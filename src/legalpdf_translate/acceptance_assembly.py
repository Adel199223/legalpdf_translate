"""Offline full-case assembly of independently verified page bundles.

Ordinary commits, packaging recoveries and explicitly reviewed derivatives keep
their own identities. Nothing here resumes a workflow, contacts a provider,
settles costs or certifies legal fidelity/rendered layout. The caller's live
evidence guard verifies historical operations and returns the exact binding
produced by ``assembly_binding``; a commit alone does not establish language.
"""
from __future__ import annotations

from collections.abc import Callable, Mapping, Sequence
from copy import deepcopy
from dataclasses import dataclass
import hashlib
from io import BytesIO
from pathlib import Path
import stat
from typing import Any
from zipfile import BadZipFile, ZipFile

from .structured_artifacts import (
    StructuredArtifactError, _decode, _hash_value, _json, _names,
    _publish_file_exclusive, _read_file, validate_commit_record,
    validate_structured_page,
)
from .types import TargetLang

ASSEMBLY_VERSION = "mixed_structured_case_assembly_v1"
_PROTOCOLS = frozenset({"legal_blocks_v2", "legal_blocks_recovery_v1",
                       "legal_blocks_reviewed_recovery_v1", "legal_blocks_legacy_reviewed_recovery_v1"})
_IGNORED_LAYOUT_POLICY = "bound_original_layout_sidecars_retained_unused_v1"


class AcceptanceAssemblyError(StructuredArtifactError):
    """Content-free failure; incomplete new output is retained, never reused."""


@dataclass(frozen=True)
class PageBundleRef:
    pages_dir: Path
    expected_commit: dict
    commit_file_sha256: str
    evidence_identity_sha256: str
    # Exact literal page filenames and hashes, never directories or patterns.
    # Empty retains the original strict rejection of both layout sidecars.
    ignored_layout_sidecars: tuple[tuple[str, str], ...] = ()


def _digest(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _layout_sidecar_pins(ref: PageBundleRef) -> dict[str, str]:
    value = ref.ignored_layout_sidecars
    allowed = {f"page_{ref.expected_commit['page_number']:04d}.{suffix}"
               for suffix in ("layout.json", "layout_eligibility.json")}
    if (type(value) is not tuple or len(value) > 2
            or any(type(row) is not tuple or len(row) != 2
                   or type(row[0]) is not str or row[0] not in allowed
                   or not _hash_value(row[1]) for row in value)
            or len({row[0] for row in value}) != len(value)
            or tuple(sorted(value)) != value):
        raise AcceptanceAssemblyError("assembly_invalid_ignored_layout_sidecars")
    return dict(value)


def _read_ignored_layout_sidecars(ref: PageBundleRef) -> dict[str, bytes]:
    pins = _layout_sidecar_pins(ref)
    result = {}
    for suffix in ("layout.json", "layout_eligibility.json"):
        name = f"page_{ref.expected_commit['page_number']:04d}.{suffix}"
        path = Path(ref.pages_dir) / name
        if name in pins:
            raw = _read_file(path)
            if _digest(raw) != pins[name]:
                raise AcceptanceAssemblyError("assembly_ignored_layout_sidecar_changed")
            result[name] = raw
        elif path.exists() or path.is_symlink():
            raise AcceptanceAssemblyError("assembly_unbound_layout_derivative")
    return result


def _committed_contents(snapshot: dict[str, bytes]) -> dict[str, bytes]:
    """Only the original four artifacts are legal writer/staging inputs."""
    return {kind: snapshot[kind] for kind in ("text", "source", "target", "commit")}


def _ignored_layout_metadata(refs: Sequence[PageBundleRef]) -> dict:
    pages = [{"page_number": ref.expected_commit["page_number"],
              "files": [{"name": name, "sha256": digest} for name, digest in _layout_sidecar_pins(ref).items()]}
             for ref in refs if ref.ignored_layout_sidecars]
    if not pages:
        return {}
    return {"ignored_layout_sidecars": {"policy": _IGNORED_LAYOUT_POLICY,
        "retained_in_original_locations": True, "used_for_assembly": False,
        "copied_to_staging": False, "pages": pages}}


def _directory(path: Path) -> None:
    """Reject links/reparse directories, including redirected ancestors."""
    if not path.is_absolute() or ".." in path.parts:
        raise AcceptanceAssemblyError("assembly_absolute_path_required")
    try:
        for item in (path, *path.parents):
            info = item.lstat()
            if not stat.S_ISDIR(info.st_mode) or getattr(info, "st_file_attributes", 0) & 0x400:
                raise AcceptanceAssemblyError("assembly_unsafe_directory")
    except OSError as exc:
        raise AcceptanceAssemblyError("assembly_directory_unavailable") from exc


def assembly_binding(page_bundles: Sequence[PageBundleRef], *, full_case_pages: Sequence[int],
                     source_file_sha256: str, lang: TargetLang,
                     preferences_sha256: str) -> dict:
    """Describe what the caller must independently verify, not an authorization.

The guard must read pinned historical source/request/response/review evidence;
returning unchecked caller values is not a substitute for that verification.
"""
    pages = list(full_case_pages)
    if (not pages or len(pages) > 10000 or any(type(n) is not int for n in pages)
            or pages != list(range(1, len(pages) + 1)) or len(page_bundles) != len(pages)):
        raise AcceptanceAssemblyError("assembly_complete_ordered_case_required")
    if (not isinstance(lang, TargetLang) or not _hash_value(source_file_sha256)
            or not _hash_value(preferences_sha256)):
        raise AcceptanceAssemblyError("assembly_invalid_case_identity")
    entries = []
    for number, ref in zip(pages, page_bundles):
        if (not isinstance(ref, PageBundleRef) or not _hash_value(ref.commit_file_sha256)
                or not _hash_value(ref.evidence_identity_sha256)):
            raise AcceptanceAssemblyError("assembly_invalid_page_reference")
        record = validate_commit_record(ref.expected_commit, page_number=number)
        if (record["source_file_sha256"] != source_file_sha256
                or record["protocol_identity"]["protocol"] not in _PROTOCOLS):
            raise AcceptanceAssemblyError("assembly_wrong_source_or_protocol")
        entries.append({"page_number": number, "commit_file_sha256": ref.commit_file_sha256,
                        "bundle_sha256": record["bundle_sha256"],
                        "evidence_identity_sha256": ref.evidence_identity_sha256})
        pins = _layout_sidecar_pins(ref)
        if pins:
            entries[-1]["ignored_layout_sidecars"] = {"policy": _IGNORED_LAYOUT_POLICY,
                "files": [{"name": name, "sha256": digest} for name, digest in pins.items()]}
    return {"version": ASSEMBLY_VERSION, "full_case_pages": pages,
            "source_file_sha256": source_file_sha256, "target_lang": lang.value,
            "preferences_sha256": preferences_sha256, "pages": entries}


def _read_bundle(ref: PageBundleRef) -> dict[str, bytes]:
    folder, record = Path(ref.pages_dir), ref.expected_commit
    _directory(folder)
    names = _names(record["page_number"])
    ignored_sidecars = _read_ignored_layout_sidecars(ref)
    commit_bytes = _read_file(folder / names["commit"], maximum=64 * 1024)
    if _digest(commit_bytes) != ref.commit_file_sha256:
        raise AcceptanceAssemblyError("assembly_commit_file_changed")
    validate_structured_page(folder, record["page_number"],
        protocol_identity=record["protocol_identity"], page_fingerprint=record["page_fingerprint"],
        expected_commit=record)
    contents = {kind: _read_file(folder / names[kind]) for kind in ("text", "source", "target")}
    for kind, content in contents.items():
        item = record["artifacts"][kind]
        if _digest(content) != item["sha256"] or len(content) != item["bytes"]:
            raise AcceptanceAssemblyError("assembly_artifact_changed")
    if _read_file(folder / names["commit"], maximum=64 * 1024) != commit_bytes:
        raise AcceptanceAssemblyError("assembly_commit_file_changed")
    if _read_ignored_layout_sidecars(ref) != ignored_sidecars:
        raise AcceptanceAssemblyError("assembly_ignored_layout_sidecar_changed")
    contents["commit"] = commit_bytes
    # Snapshot the ignored bytes, but keep them outside the four named artifacts.
    contents.update(ignored_sidecars)
    return contents


def _check_source_map(mapping: Any, records: list[dict], contents: list[dict[str, bytes]],
                      docx_sha256: str) -> None:
    if (not isinstance(mapping, dict) or mapping.get("version") != 1
            or mapping.get("docx_sha256") != docx_sha256
            or mapping.get("source_page_count") != len(records)
            or mapping.get("rendered_page_count") is not None
            or not isinstance(mapping.get("pages"), list)
            or len(mapping["pages"]) != len(records)):
        raise AcceptanceAssemblyError("assembly_source_map_mismatch")
    for record, content, page in zip(records, contents, mapping["pages"]):
        if (not isinstance(page, dict) or page.get("source_page_number") != record["page_number"]
                or page.get("structure_status") != "validated"
                or page.get("source_block_coverage_status") != "complete"
                or page.get("source_file_sha256") != record["source_file_sha256"]
                or page.get("source_sha256") != record["source_sha256"]
                or page.get("source_text_sha256") != record["source_sha256"]
                or page.get("translation_sha256") != record["translation_sha256"]):
            raise AcceptanceAssemblyError("assembly_source_map_page_mismatch")
        blocks, mapped = _decode(content["target"])["blocks"], page.get("blocks")
        if (not isinstance(mapped, list) or len(mapped) != len(blocks)
                or any(not isinstance(row, dict) or row.get("block_id") != block["id"]
                    or row.get("role") != block["role"]
                    or not isinstance(row.get("location"), dict)
                    or not isinstance(row["location"].get("kind"), str)
                    or not row["location"]["kind"]
                    for row, block in zip(mapped, blocks))):
            raise AcceptanceAssemblyError("assembly_source_map_block_mismatch")
        # This version never silently substitutes a canonical footer variant.
        if any("furniture_alias" in row for row in mapped):
            raise AcceptanceAssemblyError("assembly_furniture_alias_requires_review")


def _check_docx_text(content: bytes, mapping: dict, originals: list[dict[str, bytes]]) -> None:
    """Independent visible-text proof for the current whole-page/body cohort.

Tables, joins and furniture projection require an additional reviewed locator
contract; they are rejected, not certified by a map merely claiming coverage.
All substantive body paragraphs must have one unique source-block owner.
"""
    from docx import Document
    from docx.oxml.ns import qn
    from .docx_writer import sanitize_bidi_controls, unwrap_internal_placeholders

    try:
        with ZipFile(BytesIO(content)) as package:
            entries = package.infolist()
            if (len(entries) > 2000 or len({e.filename for e in entries}) != len(entries)
                    or sum(e.file_size for e in entries) > 128 * 1024 * 1024):
                raise AcceptanceAssemblyError("assembly_oversized_docx_package")
        document = Document(BytesIO(content))
    except (BadZipFile, KeyError, ValueError) as exc:
        raise AcceptanceAssemblyError("assembly_invalid_docx_package") from exc
    body = document._element.body
    if (len(list(body.iter(qn("w:p")))) != len(document.paragraphs)
            or any(list(body.iter(qn(tag))) for tag in ("w:tbl", "w:drawing", "w:pict", "w:fldChar"))):
        raise AcceptanceAssemblyError("assembly_unsupported_body_content")
    covered: set[int] = set()
    previous_index = -1
    for page, original in zip(mapping["pages"], originals):
        blocks = _decode(original["target"])["blocks"]
        for row, block in zip(page["blocks"], blocks):
            expected = sanitize_bidi_controls(unwrap_internal_placeholders(block["text"]))
            loc = row["location"]
            if not expected.strip() and loc == {"kind": "empty_block"}:
                continue
            index = loc.get("paragraph_index")
            if (loc.get("kind") != "body_paragraph" or set(loc) != {"kind", "paragraph_index"}
                    or type(index) is not int or not 0 <= index < len(document.paragraphs)
                    or index <= previous_index or index in covered or "joined_to_block_id" in row):
                raise AcceptanceAssemblyError("assembly_unsupported_or_duplicate_locator")
            if sanitize_bidi_controls(document.paragraphs[index].text) != expected:
                raise AcceptanceAssemblyError("assembly_docx_text_mismatch")
            covered.add(index)
            previous_index = index
    if any(sanitize_bidi_controls(p.text).strip() and i not in covered
           for i, p in enumerate(document.paragraphs)):
        raise AcceptanceAssemblyError("assembly_unmapped_docx_text")
    # Inspect package parts directly, without creating absent header/footer parts.
    for part in document.part.package.parts:
        name = str(part.partname)
        if name.startswith(("/word/header", "/word/footer")):
            element = part.element
            if list(element.iter(qn("w:tbl"))) or list(element.iter(qn("w:drawing"))):
                raise AcceptanceAssemblyError("assembly_unmapped_furniture")
            for paragraph in element.iter(qn("w:p")):
                texts = [n.text or "" for n in paragraph.iter(qn("w:t"))]
                fields = [n.text or "" for n in paragraph.iter(qn("w:instrText"))]
                if texts or fields:
                    # The unchanged writer emits exactly this generated PAGE field.
                    chars = [n.get(qn("w:fldCharType")) for n in paragraph.iter(qn("w:fldChar"))]
                    if not (name.startswith("/word/footer") and texts == ["1"]
                            and fields == [" PAGE "] and chars == ["begin", "separate", "end"]):
                        raise AcceptanceAssemblyError("assembly_unmapped_furniture")
        elif name.startswith(("/word/footnotes", "/word/endnotes", "/word/comments")):
            raise AcceptanceAssemblyError("assembly_unmapped_annotation_part")


def assemble_mixed_case(page_bundles: Sequence[PageBundleRef], *, full_case_pages: Sequence[int],
                        source_file_sha256: str, lang: TargetLang, preferences_sha256: str,
                        output_dir: Path, evidence_guard: Callable[[], Mapping], writer=None) -> Path:
    """Make one exclusive full-case review DOCX without rewriting any page.

The manifest establishes ordered bundle/copy/writer provenance, not visual or
bilingual acceptance. ``writer`` is an internal dependency-injection seam for
tests, not an arbitrary external execution hook. No original directory is used
as writer input and no partially completed output can be resumed here.
"""
    if not callable(evidence_guard):
        raise AcceptanceAssemblyError("assembly_evidence_guard_required")
    if any(not isinstance(ref, PageBundleRef) for ref in page_bundles):
        raise AcceptanceAssemblyError("assembly_invalid_page_reference")
    refs = tuple(PageBundleRef(Path(ref.pages_dir), deepcopy(ref.expected_commit),
                 ref.commit_file_sha256, ref.evidence_identity_sha256,
                 ref.ignored_layout_sidecars) for ref in page_bundles)
    binding = assembly_binding(refs, full_case_pages=full_case_pages,
        source_file_sha256=source_file_sha256, lang=lang, preferences_sha256=preferences_sha256)
    records = [ref.expected_commit for ref in refs]
    output_dir = Path(output_dir)
    _directory(output_dir.parent)
    if any(output_dir == Path(ref.pages_dir) or output_dir.is_relative_to(Path(ref.pages_dir))
           or Path(ref.pages_dir).is_relative_to(output_dir) for ref in refs):
        raise AcceptanceAssemblyError("assembly_output_overlaps_evidence")

    def guard():
        if _json(evidence_guard()) != _json(binding):
            raise AcceptanceAssemblyError("assembly_historical_binding_mismatch")

    guard()
    originals = [_read_bundle(ref) for ref in refs]
    for content in originals:
        target = _decode(content["target"])
        recovery = target.get("metadata", {}).get("recovery")
        if recovery is not None and (not isinstance(recovery, dict)
                or recovery.get("lang") != lang.value
                or recovery.get("preferences_sha256") != preferences_sha256):
            raise AcceptanceAssemblyError("assembly_recovery_language_or_preferences_mismatch")
    for record, content in zip(records, originals):
        if record["protocol_identity"]["protocol"] != "legal_blocks_v2":
            provenance = _decode(content["target"])["metadata"].get("recovery")
            if (not isinstance(provenance, dict) or provenance.get("version") != record["protocol_identity"]["protocol"]
                    or _digest(_json(provenance)) != record["protocol_identity"]["fingerprint"]
                    or _digest(_json({"recovery": record["protocol_identity"], "page": record["page_number"]}))
                        != record["page_fingerprint"] or record.get("page_result") is not None):
                raise AcceptanceAssemblyError("assembly_recovery_provenance_mismatch")
    guard()
    output_dir.mkdir(exist_ok=False)
    staging = output_dir / "pages"
    staging.mkdir()
    for ref, content in zip(refs, originals):
        guard()
        if _read_bundle(ref) != content:
            raise AcceptanceAssemblyError("assembly_original_changed")
        names = _names(ref.expected_commit["page_number"])
        for kind in ("text", "source", "target", "commit"):
            _publish_file_exclusive(staging / names[kind], content[kind])
    expected_names = {name for record in records for name in _names(record["page_number"]).values()}

    def check():
        guard()
        _directory(output_dir)
        _directory(staging)
        if {p.name for p in staging.iterdir()} != expected_names:
            raise AcceptanceAssemblyError("assembly_unexpected_staged_file")
        for ref, content in zip(refs, originals):
            if _read_bundle(ref) != content or _read_bundle(PageBundleRef(staging,
                    ref.expected_commit, ref.commit_file_sha256, ref.evidence_identity_sha256)) != _committed_contents(content):
                raise AcceptanceAssemblyError("assembly_evidence_changed")
        guard()

    check()
    output = output_dir / f"assembled_{lang.value}.docx"
    if writer is None:
        from .docx_writer import assemble_docx
        writer = assemble_docx
    stats: dict = {}
    result = writer(staging, output, lang=lang, page_breaks=False,
        page_numbers=list(binding["full_case_pages"]), partial_output=False,
        derive_source_continuations=True, stats=stats)
    check()
    if (result != output or stats.get("structured_page_count") != len(refs)
            or stats.get("structure_fallback_count") != 0):
        raise AcceptanceAssemblyError("assembly_writer_fallback_or_wrong_output")
    docx = _read_file(output, maximum=64 * 1024 * 1024)
    map_path = output.with_suffix(".source_map.json")
    map_bytes = _read_file(map_path)
    if not docx:
        raise AcceptanceAssemblyError("assembly_empty_docx")
    mapping = _decode(map_bytes)
    _check_source_map(mapping, records, originals, _digest(docx))
    _check_docx_text(docx, mapping, originals)
    if {p.name for p in output_dir.iterdir()} != {"pages", output.name, map_path.name}:
        raise AcceptanceAssemblyError("assembly_unexpected_output_file")
    check()
    if _read_file(output, maximum=64 * 1024 * 1024) != docx or _read_file(map_path) != map_bytes:
        raise AcceptanceAssemblyError("assembly_output_changed")
    manifest = {"binding": binding, "binding_sha256": _digest(_json(binding)),
        "original_commits": records, "docx_sha256": _digest(docx),
        "source_map_sha256": _digest(map_bytes), "assembly_stats": stats,
        "complete_source_page_set": True, "original_artifacts_unchanged": True,
        "visible_text_coverage": "exact_unique_body_locators",
        "workflow_resumed": False, "provider_dispatch_count": 0, "assembly_cost_usd": "0",
        "fidelity_acceptance": "not_evaluated", "rendered_layout_acceptance": "not_evaluated"}
    manifest.update(_ignored_layout_metadata(refs))
    _publish_file_exclusive(output_dir / "assembly.json", _json(manifest))
    return output
