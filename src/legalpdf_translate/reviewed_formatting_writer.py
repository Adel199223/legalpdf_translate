"""Pure FR/EN/AR DOCX builder for detached reviewed formatting projections.

This is an offline test-formatting lane, not the stored flow preference and not
render acceptance. Integration must revalidate the pinned manifest against its
actual source images and original committed bundles immediately before use.
Dataclass construction is not authority: this module defensively rechecks all
retained structures, slices, flags, identities and geometry, but cannot replace
the external historical-evidence/image-byte check with a dataclass assertion.

No file writes, provider calls, native automation, OCR, resumption or accounting.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
import hashlib
import io
import json
import math
import re
import unicodedata
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

from . import reviewed_formatting as reviewed
from .docx_writer import (
    _add_rtl_flags, _segment_rtl_placeholder_aware_runs, _set_font_properties,
    _set_ltr_run_props, _set_rtl_run_props, _set_rtl_visual_alignment,
    _wrap_ltr_run_with_lrm, sanitize_bidi_controls, unwrap_internal_placeholders,
)

WRITER_VERSION = "reviewed_formatting_writer_v1"
MAP_VERSION = "reviewed_formatting_source_map_v1"
DISPLAY_POLICY = "unwrap_placeholders_strip_bidi_crlf_v1"
AR_DISPLAY_POLICY = "unwrap_placeholders_strip_bidi_crlf_rtl_runs_lrm_v1"
SPACING_POLICY = "source_region_gap_capped_48pt_v1"
MAX_PACKAGE_BYTES = 32 * 1024 * 1024
MAX_EXPANDED_BYTES = 128 * 1024 * 1024
_ALIGNMENTS = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
               "center": WD_ALIGN_PARAGRAPH.CENTER, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
_JC = {"left": "left", "right": "right", "center": "center", "justify": "both"}
_AR_JC = {"left": "end", "right": "start", "center": "center", "justify": "both"}
_TRAILING = re.compile(r"(?:\r?\n)[ \t\r\n]*\Z")
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


class ReviewedFormattingWriterError(ValueError):
    """Content-free failure; never return document text in an exception."""


@dataclass(frozen=True, slots=True)
class ReviewedDocxArtifact:
    docx_bytes: bytes
    source_map_bytes: bytes
    projection_sha256: str


def _fail(code):
    raise ReviewedFormattingWriterError(code)


def _sha(raw):
    return hashlib.sha256(raw).hexdigest()


def _json(value):
    try:
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
                          allow_nan=False).encode("utf-8")
    except (ValueError, TypeError, UnicodeError, RecursionError):
        _fail("invalid_json_value")


def _decode(raw):
    if type(raw) is not bytes or not 0 < len(raw) <= MAX_EXPANDED_BYTES:
        _fail("invalid_json_bytes")
    try:
        return json.loads(raw.decode("utf-8"), object_pairs_hook=reviewed._unique,
                          parse_constant=lambda _: _fail("invalid_json_constant"))
    except (ValueError, UnicodeError, RecursionError):
        _fail("invalid_json_bytes")


def _fragment_input(fragment):
    return {"rendering_id": fragment.rendering_id, "parent_block_id": fragment.parent_block_id,
        "source_range": list(fragment.source_range), "target_range": list(fragment.target_range),
        "source_text_sha256": fragment.source_text_sha256, "target_text_sha256": fragment.target_text_sha256,
        "role": fragment.role, "alignment": fragment.alignment, "bold": fragment.bold, "italic": fragment.italic,
        "bbox_px": list(fragment.bbox_px), "review_note": fragment.review_note}


def _checked_projection(projection, *, expected_reviewer_kind="ai_test_review"):
    """Recheck consistency, not physical image bytes or reviewer authority."""
    if type(expected_reviewer_kind) is not str or expected_reviewer_kind not in {"ai_test_review", "operator_review"}:
        _fail("invalid_expected_reviewer_kind")
    if type(projection) is not reviewed.ReviewedFormattingProjection:
        _fail("unsupported_projection_type")
    if projection.target_lang not in {"FR", "EN", "AR"}:
        _fail("target_language_requires_separate_writer_profile")
    regions = projection.version == reviewed.REGION_VERSION
    if type(projection.boundary_policy) is not str or projection.boundary_policy not in ({reviewed.BOUNDARY_POLICY, reviewed.CELL_BOUNDARY_POLICY}
                                         if regions else {reviewed.BOUNDARY_POLICY}):
        _fail("unsupported_projection_state")
    if projection.spacing_policy is not None and (not regions or projection.spacing_policy != reviewed.SOURCE_GAP_POLICY):
        _fail("unsupported_projection_spacing")
    flags = {"version": reviewed.REGION_VERSION if regions else reviewed.VERSION,
        "policy": reviewed.REGION_POLICY if regions else reviewed.POLICY,
        "boundary_policy": projection.boundary_policy, "offset_unit": "unicode_codepoint",
        "reviewer_kind": expected_reviewer_kind, "geometry_basis": "reviewed_image_regions_not_ocr",
        "layout_review_required": True, "rendered_layout_acceptance": "not_evaluated", "rendered_page_count": None}
    if any(getattr(projection, key) != value or type(getattr(projection, key)) is not type(value)
           for key, value in flags.items()):
        _fail("unsupported_projection_state")
    if regions:
        if projection.folio_policy != reviewed.FOLIO_POLICY or type(projection.document_groups_json) is not bytes:
            _fail("incomplete_region_projection")
        groups = _decode(projection.document_groups_json)
        if reviewed._canonical(groups) != projection.document_groups_json:
            _fail("noncanonical_document_groups")
    elif projection.document_groups_json is not None or projection.folio_policy is not None:
        _fail("unsupported_projection_state")
    for value in (projection.manifest_sha256, projection.source_file_sha256, projection.preferences_sha256):
        reviewed._hash(value)
    if type(projection.pages) is not tuple or not 0 < len(projection.pages) <= reviewed.MAX_PAGES:
        _fail("invalid_projection_pages")
    descriptor = {**flags, "manifest_sha256": projection.manifest_sha256,
        "source_file_sha256": projection.source_file_sha256, "preferences_sha256": projection.preferences_sha256,
        "target_lang": projection.target_lang, "pages": []}
    if projection.spacing_policy is not None:
        descriptor["spacing_policy"] = projection.spacing_policy
    commits, bundles, total_bytes = set(), set(), 0
    for number, page in enumerate(projection.pages, 1):
        if (type(page) is not reviewed.ReviewedFormattingPage or type(page.page_number) is not int
                or type(page.section_index) is not int or page.page_number != number or page.section_index != number - 1
                or page.source_uncertain is not True or page.geometry_status != "not_verified"
                or page.paper_size_basis != "a4_assumed" or page.original_parent_separator != "\n"
                or type(page.fragments) is not tuple or not 0 < len(page.fragments) <= reviewed.MAX_FRAGMENTS):
            _fail("invalid_projection_page")
        if regions:
            if type(page.region_layout_json) is not bytes:
                _fail("missing_region_layout")
            region_layout = _decode(page.region_layout_json)
            if reviewed._canonical(region_layout) != page.region_layout_json:
                _fail("noncanonical_region_layout")
        else:
            if page.region_layout_json is not None or page.folio_fragment_id is not None:
                _fail("unsupported_projection_page_state")
            region_layout = None
        for value in (page.commit_file_sha256, page.bundle_sha256, page.source_image_sha256,
                      page.source_structure_sha256, page.target_structure_sha256):
            reviewed._hash(value)
        if page.commit_file_sha256 in commits or page.bundle_sha256 in bundles:
            _fail("duplicate_original_commit")
        commits.add(page.commit_file_sha256)
        bundles.add(page.bundle_sha256)
        source, target = _decode(page.source_structure_json), _decode(page.target_structure_json)
        total_bytes += len(page.source_structure_json) + len(page.target_structure_json)
        if total_bytes > MAX_EXPANDED_BYTES:
            _fail("oversized_projection")
        bound = reviewed.FormattingPageInput(source, target, page.commit_file_sha256, page.bundle_sha256, b"")
        left, right, source_json, target_json = reviewed._bound_pair(bound, number=number,
                                                               source_hash=projection.source_file_sha256)
        if (source_json != page.source_structure_json or target_json != page.target_structure_json
                or _sha(source_json) != page.source_structure_sha256 or _sha(target_json) != page.target_structure_sha256):
            _fail("retained_structure_hash_mismatch")
        if (type(page.image_size_px) is not tuple or len(page.image_size_px) != 2
                or any(type(n) is not int or n <= 0 for n in page.image_size_px)
                or math.prod(page.image_size_px) > reviewed.MAX_IMAGE_PIXELS
                or type(page.page_size_pt) is not tuple or len(page.page_size_pt) != 2
                or page.page_size_pt != (left["width_pt"], left["height_pt"])
                or any(abs(actual - expected) > .01 for actual, expected in zip(page.page_size_pt, (Cm(21).pt, Cm(29.7).pt)))):
            _fail("unsupported_or_changed_source_frame")
        identity = left.get("metadata", {}).get("source_page_identity", {})
        review = left.get("metadata", {}).get("reviewed_source", {})
        if (not isinstance(identity, dict) or not isinstance(review, dict)
                or identity.get("image_sha256") != page.source_image_sha256
                or identity.get("source_file_sha256") != projection.source_file_sha256
                or identity.get("source_type") != "browser_pdf_image"
                or identity.get("paper_size_basis") != "a4_assumed"
                or review.get("review_image_sha256") != page.source_image_sha256
                or review.get("geometry_status") != page.geometry_status):
            _fail("retained_source_image_binding_mismatch")
        if any(type(fragment) is not reviewed.ReviewedFormattingFragment
               or type(fragment.parent_uncertain) is not bool for fragment in page.fragments):
            _fail("invalid_fragment_type")
        rows = [_fragment_input(fragment) for fragment in page.fragments]
        recomputed = reviewed._fragments(rows, left, right, width=page.image_size_px[0],
                                         height=page.image_size_px[1], total=len(projection.pages),
                                         region_layout=region_layout, document_local_folios=regions,
                                         boundary_policy=projection.boundary_policy)
        if recomputed != page.fragments:
            _fail("forged_fragment_projection")
        if not any(fragment.role in {"body", "signature"} for fragment in page.fragments):
            _fail("source_page_without_body_is_unsupported")
        descriptor["pages"].append({"page_number": number, "section_index": number - 1,
            "commit_file_sha256": page.commit_file_sha256, "bundle_sha256": page.bundle_sha256,
            "source_structure_sha256": page.source_structure_sha256, "target_structure_sha256": page.target_structure_sha256,
            "source_image_sha256": page.source_image_sha256, "image_size_px": list(page.image_size_px),
            "page_size_pt": list(page.page_size_pt), "paper_size_basis": page.paper_size_basis,
            "source_uncertain": page.source_uncertain, "geometry_status": page.geometry_status,
            "original_parent_separator": page.original_parent_separator, "fragments": rows})
        if regions:
            descriptor["pages"][-1].update(region_layout=region_layout, folio_fragment_id=page.folio_fragment_id)
    if regions:
        folios = reviewed._validated_folios(groups, pages=projection.pages,
            source_file_sha256=projection.source_file_sha256, target_lang=projection.target_lang,
            boundary_policy=projection.boundary_policy)
        descriptor.update(document_groups=groups, folio_policy=projection.folio_policy,
                          validated_folios=[asdict(row) for row in folios])
    return descriptor, _sha(_json(descriptor))


def canonical_fragment_display(raw, *, target_lang="FR"):
    """Return display text and exact raw trailing separator ownership.

    Structural paragraph endings/spacing represent the owned trailing separator
    sequence. All other spaces, tabs and interior line boundaries are retained;
    existing placeholder/bidi display cleanup and CRLF folding are explicit.
    """
    suffix = _TRAILING.search(raw)
    cut = suffix.start() if suffix else len(raw)
    content, separators = raw[:cut], raw[cut:]
    display = sanitize_bidi_controls(unwrap_internal_placeholders(content)).replace("\r\n", "\n")
    if not display.strip():
        _fail("empty_fragment_display")
    if target_lang not in {"FR", "EN", "AR"}:
        _fail("target_language_requires_separate_writer_profile")
    if target_lang != "AR" and any(unicodedata.bidirectional(char) in {"R", "AL"} for char in display):
        _fail("rtl_text_requires_separate_writer_profile")
    if any(not (ord(char) in {9, 10} or 32 <= ord(char) <= 0xD7FF
                or 0xE000 <= ord(char) <= 0xFFFD or 0x10000 <= ord(char) <= 0x10FFFF) for char in display):
        _fail("unsupported_xml_text_character")
    return display, separators, cut


def _arabic_runs(raw):
    """Derive exact OOXML runs from raw protected text, not a cleaned copy.

    The visible display hash excludes our deterministic LRM pairs. Saved bidi
    controls are removed first; the checker permits only the exact pairs and
    directional boundaries regenerated here. No normalization of graphemes,
    rewriting of names/clocks, or geometry-based text shortening is allowed.
    """
    display, _, cut = canonical_fragment_display(raw, target_lang="AR")
    runs, mixed = _segment_rtl_placeholder_aware_runs(
        raw[:cut].replace("\r\n", "\n"), strip_bidi_controls=True,
    )
    if (not runs or any(kind not in {"rtl", "ltr"} for kind, _ in runs)
            or "".join(chunk for _, chunk in runs) != display):
        _fail("arabic_run_projection_mismatch")
    result = tuple((kind, _wrap_ltr_run_with_lrm(chunk) if mixed and kind == "ltr" else chunk)
                   for kind, chunk in runs if chunk)
    if not result or sanitize_bidi_controls("".join(chunk for _, chunk in result)) != display:
        _fail("arabic_run_projection_mismatch")
    return result


def _plan(projection, *, expected_reviewer_kind="ai_test_review"):
    descriptor, identity = _checked_projection(projection, expected_reviewer_kind=expected_reviewer_kind)
    arabic = projection.target_lang == "AR"
    mapping = {"version": MAP_VERSION, "writer_version": WRITER_VERSION,
        "policy": projection.policy, "projection_sha256": identity, "manifest_sha256": projection.manifest_sha256,
        "source_file_sha256": projection.source_file_sha256, "target_lang": projection.target_lang,
        "preferences_sha256": projection.preferences_sha256, "source_page_count": len(projection.pages),
        "rendered_page_count": None, "layout_review_required": True, "rendered_layout_acceptance": "not_evaluated",
        "source_geometry_status": "not_verified", "physical_source_evidence_recheck": "required_at_integration",
        "typography": {"font": "Arial" if arabic else "Times New Roman", "size_pt": 11 if arabic else 10.5,
                       "page_size_cm": [21, 29.7],
                       "margins_cm": {"left": 1.7, "right": 1.7, "top": 1.5, "bottom": 1.5}},
        "footer_arrangement": "separate_literal_folio_paragraph_v1", "pages": []}
    # Preserve the strict v1/v2 map shape. Explicit operator output retains its
    # provenance in addition to the reviewer-bound projection hash.
    if expected_reviewer_kind != "ai_test_review":
        mapping["reviewer_kind"] = expected_reviewer_kind
    body_index = 0
    for page, data in zip(projection.pages, descriptor["pages"]):
        row = {key: value for key, value in data.items() if key != "fragments"}
        row.update(section_parts={"header": f"/word/header{page.page_number}.xml",
                                  "footer": f"/word/footer{page.page_number}.xml"}, fragments=[])
        counts, last_bottom = {"header": 0, "footer": 0}, {}
        for fragment in page.fragments:
            part = "header" if fragment.role == "header" else "footer" if fragment.role in {"footer", "folio"} else "body"
            display, separators, cut = canonical_fragment_display(fragment.target_text, target_lang=projection.target_lang)
            if arabic:
                _arabic_runs(fragment.target_text)
            if part == "body":
                location = {"kind": "body_paragraph", "part_uri": "/word/document.xml", "paragraph_index": body_index,
                            "section_index": page.section_index}
                body_index += 1
            else:
                location = {"kind": "section_" + part, "part_uri": row["section_parts"][part],
                            "paragraph_index": counts[part], "section_index": page.section_index}
                counts[part] += 1
            scale = page.page_size_pt[1] / page.image_size_px[1]
            if part in last_bottom:
                gap = max(0, (fragment.bbox_px[1] - last_bottom[part]) * scale)
            elif part == "body":
                gap = max(0, fragment.bbox_px[1] * scale - Cm(1.5).pt)
            else:
                gap = 0.0  # Furniture anchoring needs actual render review.
            last_bottom[part] = max(last_bottom.get(part, 0), fragment.bbox_px[3])
            spacing = {"basis": SPACING_POLICY, "source_gap_pt": round(gap, 6),
                "space_before_pt": round(min(48.0, gap) * 20) / 20, "space_after_pt": 0,
                "gap_capped": gap > 48.0}
            row["fragments"].append({**_fragment_input(fragment), "parent_uncertain": fragment.parent_uncertain,
                "location": location, "display_policy": AR_DISPLAY_POLICY if arabic else DISPLAY_POLICY,
                "display_text_sha256": _sha(display.encode()), "raw_content_end": cut,
                "raw_content_sha256": _sha(fragment.target_text[:cut].encode()),
                "trailing_separators": separators, "trailing_separators_sha256": _sha(separators.encode()),
                "separator_representation": "owned_paragraph_boundary_and_source_spacing",
                "interior_crlf_count": fragment.target_text[:cut].count("\r\n"), "spacing": spacing})
        mapping["pages"].append(row)
    return mapping


def _base_document(target_lang="FR"):
    document = Document()
    normal = document.styles["Normal"]
    if target_lang == "AR":
        _set_font_properties(normal, name="Arial", size=11)
    else:
        normal.font.name, normal.font.size = "Times New Roman", Pt(10.5)
        normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    document.settings.odd_and_even_pages_header_footer = False
    for name in ("author", "last_modified_by", "comments", "title", "subject", "keywords"):
        setattr(document.core_properties, name, "")
    return document


def _configure_section(section):
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.7)
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.header_distance = section.footer_distance = Cm(.7)
    section.different_first_page_header_footer = False
    for name in ("header", "footer"):
        part = getattr(section, name)
        part.is_linked_to_previous = False
        for child in list(part._element):
            part._element.remove(child)


def _new_section(document):
    """Attach the break to the existing last paragraph; add no empty body row."""
    previous = document.paragraphs[-1]
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    empty_break = document.paragraphs[-1]
    section_properties = empty_break._p.get_or_add_pPr().find(qn("w:sectPr"))
    if empty_break.text or section_properties is None:
        _fail("unexpected_section_break_shape")
    empty_break._p.getparent().remove(empty_break._p)
    previous._p.get_or_add_pPr().append(section_properties)
    return section


def _paragraph(paragraph, fragment, row, *, target_lang="FR"):
    paragraph.style = "Normal"
    paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = "Normal"
    paragraph.alignment = _ALIGNMENTS[fragment.alignment]
    if target_lang == "AR":
        _add_rtl_flags(paragraph)
        _set_rtl_visual_alignment(paragraph, fragment.alignment)
    formatting = paragraph.paragraph_format
    formatting.space_before = Pt(row["spacing"]["space_before_pt"])
    formatting.space_after = Pt(0)
    formatting.line_spacing = 1.0
    formatting.keep_with_next = formatting.keep_together = formatting.widow_control = False
    if target_lang == "AR":
        for kind, text in _arabic_runs(fragment.target_text):
            run = paragraph.add_run(text)
            _set_font_properties(run, name="Arial", size=11)
            run.bold, run.italic = fragment.bold, fragment.italic
            run.font.cs_bold, run.font.cs_italic = fragment.bold, fragment.italic
            if kind == "ltr":
                _set_ltr_run_props(run)
            else:
                _set_rtl_run_props(run, bidi_lang="ar-SA")
        return
    display, _, _ = canonical_fragment_display(fragment.target_text)
    run = paragraph.add_run(display)
    run.font.name, run.font.size = "Times New Roman", Pt(10.5)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold, run.italic = fragment.bold, fragment.italic
    rpr = run._r.get_or_add_rPr()
    for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
        rpr.rFonts.set(qn("w:" + kind), "Times New Roman")
    size = OxmlElement("w:szCs")
    size.set(qn("w:val"), "21")
    rpr.append(size)


def _finish_empty_parts(section):
    for name in ("header", "footer"):
        part = getattr(section, name)
        if not len(part._element):
            part.add_paragraph()


def _save(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _scaffold(count, *, target_lang="FR"):
    """Trusted, text-free package infrastructure for this exact writer profile."""
    document = _base_document(target_lang)
    for index in range(count):
        section = _new_section(document) if index else document.sections[0]
        _configure_section(section)
        document.add_paragraph()  # Reference package only; never delivered.
        _finish_empty_parts(section)
    return document, _save(document)


def _package(raw):
    if type(raw) is not bytes or not 0 < len(raw) <= MAX_PACKAGE_BYTES:
        _fail("invalid_docx_bytes")
    try:
        with ZipFile(io.BytesIO(raw)) as package:
            entries = package.infolist()
            if (len(entries) > 2 * reviewed.MAX_PAGES + 100 or len({e.filename for e in entries}) != len(entries)
                    or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES
                    or any(e.flag_bits & 1 or "\\" in e.filename or e.filename.startswith("/")
                           or any(part in {"", ".", ".."} for part in e.filename.split("/")) for e in entries)):
                _fail("unsupported_docx_package")
            contents = {entry.filename: package.read(entry.filename) for entry in entries}
    except (BadZipFile, KeyError, OSError, RuntimeError):
        _fail("invalid_docx_package")
    for name, content in contents.items():
        if name.endswith((".xml", ".rels")) and (b"\x00" in content or b"<!DOCTYPE" in content or b"<!ENTITY" in content):
            _fail("unsupported_xml_declaration")
    return contents


def _xml(raw):
    try:
        parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)
        result = etree.fromstring(raw, parser=parser)
        if result.getprevious() is not None or result.getnext() is not None:
            _fail("unsupported_xml_prolog")
        return result
    except (etree.XMLSyntaxError, ValueError):
        _fail("invalid_part_xml")


def _properties(element, expected):
    if element is None or element.attrib or element.text or element.tail or len(element) != len(expected):
        _fail("unsupported_text_properties")
    actual = {}
    for child in element:
        if child.tag in actual or len(child) or child.text or child.tail:
            _fail("unsupported_text_properties")
        actual[child.tag] = dict(child.attrib)
    if actual != {qn("w:" + tag): {qn("w:" + key): value for key, value in attrs.items()}
                  for tag, attrs in expected.items()}:
        _fail("unsupported_text_properties")


def _paragraph_text(element, row, fragment, *, target_lang="FR"):
    if element.tag != qn("w:p") or element.attrib or element.text or element.tail:
        _fail("unsupported_paragraph")
    children = list(element)
    if not children or children[0].tag != qn("w:pPr") or any(child.tag != qn("w:r") for child in children[1:]):
        _fail("unsupported_paragraph_content")
    ppr = children[0]
    if ppr.text or ppr.tail:
        _fail("unowned_paragraph_property_text")
    structural = [child for child in ppr if child.tag == qn("w:sectPr")]
    if len(structural) > 1:
        _fail("duplicate_section_properties")
    properties = etree.fromstring(etree.tostring(ppr))
    for child in list(properties):
        if child.tag == qn("w:sectPr"):
            properties.remove(child)
    arabic = target_lang == "AR"
    expected_properties = {"pStyle": {"val": "Normal"}, "jc": {"val": (_AR_JC if arabic else _JC)[fragment.alignment]},
        "spacing": {"before": str(round(row["spacing"]["space_before_pt"] * 20)), "after": "0", "line": "240", "lineRule": "auto"},
        "keepNext": {"val": "0"}, "keepLines": {"val": "0"}, "widowControl": {"val": "0"}}
    if arabic:
        expected_properties["bidi"] = {"val": "1"}
    _properties(properties, expected_properties)
    expected_runs = _arabic_runs(fragment.target_text) if arabic else None
    if arabic and len(children) - 1 != len(expected_runs):
        _fail("arabic_run_count_mismatch")
    result = []
    for index, run in enumerate(children[1:]):
        if run.attrib or run.text or run.tail or len(run) < 2 or run[0].tag != qn("w:rPr"):
            _fail("unsupported_run")
        expected_run_properties = {"rFonts": {name: "Arial" if arabic else "Times New Roman"
                                               for name in ("ascii", "hAnsi", "eastAsia", "cs")},
            "b": {} if fragment.bold else {"val": "0"}, "i": {} if fragment.italic else {"val": "0"},
            "color": {"val": "000000"}, "sz": {"val": "22" if arabic else "21"},
            "szCs": {"val": "22" if arabic else "21"}}
        if arabic:
            kind, expected_text = expected_runs[index]
            expected_run_properties.update(
                bCs={} if fragment.bold else {"val": "0"}, iCs={} if fragment.italic else {"val": "0"},
                rtl={"val": "0" if kind == "ltr" else "1"},
                lang={"val": "en-US"} if kind == "ltr" else {"val": "ar-SA", "bidi": "ar-SA"})
        _properties(run[0], expected_run_properties)
        run_text = []
        for node in list(run)[1:]:
            if len(node) or node.tail:
                _fail("unsupported_nested_run_content")
            if node.tag == qn("w:t") and dict(node.attrib) in ({}, {_XML_SPACE: "preserve"}):
                if arabic and node.text and node.text.strip() != node.text and node.get(_XML_SPACE) != "preserve":
                    _fail("arabic_run_whitespace_not_preserved")
                run_text.append(node.text or "")
            elif node.tag in {qn("w:tab"), qn("w:br")} and not node.attrib and not node.text:
                run_text.append("\t" if node.tag == qn("w:tab") else "\n")
            else:
                _fail("unsupported_run_content")
        text = "".join(run_text)
        if arabic and text != expected_text:
            _fail("arabic_directional_run_text_mismatch")
        result.append(text)
    display, _, _ = canonical_fragment_display(fragment.target_text, target_lang=target_lang)
    actual_display = sanitize_bidi_controls("".join(result)) if arabic else "".join(result)
    if actual_display != display:
        _fail("fragment_display_text_mismatch")


def _validate(raw, source_map_bytes, projection, *, expected_reviewer_kind="ai_test_review"):
    expected = _plan(projection, expected_reviewer_kind=expected_reviewer_kind)
    expected["docx_sha256"] = _sha(raw)
    if _json(_decode(source_map_bytes)) != _json(expected):
        _fail("source_map_mismatch")
    actual = _package(raw)
    scaffold, scaffold_bytes = _scaffold(len(projection.pages), target_lang=projection.target_lang)
    baseline = _package(scaffold_bytes)
    text_parts = {"word/document.xml", *(uri[1:] for page in expected["pages"] for uri in page["section_parts"].values())}
    if set(actual) != set(baseline) or any(actual[name] != baseline[name] for name in actual if name not in text_parts):
        _fail("unsupported_or_changed_package_infrastructure")
    roots = {"/" + name: _xml(actual[name]) for name in text_parts}
    if any(node.tail or (node.text and node.tag != qn("w:t"))
           for root in roots.values() for node in root.iter()):
        _fail("unowned_xml_text_or_tail")
    document = roots["/word/document.xml"]
    reference_document = _xml(baseline["word/document.xml"])
    if (document.tag != qn("w:document") or len(document) != 1 or document[0].tag != qn("w:body")
            or document.text or document.tail or dict(document.attrib) != dict(reference_document.attrib)):
        _fail("unsupported_document_body")
    body = document[0]
    if (body.attrib or body.text or body.tail or not len(body) or body[-1].tag != qn("w:sectPr")
            or any(node.tag != qn("w:p") for node in list(body)[:-1])):
        _fail("unsupported_document_body")
    paragraphs = {"/word/document.xml": list(body)[:-1]}
    for uri, root in roots.items():
        if uri == "/word/document.xml":
            continue
        tag = qn("w:hdr") if uri.startswith("/word/header") else qn("w:ftr")
        reference_root = _xml(baseline[uri[1:]])
        if (root.tag != tag or root.text or root.tail or dict(root.attrib) != dict(reference_root.attrib)
                or any(child.tag != qn("w:p") for child in root)):
            _fail("unsupported_section_furniture")
        paragraphs[uri] = list(root)
    expected_counts = {uri: 0 for uri in roots}
    actual_sections, expected_ends = [], []
    for page, page_map in zip(projection.pages, expected["pages"]):
        body_locations = []
        for fragment, row in zip(page.fragments, page_map["fragments"]):
            location = row["location"]
            uri, index = location["part_uri"], location["paragraph_index"]
            if index != expected_counts[uri] or not 0 <= index < len(paragraphs[uri]):
                _fail("missing_duplicate_or_reordered_locator")
            expected_counts[uri] += 1
            _paragraph_text(paragraphs[uri][index], row, fragment, target_lang=projection.target_lang)
            if uri == "/word/document.xml":
                body_locations.append(index)
        if page.page_number < len(projection.pages):
            expected_ends.append(body_locations[-1])
    for uri, items in paragraphs.items():
        count = expected_counts[uri]
        if count == 0 and uri != "/word/document.xml":
            if len(items) != 1 or items[0].attrib or len(items[0]) or items[0].text or items[0].tail:
                _fail("unowned_empty_furniture")
        elif len(items) != count:
            _fail("unowned_paragraph")
    for index, paragraph in enumerate(paragraphs["/word/document.xml"]):
        section = paragraph.find("./" + qn("w:pPr") + "/" + qn("w:sectPr"))
        if section is not None:
            if index not in expected_ends:
                _fail("section_boundary_mismatch")
            actual_sections.append(section)
        elif index in expected_ends:
            _fail("missing_section_boundary")
    actual_sections.append(body[-1])
    if len(actual_sections) != len(scaffold.sections):
        _fail("section_count_mismatch")
    for actual_section, trusted_section in zip(actual_sections, scaffold.sections):
        if etree.tostring(actual_section, method="c14n") != etree.tostring(trusted_section._sectPr, method="c14n"):
            _fail("section_settings_or_part_binding_mismatch")


def validate_reviewed_docx(docx_bytes: bytes, source_map_bytes: bytes, *,
        projection: reviewed.ReviewedFormattingProjection, expected_reviewer_kind: str = "ai_test_review") -> None:
    """Independently inspect exact package text, properties, sections and map."""
    try:
        if type(projection) is reviewed.ReviewedFormattingProjection and projection.version == reviewed.REGION_VERSION:
            from .reviewed_region_writer import validate_region_docx
            validate_region_docx(docx_bytes, source_map_bytes, projection=projection,
                                 expected_reviewer_kind=expected_reviewer_kind)
            return
        _validate(docx_bytes, source_map_bytes, projection, expected_reviewer_kind=expected_reviewer_kind)
    except ReviewedFormattingWriterError:
        raise
    except (reviewed.ReviewedFormattingError, TypeError, ValueError, KeyError, AttributeError, OverflowError):
        _fail("invalid_reviewed_docx_contract")


def build_reviewed_docx(projection: reviewed.ReviewedFormattingProjection, *,
        expected_reviewer_kind: str = "ai_test_review") -> ReviewedDocxArtifact:
    """Build and independently validate in memory; caller owns publication."""
    try:
        if type(projection) is reviewed.ReviewedFormattingProjection and projection.version == reviewed.REGION_VERSION:
            from .reviewed_region_writer import build_region_docx
            return build_region_docx(projection, expected_reviewer_kind=expected_reviewer_kind)
        mapping = _plan(projection, expected_reviewer_kind=expected_reviewer_kind)
        document = _base_document(projection.target_lang)
        for index, (page, page_map) in enumerate(zip(projection.pages, mapping["pages"])):
            section = _new_section(document) if index else document.sections[0]
            _configure_section(section)
            for fragment, row in zip(page.fragments, page_map["fragments"]):
                kind = row["location"]["kind"]
                container = (document if kind == "body_paragraph" else section.header if kind == "section_header" else section.footer)
                paragraph = container.add_paragraph()
                _paragraph(paragraph, fragment, row, target_lang=projection.target_lang)
            _finish_empty_parts(section)
        raw = _save(document)
        mapping["docx_sha256"] = _sha(raw)
        source_map_bytes = _json(mapping)
        validate_reviewed_docx(raw, source_map_bytes, projection=projection,
                               expected_reviewer_kind=expected_reviewer_kind)
        return ReviewedDocxArtifact(raw, source_map_bytes, mapping["projection_sha256"])
    except ReviewedFormattingWriterError:
        raise
    except (reviewed.ReviewedFormattingError, TypeError, ValueError, KeyError, AttributeError, OverflowError):
        _fail("invalid_reviewed_writer_input")
