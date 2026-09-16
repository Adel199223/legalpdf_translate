"""In-memory v2 source-region DOCX profile and independent package checker.

Tables implement explicitly reviewed physical columns, never guessed layout.
All mapped text is checked before it is removed from an inspection copy; the
remaining package must equal a separately rebuilt text-free structural scaffold.
No font fitting, content shortening, native export, or render acceptance occurs.
"""
from copy import deepcopy

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

from . import reviewed_formatting as reviewed
from . import reviewed_formatting_writer as base
from .reviewed_regions import GUTTER_VERSION


WRITER_VERSION = "reviewed_region_writer_v1"
MAP_VERSION = "reviewed_region_source_map_v1"
SPACING_POLICY = "source_region_block_cell_gap_capped_48pt_v1"
TABLE_POLICY = "explicit_ltr_fixed_percent_columns_zero_cell_margins_v1"
GUTTER_TABLE_POLICY = "explicit_ltr_fixed_percent_columns_reviewed_gutters_v2"
BREAK_POLICY = "do_not_expand_shift_return_v1"


def _spacing(gap, *, preserve_source=False):
    return {"basis": reviewed.SOURCE_GAP_POLICY if preserve_source else SPACING_POLICY,
        "source_gap_pt": round(gap, 6),
        "space_before_pt": round((gap if preserve_source else min(48.0, gap)) * 20) / 20,
        "space_after_pt": 0, "gap_capped": not preserve_source and gap > 48.0}


def _bounds(ids, fragments):
    boxes = [fragments[rid].bbox_px for rid in ids]
    if not boxes:
        base._fail("empty_region_bounds")
    return (min(b[0] for b in boxes), min(b[1] for b in boxes),
            max(b[2] for b in boxes), max(b[3] for b in boxes))


def _row_ids(row):
    return [rid for cell in row["cells"] for rid in cell["fragment_ids"]]


def _block_ids(block):
    return ([block["fragment_id"]] if block["kind"] == "paragraph"
            else [rid for row in block["rows"] for rid in _row_ids(row)])


def _widths(percentages):
    """Deterministic integer twips; last column owns rounding remainder."""
    total = Cm(17.6).twips
    result = [round(total * percent / 100) for percent in percentages[:-1]]
    return [*result, total - sum(result)]


def _column_gutters(block, page):
    """Split each approved physical source gap across its two cell edges."""
    scale = page.page_size_pt[0] / page.image_size_px[0]
    gaps = [round(value * scale * 20) for value in block["column_gaps_px"]]
    margins = [{"left": 0, "right": 0} for _ in block["column_widths"]]
    for column, gap in enumerate(gaps):
        margins[column]["right"] = gap // 2
        margins[column + 1]["left"] = gap - gap // 2
    if any(row["left"] + row["right"] >= width
           for row, width in zip(margins, _widths(block["column_widths"]))):
        base._fail("column_gutter_leaves_no_text_width")
    return gaps, margins


def _plan(projection, *, expected_reviewer_kind="ai_test_review"):
    descriptor, _ = base._checked_projection(projection, expected_reviewer_kind=expected_reviewer_kind)
    if projection.version != reviewed.REGION_VERSION or projection.policy != reviewed.REGION_POLICY:
        base._fail("region_writer_requires_v2_projection")
    # Retain the established exact raw/display/slice/separator bindings. Replace
    # every v1 positional and spacing assumption with the explicit region map.
    mapping = base._plan(projection, expected_reviewer_kind=expected_reviewer_kind)
    mapping.update(version=MAP_VERSION, writer_version=WRITER_VERSION,
        table_policy=TABLE_POLICY, manual_line_break_policy=BREAK_POLICY,
        document_groups=descriptor["document_groups"], folio_policy=descriptor["folio_policy"],
        validated_folios=descriptor["validated_folios"])
    preserve_source = projection.spacing_policy == reviewed.SOURCE_GAP_POLICY
    if preserve_source:
        mapping["spacing_policy"] = projection.spacing_policy
    child_index = paragraph_index = table_index = 0
    for page, page_map in zip(projection.pages, mapping["pages"]):
        fragments = {f.rendering_id: f for f in page.fragments}
        rows = {row["rendering_id"]: row for row in page_map["fragments"]}
        layout = page_map["region_layout"]
        has_gutters = layout["version"] == GUTTER_VERSION
        if has_gutters:
            mapping["table_policy"] = GUTTER_TABLE_POLICY
        scale = page.page_size_pt[1] / page.image_size_px[1]
        page_map.update(structural_paragraphs=[], tables=[])
        assigned = set()

        def assign(rid, location, gap):
            if rid in assigned:
                base._fail("duplicate_region_locator")
            assigned.add(rid)
            rows[rid]["location"] = location
            rows[rid]["spacing"] = _spacing(gap, preserve_source=preserve_source)

        for part in ("header", "footer"):
            previous = None
            for index, rid in enumerate(layout[part]):
                box = fragments[rid].bbox_px
                gap = 0.0 if previous is None else max(0, (box[1] - previous) * scale)
                assign(rid, {"kind": "section_" + part, "part_uri": page_map["section_parts"][part],
                    "paragraph_index": index, "section_index": page.section_index}, gap)
                previous = max(previous or 0, box[3])

        previous_block_bottom = None
        for block in layout["body"]:
            block_box = _bounds(_block_ids(block), fragments)
            gap = (max(0, block_box[1] * scale - Cm(1.5).pt) if previous_block_bottom is None
                   else max(0, (block_box[1] - previous_block_bottom) * scale))
            if block["kind"] == "paragraph":
                assign(block["fragment_id"], {"kind": "body_paragraph", "part_uri": "/word/document.xml",
                    "body_child_index": child_index, "paragraph_index": paragraph_index,
                    "section_index": page.section_index}, gap)
                paragraph_index += 1
            else:
                table_location = {"part_uri": "/word/document.xml", "body_child_index": child_index,
                    "table_index": table_index, "table_id": block["table_id"], "section_index": page.section_index}
                page_map["tables"].append({**table_location, "column_widths_percent": block["column_widths"],
                    "column_widths_twips": _widths(block["column_widths"]), "bbox_px": list(block_box)})
                if has_gutters:
                    gap_twips, margins = _column_gutters(block, page)
                    page_map["tables"][-1].update(column_gaps_px=list(block["column_gaps_px"]),
                        column_gaps_twips=gap_twips, cell_margins_twips=margins)
                previous_row_bottom = None
                for row_index, table_row in enumerate(block["rows"]):
                    row_box = _bounds(_row_ids(table_row), fragments)
                    row_gap = (gap if previous_row_bottom is None
                               else max(0, (row_box[1] - previous_row_bottom) * scale))
                    for column_index, cell in enumerate(table_row["cells"]):
                        common = {"kind": "body_table_cell", **table_location,
                            "row_index": row_index, "column_index": column_index}
                        if not cell["fragment_ids"]:
                            page_map["structural_paragraphs"].append({"reason": "empty_table_cell",
                                "location": {**common, "paragraph_index": 0}})
                        previous_cell_bottom = None
                        for index, rid in enumerate(cell["fragment_ids"]):
                            box = fragments[rid].bbox_px
                            cell_gap = (row_gap + max(0, (box[1] - row_box[1]) * scale)
                                if previous_cell_bottom is None
                                else max(0, (box[1] - previous_cell_bottom) * scale))
                            assign(rid, {**common, "paragraph_index": index}, cell_gap)
                            previous_cell_bottom = box[3]
                    previous_row_bottom = row_box[3]
                table_index += 1
            previous_block_bottom = block_box[3]
            child_index += 1

        if layout["body"][-1]["kind"] == "table":
            # WML cannot attach a paragraph section break to a table. This
            # explicit, owned, text-free structural paragraph is not a text
            # carrier or a workaround for short justified prose.
            page_map["structural_paragraphs"].append({"reason": "table_end_section_anchor",
                "location": {"kind": "body_paragraph", "part_uri": "/word/document.xml",
                    "body_child_index": child_index, "paragraph_index": paragraph_index,
                    "section_index": page.section_index}, "line_height_pt": 1})
            child_index += 1
            paragraph_index += 1
        if assigned != set(rows):
            base._fail("unassigned_region_fragment")
    return mapping


def _property(parent, name, **attributes):
    child = OxmlElement("w:" + name)
    for key, value in attributes.items():
        child.set(qn("w:" + key), str(value))
    parent.append(child)
    return child


def _configure_table(table, percentages, margins=None):
    """Source physical columns remain LTR even when all paragraphs are RTL."""
    widths = _widths(percentages)
    properties = table._tbl.tblPr
    for node in list(properties):
        properties.remove(node)
    _property(properties, "bidiVisual", val=0)
    _property(properties, "tblW", w=sum(widths), type="dxa")
    _property(properties, "jc", val="left")
    borders = _property(properties, "tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _property(borders, side, val="nil")
    _property(properties, "tblLayout", type="fixed")
    table_margins = _property(properties, "tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        _property(table_margins, side, w=0, type="dxa")
    grid = table._tbl.tblGrid
    for node in list(grid):
        grid.remove(node)
    for width in widths:
        _property(grid, "gridCol", w=width)
    for row in table.rows:
        for column, (cell, width) in enumerate(zip(row.cells, widths)):
            properties = cell._tc.get_or_add_tcPr()
            for node in list(properties):
                properties.remove(node)
            _property(properties, "tcW", w=width, type="dxa")
            if margins is not None:
                cell_margins = _property(properties, "tcMar")
                for side in ("top", "left", "bottom", "right"):
                    _property(cell_margins, side, w=margins[column].get(side, 0), type="dxa")
            _property(properties, "vAlign", val="top")


def _anchor(document):
    paragraph = document.add_paragraph()
    paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = "Normal"
    fmt = paragraph.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(1)
    fmt.keep_with_next = fmt.keep_together = fmt.widow_control = False
    return paragraph


def _document(projection, mapping, *, include_text):
    document = base._base_document(projection.target_lang)
    settings = document.settings._element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = _property(settings, "compat")
    if compat.find(qn("w:doNotExpandShiftReturn")) is not None:
        base._fail("unexpected_default_manual_break_policy")
    # CT_Compat is a sequence, not an unordered property bag. This profile
    # starts from the fixed python-docx template: both existing child kinds
    # follow doNotExpandShiftReturn in the Open XML schema. A changed template
    # needs explicit review rather than guessing its insertion order.
    if any(child.tag not in {qn("w:useFELayout"), qn("w:compatSetting")} for child in compat):
        base._fail("unsupported_default_compatibility_profile")
    flag = _property(compat, "doNotExpandShiftReturn", val=1)
    compat.insert(0, flag)
    for index, (page, page_map) in enumerate(zip(projection.pages, mapping["pages"])):
        section = base._new_section(document) if index else document.sections[0]
        base._configure_section(section)
        fragments = {f.rendering_id: f for f in page.fragments}
        rows = {row["rendering_id"]: row for row in page_map["fragments"]}
        layout = page_map["region_layout"]

        def fill(paragraph, rid):
            if include_text:
                base._paragraph(paragraph, fragments[rid], rows[rid], target_lang=projection.target_lang)

        for part in ("header", "footer"):
            for rid in layout[part]:
                fill(getattr(section, part).add_paragraph(), rid)
        for block in layout["body"]:
            if block["kind"] == "paragraph":
                fill(document.add_paragraph(), block["fragment_id"])
                continue
            table = document.add_table(rows=len(block["rows"]), cols=len(block["column_widths"]))
            margins = (_column_gutters(block, page)[1]
                       if layout["version"] == GUTTER_VERSION else None)
            _configure_table(table, block["column_widths"], margins)
            for row_index, row in enumerate(block["rows"]):
                for column_index, cell_spec in enumerate(row["cells"]):
                    cell = table.cell(row_index, column_index)
                    # Every cell requires a paragraph, including an explicitly
                    # empty structural cell. Additional paragraphs require IDs.
                    for paragraph_index, rid in enumerate(cell_spec["fragment_ids"]):
                        paragraph = cell.paragraphs[0] if paragraph_index == 0 else cell.add_paragraph()
                        fill(paragraph, rid)
        if layout["body"][-1]["kind"] == "table":
            _anchor(document)
        base._finish_empty_parts(section)
    return document


def _located(roots, location):
    root = roots[location["part_uri"]]
    if location["kind"] in {"section_header", "section_footer"}:
        return root[location["paragraph_index"]]
    if len(root) != 1 or root[0].tag != qn("w:body"):
        base._fail("unsupported_region_document_body")
    body = root[0]
    element = body[location["body_child_index"]]
    if location["kind"] == "body_paragraph":
        paragraphs = [child for child in body if child.tag == qn("w:p")]
        if paragraphs[location["paragraph_index"]] is not element:
            base._fail("region_paragraph_locator_mismatch")
        return element
    if location["kind"] != "body_table_cell" or element.tag != qn("w:tbl"):
        base._fail("region_table_locator_mismatch")
    tables = [child for child in body if child.tag == qn("w:tbl")]
    if tables[location["table_index"]] is not element:
        base._fail("region_table_locator_mismatch")
    table_rows = [child for child in element if child.tag == qn("w:tr")]
    cells = [child for child in table_rows[location["row_index"]] if child.tag == qn("w:tc")]
    paragraphs = [child for child in cells[location["column_index"]] if child.tag == qn("w:p")]
    return paragraphs[location["paragraph_index"]]


def _strip_checked_text(paragraph):
    """Remove only text/properties already proved exact; keep section markup."""
    properties = paragraph.find(qn("w:pPr"))
    section = properties.find(qn("w:sectPr"))
    saved_section = deepcopy(section) if section is not None else None
    for child in list(paragraph):
        paragraph.remove(child)
    if saved_section is not None:
        properties = OxmlElement("w:pPr")
        properties.append(saved_section)
        paragraph.append(properties)


def _validate(raw, source_map_bytes, projection, *, expected_reviewer_kind="ai_test_review"):
    expected = _plan(projection, expected_reviewer_kind=expected_reviewer_kind)
    expected["docx_sha256"] = base._sha(raw)
    if base._json(base._decode(source_map_bytes)) != base._json(expected):
        base._fail("source_map_mismatch")
    actual = base._package(raw)
    scaffold = _document(projection, expected, include_text=False)
    baseline = base._package(base._save(scaffold))
    text_parts = {"word/document.xml", *(uri[1:] for page in expected["pages"]
                    for uri in page["section_parts"].values())}
    if set(actual) != set(baseline) or any(actual[name] != baseline[name] for name in actual if name not in text_parts):
        base._fail("unsupported_or_changed_package_infrastructure")
    roots = {"/" + name: base._xml(actual[name]) for name in text_parts}
    if any(node.tail or (node.text and node.tag != qn("w:t"))
           for root in roots.values() for node in root.iter()):
        base._fail("unowned_xml_text_or_tail")
    owned = set()
    for page, page_map in zip(projection.pages, expected["pages"]):
        for fragment, row in zip(page.fragments, page_map["fragments"]):
            paragraph = _located(roots, row["location"])
            if paragraph in owned:
                base._fail("duplicate_region_paragraph_ownership")
            owned.add(paragraph)
            base._paragraph_text(paragraph, row, fragment, target_lang=projection.target_lang)
            _strip_checked_text(paragraph)
    # Empty cells, table-ending anchors, grids/widths, row/cell order, section
    # placement, part relationships and all unowned nodes remain untouched.
    # Thus even an extra empty paragraph or a hidden run cannot be normalized
    # away by the checker. It does not trust the supplied map for ownership.
    for uri, root in roots.items():
        reference = base._xml(baseline[uri[1:]])
        if etree.tostring(root, method="c14n") != etree.tostring(reference, method="c14n"):
            base._fail("region_structural_scaffold_mismatch")


def validate_region_docx(docx_bytes: bytes, source_map_bytes: bytes, *,
                         projection: reviewed.ReviewedFormattingProjection,
                         expected_reviewer_kind: str = "ai_test_review") -> None:
    """Check v2 exact display, ownership, complete structure and infrastructure."""
    try:
        _validate(docx_bytes, source_map_bytes, projection, expected_reviewer_kind=expected_reviewer_kind)
    except base.ReviewedFormattingWriterError:
        raise
    except (reviewed.ReviewedFormattingError, TypeError, ValueError, KeyError,
            IndexError, AttributeError, OverflowError):
        base._fail("invalid_reviewed_region_docx_contract")


def build_region_docx(projection: reviewed.ReviewedFormattingProjection, *,
        expected_reviewer_kind: str = "ai_test_review") -> base.ReviewedDocxArtifact:
    """Build/verify in memory; publication and physical review remain external."""
    try:
        mapping = _plan(projection, expected_reviewer_kind=expected_reviewer_kind)
        raw = base._save(_document(projection, mapping, include_text=True))
        mapping["docx_sha256"] = base._sha(raw)
        source_map_bytes = base._json(mapping)
        validate_region_docx(raw, source_map_bytes, projection=projection,
                             expected_reviewer_kind=expected_reviewer_kind)
        return base.ReviewedDocxArtifact(raw, source_map_bytes, mapping["projection_sha256"])
    except base.ReviewedFormattingWriterError:
        raise
    except (reviewed.ReviewedFormattingError, TypeError, ValueError, KeyError,
            IndexError, AttributeError, OverflowError):
        base._fail("invalid_reviewed_region_writer_input")
