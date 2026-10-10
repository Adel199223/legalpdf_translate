"""Pure ownership and preservation contracts for automatic ordinary layout.

Writer locations, not translated wording or paragraph counts, establish page
ownership. The optional saved-layout snapshot is available only for packages
that its existing bounded renderer can safely rebuild.
"""
from __future__ import annotations

from dataclasses import dataclass
from copy import deepcopy
from io import BytesIO
import hashlib
import json
import re

from docx import Document

from . import saved_docx_layout as model
from .saved_docx_layout_writer import build_unreviewed_docx, validate_built_docx

VERSION = "ordinary_auto_layout_artifacts_v1"
_HASH = re.compile(r"[a-f0-9]{64}\Z")


class OrdinaryAutoArtifactError(ValueError):
    def __init__(self, code):
        self.code = "ordinary_auto_layout_" + code
        super().__init__(self.code)


def _fail(code):
    raise OrdinaryAutoArtifactError(code) from None


def _sha(raw):
    return hashlib.sha256(raw).hexdigest()


def _json(raw):
    return json.dumps(raw, sort_keys=True, ensure_ascii=False, allow_nan=False,
                      separators=(",", ":")).encode("utf-8")


@dataclass(frozen=True, slots=True)
class RawParagraph:
    id: str
    ordinal: int
    text: str
    location: str
    page_number: int | None
    control: bool
    has_page_break: bool


@dataclass(frozen=True, slots=True)
class RawOrdinarySnapshot:
    raw_docx_sha256: str
    raw_source_map_sha256: str
    source_pdf_sha256: str
    selected_pages: tuple[int, ...]
    target_lang: str
    paragraphs: tuple[RawParagraph, ...]
    page_groups: tuple[tuple[int, tuple[str, ...]], ...]
    saved_snapshot_json: bytes | None
    fingerprint: str

    @property
    def page_ids(self):
        return dict(self.page_groups)

    @property
    def saved_snapshot(self):
        return json.loads(self.saved_snapshot_json) if self.saved_snapshot_json is not None else None


@dataclass(frozen=True, slots=True)
class AutoCandidateArtifact:
    docx_bytes: bytes
    source_map_bytes: bytes
    receipt_bytes: bytes
    docx_sha256: str
    source_map_sha256: str
    receipt_sha256: str

    @property
    def source_map(self):
        return json.loads(self.source_map_bytes)

    @property
    def receipt(self):
        return json.loads(self.receipt_bytes)


_MAP_EXTENSIONS = {"raw_source_map_sha256", "source_pdf_sha256",
                   "selected_physical_pages", "raw_snapshot_fingerprint",
                   "proposal_evidence_sha256", "source_association_basis"}


def _checked_automatic_decisions(raw_snapshot, page_frames, decisions):
    checked = model.validate_decisions(raw_snapshot.saved_snapshot, page_frames, decisions,
                                       require_review=False)
    if checked["review"]["document_reviewed"] or checked["review"]["pages_reviewed"]:
        _fail("automatic_review_claim")
    owners = {row.id: row.page_number for row in raw_snapshot.paragraphs}
    if set(owners) != {row["id"] for row in raw_snapshot.saved_snapshot["paragraphs"]}:
        _fail("candidate_paragraph_mismatch")
    for choice in checked["paragraphs"]:
        owner = owners[choice["paragraph_id"]]
        if any(owner is None or region["page_number"] != owner
               or region["page_number"] not in raw_snapshot.selected_pages
               for region in choice["regions"]):
            _fail("proposal_page_owner_mismatch")
    return checked


def _extended_source_map(base_map, raw_snapshot, proposal_evidence):
    owners = {row.id: row.page_number for row in raw_snapshot.paragraphs}
    if [row["paragraph_id"] for row in base_map["paragraphs"]] != [row.id for row in raw_snapshot.paragraphs]:
        _fail("candidate_paragraph_mismatch")
    mapping = deepcopy(base_map)
    mapping["version"] = VERSION
    mapping.update({"raw_source_map_sha256": raw_snapshot.raw_source_map_sha256,
        "source_pdf_sha256": raw_snapshot.source_pdf_sha256,
        "selected_physical_pages": list(raw_snapshot.selected_pages),
        "raw_snapshot_fingerprint": raw_snapshot.fingerprint,
        "proposal_evidence_sha256": _sha(_json(proposal_evidence)),
        "source_association_basis": "raw_writer_map"})
    for row in mapping["paragraphs"]:
        row["raw_source_page_number"] = owners[row["paragraph_id"]]
    return mapping


def _candidate_receipt(raw_snapshot, checked, proposal_evidence, docx_bytes, map_bytes):
    return {"version": VERSION, "docx_sha256": _sha(docx_bytes),
        "source_map_sha256": _sha(map_bytes),
        "raw_docx_sha256": raw_snapshot.raw_docx_sha256,
        "raw_source_map_sha256": raw_snapshot.raw_source_map_sha256,
        "source_pdf_sha256": raw_snapshot.source_pdf_sha256,
        "selected_physical_pages": list(raw_snapshot.selected_pages),
        "raw_snapshot_fingerprint": raw_snapshot.fingerprint,
        "decisions_sha256": _sha(_json(checked)),
        "proposal_evidence_sha256": _sha(_json(proposal_evidence)),
        "target_lang": raw_snapshot.target_lang, "document_reviewed": False,
        "rendered_layout_acceptance": "not_evaluated",
        "source_character_coverage": "not_proven", "exact_text_preserved": True}


def _container_step(container, step):
    if type(step) is not dict or set(step) != {"table_index", "row", "col"}:
        _fail("invalid_raw_location")
    a, b, c = (step[key] for key in ("table_index", "row", "col"))
    if any(type(v) is not int or v < 0 for v in (a, b, c)):
        _fail("invalid_raw_location")
    try:
        return container.tables[a].cell(b, c)
    except (IndexError, ValueError):
        _fail("unresolved_raw_location")


def _resolve(document, location):
    if type(location) is not dict or type(location.get("kind")) is not str:
        _fail("invalid_raw_location")
    kind = location["kind"]
    if kind == "body_paragraph":
        if set(location) != {"kind", "paragraph_index"}:
            _fail("invalid_raw_location")
        container = document
    elif kind == "table_cell":
        path = location.get("table_path")
        if path is None:
            path = [{key: location[key] for key in ("table_index", "row", "col")}
                    ] if all(key in location for key in ("table_index", "row", "col")) else None
        if type(path) is not list or not path:
            _fail("invalid_raw_location")
        container = document
        for step in path:
            container = _container_step(container, step)
    elif kind == "layout_region_paragraph":
        path = location.get("table_path")
        if type(path) is not list or not path:
            _fail("invalid_raw_location")
        container = document
        for step in path:
            container = _container_step(container, step)
    else:
        return None
    index = location.get("paragraph_index")
    if type(index) is not int or index < 0:
        _fail("invalid_raw_location")
    try:
        return container.paragraphs[index]._p
    except IndexError:
        _fail("unresolved_raw_location")


def bind_raw_page_map(raw_docx_bytes: bytes, raw_source_map: dict, *, source_pdf_sha256: str,
                      selected_pages: tuple[int, ...], target_lang: str) -> RawOrdinarySnapshot:
    """Bind every visible body paragraph to a writer location or explicit control."""
    try:
        if (type(raw_docx_bytes) is not bytes or not raw_docx_bytes
                or type(raw_source_map) is not dict or type(source_pdf_sha256) is not str
                or not _HASH.fullmatch(source_pdf_sha256)
                or type(selected_pages) is not tuple or not selected_pages
                or tuple(sorted(set(selected_pages))) != selected_pages
                or any(type(n) is not int or not 1 <= n <= 10000 for n in selected_pages)
                or target_lang not in {"EN", "FR", "AR"}):
            _fail("invalid_raw_binding")
        raw_hash = _sha(raw_docx_bytes)
        if (raw_source_map.get("version") != 1 or raw_source_map.get("docx_sha256") != raw_hash
                or raw_source_map.get("source_page_count") != len(selected_pages)
                or type(raw_source_map.get("pages")) is not list
                or len(raw_source_map["pages"]) != len(selected_pages)):
            _fail("stale_raw_map")
        document = Document(BytesIO(raw_docx_bytes))
        members, _ = model._package(raw_docx_bytes)
        context = model._styles(members)
        body = document._element.body
        all_nodes = list(body.iter(model.W + "p"))
        if not all_nodes or len(all_nodes) > model.MAX_PARAGRAPHS:
            _fail("unsupported_raw_paragraphs")
        paths = {document._element.getroottree().getpath(node): node for node in all_nodes}
        owners: dict[str, int] = {}
        for expected_page, page in zip(selected_pages, raw_source_map["pages"]):
            if (type(page) is not dict or page.get("source_page_number") != expected_page
                    or type(page.get("blocks")) is not list):
                _fail("raw_page_selection_mismatch")
            source_file = page.get("source_file_sha256")
            if source_file is not None and source_file != source_pdf_sha256:
                _fail("raw_source_changed")
            for block in page["blocks"]:
                if type(block) is not dict or type(block.get("location")) is not dict:
                    _fail("invalid_raw_block")
                location = block["location"]
                if location.get("kind") in {"empty_block", "generated_footer_page_field",
                                            "section_furniture", "generated_header", "generated_footer"}:
                    continue
                node = _resolve(document, location)
                if node is None:
                    # A non-body furniture alias is not source-page paragraph text.
                    if block.get("furniture_alias") is not None:
                        continue
                    _fail("unsupported_raw_location")
                path = document._element.getroottree().getpath(node)
                if path not in paths:
                    _fail("raw_location_outside_body")
                prior = owners.setdefault(path, expected_page)
                if prior != expected_page:
                    _fail("ambiguous_raw_page")
        rows = []
        groups = {n: [] for n in selected_pages}
        for ordinal, node in enumerate(all_nodes, 1):
            path = document._element.getroottree().getpath(node)
            page = owners.get(path)
            control = page is None
            try:
                row = model._paragraph_inventory(node, context, ordinal=ordinal)
            except model.SavedDocxLayoutError:
                if not control or any((part.text or "").strip() for part in node.iter(model.W + "t")):
                    _fail("unsupported_raw_paragraph")
                row = {"id": f"p{ordinal:06d}", "text": "", "has_page_break": bool(
                    node.findall(".//" + model.W + "sectPr")
                    or any(br.get(model.W + "type") == "page"
                           for br in node.iter(model.W + "br")))}
            if control and row["text"].strip():
                _fail("unmapped_raw_text")
            pid = row["id"]
            rows.append(RawParagraph(pid, ordinal, row["text"], path, page, control,
                                     row["has_page_break"]))
            if page is not None:
                groups[page].append(pid)
        if any(not group for group in groups.values()):
            _fail("empty_raw_page")
        try:
            saved = model.inspect_docx(raw_docx_bytes, target_lang)
        except model.SavedDocxLayoutError:
            saved = None
        saved_json = _json(saved) if saved is not None else None
        source_map_hash = _sha(_json(raw_source_map))
        identity = {"version": VERSION, "raw_docx_sha256": raw_hash,
                    "raw_source_map_sha256": source_map_hash, "source_pdf_sha256": source_pdf_sha256,
                    "selected_pages": selected_pages, "target_lang": target_lang,
                    "paragraphs": [(r.id, r.location, r.page_number, r.control) for r in rows]}
        return RawOrdinarySnapshot(raw_hash, source_map_hash, source_pdf_sha256,
                                   selected_pages, target_lang, tuple(rows),
                                   tuple((n, tuple(groups[n])) for n in selected_pages), saved_json,
                                   _sha(_json(identity)))
    except OrdinaryAutoArtifactError:
        raise
    except (model.SavedDocxLayoutError, TypeError, ValueError, KeyError, IndexError,
            AttributeError, OverflowError, RecursionError):
        _fail("invalid_raw_package")


def build_unreviewed_candidate(raw_docx_bytes: bytes, raw_snapshot: RawOrdinarySnapshot,
                               page_frames: list, decisions: dict, *,
                               proposal_evidence: dict) -> AutoCandidateArtifact:
    """Build detached editable bytes; never turn a model proposal into review."""
    if (type(raw_snapshot) is not RawOrdinarySnapshot or _sha(raw_docx_bytes) != raw_snapshot.raw_docx_sha256
            or raw_snapshot.saved_snapshot is None or type(proposal_evidence) is not dict
            or not proposal_evidence or proposal_evidence.get("document_reviewed") is True
            or proposal_evidence.get("rendered_layout_acceptance") not in {None, "not_evaluated"}
            or proposal_evidence.get("source_character_coverage") not in {None, "not_proven"}):
        _fail("invalid_candidate_binding")
    checked = _checked_automatic_decisions(raw_snapshot, page_frames, decisions)
    rendered = build_unreviewed_docx(raw_docx_bytes, raw_snapshot.saved_snapshot,
                                     page_frames, checked, ordinary_context={"selected_pages":list(raw_snapshot.selected_pages),
                                         "page_groups":[[page,list(ids)] for page,ids in raw_snapshot.page_groups]})
    source_map = _extended_source_map(rendered.source_map, raw_snapshot, proposal_evidence)
    map_bytes = _json(source_map)
    receipt = _candidate_receipt(raw_snapshot, checked, proposal_evidence,
                                 rendered.docx_bytes, map_bytes)
    receipt_bytes = _json(receipt)
    artifact = AutoCandidateArtifact(rendered.docx_bytes, map_bytes, receipt_bytes,
                                     receipt["docx_sha256"], receipt["source_map_sha256"],
                                     _sha(receipt_bytes))
    verify_unreviewed_candidate(raw_docx_bytes, raw_snapshot, page_frames, checked, artifact,
                                proposal_evidence=proposal_evidence)
    return artifact


def verify_unreviewed_candidate(raw_docx_bytes: bytes, raw_snapshot: RawOrdinarySnapshot,
                                page_frames: list, decisions: dict, artifact: AutoCandidateArtifact,
                                *, proposal_evidence: dict) -> None:
    if (type(artifact) is not AutoCandidateArtifact
            or _sha(raw_docx_bytes) != raw_snapshot.raw_docx_sha256
            or _sha(artifact.docx_bytes) != artifact.docx_sha256
            or _sha(artifact.source_map_bytes) != artifact.source_map_sha256
            or _sha(artifact.receipt_bytes) != artifact.receipt_sha256):
        _fail("candidate_changed")
    checked = _checked_automatic_decisions(raw_snapshot, page_frames, decisions)
    source_map = artifact.source_map
    base_map = {k: deepcopy(v) for k, v in source_map.items() if k not in _MAP_EXTENSIONS}
    base_map["version"] = model.VERSION
    if type(base_map.get("paragraphs")) is not list:
        _fail("candidate_provenance_changed")
    for row in base_map["paragraphs"]:
        if type(row) is not dict:
            _fail("candidate_provenance_changed")
        row.pop("raw_source_page_number", None)
    validate_built_docx(artifact.docx_bytes, base_map, original_docx=raw_docx_bytes,
                        snapshot=raw_snapshot.saved_snapshot, pages=page_frames,
                        decisions=checked, require_review=False,
                        ordinary_context={"selected_pages":list(raw_snapshot.selected_pages),
                            "page_groups":[[page,list(ids)] for page,ids in raw_snapshot.page_groups]}
                        if base_map.get("writer_version") == "saved_docx_layout_writer_ordinary_presentation_v6" else None)
    expected_map = _extended_source_map(base_map, raw_snapshot, proposal_evidence)
    if artifact.source_map_bytes != _json(expected_map):
        _fail("candidate_provenance_changed")
    expected_receipt = _candidate_receipt(raw_snapshot, checked, proposal_evidence,
                                          artifact.docx_bytes, artifact.source_map_bytes)
    if artifact.receipt_bytes != _json(expected_receipt):
        _fail("candidate_provenance_changed")
