"""Fictional saved-document import and source-decision contracts, no providers."""
from copy import deepcopy
import io
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree
import pytest

from legalpdf_translate import saved_docx_layout as model


PAGES = [{"page_number": 1, "width_px": 1000, "height_px": 1400,
          "image_sha256": "1" * 64, "rotation": 0, "render_dpi": 144, "backend": "fictional"}]


def document_bytes(texts=("Example institution", "The complete obligation remains unchanged."), edit=None):
    document = Document()
    for text in texts:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(11)
    if edit:
        edit(document)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def changed_part(raw, name, edit):
    stream = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as source, ZipFile(stream, "w") as target:
        for entry in source.infolist():
            value = source.read(entry.filename)
            if entry.filename == name:
                root = etree.fromstring(value)
                edit(root)
                value = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(entry, value)
    return stream.getvalue()


def reviewed(raw, lang="EN", pages=None):
    pages = deepcopy(PAGES if pages is None else pages)
    snapshot = model.inspect_docx(raw, lang)
    decisions = model.default_decisions(snapshot, pages)
    for index, choice in enumerate(decisions["paragraphs"]):
        choice["regions"] = [{"page_number": 1, "bbox_px": [20, 20 + index * 30, 800, 45 + index * 30]}]
    decisions["review"].update(reviewer="Fictional reviewer", note="All fictional source groups reviewed.",
                               pages_reviewed=[p["page_number"] for p in pages], document_reviewed=True)
    return snapshot, pages, decisions


@pytest.mark.parametrize("lang,text", [("EN", "Example text"), ("FR", "La déclaration, sans modification."),
    ("AR", "نَصّ محفوظ مع الاسم \u200ePerson Example\u200e")])
def test_import_preserves_complete_typed_offsets_and_blank_paragraphs(lang, text):
    def edit(doc):
        run = doc.paragraphs[0].add_run()
        run.add_tab()
        run.add_text("after tab")
        run.add_break()
        run.add_text("after break")
        run.add_break(WD_BREAK.PAGE)
    raw = document_bytes((text, "", "  "), edit)
    before = bytes(raw)
    snapshot = model.inspect_docx(raw, lang)
    assert raw == before
    rows = snapshot["paragraphs"]
    assert [p["id"] for p in rows] == ["p000001", "p000002", "p000003"]
    assert rows[0]["text"] == text + "\tafter tab\nafter break\f"
    assert rows[0]["has_page_break"] and not rows[1]["has_page_break"]
    assert [p["text"] for p in rows[1:]] == ["", "  "]
    assert [t["kind"] for t in rows[0]["tokens"]] == ["t", "tab", "t", "line_break", "t", "page_break"]
    for token in rows[0]["tokens"]:
        assert rows[0]["text"][token["start"]:token["end"]] == token["text"]


@pytest.mark.parametrize("feature", ["table", "drawing", "hyperlink", "field", "revision", "bookmark",
    "comment", "content_control", "numbering", "hidden", "section", "midpage", "second_pagebreak", "frame"])
def test_unsupported_body_features_decline_without_flattening(feature):
    raw = document_bytes()
    def change(root):
        body, p = root[0], root[0][0]
        run = p.find(model.W + "r")
        if feature == "table":
            body.insert(0, etree.Element(model.W + "tbl"))
        elif feature in {"drawing", "field"}:
            run.append(etree.Element(model.W + ("drawing" if feature == "drawing" else "fldChar")))
        elif feature in {"hyperlink", "revision", "bookmark", "comment", "content_control"}:
            p.append(etree.Element(model.W + {"hyperlink": "hyperlink", "revision": "ins", "bookmark": "bookmarkStart",
                "comment": "commentRangeStart", "content_control": "sdt"}[feature]))
        elif feature == "hidden":
            run.find(model.W + "rPr").append(etree.Element(model.W + "vanish"))
        elif feature in {"numbering", "section", "frame"}:
            props = etree.Element(model.W + "pPr")
            props.append(etree.Element(model.W + {"numbering": "numPr", "section": "sectPr", "frame": "framePr"}[feature]))
            p.insert(0, props)
        else:
            br = etree.Element(model.W + "br", {model.W + "type": "page"})
            run.append(br)
            if feature == "midpage":
                later = etree.Element(model.W + "t")
                later.text = "still same paragraph"
            else:
                later = deepcopy(br)
            run.append(later)
    broken = changed_part(raw, "word/document.xml", change)
    with pytest.raises(model.SavedDocxLayoutError) as exc:
        model.inspect_docx(broken, "EN")
    assert "Example" not in str(exc.value)


def test_inherited_hidden_and_numbered_styles_are_declined():
    for tag, location in (("vanish", "rPr"), ("numPr", "pPr")):
        def mutate(root):
            normal = root.find(model.W + "style[@" + model.W + "styleId='Normal']")
            props = normal.find(model.W + location)
            if props is None:
                props = etree.SubElement(normal, model.W + location)
            etree.SubElement(props, model.W + tag)
        with pytest.raises(model.SavedDocxLayoutError):
            model.inspect_docx(changed_part(document_bytes(), "word/styles.xml", mutate), "EN")


def test_static_footer_and_page_field_are_inventory_supported_but_other_fields_decline():
    def footer(doc):
        run = doc.sections[0].footer.paragraphs[0].add_run()
        for tag, attrs, text in (("fldChar", {"fldCharType": "begin"}, None),
                ("instrText", {}, " PAGE "), ("fldChar", {"fldCharType": "separate"}, None),
                ("t", {}, "1"), ("fldChar", {"fldCharType": "end"}, None)):
            node = OxmlElement("w:" + tag)
            for key, value in attrs.items():
                node.set(qn("w:" + key), value)
            node.text = text
            run._r.append(node)
    raw = document_bytes(edit=footer)
    assert "word/footer1.xml" in model.inspect_docx(raw, "EN")["package_sha256"]
    def malicious(root):
        next(root.iter(model.W + "instrText")).text = " INCLUDETEXT remote "
    with pytest.raises(model.SavedDocxLayoutError, match="unsupported_footer_field"):
        model.inspect_docx(changed_part(raw, "word/footer1.xml", malicious), "EN")


@pytest.mark.parametrize("kind", ["unknown_part", "external", "traversal", "duplicate", "doctype", "macro"])
def test_package_declines_unsupported_or_ambiguous_members(kind):
    raw = document_bytes()
    if kind in {"unknown_part", "traversal", "duplicate"}:
        stream = io.BytesIO(raw)
        with ZipFile(stream, "a") as archive:
            name = {"unknown_part": "word/embeddings/object1.bin", "traversal": "../outside.xml", "duplicate": "word/document.xml"}[kind]
            with pytest.warns(UserWarning) if kind == "duplicate" else _no_warning():
                archive.writestr(name, b"empty")
        raw = stream.getvalue()
    elif kind == "external":
        raw = changed_part(raw, "word/_rels/document.xml.rels", lambda r: r[0].set("TargetMode", "External"))
    elif kind == "macro":
        raw = changed_part(raw, "[Content_Types].xml", lambda r: next(n for n in r if n.get("PartName") == "/word/document.xml").set("ContentType", "application/vnd.ms-word.document.macroEnabled.main+xml"))
    else:
        stream = io.BytesIO()
        with ZipFile(io.BytesIO(raw)) as source, ZipFile(stream, "w") as target:
            for entry in source.infolist():
                data = source.read(entry.filename)
                target.writestr(entry, b'<!DOCTYPE x [<!ENTITY leak "x">]><x/>' if entry.filename == "customXml/item1.xml" else data)
        raw = stream.getvalue()
    with pytest.raises(model.SavedDocxLayoutError):
        model.inspect_docx(raw, "EN")


from contextlib import nullcontext as _no_warning


@pytest.mark.parametrize("limit,value", [("DOCX_MAX_BYTES", 10), ("EXPANDED_MAX_BYTES", 10),
    ("MEMBER_MAX_BYTES", 10), ("MAX_MEMBERS", 1), ("MAX_PARAGRAPHS", 1), ("MAX_CODEPOINTS", 1)])
def test_import_resource_limits_are_enforced(monkeypatch, limit, value):
    raw = document_bytes()
    monkeypatch.setattr(model, limit, value)
    with pytest.raises(model.SavedDocxLayoutError):
        model.inspect_docx(raw, "EN")


@pytest.mark.parametrize("mutation", ["missing", "duplicate", "reorder", "replacement_text", "bad_box", "nan_box",
    "unassigned_review", "missed_source_page", "missing_reviewer", "invalid_size", "bad_columns", "unknown_role"])
def test_decisions_require_exact_complete_order_and_honest_review(mutation):
    snapshot, pages, dec = reviewed(document_bytes())
    ids = dec["bands"][0]["groups"][0]["paragraph_ids"]
    if mutation == "missing": ids.pop()
    elif mutation == "duplicate": ids.append(ids[0])
    elif mutation == "reorder": ids.reverse()
    elif mutation == "replacement_text": dec["paragraphs"][0]["text"] = "Replacement"
    elif mutation == "bad_box": dec["paragraphs"][0]["regions"][0]["bbox_px"] = [0, 0, 1001, 50]
    elif mutation == "nan_box": dec["paragraphs"][0]["regions"][0]["bbox_px"][0] = float("nan")
    elif mutation == "unassigned_review": dec["paragraphs"][0]["regions"] = []
    elif mutation == "missed_source_page": dec["review"]["pages_reviewed"] = []
    elif mutation == "missing_reviewer": dec["review"]["reviewer"] = ""
    elif mutation == "invalid_size": dec["paragraphs"][0]["heading_size_pt"] = 100
    elif mutation == "unknown_role": dec["paragraphs"][0]["role"] = "unknown"
    else: dec["bands"] = [{"kind": "columns", "widths_pct": [40, 40], "gutter_pt": 8,
                           "cells": [{"groups": [{"paragraph_ids": [identifier], "panel": False}]} for identifier in ids]}]
    with pytest.raises(model.SavedDocxLayoutError):
        model.validate_decisions(snapshot, pages, dec, require_review=True)


def test_explicit_unmapped_retention_is_honest_and_drafts_remain_unreviewed():
    raw = document_bytes()
    snapshot = model.inspect_docx(raw, "EN")
    dec = model.default_decisions(snapshot, PAGES)
    assert model.validate_decisions(snapshot, PAGES, dec) == dec
    with pytest.raises(model.SavedDocxLayoutError, match="incomplete_source_review"):
        model.validate_decisions(snapshot, PAGES, dec, require_review=True)
    snapshot, pages, dec = reviewed(raw)
    dec["paragraphs"][0].update(regions=[], unmapped_reason="Current target-only note retained.")
    actual = model.validate_decisions(snapshot, pages, dec, require_review=True)
    assert actual == dec and actual is not dec


@pytest.mark.parametrize("kind", ["terminal", "before", "inherited_before"])
@pytest.mark.parametrize("container", ["columns", "panel"])
def test_page_breaks_cannot_be_hidden_in_a_layout_container(kind, container):
    def edit(doc):
        if kind == "terminal": doc.paragraphs[0].runs[0].add_break(WD_BREAK.PAGE)
        elif kind == "before": doc.paragraphs[0].paragraph_format.page_break_before = True
        else: doc.styles["Normal"].paragraph_format.page_break_before = True
    snapshot, pages, dec = reviewed(document_bytes(edit=edit))
    ids = [p["id"] for p in snapshot["paragraphs"]]
    if container == "panel": dec["bands"][0]["groups"][0]["panel"] = True
    else: dec["bands"] = [{"kind": "columns", "widths_pct": [50, 50], "gutter_pt": 12,
                           "cells": [{"groups": [{"paragraph_ids": [identifier], "panel": False}]} for identifier in ids]}]
    with pytest.raises(model.SavedDocxLayoutError, match="page_break_requires_flow"):
        model.validate_decisions(snapshot, pages, dec)


@pytest.mark.parametrize("text,selected,accepted", [
    ("Keep this whole phrase exactly.", "this whole phrase", True),
    ("Une de\u0301claration reste complète.", "de\u0301claration", True),
    ("Une de\u0301claration reste complète.", "de", False),
    ("نَصّ محفوظ دون تغيير", "محفوظ دون", True),
    ("نَصّ محفوظ دون تغيير", "ن", False),
    ("Clock 09:30 remains", "09", False),
    ("Call 123 456 789 today", "123", False),
    ("Visit https://example.invalid/a today", "https://example.invalid/a", True),
    ("Value 12/2026/AB remains", "12/2026", False),
    ("Amount 12.50 € remains", "12.50", False),
    ("نص \u200eName Example\u200e ثم نص", "Name", False),
    ("Text [[Name Example]] retained", "Name", False),
    ("Text \u2066Name Example\u2069 retained", "Name", False),
])
def test_phrase_edges_do_not_split_literals_or_graphemes(text, selected, accepted):
    snapshot, pages, dec = reviewed(document_bytes((text,)), "AR" if "ن" in text else "FR")
    start = text.index(selected)
    dec["paragraphs"][0]["emphasis"] = [{"start": start, "end": start + len(selected), "bold": True, "italic": False, "underline": False}]
    if accepted:
        assert model.validate_decisions(snapshot, pages, dec)
    else:
        with pytest.raises(model.SavedDocxLayoutError, match="unsafe_emphasis_boundary"):
            model.validate_decisions(snapshot, pages, dec)


def test_span_cannot_cross_typed_break_even_if_projected_string_matches():
    snapshot, pages, dec = reviewed(document_bytes(("One\nTwo",)))
    dec["paragraphs"][0]["emphasis"] = [{"start": 0, "end": 7, "bold": True, "italic": False, "underline": False}]
    with pytest.raises(model.SavedDocxLayoutError, match="unsafe_emphasis_boundary"):
        model.validate_decisions(snapshot, pages, dec)


@pytest.mark.parametrize("feature", ["fit_direct", "fit_inherited", "gutter", "mirror_margins"])
def test_run_fitting_and_binding_gutter_profiles_are_explicitly_declined(feature):
    raw = document_bytes()
    def change(root):
        if feature == "fit_direct":
            next(root.iter(model.W + "rPr")).append(etree.Element(model.W + "fitText", {model.W + "val": "500"}))
        elif feature == "fit_inherited":
            normal = root.find(model.W + "style[@" + model.W + "styleId='Normal']")
            props = normal.find(model.W + "rPr")
            if props is None: props = etree.SubElement(normal, model.W + "rPr")
            etree.SubElement(props, model.W + "fitText", {model.W + "val": "500"})
        elif feature == "gutter":
            next(root.iter(model.W + "pgMar")).set(model.W + "gutter", "300")
        else: etree.SubElement(root, model.W + "mirrorMargins")
    name = "word/styles.xml" if feature == "fit_inherited" else "word/settings.xml" if feature == "mirror_margins" else "word/document.xml"
    with pytest.raises(model.SavedDocxLayoutError):
        model.inspect_docx(changed_part(raw, name, change), "EN")
