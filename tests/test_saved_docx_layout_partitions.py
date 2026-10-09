"""Exact parent partitions produce editable siblings, never rewritten content."""
from copy import deepcopy
import io
from zipfile import ZipFile
from lxml import etree
import pytest

from legalpdf_translate import saved_docx_layout as model
from legalpdf_translate import saved_docx_layout_writer as writer
from legalpdf_translate.docx_writer import _add_rtl_flags, _set_rtl_run_props, _set_ltr_run_props
from tests.test_saved_docx_layout import document_bytes, reviewed, changed_part


def packet(lang="EN"):
    text = "First clause. Second clause. Third clause." if lang != "AR" else "الجملة الأولى. الجملة الثانية. الجملة الثالثة."
    def edit(doc):
        if lang == "AR":
            for p in doc.paragraphs:
                _add_rtl_flags(p)
                for r in p.runs:
                    _set_rtl_run_props(r, bidi_lang="ar-SA")
    raw = document_bytes([text, "Final original paragraph"], edit)
    snapshot, pages, dec = reviewed(raw, lang)
    dec["version"] = model.PARTITION_DECISIONS_VERSION
    cuts = [text.index("Second"), text.index("Third")] if lang != "AR" else [text.index("الجملة", 1), text.index("الجملة", text.index("الجملة", 1) + 1)]
    dec["paragraph_partitions"] = [{"paragraph_id": "p000001", "offsets": cuts}]
    dec["review"].update(document_reviewed=False, pages_reviewed=[], reviewer="", note="")
    return raw, snapshot, pages, dec


@pytest.mark.parametrize("lang", ["EN", "FR", "AR"])
def test_exact_unicode_ranges_sibling_order_properties_and_source_map(lang):
    raw, snapshot, pages, dec = packet(lang)
    built = writer.build_unreviewed_docx(raw, snapshot, pages, dec)
    assert built.source_map["writer_version"] == writer.AUTOMATIC_MODERN_WRITER_VERSION
    assert len(built.source_map["paragraphs"]) == 2
    parts = built.source_map["paragraphs"][0]["parts"]
    assert [p["part_id"] for p in parts] == ["p000001.part001", "p000001.part002", "p000001.part003"]
    assert parts[0]["start"] == 0 and parts[-1]["end"] == len(snapshot["paragraphs"][0]["text"])
    assert built.source_map["paragraphs"][0]["location"] == parts[0]["location"]
    with ZipFile(io.BytesIO(built.docx_bytes)) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    owned = [root.getroottree().xpath(p["location"], namespaces=root.nsmap)[0] for p in parts]
    assert all(p.getparent() is root[0] for p in owned)
    texts = ["".join(n.text or "" for n in p.iter(model.W + "t")) for p in owned]
    assert "".join(texts) == snapshot["paragraphs"][0]["text"]
    assert texts[0].endswith(" ") and texts[1].endswith(" ")
    assert all(p.find(model.W + "r/" + model.W + "t").get(model.XML_SPACE) == "preserve" for p in owned[:2])
    writer.validate_built_docx(built.docx_bytes, built.source_map, original_docx=raw,
        snapshot=snapshot, pages=pages, decisions=dec, require_review=False)


@pytest.mark.parametrize("offsets", [[0], [999], [True], [4], [14, 14], [28, 14], list(range(1, 9))])
def test_invalid_and_nonphrase_offsets_rejected(offsets):
    raw, snapshot, pages, dec = packet()
    dec["paragraph_partitions"][0]["offsets"] = offsets
    with pytest.raises(model.SavedDocxLayoutError):
        writer.build_unreviewed_docx(raw, snapshot, pages, dec)


@pytest.mark.parametrize("role", ["list", "reference", "institution", "recipient", "signature"])
def test_only_plain_body_flow_can_partition(role):
    _, snapshot, pages, dec = packet()
    dec["paragraphs"][0]["role"] = role
    with pytest.raises(model.SavedDocxLayoutError, match="unsupported_partition_parent"):
        model.validate_decisions(snapshot, pages, dec)


def test_explicit_ar_ltr_run_interior_is_protected():
    def edit(doc):
        p = doc.paragraphs[0]
        _add_rtl_flags(p)
        _set_rtl_run_props(p.runs[0], bidi_lang="ar-SA")
        _set_ltr_run_props(p.add_run(" Example Person "))
        _set_rtl_run_props(p.add_run("نص آخر"), bidi_lang="ar-SA")
    raw = document_bytes(["نص أول"], edit)
    snapshot, pages, dec = reviewed(raw, "AR")
    dec.update(version=model.PARTITION_DECISIONS_VERSION, paragraph_partitions=[{
        "paragraph_id": "p000001", "offsets": [snapshot["paragraphs"][0]["text"].index("Person")]}])
    with pytest.raises(model.SavedDocxLayoutError, match="unsafe_partition_boundary"):
        model.validate_decisions(snapshot, pages, dec)


def test_latin_only_ar_paragraph_keeps_ordinary_whitespace_partition_boundaries():
    raw = document_bytes(["First clause. Second clause."],
        lambda doc: _set_ltr_run_props(doc.paragraphs[0].runs[0]))
    snapshot, pages, dec = reviewed(raw, "AR")
    dec.update(version=model.PARTITION_DECISIONS_VERSION, paragraph_partitions=[{
        "paragraph_id": "p000001", "offsets": [14]}])
    model.validate_decisions(snapshot, pages, dec)


def test_reviewed_v2_retains_parts_and_map_tampering_or_child_deletion_rejected():
    raw, snapshot, pages, dec = packet()
    dec["review"].update(document_reviewed=True, reviewer="Fictional reviewer", note="Reviewed", pages_reviewed=[1])
    built = writer.build_docx(raw, snapshot, pages, dec)
    assert len(built.source_map["paragraphs"][0]["parts"]) == 3
    changed = deepcopy(built.source_map)
    changed["paragraphs"][0]["parts"][1]["start"] += 1
    with pytest.raises(model.SavedDocxLayoutError, match="source_map_mismatch"):
        writer.validate_built_docx(built.docx_bytes, changed, original_docx=raw,
            snapshot=snapshot, pages=pages, decisions=dec)
    damaged = changed_part(built.docx_bytes, "word/document.xml", lambda root: root[0].remove(root[0][1]))
    with pytest.raises(model.SavedDocxLayoutError):
        writer.validate_built_docx(damaged, built.source_map, original_docx=raw,
            snapshot=snapshot, pages=pages, decisions=dec)


@pytest.mark.parametrize("control", ["tab", "line", "page_before"])
def test_controls_and_page_break_before_reject_partition(control):
    from docx import Document
    from docx.enum.text import WD_BREAK
    raw, _, pages, dec = packet()
    doc = Document(io.BytesIO(raw))
    if control == "page_before":
        doc.paragraphs[0].paragraph_format.page_break_before = True
    elif control == "tab":
        doc.paragraphs[0].add_run().add_tab()
    else:
        doc.paragraphs[0].add_run().add_break(WD_BREAK.LINE)
    stream = io.BytesIO()
    doc.save(stream)
    raw = stream.getvalue()
    snapshot = model.inspect_docx(raw, "EN")
    with pytest.raises(model.SavedDocxLayoutError, match="unsupported_partition_parent"):
        model.validate_decisions(snapshot, pages, dec)


def test_grapheme_boundary_and_blank_children_rejected():
    raw = document_bytes(["First a \u0301combining. Last part."])
    snapshot, pages, dec = reviewed(raw)
    dec.update(version=model.PARTITION_DECISIONS_VERSION, paragraph_partitions=[{
        "paragraph_id": "p000001", "offsets": [snapshot["paragraphs"][0]["text"].index("\u0301")]}])
    with pytest.raises(model.SavedDocxLayoutError, match="unsafe_partition_boundary"):
        model.validate_decisions(snapshot, pages, dec)
    raw = document_bytes(["   Actual body text"])
    snapshot, pages, dec = reviewed(raw)
    dec.update(version=model.PARTITION_DECISIONS_VERSION, paragraph_partitions=[{
        "paragraph_id": "p000001", "offsets": [3]}])
    with pytest.raises(model.SavedDocxLayoutError, match="blank_partition_child"):
        model.validate_decisions(snapshot, pages, dec)


def test_multiple_run_styles_and_emphasis_preserved_across_child_ranges():
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    p = doc.add_paragraph()
    first = p.add_run("First clause. ")
    first.bold = True
    first.font.name = "Arial"
    second = p.add_run("Second clause. Third clause.")
    second.italic = True
    second.font.size = Pt(13)
    stream = io.BytesIO()
    doc.save(stream)
    raw = stream.getvalue()
    snapshot, pages, dec = reviewed(raw)
    text = snapshot["paragraphs"][0]["text"]
    dec.update(version=model.PARTITION_DECISIONS_VERSION, paragraph_partitions=[{
        "paragraph_id": "p000001", "offsets": [text.index("Second"), text.index("Third")]}])
    dec["paragraphs"][0]["emphasis"] = [{"start": text.index("Second"), "end": text.index("Third"),
        "bold": False, "italic": False, "underline": True}]
    built = writer.build_docx(raw, snapshot, pages, dec)
    writer.validate_built_docx(built.docx_bytes, built.source_map, original_docx=raw,
        snapshot=snapshot, pages=pages, decisions=dec)


def test_independent_verifier_rejects_splitter_shift_with_exact_parent_concat(monkeypatch):
    raw, snapshot, pages, dec = packet()
    real = writer._partition_children
    def shifted(paragraph, offsets):
        return real(paragraph, [offsets[0] + 1, *offsets[1:]])
    monkeypatch.setattr(writer, "_partition_children", shifted)
    with pytest.raises(model.SavedDocxLayoutError, match="output_partition_text_range_changed"):
        writer.build_unreviewed_docx(raw, snapshot, pages, dec)


def test_independent_verifier_rejects_splitter_dropped_later_child_bidi(monkeypatch):
    raw, snapshot, pages, dec = packet("AR")
    real = writer._partition_children
    def dropped(paragraph, offsets):
        children = real(paragraph, offsets)
        props = children[1].find(model.W + "pPr")
        props.remove(props.find(model.W + "bidi"))
        return children
    monkeypatch.setattr(writer, "_partition_children", dropped)
    with pytest.raises(model.SavedDocxLayoutError, match="output_partition_paragraph_properties_changed"):
        writer.build_unreviewed_docx(raw, snapshot, pages, dec)
