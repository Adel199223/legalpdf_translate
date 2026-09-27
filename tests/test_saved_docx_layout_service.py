"""Fictional local snapshots: persistence, ownership, retries and owned workers."""
from copy import deepcopy
from io import BytesIO
import json
import os
from pathlib import Path
import threading
from types import SimpleNamespace
import uuid

from docx import Document
import fitz
from PIL import Image
import pytest

from legalpdf_translate import saved_docx_layout_service as module
from legalpdf_translate.saved_docx_layout import inspect_docx
from legalpdf_translate.saved_docx_layout_service import SavedDocxLayoutService, SavedDocxLayoutServiceError
from legalpdf_translate.run_workspace_lock import run_workspace_slot


def nonce():
    return uuid.uuid4().hex


def docx_bytes(language="EN"):
    doc = Document()
    doc.add_paragraph({"EN": "Fictional court notice", "FR": "Avis fictif du tribunal",
                       "AR": "إشعار خيالي للمحكمة"}[language])
    doc.add_paragraph({"EN": "Case 121/26. Preserve every word.", "FR": "Affaire 121/26. Préserver chaque mot.",
                       "AR": "القضية \u200e121/26\u200e. الحفاظ على كل كلمة."}[language])
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def pdf_bytes():
    with fitz.open() as document:
        page = document.new_page(width=240, height=320)
        page.insert_text((20, 40), "Fictional court notice")
        return document.tobytes()


def fake_prepare(folder, language):
    """Replace raster work, not the snapshot parser/decision validator."""
    snapshot = inspect_docx(module._read(folder / "original.docx", module.DOCX_MAX_BYTES), language)
    module._write(folder / "snapshot.json", snapshot, module.SNAPSHOT_MAX_BYTES)
    image = BytesIO()
    Image.new("RGB", (240, 320), "white").save(image, format="PNG")
    raw = image.getvalue()
    module._mkdir(folder / "images")
    module._atomic(folder / "images" / "page-0001.png", raw)
    module._write(folder / "pages.json", [{"page_number": 1, "width_px": 240, "height_px": 320,
        "image_sha256": module._sha(raw), "rotation": 0, "render_dpi": 144,
        "backend": "fictional_test_renderer", "backend_version": "1"}])


@pytest.fixture
def local_case(tmp_path, monkeypatch):
    monkeypatch.setattr(module, "_prepare_import", fake_prepare)
    root = tmp_path / "isolated"
    service = SavedDocxLayoutService(root, mode="shadow", workspace_id="fictional")
    source, word, import_nonce = pdf_bytes(), docx_bytes(), nonce()
    imported = service.import_document(source, word, "EN", import_nonce)
    return SimpleNamespace(service=service, root=root, source=source, word=word,
        import_nonce=import_nonce, view=imported, folder=service.root / "reviews" / imported["review_id"])


def reviewed(view):
    decisions = deepcopy(view["decisions"])
    for index, row in enumerate(decisions["paragraphs"]):
        row["regions"] = [{"page_number": 1, "bbox_px": [10, 10 + index * 40, 220, 40 + index * 40]}]
    decisions["review"] = {"reviewer_kind": "operator_review", "reviewer": "Fictional reviewer",
        "note": "Fictional source comparison of all material.", "pages_reviewed": [1], "document_reviewed": True}
    return decisions


def saved_review(case):
    return case.service.save_decisions(case.view["review_id"], 1, nonce(), reviewed(case.view))


def test_import_preserves_bytes_and_restart_owns_the_snapshot(local_case):
    case = local_case
    reopened = SavedDocxLayoutService(case.root, mode="shadow", workspace_id="fictional")
    assert reopened.read(case.view["review_id"]) == case.view
    assert (case.folder / "original.docx").read_bytes() == case.word
    assert (case.folder / "source.pdf").read_bytes() == case.source
    assert reopened.image(case.view["review_id"], 1).startswith(b"\x89PNG")
    summaries = reopened.list_reviews()
    assert len(summaries) == 1 and summaries[0]["paragraph_count"] == 2
    assert "paragraphs" not in summaries[0] and "Fictional court" not in json.dumps(summaries)
    assert str(case.root) not in json.dumps(case.view)
    assert set(case.view["paragraphs"][0]) == {"id", "ordinal", "text", "has_page_break"}


def test_import_nonce_retry_and_changed_bytes_conflict(local_case):
    case = local_case
    assert case.service.import_document(case.source, case.word, "EN", case.import_nonce) == case.view
    for source, word, language in ((case.source + b"\n", case.word, "EN"),
                                  (case.source, case.word + b"\n", "EN"),
                                  (case.source, case.word, "FR")):
        with pytest.raises(SavedDocxLayoutServiceError, match="nonce_conflict") as caught:
            case.service.import_document(source, word, language, case.import_nonce)
        assert caught.value.status == 409
    assert len(case.service.list_reviews()) == 1


def test_workspace_and_mode_isolation(local_case):
    case = local_case
    other = SavedDocxLayoutService(case.root, mode="shadow", workspace_id="other")
    assert other.list_reviews() == []
    with pytest.raises(SavedDocxLayoutServiceError, match="not_found"):
        other.read(case.view["review_id"])
    wrong_mode = SavedDocxLayoutService(case.root, mode="live", workspace_id="fictional")
    with pytest.raises(SavedDocxLayoutServiceError, match="owner_changed"):
        wrong_mode.read(case.view["review_id"])


def test_save_nonce_precedes_generation_and_does_not_discard_later_edits(local_case):
    case = local_case
    first_nonce = nonce()
    decision = reviewed(case.view)
    first = case.service.save_decisions(case.view["review_id"], 1, first_nonce, decision)
    later = deepcopy(decision)
    later["paragraphs"][0]["bold"] = True
    second = case.service.save_decisions(case.view["review_id"], 2, nonce(), later)
    retry = case.service.save_decisions(case.view["review_id"], 1, first_nonce, decision)
    assert first["generation"] == 2 and retry == second and retry["generation"] == 3
    assert retry["decisions"]["paragraphs"][0]["bold"] is True
    with pytest.raises(SavedDocxLayoutServiceError, match="nonce_conflict"):
        case.service.save_decisions(case.view["review_id"], 1, first_nonce, later)
    with pytest.raises(SavedDocxLayoutServiceError, match="generation_conflict"):
        case.service.save_decisions(case.view["review_id"], 1, nonce(), later)


def test_lost_save_receipt_recovers_from_generation(local_case, monkeypatch):
    case = local_case
    original = case.service._save_receipt
    def fail_once(folder, row):
        raise OSError("Simulated private wording must never escape")
    monkeypatch.setattr(case.service, "_save_receipt", fail_once)
    save_nonce = nonce()
    with pytest.raises(SavedDocxLayoutServiceError, match="operation_failed"):
        case.service.save_decisions(case.view["review_id"], 1, save_nonce, reviewed(case.view))
    monkeypatch.setattr(case.service, "_save_receipt", original)
    reopened = SavedDocxLayoutService(case.root, mode="shadow", workspace_id="fictional")
    read = reopened.read(case.view["review_id"])
    assert read["generation"] == 2
    assert read["saves"] == [{"save_nonce": save_nonce, "generation": 2}]
    assert (case.folder / "saves" / (save_nonce + ".json")).is_file()


def test_interrupted_atomic_generation_preserves_temporary_and_previous_state(local_case, monkeypatch):
    case = local_case
    original = module.os.rename
    def interrupted(source, destination):
        if Path(destination).parent.name == "generations":
            raise OSError("Simulated interrupted save")
        return original(source, destination)
    monkeypatch.setattr(module.os, "rename", interrupted)
    with pytest.raises(SavedDocxLayoutServiceError, match="operation_failed"):
        case.service.save_decisions(case.view["review_id"], 1, nonce(), reviewed(case.view))
    assert len(list((case.folder / "generations").glob(".pending-*"))) == 1
    assert case.service.read(case.view["review_id"])["generation"] == 1


def test_invalid_review_cannot_be_attested_by_build(local_case):
    case = local_case
    with pytest.raises(SavedDocxLayoutServiceError, match="review"):
        case.service.build(case.view["review_id"], 1, nonce(), True)
    assert not (case.folder / "builds").exists()


def test_build_and_download_are_separate_verified_artifacts(local_case):
    case = local_case
    current = saved_review(case)
    operation_nonce = nonce()
    built = case.service.build(case.view["review_id"], current["generation"], operation_nonce, True)
    assert built["status"] == "built"
    artifact_id = built["artifact_id"]
    raw = case.service.artifact(case.view["review_id"], artifact_id, "docx")
    assert raw.startswith(b"PK") and (case.folder / "original.docx").read_bytes() == case.word
    receipt = json.loads(case.service.artifact(case.view["review_id"], artifact_id, "receipt"))
    assert receipt["exact_text_preserved"] is True and receipt["provider_dispatch_count"] == 0
    assert receipt["rendered_layout_acceptance"] == "not_evaluated"
    reopened = SavedDocxLayoutService(case.root, mode="shadow", workspace_id="fictional")
    assert reopened.build(case.view["review_id"], 2, operation_nonce, True) == built
    assert reopened.artifact(case.view["review_id"], artifact_id, "docx") == raw


def test_build_nonce_recovery_precedes_later_generation(local_case):
    case = local_case
    view = saved_review(case)
    operation_nonce = nonce()
    built = case.service.build(case.view["review_id"], 2, operation_nonce, True)
    decision = deepcopy(view["decisions"])
    decision["paragraphs"][0]["bold"] = True
    case.service.save_decisions(case.view["review_id"], 2, nonce(), decision)
    retry = case.service.build(case.view["review_id"], 2, operation_nonce, True)
    assert retry["generation"] == 3 and retry["artifact_id"] == built["artifact_id"]
    assert len(retry["builds"]) == 1
    with pytest.raises(SavedDocxLayoutServiceError, match="nonce_conflict"):
        case.service.build(case.view["review_id"], 3, operation_nonce, True)


def test_interrupted_build_stays_incomplete_until_explicit_new_attempt(local_case, monkeypatch):
    from legalpdf_translate import saved_docx_layout_writer as writer
    case = local_case
    saved_review(case)
    original = writer.build_docx
    def fail(*_):
        raise OSError("Private text should not be in the public error")
    monkeypatch.setattr(writer, "build_docx", fail)
    failed_nonce = nonce()
    with pytest.raises(SavedDocxLayoutServiceError, match="operation_failed"):
        case.service.build(case.view["review_id"], 2, failed_nonce, True)
    monkeypatch.setattr(writer, "build_docx", original)
    reopened = SavedDocxLayoutService(case.root, mode="shadow", workspace_id="fictional")
    view = reopened.read(case.view["review_id"])
    assert view["builds"][0]["status"] == "incomplete"
    with pytest.raises(SavedDocxLayoutServiceError, match="build_incomplete"):
        reopened.build(case.view["review_id"], 2, failed_nonce, True)
    built = reopened.build(case.view["review_id"], 2, nonce(), True)
    assert len(built["builds"]) == 2
    assert (case.folder / "builds" / failed_nonce / "intent.json").exists()


def test_complete_evidence_recovers_missing_completion_marker_without_rebuild(local_case, monkeypatch):
    from legalpdf_translate import saved_docx_layout_writer as writer
    case = local_case
    saved_review(case)
    original_atomic = module._atomic
    def stop_before_marker(path, raw):
        if path.name == "complete.json":
            raise OSError("Simulated response-loss boundary")
        return original_atomic(path, raw)
    monkeypatch.setattr(module, "_atomic", stop_before_marker)
    build_nonce = nonce()
    with pytest.raises(SavedDocxLayoutServiceError, match="operation_failed"):
        case.service.build(case.view["review_id"], 2, build_nonce, True)
    monkeypatch.setattr(module, "_atomic", original_atomic)
    monkeypatch.setattr(writer, "build_docx", lambda *_: pytest.fail("Recovery must not rebuild"))
    reopened = SavedDocxLayoutService(case.root, mode="shadow", workspace_id="fictional")
    assert reopened.read(case.view["review_id"])["builds"][0]["status"] == "built"
    assert (case.folder / "builds" / build_nonce / "complete.json").is_file()


@pytest.mark.parametrize("part", ["original.docx", "source.pdf", "snapshot.json", "pages.json"])
def test_snapshot_tampering_blocks_read(local_case, part):
    path = local_case.folder / part
    path.write_bytes(path.read_bytes() + b"\n")
    with pytest.raises(SavedDocxLayoutServiceError, match="snapshot_changed"):
        local_case.service.read(local_case.view["review_id"])


@pytest.mark.parametrize("first", [False, True])
def test_last_generation_tampering_cannot_reuse_request_hash(local_case, first):
    case = local_case
    if not first:
        saved_review(case)
    path = case.folder / "generations" / ("000001.json" if first else "000002.json")
    record = json.loads(path.read_bytes())
    record["decisions"]["paragraphs"][0]["bold"] = True
    path.write_bytes(module._encode(record))
    with pytest.raises(SavedDocxLayoutServiceError, match="generation_changed"):
        case.service.read(case.view["review_id"])


def test_raster_and_artifact_tampering_blocks_reuse(local_case):
    case = local_case
    saved_review(case)
    built = case.service.build(case.view["review_id"], 2, nonce(), True)
    row = built["builds"][0]
    docx = case.folder / "builds" / row["operation_nonce"] / "output.docx"
    docx.write_bytes(docx.read_bytes() + b"changed")
    with pytest.raises(SavedDocxLayoutServiceError, match="artifact_changed"):
        case.service.artifact(case.view["review_id"], built["artifact_id"], "docx")
    image = case.folder / "images" / "page-0001.png"
    image.write_bytes(image.read_bytes() + b"changed")
    with pytest.raises(SavedDocxLayoutServiceError, match="source_image_changed"):
        case.service.image(case.view["review_id"], 1)


def test_concurrent_review_lock_refuses_second_thread(local_case):
    case = local_case
    errors = []
    def read():
        try:
            case.service.read(case.view["review_id"])
        except SavedDocxLayoutServiceError as exc:
            errors.append((exc.code, exc.status))
    with run_workspace_slot(case.folder):
        thread = threading.Thread(target=read)
        thread.start()
        thread.join(timeout=5)
        assert not thread.is_alive()
    assert errors == [("saved_docx_layout_busy", 409)]


def test_failed_import_is_retained_but_never_reopened(tmp_path, monkeypatch):
    def fail(folder, language):
        module._fail("import_timeout", 409)
    monkeypatch.setattr(module, "_prepare_import", fail)
    service = SavedDocxLayoutService(tmp_path, mode="shadow", workspace_id="failed")
    source, word, identity = pdf_bytes(), docx_bytes(), nonce()
    with pytest.raises(SavedDocxLayoutServiceError, match="import_timeout"):
        service.import_document(source, word, "EN", identity)
    assert service.list_reviews() == []
    folder = next((service.root / "reviews").iterdir())
    assert (folder / "original.docx").read_bytes() == word
    assert (folder / "import_failed.json").is_file()
    with pytest.raises(SavedDocxLayoutServiceError, match="import_incomplete"):
        service.import_document(source, word, "EN", identity)


def test_source_changed_during_preparation_cannot_bind_old_rasters_to_new_pdf(tmp_path, monkeypatch):
    def changing_prepare(folder, language):
        fake_prepare(folder, language)
        source = folder / "source.pdf"
        source.write_bytes(source.read_bytes() + b"\n%changed during preparation")
    monkeypatch.setattr(module, "_prepare_import", changing_prepare)
    service = SavedDocxLayoutService(tmp_path, mode="shadow", workspace_id="changed-import")
    with pytest.raises(SavedDocxLayoutServiceError, match="import_inputs_changed"):
        service.import_document(pdf_bytes(), docx_bytes(), "EN", nonce())
    assert service.list_reviews() == []
    folder = next((service.root / "reviews").iterdir())
    assert not (folder / "import_complete.json").exists()


def test_limits_apply_before_any_import_storage(tmp_path, monkeypatch):
    service = SavedDocxLayoutService(tmp_path, mode="shadow", workspace_id="limits")
    monkeypatch.setattr(module, "DOCX_MAX_BYTES", 2)
    with pytest.raises(SavedDocxLayoutServiceError, match="input_limit") as caught:
        service.import_document(b"pdf", b"abc", "EN", nonce())
    assert caught.value.status == 413 and not service.root.exists()


def test_generation_and_registry_limits_preserve_existing(local_case, monkeypatch):
    case = local_case
    monkeypatch.setattr(module, "MAX_GENERATIONS", 1)
    with pytest.raises(SavedDocxLayoutServiceError, match="generation_limit"):
        case.service.save_decisions(case.view["review_id"], 1, nonce(), reviewed(case.view))
    monkeypatch.setattr(module, "MAX_REVIEWS", 1)
    with pytest.raises(SavedDocxLayoutServiceError, match="registry_limit"):
        case.service.import_document(case.source, case.word, "EN", nonce())
    assert case.service.read(case.view["review_id"])["generation"] == 1


def test_capabilities_and_no_settings_or_network_dependencies(tmp_path, monkeypatch):
    import legalpdf_translate.user_settings as settings
    monkeypatch.setattr(settings, "load_gui_settings_from_path", lambda *_a, **_k: pytest.fail("No settings reads"))
    service = SavedDocxLayoutService(tmp_path, mode="shadow", workspace_id="caps")
    caps = service.capabilities()
    assert caps["languages"] == ["EN", "AR", "FR"] and caps["provider_required"] is False
    assert caps["limits"]["source_pages"] == 100 and not service.root.exists()


def test_actual_owned_worker_imports_fictional_pdf_and_docx(tmp_path):
    service = SavedDocxLayoutService(tmp_path, mode="shadow", workspace_id="actual-worker")
    result = service.import_document(pdf_bytes(), docx_bytes("FR"), "FR", nonce())
    assert result["pages"][0]["backend"] == "pymupdf"
    assert result["pages"][0]["width_px"] == 480
    assert result["paragraphs"][0]["text"] == "Avis fictif du tribunal"
    image = Image.open(BytesIO(service.image(result["review_id"], 1)))
    assert image.size == (480, 640)


def test_actual_owned_worker_rejects_image_disguised_as_source_pdf(tmp_path):
    image = BytesIO()
    Image.new("RGB", (24, 24), "white").save(image, format="PNG")
    service = SavedDocxLayoutService(tmp_path, mode="shadow", workspace_id="not-pdf")
    with pytest.raises(SavedDocxLayoutServiceError):
        service.import_document(image.getvalue(), docx_bytes(), "EN", nonce())
    assert service.list_reviews() == []
    folder = next((service.root / "reviews").iterdir())
    assert (folder / "source.pdf").read_bytes() == image.getvalue()
    assert not (folder / "import_complete.json").exists()


@pytest.mark.parametrize("per_page", [False, True])
def test_owned_worker_deadlines_terminate_only_that_worker(tmp_path, monkeypatch, per_page):
    class Channel:
        def __init__(self):
            self.first = True
        def poll(self, timeout):
            if per_page and self.first:
                self.first = False
                return True
            return False
        def recv(self):
            return ("page", 1)
        def close(self):
            pass
    class Worker:
        exitcode = None
        alive = True
        terminated = closed = False
        def start(self):
            pass
        def is_alive(self):
            return self.alive
        def terminate(self):
            self.terminated = True
            self.alive = False
            self.exitcode = -1
        def join(self, timeout):
            pass
        def close(self):
            self.closed = True
        def kill(self):
            pytest.fail("Terminate already stopped the owned worker")
    worker = Worker()
    context = SimpleNamespace(Pipe=lambda **_: (Channel(), Channel()), Process=lambda **_: worker)
    monkeypatch.setattr(module.multiprocessing, "get_context", lambda _: context)
    times = iter(range(20))
    monkeypatch.setattr(module.time, "monotonic", lambda: float(next(times)))
    monkeypatch.setattr(module, "IMPORT_TIMEOUT_SECONDS", 100 if per_page else 1.5)
    monkeypatch.setattr(module, "PAGE_TIMEOUT_SECONDS", 0.5)
    with pytest.raises(SavedDocxLayoutServiceError, match="import_timeout"):
        module._prepare_import(tmp_path, "EN")
    assert worker.terminated and worker.closed


def test_symlink_root_is_rejected(tmp_path, local_case):
    link = tmp_path / "linked-root"
    try:
        link.symlink_to(local_case.root, target_is_directory=True)
    except OSError:
        pytest.skip("Windows symlink privilege unavailable")
    with pytest.raises(SavedDocxLayoutServiceError, match="indirect_path"):
        SavedDocxLayoutService(link, mode="shadow", workspace_id="fictional")


def test_hardlinked_snapshot_is_rejected(tmp_path, local_case):
    os.link(local_case.folder / "original.docx", tmp_path / "extra-word-link.docx")
    with pytest.raises(SavedDocxLayoutServiceError, match="indirect_path"):
        local_case.service.read(local_case.view["review_id"])
