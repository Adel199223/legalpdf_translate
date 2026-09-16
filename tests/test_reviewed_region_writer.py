"""Synthetic in-memory region contracts; no native or rendered acceptance."""
from copy import deepcopy
import hashlib
import io
import json

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree
import pytest

from legalpdf_translate.document_structure import text_sha256
from legalpdf_translate.reviewed_formatting_writer import (
    ReviewedFormattingWriterError, build_reviewed_docx, validate_reviewed_docx,
)
from legalpdf_translate.reviewed_region_writer import MAP_VERSION, WRITER_VERSION
from tests.test_reviewed_formatting import packet, rebind, validate
from tests.test_reviewed_formatting_v2 import grouped_packet, region_packet
from tests.test_reviewed_formatting_writer import changed_document, changed_package, decoded


def projection(lang="FR", *, empty_cells=False, paragraph_end=False):
    manifest, page = region_packet(columns=not paragraph_end)
    manifest["target_lang"] = lang
    if lang == "AR":
        parts = ["المحكمة القضائية\nالدائرة الجنائية\n\n",
                 "يُحافَظ على مهلة عشرة أيام.\nيحضر [[João Guerreiro]] في [[09]]:[[30]].\n\n",
                 "القاضية\n[[Nome Exemplo]]\n\n",
                 "[[Largo do Exemplo - 1234-567 Exemplo]]\nالهاتف: [[123456789]]\n", "1 / 1"]
        page.target_structure["blocks"][0]["text"] = "".join(parts)
        page.target_structure["translation_sha256"] = text_sha256("".join(parts))
        offset = 0
        for row, part in zip(manifest["pages"][0]["fragments"], parts):
            row["target_range"] = [offset, offset + len(part)]
            row["alignment"] = "right"
            offset += len(part)
        rebind(manifest, page)
    if empty_cells:
        table = manifest["pages"][0]["region_layout"]["body"][0]
        left, right = table["rows"][0]["cells"]
        table["column_widths"] = [30, 10, 40, 20]
        table["rows"][0]["cells"] = [left, {"fragment_ids": []}, right, {"fragment_ids": []}]
    return validate(manifest, page, target_lang=lang)


def located(document, row):
    loc = row["location"]
    if loc["kind"] == "body_table_cell":
        return document.tables[loc["table_index"]].cell(loc["row_index"], loc["column_index"]).paragraphs[loc["paragraph_index"]]
    if loc["kind"] == "body_paragraph":
        return document.paragraphs[loc["paragraph_index"]]
    part = "header" if loc["kind"] == "section_header" else "footer"
    return getattr(document.sections[loc["section_index"]], part).paragraphs[loc["paragraph_index"]]


@pytest.mark.parametrize("lang", ["FR", "EN", "AR"])
@pytest.mark.parametrize("empty_cells", [False, True])
def test_regions_preserve_exact_source_partition_and_fixed_typography(lang, empty_cells):
    reviewed = projection(lang, empty_cells=empty_cells)
    before = deepcopy(reviewed)
    artifact = build_reviewed_docx(reviewed)
    doc, mapping = decoded(artifact)
    validate_reviewed_docx(artifact.docx_bytes, artifact.source_map_bytes, projection=reviewed)
    assert reviewed == before
    assert mapping["version"] == MAP_VERSION and mapping["writer_version"] == WRITER_VERSION
    assert mapping["document_groups"] == [{"start_page": 1, "end_page": 1}]
    assert len(mapping["validated_folios"]) == 1
    assert mapping["rendered_layout_acceptance"] == "not_evaluated"
    assert mapping["rendered_page_count"] is None and mapping["layout_review_required"]
    assert mapping["pages"][0]["region_layout"] == json.loads(reviewed.pages[0].region_layout_json)
    assert len(doc.tables) == 1 and not doc.tables[0].autofit
    assert doc.tables[0]._tbl.tblPr.find(qn("w:bidiVisual")).get(qn("w:val")) == "0"
    widths = [int(node.get(qn("w:w"))) for node in doc.tables[0]._tbl.tblGrid]
    assert sum(widths) == Cm(17.6).twips
    assert len(widths) == (4 if empty_cells else 2)
    assert [int(cell._tc.tcPr.find(qn("w:tcW")).get(qn("w:w"))) for cell in doc.tables[0].rows[0].cells] == widths
    assert not list(doc._element.iter(qn("w:trHeight")))
    assert not list(doc._element.iter(qn("w:vanish")))
    assert not list(doc._element.iter(qn("w:fldChar")))
    assert doc.settings._element.find(".//" + qn("w:doNotExpandShiftReturn")).get(qn("w:val")) == "1"
    compat = doc.settings._element.find(qn("w:compat"))
    names = [child.tag for child in compat]
    assert names.index(qn("w:doNotExpandShiftReturn")) < names.index(qn("w:useFELayout"))
    assert names.index(qn("w:doNotExpandShiftReturn")) < names.index(qn("w:compatSetting"))
    for row, original in zip(mapping["pages"][0]["fragments"], reviewed.pages[0].fragments):
        p = located(doc, row)
        visible = p.text.replace("\u200e", "") if lang == "AR" else p.text
        assert hashlib.sha256(visible.encode()).hexdigest() == row["display_text_sha256"]
        assert row["source_text_sha256"] == original.source_text_sha256
        assert row["target_text_sha256"] == original.target_text_sha256
        assert row["target_range"] == list(original.target_range)
        assert 0 <= row["spacing"]["space_before_pt"] <= 48
        for run in p.runs:
            assert run.font.name == ("Arial" if lang == "AR" else "Times New Roman")
            assert run.font.size.pt == (11 if lang == "AR" else 10.5)
        if lang == "AR":
            assert p._p.pPr.find(qn("w:bidi")).get(qn("w:val")) == "1"
    if lang == "AR":
        body = located(doc, mapping["pages"][0]["fragments"][1])
        assert "João Guerreiro" in body.text and "09:30" in body.text.replace("\u200e", "")
        assert {run._r.rPr.find(qn("w:rtl")).get(qn("w:val")) for run in body.runs} == {"0", "1"}
    anchors = mapping["pages"][0]["structural_paragraphs"]
    assert sum(row["reason"] == "table_end_section_anchor" for row in anchors) == 1
    assert sum(row["reason"] == "empty_table_cell" for row in anchors) == (2 if empty_cells else 0)
    assert len(doc.paragraphs) == 1 and doc.paragraphs[0].text == ""
    assert not doc.paragraphs[0].runs and doc.paragraphs[0].paragraph_format.line_spacing.pt == 1


@pytest.mark.parametrize("lang", ["FR", "EN", "AR"])
def test_multiple_table_ended_pages_keep_local_folios_and_independent_sections(lang):
    manifest, pages = grouped_packet(lang)
    for number, row in enumerate(manifest["pages"], 1):
        body_id = row["region_layout"]["body"][0]["fragment_id"]
        row["region_layout"]["body"] = [{"kind": "table", "table_id": f"p{number:04d}_t0001",
            "column_widths": [100], "rows": [{"cells": [{"fragment_ids": [body_id]}]}]}]
    reviewed = validate(manifest, pages[0], pages=pages, target_lang=lang)
    artifact = build_reviewed_docx(reviewed)
    doc, mapping = decoded(artifact)
    validate_reviewed_docx(artifact.docx_bytes, artifact.source_map_bytes, projection=reviewed)
    assert len(doc.sections) == len(doc.tables) == len(doc.paragraphs) == 9
    assert mapping["document_groups"] == [{"start_page": 1, "end_page": 2},
        {"start_page": 3, "end_page": 3}, {"start_page": 4, "end_page": 9}]
    for index, (section, page) in enumerate(zip(doc.sections, mapping["pages"])):
        assert not section.header.is_linked_to_previous and not section.footer.is_linked_to_previous
        assert section.page_width.twips == Cm(21).twips and section.page_height.twips == Cm(29.7).twips
        assert section.left_margin.twips == section.right_margin.twips == Cm(1.7).twips
        assert section.top_margin.twips == section.bottom_margin.twips == Cm(1.5).twips
        assert len(section.header.paragraphs) == 1 and section.header.paragraphs[0].text == ""
        anchor = page["structural_paragraphs"][-1]
        assert anchor["location"]["body_child_index"] == index * 2 + 1
        assert (doc.paragraphs[index]._p.pPr.find(qn("w:sectPr")) is not None) == (index < 8)
        assert not list(section._sectPr.iter(qn("w:pgNumType")))


def test_paragraph_end_never_gets_table_anchor_and_v1_keeps_original_profile():
    reviewed = projection(paragraph_end=True)
    doc, mapping = decoded(build_reviewed_docx(reviewed))
    assert not doc.tables and all(p.text for p in doc.paragraphs)
    assert mapping["pages"][0]["structural_paragraphs"] == []
    manifest, page = packet()
    doc, old = decoded(build_reviewed_docx(validate(manifest, page)))
    assert old["version"] == "reviewed_formatting_source_map_v1"
    assert old["writer_version"] == "reviewed_formatting_writer_v1"
    assert doc.settings._element.find(".//" + qn("w:doNotExpandShiftReturn")) is None


def test_first_table_can_own_letterhead_and_multiple_rows_with_structural_empty_cell():
    manifest, page = region_packet(columns=True)
    row = manifest["pages"][0]
    row["fragments"][0]["bbox_px"] = [20, 20, 270, 70]
    row["fragments"][1]["bbox_px"] = [310, 20, 550, 100]
    row["fragments"][2]["bbox_px"] = [310, 150, 550, 250]
    row["region_layout"]["header"] = []
    table = row["region_layout"]["body"][0]
    table["rows"] = [
        {"cells": [{"fragment_ids": ["p0001_f0001"]}, {"fragment_ids": ["p0001_f0002"]}]},
        {"cells": [{"fragment_ids": []}, {"fragment_ids": ["p0001_f0003"]}]}]
    reviewed = validate(manifest, page)
    artifact = build_reviewed_docx(reviewed)
    doc, mapping = decoded(artifact)
    assert len(doc.tables[0].rows) == 2
    assert doc.tables[0].cell(1, 0).text == ""
    assert doc.sections[0].header.paragraphs[0].text == ""
    for fragment, mapped in zip(reviewed.pages[0].fragments, mapping["pages"][0]["fragments"]):
        assert hashlib.sha256(located(doc, mapped).text.encode()).hexdigest() == mapped["display_text_sha256"]
    validate_reviewed_docx(artifact.docx_bytes, artifact.source_map_bytes, projection=reviewed)


def test_interleaved_source_ids_in_columns_keep_multiple_paragraph_locators():
    manifest, page = region_packet(columns=True)
    row = manifest["pages"][0]
    row["fragments"][0]["bbox_px"] = [20, 20, 270, 70]
    row["fragments"][1]["bbox_px"] = [310, 20, 550, 100]
    row["fragments"][2]["bbox_px"] = [20, 150, 270, 250]
    row["region_layout"]["header"] = []
    row["region_layout"]["body"][0]["rows"] = [{"cells": [
        {"fragment_ids": ["p0001_f0001", "p0001_f0003"]},
        {"fragment_ids": ["p0001_f0002"]}]}]
    reviewed = validate(manifest, page)
    artifact = build_reviewed_docx(reviewed)
    doc, mapping = decoded(artifact)
    assert len(doc.tables[0].cell(0, 0).paragraphs) == 2
    third = mapping["pages"][0]["fragments"][2]
    assert third["location"]["column_index"] == 0
    assert third["location"]["paragraph_index"] == 1
    assert third["spacing"]["space_before_pt"] == 48
    assert hashlib.sha256(located(doc, third).text.encode()).hexdigest() == third["display_text_sha256"]
    validate_reviewed_docx(artifact.docx_bytes, artifact.source_map_bytes, projection=reviewed)


@pytest.mark.parametrize("mutate", [
    lambda d: setattr(d.tables[0], "autofit", True),
    lambda d: d.tables[0]._tbl.tblGrid[0].set(qn("w:w"), "1"),
    lambda d: d.tables[0].cell(0, 0)._tc.tcPr.find(qn("w:tcW")).set(qn("w:w"), "1"),
    lambda d: d.tables[0]._tbl.tblPr.find(qn("w:bidiVisual")).set(qn("w:val"), "1"),
    lambda d: d.tables[0].add_row(),
    lambda d: d.tables[0].cell(0, 0).add_paragraph("extra"),
    lambda d: d.tables[0].cell(0, 0).add_table(rows=1, cols=1),
    lambda d: d.tables[0].cell(0, 1).paragraphs[0].add_run("unowned empty-cell words"),
    lambda d: d.tables[0].cell(0, 0).paragraphs[0].add_run("extra"),
    lambda d: setattr(d.tables[0].cell(0, 0).paragraphs[0].runs[0].font, "hidden", True),
    lambda d: setattr(d.tables[0].cell(0, 0).paragraphs[0].runs[0].font, "size", Pt(8)),
    lambda d: d.tables[0].cell(0, 0).paragraphs[0]._p.getparent().remove(d.tables[0].cell(0, 0).paragraphs[0]._p),
    lambda d: d.tables[0]._tbl.tr_lst[0].append(d.tables[0]._tbl.tr_lst[0].tc_lst[0]),
    lambda d: d.tables[0]._tbl.getparent().remove(d.tables[0]._tbl),
    lambda d: d.paragraphs[0].add_run("unowned anchor words"),
    lambda d: setattr(d.paragraphs[0].paragraph_format, "line_spacing", Pt(2)),
    lambda d: d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p),
    lambda d: d.add_paragraph(),
    lambda d: setattr(d.sections[0], "bottom_margin", Cm(.2)),
    lambda d: d.settings._element.find(".//" + qn("w:doNotExpandShiftReturn")).set(qn("w:val"), "0"),
])
def test_table_content_geometry_empty_cells_and_anchor_tampering_rejected_after_rehash(mutate):
    reviewed = projection(empty_cells=True)
    artifact = build_reviewed_docx(reviewed)
    raw, mapping = changed_document(artifact, mutate)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(raw, mapping, projection=reviewed)


def test_wrong_arabic_run_direction_rejected_after_rehash():
    reviewed = projection("AR")
    artifact = build_reviewed_docx(reviewed)
    def mutate(doc):
        doc.tables[0].cell(0, 0).paragraphs[0].runs[0]._r.rPr.find(qn("w:rtl")).set(qn("w:val"), "0")
    raw, mapping = changed_document(artifact, mutate)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(raw, mapping, projection=reviewed)


@pytest.mark.parametrize("change", [
    lambda m: m.update(version="reviewed_formatting_source_map_v1"),
    lambda m: m.update(document_groups=[{"start_page": 1, "end_page": 2}]),
    lambda m: m.update(validated_folios=[]),
    lambda m: m["pages"][0]["fragments"][1]["location"].update(column_index=1),
    lambda m: m["pages"][0]["fragments"][1]["spacing"].update(space_before_pt=0),
    lambda m: m["pages"][0].update(structural_paragraphs=[]),
])
def test_region_map_and_structural_ownership_cannot_be_rebound(change):
    reviewed = projection()
    artifact = build_reviewed_docx(reviewed)
    mapping = json.loads(artifact.source_map_bytes)
    change(mapping)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(artifact.docx_bytes, json.dumps(mapping).encode(), projection=reviewed)


@pytest.mark.parametrize("path,attribute", [(".//w:tbl", "text"), (".//w:tc", "tail"),
    (".//w:tblPr", "text"), (".//w:tr", "tail")])
def test_unowned_xml_text_around_tables_rejected(path, attribute):
    reviewed = projection()
    artifact = build_reviewed_docx(reviewed)
    def mutate(raw):
        root = etree.fromstring(raw)
        node = root.find(path, {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
        setattr(node, attribute, "unowned XML words")
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    raw, mapping = changed_package(artifact, "word/document.xml", mutate)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(raw, mapping, projection=reviewed)
