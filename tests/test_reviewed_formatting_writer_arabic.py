"""Synthetic Arabic partition contracts; no native, paid or visual acceptance."""
from copy import deepcopy
from dataclasses import replace
import hashlib
import io
import json

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from lxml import etree
import pytest

from legalpdf_translate.document_structure import text_sha256
from legalpdf_translate.docx_writer import sanitize_bidi_controls, unwrap_internal_placeholders
from legalpdf_translate.reviewed_formatting_writer import (
    ReviewedFormattingWriterError, build_reviewed_docx, validate_reviewed_docx,
)
from tests.test_reviewed_formatting import packet, rebind, validate
from tests.test_reviewed_formatting_writer import changed_document, changed_package


AR_DISPLAY_POLICY = "unwrap_placeholders_strip_bidi_crlf_rtl_runs_lrm_v1"
BODY = ("يُستدعى \u2066[[João Guerreiro]]\u2069 و[[Cafe\u0301]] عند [[09]]:[[30]].\r\n"
        "نص كامل لا يُختصر، مع وصل\u200dة و فصل\u200cة.\r\n \r\n")


def arabic_projection(*, pages=1, body=BODY, alignment="justify", bold=False, italic=False,
                      absent_middle_footer=False):
    inputs, rows = [], []
    for number in range(1, pages + 1):
        target_parts = [f"المحكمة القضائية رقم {number}\r\nالدائرة الجنائية\r\n\r\n",
                        body, "القاضية\r\n[[João Guerreiro]]\r\n\r\n",
                        "[[Rua João de Deus]], [[20]]\r\n[[7800-001]] [[Beja]]\r\n",
                        f"{number} / {pages}"]
        manifest, page = packet(target_parts=target_parts)
        row = manifest["pages"][0]
        absent = absent_middle_footer and number == 2
        if absent:
            for side in ("source", "target"):
                original = getattr(page, side + "_structure")["blocks"][0]
                original["text"] = original["text"][:row["fragments"][3][side + "_range"][0]]
            row["fragments"] = row["fragments"][:3]
        else:
            page.source_structure["blocks"][0]["text"] = (
                page.source_structure["blocks"][0]["text"][:-5] + f"{number} / {pages}")
        for structure in (page.source_structure, page.target_structure):
            structure["page_number"] = number
            structure["blocks"][0]["id"] = f"p{number:04d}_b100000001"
            structure["blocks"][0]["document_start"] = number in {1, 3}
        source_hash = text_sha256(page.source_structure["blocks"][0]["text"])
        for structure in (page.source_structure, page.target_structure):
            structure["source_sha256"] = structure["source_text_sha256"] = source_hash
        page.target_structure["translation_sha256"] = text_sha256(page.target_structure["blocks"][0]["text"])
        page = replace(page, commit_file_sha256=text_sha256(f"arabic-commit-{number}"),
                       bundle_sha256=text_sha256(f"arabic-bundle-{number}"))
        row.update(page_number=number, commit_file_sha256=page.commit_file_sha256,
                   bundle_sha256=page.bundle_sha256)
        for index, fragment in enumerate(row["fragments"], 1):
            fragment.update(rendering_id=f"p{number:04d}_f{index:04d}",
                            parent_block_id=f"p{number:04d}_b100000001")
            fragment["alignment"] = alignment if fragment["role"] == "body" else (
                "center" if fragment["role"] in {"header", "folio"} else "right")
            fragment.update(bold=bold, italic=italic)
        rebind(manifest, page)
        rows.append(row)
        inputs.append(page)
    manifest.update(target_lang="AR", full_case_pages=list(range(1, pages + 1)), pages=rows)
    return validate(manifest, inputs[0], pages=inputs, target_lang="AR")


def decoded(artifact):
    return Document(io.BytesIO(artifact.docx_bytes)), json.loads(artifact.source_map_bytes)


def located(document, page, row):
    location = row["location"]
    if location["kind"] == "body_paragraph":
        return document.paragraphs[location["paragraph_index"]]
    section = document.sections[page["section_index"]]
    part = section.header if location["kind"] == "section_header" else section.footer
    return part.paragraphs[location["paragraph_index"]]


def visible(raw):
    return sanitize_bidi_controls(unwrap_internal_placeholders(raw)).replace("\r\n", "\n")


@pytest.mark.parametrize("bold,italic", [(False, False), (True, False), (False, True), (True, True)])
def test_arabic_profile_preserves_raw_partitions_and_exact_mixed_runs(bold, italic):
    projection = arabic_projection(bold=bold, italic=italic)
    before = deepcopy(projection)
    artifact = build_reviewed_docx(projection)
    document, mapping = decoded(artifact)
    validate_reviewed_docx(artifact.docx_bytes, artifact.source_map_bytes, projection=projection)
    assert projection == before
    assert mapping["typography"]["font"] == "Arial"
    assert mapping["typography"]["size_pt"] == 11
    assert mapping["rendered_page_count"] is None
    assert mapping["rendered_layout_acceptance"] == "not_evaluated"
    normal = document.styles["Normal"]
    assert normal.font.name == "Arial" and normal.font.size.pt == 11
    assert normal._element.rPr.rFonts.get(qn("w:cs")) == "Arial"
    assert normal._element.rPr.find(qn("w:szCs")).get(qn("w:val")) == "22"
    page = mapping["pages"][0]
    for source, row in zip(projection.pages[0].fragments, page["fragments"]):
        paragraph = located(document, page, row)
        assert row["display_policy"] == AR_DISPLAY_POLICY
        assert row["target_text_sha256"] == source.target_text_sha256
        assert row["target_range"] == list(source.target_range)
        display = visible(source.target_text[:row["raw_content_end"]])
        assert sanitize_bidi_controls(paragraph.text) == display
        assert hashlib.sha256(display.encode()).hexdigest() == row["display_text_sha256"]
        assert source.target_text[row["raw_content_end"]:] == row["trailing_separators"]
        assert paragraph._p.pPr.find(qn("w:bidi")).get(qn("w:val")) == "1"
        assert paragraph.paragraph_format.line_spacing == 1.0
        for run in paragraph.runs:
            rpr = run._r.rPr
            assert run.font.name == "Arial" and run.font.size.pt == 11
            assert all(rpr.rFonts.get(qn("w:" + slot)) == "Arial" for slot in ("ascii", "hAnsi", "eastAsia", "cs"))
            assert rpr.find(qn("w:szCs")).get(qn("w:val")) == "22"
            assert run.bold is bold and run.italic is italic
            assert run.font.cs_bold is bold and run.font.cs_italic is italic
            direction = rpr.find(qn("w:rtl")).get(qn("w:val"))
            lang = rpr.find(qn("w:lang"))
            assert direction in {"0", "1"}
            assert lang.get(qn("w:val")) == ("en-US" if direction == "0" else "ar-SA")
            assert lang.get(qn("w:bidi")) == (None if direction == "0" else "ar-SA")
    body = document.paragraphs[0]
    for token in ("João Guerreiro", "Cafe\u0301", "09:30"):
        matching = [run for run in body.runs if token in run.text]
        assert len(matching) == 1
        assert matching[0].text == "\u200e" + token + "\u200e"
        assert matching[0]._r.rPr.find(qn("w:rtl")).get(qn("w:val")) == "0"
    assert "\u200d" in body.text and "\u200c" in body.text
    assert page["fragments"][1]["trailing_separators"] == "\r\n \r\n"
    footer = document.sections[0].footer.paragraphs[0]
    assert [run.text for run in footer.runs] == ["Rua João de Deus, 20\n", "7800-001 Beja"]
    assert document.sections[0].footer.paragraphs[-1].text == "1 / 1"


@pytest.mark.parametrize("alignment,jc", [("right", "start"), ("left", "end"),
                                          ("center", "center"), ("justify", "both")])
def test_arabic_alignment_is_explicit_physical_alignment(alignment, jc):
    artifact = build_reviewed_docx(arabic_projection(alignment=alignment))
    document, _ = decoded(artifact)
    assert document.paragraphs[0]._p.pPr.jc.get(qn("w:val")) == jc
    assert document.sections[0].header.paragraphs[0]._p.pPr.jc.get(qn("w:val")) == "center"


def test_full_case_retains_sections_gaps_document_boundaries_and_absent_footer():
    projection = arabic_projection(pages=3, absent_middle_footer=True)
    artifact = build_reviewed_docx(projection)
    document, mapping = decoded(artifact)
    assert len(document.sections) == 3 and len(document.paragraphs) == 6
    assert all(paragraph.text for paragraph in document.paragraphs)
    assert [json.loads(page.source_structure_json)["blocks"][0]["document_start"]
            for page in projection.pages] == [True, False, True]
    parts = set()
    for section, page in zip(document.sections, mapping["pages"]):
        assert not section.header.is_linked_to_previous and not section.footer.is_linked_to_previous
        assert section.page_width.twips == Cm(21).twips and section.page_height.twips == Cm(29.7).twips
        assert section.left_margin.twips == Cm(1.7).twips and section.top_margin.twips == Cm(1.5).twips
        for uri in page["section_parts"].values():
            assert uri not in parts
            parts.add(uri)
        for row in page["fragments"]:
            assert 0 <= row["spacing"]["space_before_pt"] <= 48
            assert located(document, page, row).paragraph_format.space_before.pt == row["spacing"]["space_before_pt"]
    assert document.sections[0].footer.paragraphs[-1].text == "1 / 3"
    assert len(document.sections[1].footer.paragraphs) == 1
    assert document.sections[1].footer.paragraphs[0].text == ""
    assert document.sections[2].footer.paragraphs[-1].text == "3 / 3"
    validate_reviewed_docx(artifact.docx_bytes, artifact.source_map_bytes, projection=projection)


def _tamper(document, case):
    paragraph = document.paragraphs[0]
    rtl = next(run for run in paragraph.runs if run._r.rPr.find(qn("w:rtl")).get(qn("w:val")) == "1")
    ltr = next(run for run in paragraph.runs if "João Guerreiro" in run.text)
    if case == "paragraph_bidi":
        paragraph._p.pPr.find(qn("w:bidi")).set(qn("w:val"), "0")
    elif case == "legacy_alignment":
        paragraph._p.pPr.jc.set(qn("w:val"), "right")
    elif case == "run_direction":
        ltr._r.rPr.find(qn("w:rtl")).set(qn("w:val"), "1")
    elif case == "language":
        rtl._r.rPr.find(qn("w:lang")).set(qn("w:bidi"), "en-US")
    elif case == "complex_font":
        rtl._r.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    elif case == "complex_size":
        rtl._r.rPr.find(qn("w:szCs")).set(qn("w:val"), "21")
    elif case in {"complex_bold", "complex_italic"}:
        tag = qn("w:bCs" if case == "complex_bold" else "w:iCs")
        rtl._r.rPr.remove(rtl._r.rPr.find(tag))
    elif case == "extra_control":
        ltr.text = "\u202e" + ltr.text
    elif case == "missing_lrm":
        ltr.text = ltr.text.replace("\u200e", "")
    elif case == "moved_lrm":
        ltr.text = ltr.text.replace("\u200eJoão", "Jo\u200eão")
    elif case == "missing_space_preservation":
        space = "{http://www.w3.org/XML/1998/namespace}space"
        node = next(node for node in paragraph._p.iter(qn("w:t")) if node.get(space) == "preserve")
        del node.attrib[space]
    elif case == "split_name":
        duplicate = deepcopy(ltr._r)
        ltr.text = "\u200eJoão "
        from docx.text.run import Run
        Run(duplicate, paragraph).text = "Guerreiro\u200e"
        ltr._r.addnext(duplicate)
    elif case == "extra_empty_run":
        paragraph.add_run()
    elif case == "hidden_run":
        rtl.font.hidden = True
    elif case == "linked_footer":
        document.sections[1].footer.is_linked_to_previous = True
    else:
        raise AssertionError(case)


@pytest.mark.parametrize("case", ["paragraph_bidi", "legacy_alignment", "run_direction", "language",
    "complex_font", "complex_size", "complex_bold", "complex_italic", "extra_control",
    "missing_lrm", "moved_lrm", "missing_space_preservation", "split_name", "extra_empty_run", "hidden_run", "linked_footer"])
def test_arabic_direction_and_exact_run_policy_cannot_be_rebound(case):
    projection = arabic_projection(pages=2)
    artifact = build_reviewed_docx(projection)
    raw, mapping = changed_document(artifact, lambda document: _tamper(document, case))
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(raw, mapping, projection=projection)


@pytest.mark.parametrize("field,value", [("display_policy", "unwrap_placeholders_strip_bidi_crlf_v1"),
    ("raw_content_end", 1), ("display_text_sha256", "f" * 64)])
def test_arabic_map_cannot_hide_display_or_partition_changes(field, value):
    projection = arabic_projection()
    artifact = build_reviewed_docx(projection)
    mapping = json.loads(artifact.source_map_bytes)
    mapping["pages"][0]["fragments"][1][field] = value
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(artifact.docx_bytes, json.dumps(mapping).encode(), projection=projection)


def test_arabic_normal_style_complex_font_is_checked_independently():
    projection = arabic_projection()
    artifact = build_reviewed_docx(projection)
    def mutate(raw):
        root = etree.fromstring(raw)
        normal = root.find("./" + qn("w:style") + "[@" + qn("w:styleId") + "='Normal']")
        normal.find("./" + qn("w:rPr") + "/" + qn("w:rFonts")).set(qn("w:cs"), "Calibri")
        return etree.tostring(root)
    raw, mapping = changed_package(artifact, "word/styles.xml", mutate)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(raw, mapping, projection=projection)
