"""Versioned package inspection keeps selection and historical notices honest."""
from copy import deepcopy
from dataclasses import replace
import hashlib
import json
from types import SimpleNamespace

from docx import Document
import pytest

from legalpdf_translate.formatting_support import digest_text
from legalpdf_translate.layout_integration import collect_docx_layout_review, merge_layout_review_queue
from legalpdf_translate.reviewed_formatting_writer import build_reviewed_docx
from tests.test_reviewed_formatting_writer import projection as v1_projection, changed_document
from tests.test_reviewed_region_writer import projection as v2_projection


def encode(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


@pytest.fixture(params=["reviewed_v1", "reviewed_v2"])
def reviewed_case(tmp_path, request):
    projection = v1_projection() if request.param == "reviewed_v1" else v2_projection()
    artifact = build_reviewed_docx(projection)
    pages_dir = tmp_path / "pages"
    pages_dir.mkdir()
    for page in projection.pages:
        path = pages_dir / f"page_{page.page_number:04d}.txt"
        target = json.loads(page.target_structure_json)
        text = page.original_parent_separator.join(block["text"] for block in target["blocks"])
        path.write_bytes(text.encode("utf-8"))
        path.with_suffix(".source_structure.json").write_bytes(page.source_structure_json)
        path.with_suffix(".structure.json").write_bytes(page.target_structure_json)
    output = tmp_path / "reviewed.docx"
    output.write_bytes(artifact.docx_bytes)
    output.with_suffix(".source_map.json").write_bytes(artifact.source_map_bytes)
    return SimpleNamespace(projection=projection, artifact=artifact, pages_dir=pages_dir, output=output)


def collect(case, *, pages=(1,), projection=True, preparation=None):
    return collect_docx_layout_review(case.output, case.pages_dir, preparation or {}, page_numbers=pages,
        reviewed_projection=case.projection if projection else None)


def test_reviewed_versions_use_actual_package_and_keep_render_review_required(reviewed_case):
    case = reviewed_case
    originals = {path: path.read_bytes() for path in case.pages_dir.iterdir()}
    assert collect(case) == {1: ["layout_review_required"]}
    assert collect(case, pages=None) == {1: ["layout_review_required"]}
    assert all(path.read_bytes() == raw for path, raw in originals.items())


def test_explicit_done_selection_ignores_unselected_txt_and_preparation_notices(reviewed_case):
    case = reviewed_case
    # A stale/failed unselected page must not be required, decoded or returned.
    (case.pages_dir / "page_0009.txt").write_bytes(b"\xff UNSELECTED")
    preparation = {"legacy_pages": [9], "review_required_pages": [9]}
    assert collect(case, preparation=preparation) == {1: ["layout_review_required"]}
    all_pages = collect(case, pages=None, preparation=preparation)
    assert set(all_pages) == {1, 9}
    assert all("layout_mapping_unavailable" in codes for codes in all_pages.values())


def test_selected_missing_txt_is_reported_instead_of_removed(reviewed_case):
    case = reviewed_case
    (case.pages_dir / "page_0001.txt").unlink()
    assert "layout_mapping_unavailable" in collect(case)[1]


def test_reviewed_version_without_projection_is_not_accepted_as_a_flag(reviewed_case):
    assert collect(reviewed_case, projection=False) == {1: ["layout_mapping_unavailable"]}


def test_checker_inspects_actual_docx_even_when_updated_hash_matches(reviewed_case):
    case = reviewed_case
    docx_bytes, map_bytes = changed_document(case.artifact,
        lambda document: document.sections[0].header.paragraphs[0].add_run(" Unmapped words"))
    case.output.write_bytes(docx_bytes)
    case.output.with_suffix(".source_map.json").write_bytes(map_bytes)
    assert collect(case) == {1: ["layout_mapping_unavailable"]}


@pytest.mark.parametrize("kind", ["text", "source", "target", "line_endings"])
def test_reviewed_output_requires_exact_current_source_target_and_txt(reviewed_case, kind):
    case = reviewed_case
    path = case.pages_dir / "page_0001.txt"
    if kind == "text":
        path.write_bytes(path.read_bytes() + b" changed")
    elif kind == "line_endings":
        raw = path.read_bytes()
        assert b"\n" in raw
        path.write_bytes(raw.replace(b"\n", b"\r\n"))
    else:
        path = path.with_suffix(".source_structure.json" if kind == "source" else ".structure.json")
        data = json.loads(path.read_bytes())
        data["blocks"][0]["text"] += " changed"
        path.write_bytes(encode(data))
    assert collect(case) == {1: ["layout_mapping_unavailable"]}


@pytest.mark.parametrize("change", ["wrong_page", "duplicate", "wrong_docx", "render_claim", "version"])
def test_reviewed_map_schema_and_claims_are_checked_against_package(reviewed_case, change):
    case = reviewed_case
    path = case.output.with_suffix(".source_map.json")
    data = json.loads(path.read_bytes())
    if change == "wrong_page":
        data["pages"][0]["page_number"] = 9
    elif change == "duplicate":
        data["pages"].append(deepcopy(data["pages"][0]))
    elif change == "wrong_docx":
        data["docx_sha256"] = "0" * 64
    elif change == "render_claim":
        data.update(layout_review_required=False, rendered_layout_acceptance="accepted")
    else:
        data["version"] = "reviewed_region_source_map_v999"
    path.write_bytes(encode(data))
    assert collect(case) == {1: ["layout_mapping_unavailable"]}


def test_forged_projection_is_rechecked_by_actual_writer_contract(reviewed_case):
    case = reviewed_case
    case.projection = replace(case.projection, layout_review_required=False)
    assert collect(case) == {1: ["layout_mapping_unavailable"]}


def test_reviewed_selection_cannot_rename_or_omit_source_pages(reviewed_case):
    case = reviewed_case
    assert collect(case, pages=(9,)) == {9: ["layout_mapping_unavailable"]}
    result = collect(case, pages=(1, 9))
    assert set(result) == {1, 9}
    assert all("layout_mapping_unavailable" in codes for codes in result.values())


def test_reviewed_notices_merge_without_replacing_historical_accounting_or_semantics(reviewed_case):
    case = reviewed_case
    preparation = {"legacy_pages": [1], "review_required_pages": [1, 8]}
    before_preparation = deepcopy(preparation)
    notices = collect(case, preparation=preparation)
    assert notices == {1: ["layout_review_required", "source_block_layout_unavailable"]}
    payload = {"model": "historical", "usage_records": [{"input_tokens": 41}],
        "totals": {"cost_usd": "0.0123"}, "quality_risk_score": 0.8,
        "review_queue": [{"page_number": 1, "score": 0.8, "recommended_action": "rerun_page",
            "reasons": ["semantic_issue", "section_furniture_target_variant_standardized"], "retained": True}]}
    before = deepcopy(payload)
    page_records = {"1": {"status": "done", "usage": {"output_tokens": 17},
        "layout_review_required": True, "layout_review_reasons": notices[1]}}
    original_records = deepcopy(page_records)
    merge_layout_review_queue(payload, page_records)
    for key in ("model", "usage_records", "totals", "quality_risk_score"):
        assert payload[key] == before[key]
    row = payload["review_queue"][0]
    assert row["score"] == 0.8 and row["recommended_action"] == "rerun_page" and row["retained"]
    assert row["reasons"] == ["semantic_issue", "section_furniture_target_variant_standardized",
                              "layout_review_required", "source_block_layout_unavailable"]
    assert page_records == original_records and preparation == before_preparation


@pytest.fixture
def ordinary_case(tmp_path):
    pages = tmp_path / "pages"
    pages.mkdir()
    path = pages / "page_0002.txt"
    path.write_bytes("Texte conservé.\n".encode("utf-8"))
    output = tmp_path / "ordinary.docx"
    document = Document()
    document.add_paragraph("Texte conservé.")
    document.save(output)
    mapping = {"version": 1, "docx_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "pages": [{"source_page_number": 2, "translation_sha256": digest_text(path.read_text(encoding="utf-8")),
            "structure_status": "validated", "layout_review_required": False}]}
    output.with_suffix(".source_map.json").write_bytes(encode(mapping))
    return output, pages


def test_integer_v1_without_projection_keeps_existing_behavior_and_supports_done_selection(ordinary_case):
    output, pages = ordinary_case
    assert collect_docx_layout_review(output, pages, {}) == {}
    (pages / "page_0001.txt").write_bytes(b"Unselected stale text")
    assert collect_docx_layout_review(output, pages, {}, page_numbers=[2]) == {}
    all_pages = collect_docx_layout_review(output, pages, {})
    assert all_pages == {1: ["layout_mapping_unavailable"], 2: ["layout_mapping_unavailable"]}


@pytest.mark.parametrize("value", [True, "1", 1.0, {}, None])
def test_integer_v1_does_not_accept_other_json_types(ordinary_case, value):
    output, pages = ordinary_case
    path = output.with_suffix(".source_map.json")
    mapping = json.loads(path.read_bytes())
    mapping["version"] = value
    path.write_bytes(encode(mapping))
    assert collect_docx_layout_review(output, pages, {}) == {2: ["layout_mapping_unavailable"]}


@pytest.mark.parametrize("selection", [[True], [0], [-1], [2, 2], [2, 1], ["2"]])
def test_invalid_selection_never_silently_changes_source_page_numbers(ordinary_case, selection):
    output, pages = ordinary_case
    with pytest.raises(ValueError, match="source-page selection"):
        collect_docx_layout_review(output, pages, {}, page_numbers=selection)
