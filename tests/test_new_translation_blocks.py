from __future__ import annotations

from dataclasses import replace
import json

import fitz
import pytest
from docx import Document

from legalpdf_translate import workflow as workflow_module
from legalpdf_translate.checkpoint import load_run_state
from legalpdf_translate.new_translation_blocks import validate_block
from legalpdf_translate.openai_client import ApiCallResult
from legalpdf_translate.translation_structure import BlockCoverageError
from legalpdf_translate.types import ImageMode, OcrMode, RunConfig, TargetLang
from legalpdf_translate.workflow import TranslationWorkflow

SOURCE = ("O documento apresenta os factos e as circunstancias do processo. "
          "O arguido deve comparecer no tribunal e cumprir todas as obrigacoes indicadas. "
          "As condicoes continuam a aplicar-se durante o periodo determinado na decisao.")
TARGETS = {
    TargetLang.EN: "The document sets out the facts and circumstances of the case. The defendant must attend court and comply with all stated obligations. The conditions continue to apply during the period specified in the decision.",
    TargetLang.FR: "Le document expose les faits et les circonstances de la procedure. Le prevenu doit comparaitre devant le tribunal et respecter toutes les obligations indiquees. Les conditions continuent de s'appliquer pendant la periode fixee dans la decision.",
    TargetLang.AR: "يعرض المستند وقائع القضية وظروفها. يجب على المتهم الحضور أمام المحكمة والوفاء بجميع الالتزامات المذكورة. وتستمر الشروط في السريان طوال المدة المحددة في القرار.",
}


def configuration(tmp_path, *, pages=1, lang=TargetLang.EN, page_breaks=False):
    path = tmp_path / "synthetic.pdf"
    (tmp_path / "output").mkdir(exist_ok=True)
    with fitz.open() as document:
        for _ in range(pages):
            page = document.new_page()
            page.insert_textbox(fitz.Rect(48, 80, 540, 250), SOURCE, fontsize=11)
        document.save(path)
    return RunConfig(pdf_path=path, output_dir=tmp_path / "output", target_lang=lang,
        image_mode=ImageMode.OFF, ocr_mode=OcrMode.OFF, workers=1, resume=False, page_breaks=page_breaks)


class FakeClient:
    def __init__(self, lang=TargetLang.EN, transform=None):
        self.lang, self.transform = lang, transform
        self.calls = []

    def create_page_response(self, **kwargs):
        self.calls.append(kwargs)
        payload, _ = json.JSONDecoder().raw_decode(kwargs["prompt_text"])
        result = {"blocks": [{"id": row["id"], "text": TARGETS[self.lang]} for row in payload["blocks"]]}
        response = ApiCallResult(raw_output=json.dumps(result, ensure_ascii=False),
            usage={"input_tokens": 12, "output_tokens": 14, "reasoning_tokens": 4, "total_tokens": 26},
            response_id=f"synthetic-{len(self.calls)}", response_status="completed")
        return self.transform(response, result) if self.transform else response


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.delenv("LEGALPDF_TRANSLATION_PROTOCOL", raising=False)
    monkeypatch.setattr(workflow_module, "load_gui_settings", lambda: {})
    monkeypatch.setattr(workflow_module, "load_environment", lambda: None)
    monkeypatch.setattr(workflow_module, "run_translation_auth_test",
                        lambda *a, **k: pytest.fail("No authentication calls in offline test"))


@pytest.mark.parametrize("lang", list(TargetLang))
@pytest.mark.parametrize("page_breaks", [False, True])
def test_real_new_source_to_docx(tmp_path, lang, page_breaks):
    config = configuration(tmp_path, lang=lang, page_breaks=page_breaks)
    client = FakeClient(lang)
    workflow = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2")
    result = workflow.run(config)
    assert result.success, result.error
    assert len(client.calls) == 1
    pages = result.run_dir / "pages"
    source = json.loads((pages / "page_0001.source_structure.json").read_text("utf-8"))
    target = json.loads((pages / "page_0001.structure.json").read_text("utf-8"))
    assert source["provenance"] == "digital_pdf"
    assert [b["id"] for b in source["blocks"]] == [b["id"] for b in target["blocks"]]
    assert (pages / "page_0001.commit.json").exists()
    assert "p0001_b" not in " ".join(p.text for p in Document(result.output_docx).paragraphs)
    state = load_run_state(result.run_dir / "run_state.json")
    assert state.pages["1"]["fidelity_review_status"] == "not_evaluated"
    assert state.pages["1"]["output_tokens"] == 14
    assert state.pages["1"]["total_tokens"] == 26
    summary = json.loads(result.run_summary_path.read_text("utf-8"))
    assert summary["review_queue_count"] >= 1
    assert summary["cost_estimation_status"] == "not_evaluated_all_calls"
    assert summary["totals"]["total_cost_estimate_if_available"] is None
    assert summary["budget_post_run"]["total_tokens"] == 26
    assert summary["budget_post_run"]["reasoning_tokens_included_in_output"] is True


@pytest.mark.parametrize("kind", ["incomplete", "refused", "missing", "duplicate", "foreign"])
def test_rejected_output_never_committed(tmp_path, kind):
    config = configuration(tmp_path)
    def change(response, payload):
        if kind == "incomplete":
            response.response_status = "incomplete"
        elif kind == "refused":
            response.refused = True
        else:
            if kind == "missing": payload["blocks"] = []
            if kind == "duplicate": payload["blocks"] *= 2
            if kind == "foreign": payload["blocks"][0]["id"] = "p9999_b0001"
            response.raw_output = json.dumps(payload)
        return response
    client = FakeClient(transform=change)
    result = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2").run(config)
    assert not result.success
    assert len(client.calls) == (1 if kind in {"incomplete", "refused"} else 2)
    assert not list((result.run_dir / "pages").glob("page_*.txt"))
    state = load_run_state(result.run_dir / "run_state.json")
    assert state.done_count == 0
    assert state.pages["1"]["input_tokens"] == 12 * len(client.calls)


def test_resume_is_provider_free_and_fingerprint_changes_fail_before_writes(tmp_path):
    config = configuration(tmp_path)
    client = FakeClient()
    workflow = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2")
    result = workflow.run(config)
    assert result.success
    resumed = workflow.run(replace(config, resume=True, page_breaks=True))
    assert resumed.success
    assert len(client.calls) == 1
    before = {p.name: p.read_bytes() for p in result.run_dir.iterdir() if p.is_file()}
    with pytest.raises(ValueError, match="incompatible"):
        workflow.run(replace(config, resume=True, context_text="Changed translation instructions"))
    assert before == {p.name: p.read_bytes() for p in result.run_dir.iterdir() if p.is_file()}
    assert len(client.calls) == 1


def test_cancel_keeps_verified_partial_only(tmp_path):
    config = configuration(tmp_path, pages=2)
    client = FakeClient()
    workflow = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2")
    def progress(done, total, status):
        if done == 1 and "finished" in status:
            workflow.cancel()
    workflow._progress_callback = progress
    result = workflow.run(config)
    assert not result.success and result.error == "cancelled"
    assert len(client.calls) == 1
    assert result.completed_pages == 1 and result.partial_docx is not None
    assert TARGETS[TargetLang.EN] in " ".join(p.text for p in Document(result.partial_docx).paragraphs)


def test_complete_commit_before_checkpoint_recovers_without_dispatch(tmp_path, monkeypatch):
    config = configuration(tmp_path)
    client = FakeClient()
    original = workflow_module.mark_page_done
    def crash(*args, **kwargs):
        raise RuntimeError("synthetic crash before checkpoint DONE")
    monkeypatch.setattr(workflow_module, "mark_page_done", crash)
    with pytest.raises(RuntimeError, match="synthetic crash"):
        TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2").run(config)
    monkeypatch.setattr(workflow_module, "mark_page_done", original)
    result = TranslationWorkflow(client=client).run(replace(config, resume=True))
    assert result.success and len(client.calls) == 1
    state = load_run_state(result.run_dir / "run_state.json")
    assert state.pages["1"]["input_tokens"] == 12


def test_manual_txt_edit_blocks_resume_but_rebuild_preserves_work(tmp_path):
    config = configuration(tmp_path)
    client = FakeClient()
    workflow = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2")
    result = workflow.run(config)
    assert result.success
    original_docx = result.output_docx.read_bytes()
    page = result.run_dir / "pages" / "page_0001.txt"
    page.write_text(TARGETS[TargetLang.EN] + " Human amendment.", encoding="utf-8")
    with pytest.raises(ValueError):
        workflow.run(replace(config, resume=True))
    output = workflow.rebuild_docx(config)
    assert output != result.output_docx
    assert "Human amendment." in " ".join(p.text for p in Document(output).paragraphs)
    assert result.output_docx.read_bytes() == original_docx
    assert len(client.calls) == 1


def test_event_missing_is_not_assumed_checked_and_swapped_dates_reject():
    source = {"id": "p0001_b0001", "text": "A audiência será em 12.09.2026 e o interrogatório em 14.09.2026."}
    review = set()
    validate_block(source, "The appointments are on 12.09.2026 and 14.09.2026.", TargetLang.EN, None, review)
    assert "event_association_not_evaluated" in review
    with pytest.raises(BlockCoverageError, match="block_event_date_defect"):
        validate_block(source, "The hearing is on 14.09.2026 and questioning on 12.09.2026.", TargetLang.EN, None, set())


def test_same_id_clause_omission_does_not_claim_legal_acceptance(tmp_path):
    config = configuration(tmp_path)
    def omit(response, payload):
        payload["blocks"][0]["text"] = "The document sets out the facts and circumstances of the case."
        response.raw_output = json.dumps(payload)
        return response
    result = TranslationWorkflow(client=FakeClient(transform=omit), translation_protocol="legal_blocks_v2").run(config)
    assert result.success
    state = load_run_state(result.run_dir / "run_state.json")
    assert state.pages["1"]["fidelity_review_required"] is True
    assert state.pages["1"]["fidelity_review_status"] == "not_evaluated"


@pytest.mark.parametrize("checkpoint", ["missing", "corrupt", "legacy"])
@pytest.mark.parametrize("operation", ["resume", "rebuild"])
def test_orphan_or_mixed_bundle_cannot_fall_through_legacy(tmp_path, checkpoint, operation):
    config = configuration(tmp_path)
    client = FakeClient()
    result = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2").run(config)
    assert result.success
    path = result.run_dir / "run_state.json"
    if checkpoint == "missing":
        path.unlink()
    elif checkpoint == "corrupt":
        path.write_text("{broken", encoding="utf-8")
    else:
        data = json.loads(path.read_text("utf-8"))
        data.pop("protocol_identity", None)
        for page in data["pages"].values():
            page.pop("structured_commit", None)
            page.pop("structured_evidence_state", None)
        path.write_text(json.dumps(data), encoding="utf-8")
    (result.run_dir / "pages" / "page_0002.txt").write_text("Uncommitted pending material.", encoding="utf-8")
    before = {str(p.relative_to(config.output_dir)): p.read_bytes()
              for p in config.output_dir.rglob("*") if p.is_file()}
    workflow = TranslationWorkflow(client=client)
    with pytest.raises(ValueError, match="checkpoint|legacy"):
        if operation == "resume":
            workflow.run(replace(config, resume=True))
        else:
            workflow.rebuild_docx(config)
    assert before == {str(p.relative_to(config.output_dir)): p.read_bytes()
                      for p in config.output_dir.rglob("*") if p.is_file()}
    assert len(client.calls) == 1


def test_retention_opt_out_marks_evidence_unavailable_without_losing_docx(tmp_path):
    config = replace(configuration(tmp_path), keep_intermediates=False)
    client = FakeClient()
    workflow = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2")
    result = workflow.run(config)
    assert result.success and result.output_docx.exists()
    assert not (result.run_dir / "pages").exists()
    state = load_run_state(result.run_dir / "run_state.json")
    assert state.pages["1"]["structured_evidence_state"] == "purged"
    original = result.output_docx.read_bytes()
    with pytest.raises(ValueError):
        workflow.run(replace(config, resume=True))
    with pytest.raises(ValueError):
        workflow.rebuild_docx(config)
    assert result.output_docx.read_bytes() == original
    assert len(client.calls) == 1


def test_existing_zero_budget_block_precedes_provider(tmp_path):
    from legalpdf_translate.types import BudgetExceedPolicy
    config = replace(configuration(tmp_path), budget_cap_usd=0.0, budget_on_exceed=BudgetExceedPolicy.BLOCK)
    client = FakeClient()
    result = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2").run(config)
    assert not result.success and result.error == "budget_cap_exceeded"
    assert client.calls == []


FRENCH_LEGAL_SOURCE = (
    "Ao abrigo do artigo 12.º do Código de Processo Penal português, José Nobre deve cumprir "
    "a obrigação indicada no processo 12/26.4TEST. A referência electrónica permite consultar o documento."
)
FRENCH_LEGAL_TARGET = (
    "En vertu de l'article 12.º du Code de procédure pénale portugais, José Nobre doit respecter "
    "l'obligation indiquée dans l'affaire 12/26.4TEST. La référence électronique permet de consulter le document."
)


def french_legal_configuration(tmp_path):
    config = configuration(tmp_path, lang=TargetLang.FR)
    with fitz.open() as document:
        page = document.new_page()
        assert page.insert_textbox(fitz.Rect(48, 80, 540, 250), FRENCH_LEGAL_SOURCE, fontsize=11) >= 0
        document.save(config.pdf_path)
    return config


@pytest.mark.parametrize("corrected", [True, False])
def test_french_legal_title_translation_keeps_strict_existing_correction_contract(tmp_path, corrected):
    from legalpdf_translate.translation_structure import structured_system_instructions

    config = french_legal_configuration(tmp_path)
    calls = []
    leaked = FRENCH_LEGAL_TARGET.replace("Code de procédure pénale", "Código de Processo Penal")

    class FrenchClient:
        def create_page_response(self, **kwargs):
            calls.append(kwargs)
            payload, _ = json.JSONDecoder().raw_decode(kwargs["prompt_text"])
            text = FRENCH_LEGAL_TARGET if corrected and len(calls) == 2 else leaked
            return ApiCallResult(
                raw_output=json.dumps({"blocks": [{"id": row["id"], "text": text}
                                                  for row in payload["blocks"]]}, ensure_ascii=False),
                usage={}, response_id=f"fictional-fr-{len(calls)}", response_status="completed")

    result = TranslationWorkflow(client=FrenchClient(), translation_protocol="legal_blocks_v2").run(config)
    assert len(calls) == 2  # Only the existing primary and one compliance correction.
    assert all(call["instructions"] == structured_system_instructions(TargetLang.FR) for call in calls)
    instructions = calls[0]["instructions"]
    assert "translate legal terms, law and code titles" in instructions
    assert "Portuguese jurisdiction" in instructions
    assert "Do not substitute French law or institutions" in instructions
    assert "Preserve accented proper names and identifiers verbatim" in instructions
    assert "block_language_or_token_defect" in calls[1]["prompt_text"]
    assert calls[1]["effort"] == calls[0]["effort"]
    assert calls[1]["response_format"] == calls[0]["response_format"]
    assert calls[1]["max_output_tokens"] == calls[0]["max_output_tokens"]
    state = load_run_state(result.run_dir / "run_state.json")
    if corrected:
        assert result.success, result.error
        target = (result.run_dir / "pages" / "page_0001.txt").read_text("utf-8")
        assert target == FRENCH_LEGAL_TARGET
        assert state.pages["1"]["fidelity_review_status"] == "not_evaluated"
    else:
        assert not result.success
        assert state.pages["1"]["validator_defect_reason"] == "block_language_or_token_defect"
        assert state.done_count == 0
        assert not list((result.run_dir / "pages").glob("page_*.txt"))
        assert not list((result.run_dir / "pages").glob("page_*.commit.json"))


def test_french_old_instruction_resume_rejects_before_dispatch_or_artifact_write(tmp_path, monkeypatch):
    from legalpdf_translate import new_translation_blocks as block_module
    from legalpdf_translate.translation_structure import structured_system_instructions

    config = configuration(tmp_path, lang=TargetLang.FR)
    current = structured_system_instructions(TargetLang.FR)
    old = current[:current.index("Use formal legal French")] + (
        "Use formal legal French; preserve Portuguese legal concepts and accented proper names verbatim.")
    current_by_lang = {lang: structured_system_instructions(lang) for lang in TargetLang}
    monkeypatch.setattr(block_module, "structured_system_instructions",
                        lambda lang: old if lang == TargetLang.FR else current_by_lang[lang])
    client = FakeClient(TargetLang.FR)
    workflow = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2")
    result = workflow.run(config)
    assert result.success, result.error
    before = {str(p.relative_to(config.output_dir)): p.read_bytes()
              for p in config.output_dir.rglob("*") if p.is_file()}
    monkeypatch.setattr(block_module, "structured_system_instructions", structured_system_instructions)
    with pytest.raises(ValueError, match="incompatible"):
        workflow.run(replace(config, resume=True))
    assert len(client.calls) == 1
    assert before == {str(p.relative_to(config.output_dir)): p.read_bytes()
                      for p in config.output_dir.rglob("*") if p.is_file()}
