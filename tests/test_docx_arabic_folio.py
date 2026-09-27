"""Complete Arabic folios preserve exact text in one RTL sequence."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
import pytest

from legalpdf_translate import docx_writer as writer
from legalpdf_translate.types import TargetLang


def _visible(text):
    return writer.sanitize_bidi_controls(writer.unwrap_internal_placeholders(text))


@pytest.mark.parametrize("label", ["الصفحة", "صفحة", "ص.", "ص"])
@pytest.mark.parametrize("numbers", [("2", "9"), ("[[2]]", "[[9]]"), ("[[2]]", "9")])
@pytest.mark.parametrize("strip", [True, False])
def test_complete_plain_or_protected_folio_stays_one_rtl_sequence(label, numbers, strip):
    line = f" {label}\u00a0{numbers[0]} من  {numbers[1]} "
    runs, mixed = writer._segment_rtl_placeholder_aware_runs(line, strip_bidi_controls=strip)
    assert runs == [("rtl", _visible(line))]
    assert mixed is False


def test_default_stripping_preserves_isolated_folio_text_and_explicit_line_endings():
    text = "الصفحة \u2066[[12]]\u2069 من \u2066[[40]]\u2069\r\nص. [[13]] من [[40]]\n"
    runs, mixed = writer._segment_rtl_placeholder_aware_runs(text, strip_bidi_controls=True)
    assert runs == [("rtl", "الصفحة 12 من 40\r\n"), ("rtl", "ص. 13 من 40\n")]
    assert "".join(chunk for _, chunk in runs) == _visible(text)
    assert mixed is False


@pytest.mark.parametrize("line", [
    "راجع الصفحة [[2]] من [[9]]", "الصفحة [[2]] من [[9]] مرفقة",
    "الصفحة [[2]] من [[9]]/[[12]]", "الصفحة [[2]] من [[9]]A",
    "الصفحة [[2]] من [[09/2026]]", "الصفحة [[2]] من [[123 456 789]]",
    "الصفحة [[2]] من [[CASE-9]]", "الصفحة [[2]] من [[1.2]]",
    "الصفحة [[2]] من [[9]]،", "Page [[2]] من [[9]]", "[[الصفحة 2 من 9]]",
    "الصفحة [[٢]] من [[٩]]", "الصفحة [[2]] من [[1234567]]",
    "الصفحة\t[[2]] من [[9]]", "الصفحة [[2]]\tمن [[9]]",
    "الصفحة [[2]]\nمن [[9]]", "الصفحة [[2]]\vمن [[9]]",
    "الصفحة [[2]] من [[9]", "الصفحة [[[2]]] من [[9]]",
    "الصفحة \u2066[[2]] من [[9]]", "الصفحة [[2]]\u2069 من [[9]]",
    "\u202eالصفحة [[2]] من [[9]]\u202c", "الصفحة \u2067[[2]]\u2069 من [[9]]",
    "الصفحة [[2]] من \u200e[[9]]\u200e", "الصفحة [[2]] من [[9]]\u0301",
])
@pytest.mark.parametrize("strip", [True, False])
def test_non_folio_and_untrusted_control_scopes_keep_existing_segmentation(line, strip):
    # Compare with the existing mixed-line algorithm rather than assuming how a
    # particular number, identifier, punctuation mark or control should display.
    expected = []
    for part in line.splitlines(keepends=True):
        segments, _ = writer._segment_rtl_placeholder_aware_line(part, strip_bidi_controls=strip)
        expected.extend(segments)
    expected = writer._keep_phone_triplet_spaces_ltr(writer._keep_numeric_slash_separators_ltr(expected))
    actual, _ = writer._segment_rtl_placeholder_aware_runs(line, strip_bidi_controls=strip)
    assert [(kind, char) for kind, text in actual for char in text] == [
        (kind, char) for kind, text in expected for char in text
    ]


def test_retained_numeric_isolates_keep_their_original_scope():
    line = "الصفحة \u2066[[2]]\u2069 من \u2066[[9]]\u2069"
    baseline, _ = writer._segment_rtl_placeholder_aware_line(line, strip_bidi_controls=False)
    runs, _ = writer._segment_rtl_placeholder_aware_runs(line, strip_bidi_controls=False)
    assert runs == baseline
    assert "".join(text for _, text in runs) == writer.unwrap_internal_placeholders(line)


def test_folio_among_mixed_lines_keeps_other_runs_and_breaks():
    address = "العنوان: [[Rua Exemplo, 8, 1234-567 Vila Exemplo]]"
    text = address + "\nالصفحة [[2]] من [[9]]\nالهاتف [[321]] [[654]] [[987]]"
    runs, mixed = writer._segment_rtl_placeholder_aware_runs(text, strip_bidi_controls=True)
    assert mixed
    assert "".join(chunk for _, chunk in runs) == _visible(text)
    assert ("rtl", "الصفحة 2 من 9\n") in runs
    assert ("ltr", "Rua Exemplo, 8, 1234-567 Vila Exemplo") in runs
    assert ("ltr", "321 654 987") in runs


def test_docx_folio_is_one_arabic_run_without_added_marks_and_page_break_survives(tmp_path: Path):
    pages = tmp_path / "pages"
    pages.mkdir()
    first = "الصفحة \u2066[[2]]\u2069 من \u2066[[9]]\u2069"
    second = "العنوان: [[Rua Exemplo, 8, 1234-567 Vila Exemplo]]"
    (pages / "page_0001.txt").write_text(first, encoding="utf-8")
    (pages / "page_0002.txt").write_text(second, encoding="utf-8")
    path = writer.assemble_docx(pages, tmp_path / "fictional.docx", lang=TargetLang.AR, page_breaks=True)
    doc = Document(path)
    text_runs = [run for run in doc.paragraphs[0].runs if run.text.strip()]
    assert len(text_runs) == 1
    assert text_runs[0].text == "الصفحة 2 من 9"
    assert text_runs[0]._r.find(qn("w:rPr")).find(qn("w:rtl")).get(qn("w:val")) == "1"
    assert text_runs[0]._r.find(qn("w:rPr")).find(qn("w:lang")).get(qn("w:val")) == "ar-SA"
    assert len(doc.element.xpath('.//w:br[@w:type="page"]')) == 1
    assert _visible(doc.paragraphs[1].text) == _visible(second)


@pytest.mark.parametrize("lang", [TargetLang.EN, TargetLang.FR])
def test_non_arabic_output_path_does_not_add_rtl_flags(tmp_path: Path, lang):
    pages = tmp_path / "pages"
    pages.mkdir()
    (pages / "page_0001.txt").write_text("Page 2 of 9", encoding="utf-8")
    output = writer.assemble_docx(pages, tmp_path / "fictional.docx", lang=lang, page_breaks=False)
    doc = Document(output)
    assert doc.paragraphs[0].text == "Page 2 of 9"
    assert not doc.paragraphs[0]._p.xpath('.//w:rtl | .//w:bidi')
