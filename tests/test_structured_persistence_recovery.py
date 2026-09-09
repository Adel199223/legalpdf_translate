"""Offline new-run persistence, crash recovery and manual-edit boundaries."""
from copy import deepcopy
from dataclasses import replace
import json

from docx import Document
import pytest

from legalpdf_translate import checkpoint
from legalpdf_translate.docx_writer import assemble_docx
from legalpdf_translate.document_structure import PageStructure, StructureBlock, text_sha256
from legalpdf_translate import structured_artifacts as artifacts
from legalpdf_translate.types import RunConfig, TargetLang


IDENTITY = {"protocol": "legal_blocks_v2", "fingerprint": "a" * 64}
RECEIPT = {"image_used": False, "retry_used": True,
           "usage": {"attempt_1": {"input_tokens": 9, "output_tokens": 5},
                     "attempt_2": {"input_tokens": 12, "output_tokens": 6}},
           "metadata": {"api_calls_count": 2, "review_required": True,
                        "fidelity_status": "not_evaluated"}}


def case(tmp_path, pages=1):
    config = RunConfig(tmp_path / "source.pdf", tmp_path / "out", TargetLang.EN)
    paths = checkpoint.build_run_paths(config.output_dir, config.pdf_path, config.target_lang,
                                       run_started_at="20260909_180000")
    state = checkpoint.new_run_state(config=config, paths=paths, pdf_fingerprint="b" * 64,
        context_hash="context", total_pages=pages, selected_pages=list(range(1, pages + 1)),
        protocol_identity=IDENTITY)
    args = dict(config=config, paths=paths, pdf_fingerprint="b" * 64, context_hash="context",
                selection_start_page=1, selection_end_page=pages, selection_page_count=pages,
                max_pages_effective=pages, protocol_identity=IDENTITY)
    return state, paths, args


def publish(paths, number=1, **changes):
    source = PageStructure(number, text_sha256("Aviso"),
        [StructureBlock(f"p{number:04d}_b0001", "Aviso")], source_file_sha256="b" * 64,
        source_text_sha256=text_sha256("Aviso"))
    target = source.to_dict()
    target["blocks"][0]["text"] = "Notice"
    target["translation_sha256"] = text_sha256("Notice")
    options = dict(source_structure=source, translated_structure=target, translated_text="Notice",
                   protocol_identity=IDENTITY, page_fingerprint="c" * 64, page_result=deepcopy(RECEIPT))
    options.update(changes)
    return artifacts.publish_structured_page(paths.pages_dir, **options)


def snapshot(paths):
    return {str(path.relative_to(paths.run_dir)): path.read_bytes()
            for path in paths.run_dir.rglob("*") if path.is_file()}


@pytest.mark.parametrize("status", ["pending", "failed"])
def test_commit_before_done_recovers_usage_and_review_once_without_disk_mutation(tmp_path, status):
    state, paths, args = case(tmp_path)
    state.pages["1"].update(status=status, error="prior_failure", existing_findings=["keep"])
    record = publish(paths)
    before = snapshot(paths)
    before_state = deepcopy(state.to_dict())
    assert checkpoint.resume_incompatibility_reason(state, **args) is None
    assert state.to_dict() == before_state
    assert checkpoint.recover_structured_commits(state, paths, protocol_identity=IDENTITY) == [1]
    assert state.pages["1"]["usage"] == RECEIPT["usage"]
    assert state.pages["1"]["api_calls_count"] == 2
    assert state.pages["1"]["review_required"] is True
    assert state.pages["1"]["existing_findings"] == ["keep"]
    assert state.pages["1"]["error"] is None
    assert state.pages["1"]["structured_commit"] == record
    assert state.done_count == 1 and state.pending_count == state.failed_count == 0
    assert checkpoint.recover_structured_commits(state, paths, protocol_identity=IDENTITY) == []
    assert before == snapshot(paths)
    checkpoint.save_run_state_atomic(paths.run_state_path, state)
    restored = checkpoint.load_run_state(paths.run_state_path)
    assert checkpoint.resume_incompatibility_reason(restored, **args) is None


@pytest.mark.parametrize("damage", ["incomplete", "manual_text", "target", "receipt", "wrong_protocol"])
def test_recovery_checks_all_pages_before_mutating_first(tmp_path, damage):
    state, paths, args = case(tmp_path, pages=2)
    publish(paths, 1)
    publish(paths, 2)
    if damage == "incomplete":
        (paths.pages_dir / "page_0002.commit.json").unlink()
    elif damage == "manual_text":
        (paths.pages_dir / "page_0002.txt").write_text("My edited translation", encoding="utf-8")
    elif damage == "target":
        (paths.pages_dir / "page_0002.structure.json").write_text("{}", encoding="utf-8")
    elif damage == "receipt":
        path = paths.pages_dir / "page_0002.commit.json"
        data = json.loads(path.read_text(encoding="utf-8"))
        data["page_result"]["usage"] = {}
        path.write_text(json.dumps(data), encoding="utf-8")
    else:
        args["protocol_identity"] = {**IDENTITY, "fingerprint": "d" * 64}
    before_state, before_files = deepcopy(state.to_dict()), snapshot(paths)
    assert checkpoint.resume_incompatibility_reason(state, **args) is not None
    with pytest.raises(artifacts.StructuredArtifactError):
        checkpoint.recover_structured_commits(state, paths, protocol_identity=args["protocol_identity"])
    assert state.to_dict() == before_state and snapshot(paths) == before_files


def test_complete_unselected_artifact_blocks_recovery_and_is_not_globbed_into_state(tmp_path):
    state, paths, args = case(tmp_path)
    publish(paths, 1)
    publish(paths, 2)
    before = deepcopy(state.to_dict())
    assert "unselected" in checkpoint.resume_incompatibility_reason(state, **args)
    with pytest.raises(artifacts.StructuredArtifactError, match="unselected"):
        checkpoint.recover_structured_commits(state, paths, protocol_identity=IDENTITY)
    assert state.to_dict() == before


def test_manual_txt_can_rebuild_as_review_fallback_but_cannot_resume_old_binding(tmp_path):
    state, paths, args = case(tmp_path)
    record = publish(paths)
    checkpoint.mark_page_done(state, 1, structured_commit=record, **deepcopy(RECEIPT))
    original_docx = assemble_docx(paths.pages_dir, paths.frozen_outdir / "reviewed.docx",
                                  lang=TargetLang.EN, page_breaks=False)
    document = Document(original_docx)
    document.add_paragraph("Human Word-only correction")
    document.save(original_docx)
    reviewed_bytes = original_docx.read_bytes()
    (paths.pages_dir / "page_0001.txt").write_text("Human TXT correction", encoding="utf-8")
    before = snapshot(paths)
    assert checkpoint.resume_incompatibility_reason(state, **args) is not None
    rebuilt = assemble_docx(paths.pages_dir, paths.frozen_outdir / "rebuilt-separately.docx",
                            lang=TargetLang.EN, page_breaks=False)
    assert "Human TXT correction" in "\n".join(p.text for p in Document(rebuilt).paragraphs)
    assert original_docx.read_bytes() == reviewed_bytes
    assert snapshot(paths) == before
    mapping = json.loads(rebuilt.with_suffix(".source_map.json").read_text(encoding="utf-8"))
    assert mapping["pages"][0]["structure_status"] == "invalid_sidecar_txt_fallback"


@pytest.mark.parametrize("field,value", [("usage", {}), ("image_used", True), ("retry_used", False)])
def test_done_cannot_disagree_with_durable_paid_outcome(tmp_path, field, value):
    state, paths, _ = case(tmp_path)
    record = publish(paths)
    before = deepcopy(state.to_dict())
    result = deepcopy(RECEIPT)
    result[field] = value
    with pytest.raises(artifacts.StructuredArtifactError, match="outcome_mismatch"):
        checkpoint.mark_page_done(state, 1, structured_commit=record, **result)
    assert state.to_dict() == before


@pytest.mark.parametrize("bad", [[], {}, {**RECEIPT, "image_used": 1},
    {**RECEIPT, "usage": {"input_tokens": float("nan")}},
    {**RECEIPT, "metadata": {"bad": "x" * 40000}}, {**RECEIPT, "metadata": {1: "bad"}}])
def test_invalid_outcome_is_rejected_before_any_artifact_write(tmp_path, bad):
    _, paths, _ = case(tmp_path)
    with pytest.raises(artifacts.StructuredArtifactError):
        publish(paths, page_result=bad)
    assert not paths.pages_dir.exists()


@pytest.mark.parametrize("bad", [None, [], False, "legal_blocks_v2", {"protocol": "legal_blocks_v2"}])
def test_invalid_protocol_in_saved_state_never_loads_as_legacy(tmp_path, bad):
    state, paths, _ = case(tmp_path)
    checkpoint.save_run_state_atomic(paths.run_state_path, state)
    raw = json.loads(paths.run_state_path.read_text(encoding="utf-8"))
    raw["protocol_identity"] = bad
    paths.run_state_path.write_text(json.dumps(raw), encoding="utf-8")
    assert checkpoint.load_run_state(paths.run_state_path) is None


def test_duplicate_saved_protocol_identity_is_not_last_value_wins(tmp_path):
    state, paths, _ = case(tmp_path)
    checkpoint.save_run_state_atomic(paths.run_state_path, state)
    raw = paths.run_state_path.read_text(encoding="utf-8")
    raw = raw.replace('"protocol_identity": {', '"protocol_identity": {}, "protocol_identity": {')
    paths.run_state_path.write_text(raw, encoding="utf-8")
    assert checkpoint.load_run_state(paths.run_state_path) is None


def test_v1_commit_is_readable_but_missing_receipt_needs_explicit_recovery(tmp_path):
    state, paths, args = case(tmp_path)
    record = publish(paths)
    record["version"] = 1
    record.pop("page_result")
    record["bundle_sha256"] = artifacts._hash(artifacts._json(
        {key: value for key, value in record.items() if key != "bundle_sha256"}))
    (paths.pages_dir / "page_0001.commit.json").write_bytes(artifacts._json(record))
    assert artifacts.validate_structured_page(paths.pages_dir, 1) == record
    assert "orphan_commit_result_unavailable" in checkpoint.resume_incompatibility_reason(state, **args)


@pytest.mark.parametrize("changes", [{"page_breaks": False}, {"strip_bidi_controls": False},
    {"workers": 1}, {"page_breaks": False, "strip_bidi_controls": False, "workers": 6}])
def test_structured_resume_local_format_and_worker_changes_preserve_translation_evidence(tmp_path, changes):
    state, paths, args = case(tmp_path)
    record = publish(paths)
    checkpoint.mark_page_done(state, 1, structured_commit=record, **deepcopy(RECEIPT))
    before_state, before_files = deepcopy(state.to_dict()), snapshot(paths)
    args["config"] = replace(args["config"], **changes)
    assert checkpoint.resume_incompatibility_reason(state, **args) is None
    assert checkpoint.recover_structured_commits(state, paths, protocol_identity=IDENTITY) == []
    assert state.to_dict() == before_state and snapshot(paths) == before_files


def test_structured_resume_retention_change_remains_explicit(tmp_path):
    state, paths, args = case(tmp_path)
    publish(paths)
    args["config"] = replace(args["config"], keep_intermediates=False)
    assert "settings mismatch" in checkpoint.resume_incompatibility_reason(state, **args)


@pytest.mark.parametrize("changes", [{"page_breaks": False}, {"strip_bidi_controls": False}, {"workers": 1}])
def test_legacy_resume_settings_contract_is_not_relaxed(tmp_path, changes):
    state, _, args = case(tmp_path)
    state.protocol_identity = {}
    args["protocol_identity"] = None
    args["config"] = replace(args["config"], **changes)
    assert "settings mismatch" in checkpoint.resume_incompatibility_reason(state, **args)
