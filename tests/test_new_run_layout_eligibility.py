"""Real bundle/OCR-binding/workflow assembly; recognition and translation are fake."""
import hashlib
import io
import json
import subprocess

from docx import Document
from PIL import Image, ImageDraw, ImageFont
import pytest

from legalpdf_translate import ocr_engine as oe, workflow as wf
from legalpdf_translate.browser_pdf_bundle import write_browser_pdf_bundle
from legalpdf_translate.document_structure import PageStructure
from legalpdf_translate.docx_writer import assemble_docx
from legalpdf_translate.layout_integration import prepare_layout_rebuild, load_layout_eligibility, derive_source_layout
from legalpdf_translate.openai_client import ApiCallResult
from legalpdf_translate.types import ImageMode, OcrMode, RunConfig, TargetLang
from legalpdf_translate.workflow import TranslationWorkflow
from .test_new_run_continuations import offline, page_hashes, mapping


HEADER = "Tribunal Judicial da Comarca"
FOOTERS = ("Largo da Justica, 1", "Telef: 210000000 - E-mail: court@example.invalid")
FRAGMENTS = (
    "O arguido deve comparecer perante o tribunal e cumprir todas as obrigacoes, mantendo a condicao de",
    "se apresentar nos termos da notificacao e informar o tribunal sobre qualquer alteracao relevante.",
)
TARGET = (
    "The defendant must attend court and comply with every obligation, maintaining the requirement to",
    "appear as stated in the notice and inform the court about any relevant change.",
)


def raster_page(index, *, table=False):
    image = Image.new("RGB", (600, 850), "white")
    draw, font = ImageDraw.Draw(image), ImageFont.load_default()
    rows = ["level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext",
            "1\t1\t0\t0\t0\t0\t0\t0\t600\t850\t-1\t"]
    lines = [(HEADER, 35), (FRAGMENTS[index], 660 if index == 0 else 100),
             (FOOTERS[0], 800), (FOOTERS[1], 817)]
    for block, (text, y) in enumerate(lines, 1):
        x, line_number, word_number = 45, 1, 0
        for word in text.split():
            box = draw.textbbox((x, y), word, font=font)
            if box[2] > 550:
                x, y, line_number, word_number = 45, y + 16, line_number + 1, 0
                box = draw.textbbox((x, y), word, font=font)
            word_number += 1
            draw.text((x, y), word, font=font, fill="black")
            left, top, right, bottom = box
            rows.append(f"5\t1\t{block}\t1\t{line_number}\t{word_number}\t{left}\t{top}\t{right-left}\t{bottom-top}\t99\t{word}")
            x = right + 5
    if table:
        draw.rectangle((30, 400, 560, 480), outline="black", width=2)
        draw.line((300, 400, 300, 480), fill="black", width=2)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue(), "\n".join(rows), "\n".join(text for text, _ in lines)


def bundle_run(tmp_path, monkeypatch, *, table=False):
    monkeypatch.setattr(subprocess, "Popen", lambda *a, **k: pytest.fail("Native process forbidden"))
    monkeypatch.setattr(oe, "which", lambda _: "fake-tesseract")
    monkeypatch.setattr(oe, "_text_quality_score", lambda _: .99)
    source = tmp_path / "source.pdf"
    source.write_bytes(b"%PDF-1.4\n% synthetic bundle; native extraction forbidden")
    pages = [raster_page(i, table=table) for i in range(2)]
    write_browser_pdf_bundle(source_path=source, page_count=2, pages=[
        {"page_number": i + 1, "mime_type": "image/png", "width_px": 600,
         "height_px": 850, "image_bytes": p[0]}
        for i, p in enumerate(pages)])
    engine, calls = oe.LocalTesseractEngine(), []
    def recognize(*, input_path, pass_spec, preserve_structure=False):
        assert preserve_structure
        pixels = input_path.read_bytes()
        entry = next(p for p in pages if p[0] == pixels)
        calls.append(hashlib.sha256(pixels).hexdigest())
        return 0, entry[2], "", entry[1]
    monkeypatch.setattr(engine, "_run_pass", recognize)
    monkeypatch.setattr(TranslationWorkflow, "_resolve_ocr_engine_for_reason", lambda *a, **k: (engine, True))
    class Client:
        def __init__(self):
            self.calls = []
        def create_page_response(self, **kwargs):
            self.calls.append(kwargs)
            payload, _ = json.JSONDecoder().raw_decode(kwargs["prompt_text"])
            rows = []
            for block in payload["blocks"]:
                text = block["text"]
                if text == HEADER:
                    text = "Judicial Court"
                elif text in FRAGMENTS:
                    text = TARGET[FRAGMENTS.index(text)]
                rows.append({"id": block["id"], "text": text})
            return ApiCallResult(raw_output=json.dumps({"blocks": rows}),
                usage={"input_tokens": 10, "output_tokens": 10}, response_id="offline-bundle",
                response_status="completed")
    client = Client()
    output = tmp_path / "output"
    output.mkdir()
    config = RunConfig(source, output, TargetLang.EN, image_mode=ImageMode.OFF,
                       ocr_mode=OcrMode.ALWAYS, workers=1, resume=False, page_breaks=False)
    workflow = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2")
    result = workflow.run(config)
    assert result.success, result.error
    return workflow, config, result, calls, client


def test_bundle_same_pass_evidence_admits_real_furniture_and_split_sentence(tmp_path, monkeypatch, offline):
    workflow, config, result, calls, client = bundle_run(tmp_path, monkeypatch)
    pages = result.run_dir / "pages"
    assert len(calls) == len(client.calls) == 2
    assert len(list(pages.glob("*.layout_eligibility.json"))) == 2
    for path in pages.glob("*.source_structure.json"):
        assert json.loads(path.read_text("utf-8"))["uncertain"] is True
    document = Document(result.output_docx)
    assert [p.text for p in document.paragraphs if p.text.strip()] == [" ".join(TARGET)]
    assert "Judicial Court" in " ".join(p.text for p in document.sections[0].header.paragraphs)
    assert all(text in " ".join(p.text for p in document.sections[0].footer.paragraphs) for text in FOOTERS)
    assert all(page["section_furniture_adopted"] for page in mapping(result.output_docx))
    before = page_hashes(pages)
    rebuilt = workflow.rebuild_docx(config)
    assert [p.text for p in Document(rebuilt).paragraphs if p.text.strip()] == [" ".join(TARGET)]
    assert page_hashes(pages) == before and len(calls) == len(client.calls) == 2


def test_bundle_unresolved_table_retains_all_body_text_for_review(tmp_path, monkeypatch, offline):
    _, _, result, calls, client = bundle_run(tmp_path, monkeypatch, table=True)
    assert len(calls) == len(client.calls) == 2
    pages = mapping(result.output_docx)
    assert all(page["layout_review_required"] for page in pages)
    assert not any(page.get("section_furniture_adopted") for page in pages)
    assert not any("joined_to_block_id" in block for page in pages for block in page["blocks"])
    body = " ".join(p.text for p in Document(result.output_docx).paragraphs)
    assert body.count("Judicial Court") == 2
    assert all(text in body for text in TARGET + FOOTERS)


@pytest.mark.parametrize("corruption", ["translation", "version", "source_text", "word_confidence", "page_number", "missing"])
def test_stale_eligibility_cannot_admit_furniture_or_join(tmp_path, monkeypatch, offline, corruption):
    _, _, result, _, _ = bundle_run(tmp_path, monkeypatch)
    pages = result.run_dir / "pages"
    path = pages / "page_0001.txt"
    source = json.loads(path.with_suffix(".source_structure.json").read_text("utf-8"))
    target = json.loads(path.with_suffix(".structure.json").read_text("utf-8"))
    record_path = path.with_suffix(".layout_eligibility.json")
    assert load_layout_eligibility(path, source, target) is not None
    if corruption == "missing":
        record_path.unlink()
    elif corruption in {"translation", "version"}:
        record = json.loads(record_path.read_text("utf-8"))
        record["translation_sha256" if corruption == "translation" else "version"] = "changed"
        record_path.write_text(json.dumps(record), encoding="utf-8")
    else:
        if corruption == "source_text":
            source["blocks"][1]["text"] += " extra source text"
        elif corruption == "word_confidence":
            source["metadata"]["ocr_word_evidence"]["words"][0]["confidence"] = 30
        else:
            source["page_number"] = 3
        path.with_suffix(".source_structure.json").write_text(json.dumps(source), encoding="utf-8")
    assert load_layout_eligibility(path, source, target) is None
    before = page_hashes(pages)
    output = assemble_docx(pages, tmp_path / "stale.docx", lang=TargetLang.EN,
        page_breaks=False, page_numbers=[1, 2], derive_source_continuations=True)
    assert not any(page.get("section_furniture_adopted") for page in mapping(output))
    assert not any("joined_to_block_id" in b for p in mapping(output) for b in p["blocks"])
    assert page_hashes(pages) == before


@pytest.mark.parametrize("mode", ["matching", "partial", "manual_edit", "changed_raster"])
def test_bundle_rebuild_keeps_selection_and_saved_text_barriers(tmp_path, monkeypatch, offline, mode):
    workflow, config, result, calls, client = bundle_run(tmp_path, monkeypatch)
    pages = result.run_dir / "pages"
    if mode == "manual_edit":
        path = pages / "page_0002.txt"
        path.write_text(path.read_text("utf-8") + " Human note.", encoding="utf-8")
    if mode == "changed_raster":
        from legalpdf_translate.browser_pdf_bundle import browser_pdf_bundle_page_image_path
        path = browser_pdf_bundle_page_image_path(config.pdf_path, 1)
        with Image.open(path) as original:
            image = original.copy()
        ImageDraw.Draw(image).rectangle((5, 5, 12, 12), fill="black")
        image.save(path)
    before = page_hashes(pages)
    if mode in {"manual_edit", "changed_raster"}:
        output = workflow.rebuild_docx(config)
    else:
        output = assemble_docx(pages, tmp_path / (mode + ".docx"), lang=TargetLang.EN,
            page_breaks=mode == "matching", page_numbers=[1] if mode == "partial" else [1, 2],
            partial_output=mode == "partial", derive_source_continuations=True)
    assert not any(page.get("section_furniture_adopted") for page in mapping(output))
    assert not any("joined_to_block_id" in b for p in mapping(output) for b in p["blocks"])
    assert page_hashes(pages) == before and len(calls) == len(client.calls) == 2
    if mode == "manual_edit":
        assert "Human note." in " ".join(p.text for p in Document(output).paragraphs)


def test_direct_layout_derivation_rechecks_current_browser_raster(tmp_path, monkeypatch, offline):
    _, config, result, _, _ = bundle_run(tmp_path, monkeypatch)
    from legalpdf_translate.browser_pdf_bundle import browser_pdf_bundle_page_image_path
    path = result.run_dir / "pages" / "page_0001.txt"
    source = json.loads(path.with_suffix(".source_structure.json").read_text("utf-8"))
    target = json.loads(path.with_suffix(".structure.json").read_text("utf-8"))
    proof = load_layout_eligibility(path, source, target)
    assert derive_source_layout(source, config.pdf_path, layout_eligibility=proof)["status"] == "flow"
    raster = browser_pdf_bundle_page_image_path(config.pdf_path, 1)
    with raster.open("ab") as handle:
        handle.write(b"changed-source-raster")
    assert derive_source_layout(source, config.pdf_path, layout_eligibility=proof)["status"] == "needs_review"


def test_oversize_image_header_falls_back_without_losing_saved_text(tmp_path, monkeypatch, offline):
    _, config, result, calls, client = bundle_run(tmp_path, monkeypatch)
    pages = result.run_dir / "pages"
    before = page_hashes(pages)
    def bomb(*a, **k):
        raise Image.DecompressionBombError("synthetic image-header limit")
    monkeypatch.setattr(Image, "open", bomb)
    preparation = prepare_layout_rebuild(pages, config.pdf_path)
    assert preparation["review_required_pages"] == [1, 2]
    assert page_hashes(pages) == before and len(calls) == len(client.calls) == 2


def test_malformed_saved_layout_keeps_text_instead_of_aborting_assembly(tmp_path, monkeypatch, offline):
    _, _, result, _, _ = bundle_run(tmp_path, monkeypatch)
    pages = result.run_dir / "pages"
    path = pages / "page_0001.structure.json"
    target = json.loads(path.read_text("utf-8"))
    target["metadata"]["layout"] = "invalid"
    path.write_text(json.dumps(target), encoding="utf-8")
    (pages / "page_0001.layout.json").unlink()
    before = page_hashes(pages)
    output = assemble_docx(pages, tmp_path / "malformed.docx", lang=TargetLang.EN,
        page_breaks=False, page_numbers=[1, 2], derive_source_continuations=True)
    body = " ".join(p.text for p in Document(output).paragraphs)
    assert all(text in body for text in TARGET + FOOTERS)
    assert mapping(output)[0]["layout_review_required"]
    assert page_hashes(pages) == before
