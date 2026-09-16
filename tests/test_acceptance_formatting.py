"""Synthetic exact-bundle to detached-formatting assembly; no live artifacts."""
from copy import deepcopy
from dataclasses import replace
import hashlib
import json

import pytest

from legalpdf_translate import acceptance_formatting as assembly
from legalpdf_translate.acceptance_assembly import PageBundleRef
from legalpdf_translate.formatting_support import fingerprint
from legalpdf_translate.structured_artifacts import publish_structured_page, StructuredArtifactError
from legalpdf_translate.types import TargetLang
from tests.test_reviewed_formatting import two_page_packet


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


@pytest.fixture(params=[(protocol, lang) for protocol in (
    'legal_blocks_reviewed_recovery_v1', 'legal_blocks_legacy_reviewed_recovery_v1')
    for lang in (TargetLang.FR, TargetLang.AR)])
def case(tmp_path, request):
    protocol, lang = request.param
    target_parts = None if lang == TargetLang.FR else [
        'المحكمة القضائية\nالدائرة الجنائية\n\n',
        'يُستدعى [[João Guerreiro]] عند [[09]]:[[30]] مع بقاء المهلة 10 أيام.\n\n',
        'القاضية\n[[Nome Exemplo]]\n\n',
        '[[Largo do Exemplo]] - [[1234-567]] [[Exemplo]]\nهاتف: [[123456789]]\n', '1 / 1',
    ]
    manifest, pages = two_page_packet(target_parts=target_parts)
    manifest['target_lang'] = lang.value
    refs = []
    for number, (page, row) in enumerate(zip(pages, manifest['pages']), 1):
        source, target = deepcopy(page.source_structure), deepcopy(page.target_structure)
        provenance = {'version': protocol, 'lang': lang.value,
                      'preferences_sha256': 'd' * 64}
        target['metadata']['recovery'] = provenance
        identity = {'protocol': provenance['version'], 'fingerprint': fingerprint(provenance)}
        folder = tmp_path / f'original_{number}'
        record = publish_structured_page(folder, source_structure=source, translated_structure=target,
            translated_text=target['blocks'][0]['text'], protocol_identity=identity,
            page_fingerprint=fingerprint({'recovery': identity, 'page': number}), page_result=None)
        commit_sha = sha((folder / f'page_{number:04d}.commit.json').read_bytes())
        refs.append(PageBundleRef(folder, record, commit_sha, 'e' * 64))
        row.update(commit_file_sha256=commit_sha, bundle_sha256=record['bundle_sha256'],
                   source_structure_sha256=fingerprint(source), target_structure_sha256=fingerprint(target))
    raw = json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()
    args = dict(page_bundles=refs, full_case_pages=[1, 2], source_file_sha256='a' * 64,
        lang=lang, preferences_sha256='d' * 64, output_dir=tmp_path / 'formatted',
        formatting_manifest=raw, expected_manifest_sha256=sha(raw),
        source_images=tuple(page.source_image_bytes for page in pages))
    binding = assembly.partition_assembly_binding(page_bundles=refs, full_case_pages=[1, 2],
        source_file_sha256='a' * 64, lang=lang, preferences_sha256='d' * 64,
        formatting_manifest_sha256=sha(raw))
    args['evidence_guard'] = lambda: deepcopy(binding)
    return args


def snapshot(case):
    return {path: path.read_bytes() for ref in case['page_bundles'] for path in ref.pages_dir.iterdir()}


@pytest.mark.parametrize('columns', [False, True])
def test_explicit_v2_region_profile_through_original_bundle_assembly(tmp_path, columns):
    from tests.test_reviewed_formatting_v2 import region_packet
    manifest, page = region_packet(columns=columns)
    source, target = deepcopy(page.source_structure), deepcopy(page.target_structure)
    provenance = {'version': 'legal_blocks_reviewed_recovery_v1', 'lang': 'FR',
                  'preferences_sha256': 'd' * 64}
    target['metadata']['recovery'] = provenance
    identity = {'protocol': provenance['version'], 'fingerprint': fingerprint(provenance)}
    folder = tmp_path / 'original'
    record = publish_structured_page(folder, source_structure=source, translated_structure=target,
        translated_text=target['blocks'][0]['text'], protocol_identity=identity,
        page_fingerprint=fingerprint({'recovery': identity, 'page': 1}), page_result=None)
    commit_sha = sha((folder / 'page_0001.commit.json').read_bytes())
    refs = [PageBundleRef(folder, record, commit_sha, 'e' * 64)]
    manifest['pages'][0].update(commit_file_sha256=commit_sha, bundle_sha256=record['bundle_sha256'],
        source_structure_sha256=fingerprint(source), target_structure_sha256=fingerprint(target))
    raw = json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()
    binding = assembly.partition_assembly_binding(page_bundles=refs, full_case_pages=[1],
        source_file_sha256='a' * 64, lang=TargetLang.FR, preferences_sha256='d' * 64,
        formatting_manifest_sha256=sha(raw))
    before = {path: path.read_bytes() for path in folder.iterdir()}
    output = assembly.assemble_reviewed_partition_case(page_bundles=refs, full_case_pages=[1],
        source_file_sha256='a' * 64, lang=TargetLang.FR, preferences_sha256='d' * 64,
        output_dir=tmp_path / 'formatted_v2', formatting_manifest=raw,
        expected_manifest_sha256=sha(raw), source_images=(page.source_image_bytes,),
        evidence_guard=lambda: deepcopy(binding))
    mapping = json.loads(output.with_suffix('.source_map.json').read_bytes())
    assert mapping['policy'] == 'source_page_matched_reviewed_regions_v2'
    assert mapping['rendered_layout_acceptance'] == 'not_evaluated'
    assert all(path.read_bytes() == content for path, content in before.items())
    assert (output.parent / 'reviewed_formatting.json').read_bytes() == raw
    summary = json.loads((output.parent / 'assembly.json').read_bytes())
    assert summary['provider_dispatch_count'] == 0
    assert summary['assembly_cost_usd'] == '0'


def test_complete_formatting_keeps_originals_and_exact_pinned_manifest(case):
    before = snapshot(case)
    result = assembly.assemble_reviewed_partition_case(**case)
    assert result.name == f'formatted_{case["lang"].value}.docx'
    assert result.read_bytes()
    manifest = json.loads((result.parent / 'assembly.json').read_bytes())
    assert manifest['binding'] == case['evidence_guard']()
    assert manifest['docx_sha256'] == sha(result.read_bytes())
    assert manifest['source_map_sha256'] == sha(result.with_suffix('.source_map.json').read_bytes())
    assert manifest['formatting_manifest_sha256'] == case['expected_manifest_sha256']
    assert manifest['complete_source_page_set'] is True
    assert manifest['original_artifacts_unchanged'] is True
    assert manifest['provider_dispatch_count'] == 0
    assert manifest['assembly_cost_usd'] == '0'
    assert manifest['rendered_page_count'] is None
    assert manifest['rendered_layout_acceptance'] == 'not_evaluated'
    assert manifest['source_geometry_status'] == 'not_verified'
    assert (result.parent / 'reviewed_formatting.json').read_bytes() == case['formatting_manifest']
    if case['lang'] == TargetLang.AR:
        import io
        from docx import Document
        from docx.oxml.ns import qn
        document = Document(io.BytesIO(result.read_bytes()))
        assert len(document.sections) == 2
        assert all(not section.header.is_linked_to_previous
                   and not section.footer.is_linked_to_previous for section in document.sections)
        for text in ('João Guerreiro', '09:30'):
            assert sum(text in run.text for paragraph in document.paragraphs
                       for run in paragraph.runs) == 2
        source_map = json.loads(result.with_suffix('.source_map.json').read_bytes())
        assert source_map['typography']['font'] == 'Arial'
        assert source_map['typography']['size_pt'] == 11
        assert source_map['target_lang'] == 'AR'
        for page in source_map['pages']:
            for fragment in page['fragments']:
                if fragment['location']['kind'] == 'body_paragraph':
                    paragraph = document.paragraphs[fragment['location']['paragraph_index']]
                    assert paragraph._p.find('./' + qn('w:pPr') + '/' + qn('w:bidi')).get(qn('w:val')) == '1'
    for path, raw in before.items():
        assert path.read_bytes() == raw
        assert (result.parent / 'pages' / path.name).read_bytes() == raw


@pytest.mark.parametrize('field,value', [('evidence_guard', None), ('expected_manifest_sha256', 'f' * 64),
    ('full_case_pages', [2, 1]), ('source_file_sha256', 'f' * 64), ('preferences_sha256', 'f' * 64),
    ('lang', TargetLang.EN), ('source_images', (b'not image',)), ('formatting_manifest', b'{}')])
def test_wrong_binding_fails_before_output(case, field, value):
    case[field] = value
    with pytest.raises((StructuredArtifactError, ValueError)):
        assembly.assemble_reviewed_partition_case(**case)
    assert not case['output_dir'].exists()


@pytest.mark.parametrize('field', ['commit_file_sha256', 'evidence_identity_sha256'])
def test_changed_original_reference_rejected(case, field):
    case['page_bundles'][0] = replace(case['page_bundles'][0], **{field: 'f' * 64})
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_reviewed_partition_case(**case)
    assert not case['output_dir'].exists()


@pytest.mark.parametrize('suffix', ['txt', 'source_structure.json', 'structure.json', 'commit.json',
                                  'layout.json', 'layout_eligibility.json'])
def test_changed_or_unbound_original_artifact_rejected(case, suffix):
    path = case['page_bundles'][0].pages_dir / ('page_0001.' + suffix)
    path.write_bytes(b'changed')
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_reviewed_partition_case(**case)
    assert not case['output_dir'].exists()


def test_wrong_image_hash_rejected_before_output(case):
    case['source_images'] = (case['source_images'][0] + b'changed', case['source_images'][1])
    with pytest.raises(ValueError):
        assembly.assemble_reviewed_partition_case(**case)
    assert not case['output_dir'].exists()


def test_output_cannot_overlap_original_or_replace_existing(case):
    case['output_dir'] = case['page_bundles'][0].pages_dir / 'nested'
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_reviewed_partition_case(**case)
    case['output_dir'] = case['page_bundles'][0].pages_dir.parent / 'existing'
    case['output_dir'].mkdir()
    sentinel = case['output_dir'] / 'keep.txt'
    sentinel.write_text('keep')
    with pytest.raises((StructuredArtifactError, FileExistsError)):
        assembly.assemble_reviewed_partition_case(**case)
    assert sentinel.read_text() == 'keep'


def test_second_invocation_is_never_a_resume(case):
    output = assembly.assemble_reviewed_partition_case(**case)
    old = {p: p.read_bytes() for p in output.parent.rglob('*') if p.is_file()}
    with pytest.raises((StructuredArtifactError, FileExistsError)):
        assembly.assemble_reviewed_partition_case(**case)
    assert all(p.read_bytes() == raw for p, raw in old.items())


@pytest.mark.parametrize('phase', [1, 2, 4, 6])
def test_original_mutation_at_guard_phase_never_publishes_success(case, phase):
    guard = case['evidence_guard']
    calls = 0
    def changed_guard():
        nonlocal calls
        calls += 1
        if calls == phase:
            (case['page_bundles'][0].pages_dir / 'page_0001.txt').write_bytes(b'changed')
        return guard()
    case['evidence_guard'] = changed_guard
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_reviewed_partition_case(**case)
    assert not (case['output_dir'] / 'assembly.json').exists()


def test_writer_is_one_pure_call_and_tampered_result_cannot_publish(case):
    from legalpdf_translate.reviewed_formatting_writer import build_reviewed_docx
    calls = []
    def writer(projection):
        calls.append(projection)
        artifact = build_reviewed_docx(projection)
        return replace(artifact, docx_bytes=b'not a DOCX')
    with pytest.raises(ValueError):
        assembly.assemble_reviewed_partition_case(**case, writer=writer)
    assert len(calls) == 1
    assert not (case['output_dir'] / 'assembly.json').exists()


def test_writer_mutation_of_evidence_rejected(case):
    from legalpdf_translate.reviewed_formatting_writer import build_reviewed_docx
    def writer(projection):
        artifact = build_reviewed_docx(projection)
        (case['page_bundles'][0].pages_dir / 'page_0001.txt').write_bytes(b'changed')
        return artifact
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_reviewed_partition_case(**case, writer=writer)
    assert not (case['output_dir'] / 'assembly.json').exists()
