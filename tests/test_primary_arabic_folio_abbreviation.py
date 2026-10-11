"""Arabic terminal labels use a persisted primary-only recipe."""
import json
from copy import deepcopy
from types import SimpleNamespace
from zipfile import ZipFile
from io import BytesIO
import pytest
from docx import Document
from legalpdf_translate.docx_writer import assemble_docx, _terminal_labeled_folios_cover_v2
from legalpdf_translate.types import TargetLang
from legalpdf_translate.checkpoint import primary_assembly_recipe
from legalpdf_translate.ordinary_auto_layout import frozen_automatic_layout_policy, _optional_hint_policy


def _raw(tmp_path, text, version="terminal_labeled_folio_v2", **options):
    pages=tmp_path/"pages";pages.mkdir()
    (pages/"page_0001.txt").write_text(text,encoding="utf-8")
    path=assemble_docx(pages,tmp_path/"raw.docx",lang=TargetLang.AR,page_breaks=False,
        terminal_folio_policy=version,page_numbers=[1],**options)
    return path


def _has_page(raw):
    with ZipFile(BytesIO(raw)) as package:
        return any(b"PAGE" in package.read(n) for n in package.namelist() if n.startswith("word/footer"))


@pytest.mark.parametrize("label", ["ص. 1","ص.١","ص. ۱","\u200eص. 1\u200e","ص. 1 / 7"])
def test_explicit_terminal_abbreviation_is_preserved_without_generated_page(tmp_path,label):
    path=_raw(tmp_path,"Body text.\n"+label)
    assert not _has_page(path.read_bytes())
    text="\n".join(p.text for p in Document(path).paragraphs).replace("\u200e","")
    assert label.replace("\u200e","") in text
    assert json.loads(path.with_suffix(".source_map.json").read_text(encoding="utf-8"))["pages"][0]["primary_folio_policy"]=="terminal_labeled_folio_v2"


@pytest.mark.parametrize("text", ["Body.\nص", "Body.\n1", "Body.\nص. 2", "Body.\nانظر ص. 1", "ص. 1\nBody continues.", "Body.\nص. 1 / 0"])
def test_unsafe_or_ambiguous_abbreviation_does_not_suppress_page(tmp_path,text):
    assert _has_page(_raw(tmp_path,text).read_bytes())


@pytest.mark.parametrize("version", [None,"terminal_labeled_folio_v1"])
def test_historical_recipes_keep_their_page_field(tmp_path,version):
    assert _has_page(_raw(tmp_path,"Body.\nص. 1",version).read_bytes())


def test_partial_and_structured_inputs_preserve_existing_numbering(tmp_path):
    assert _has_page(_raw(tmp_path,"Body.\nص. 1",partial_output=True).read_bytes())
    assert not _terminal_labeled_folios_cover_v2([SimpleNamespace(number=1,text="Body.\nص. 1",structure={})])
    assert not _terminal_labeled_folios_cover_v2([SimpleNamespace(number=2,text="Body.\nص. 2",structure=None),SimpleNamespace(number=7,text="Body.",structure=None)])


@pytest.mark.parametrize("version,expected", [("terminal_labeled_folio_v1",True),("terminal_labeled_folio_v2",False),(None,True)])
def test_offline_assembly_dispatches_saved_version_not_current_catalog(tmp_path,monkeypatch,version,expected):
    from legalpdf_translate.workflow import TranslationWorkflow
    from legalpdf_translate import ordinary_auto_layout as policies
    old_policy=_optional_hint_policy("ordinary_optional_layout_hint_normalization_v2",emphasis_normalization=True,metadata_normalization=True)
    settings={"ordinary_auto_layout_policy":old_policy}
    if version:settings["primary_assembly_recipe"]={"version":version,"ordinary_auto_layout_policy":old_policy}
    state=SimpleNamespace(settings=settings,pages={"1":{"status":"done"}},selection_page_count=1)
    workflow=object.__new__(TranslationWorkflow)
    workflow._prepare_docx_layout=lambda *a,**k:{}
    workflow._record_docx_layout_review=lambda *a,**k:{}
    monkeypatch.setattr(policies,"frozen_automatic_layout_policy",lambda:pytest.fail("rebuild consulted current catalog"))
    pages=tmp_path/"pages";pages.mkdir();(pages/"page_0001.txt").write_text("Body.\nص. 1",encoding="utf-8")
    config=SimpleNamespace(target_lang=TargetLang.AR,page_breaks=False,strip_bidi_controls=True)
    path,_=workflow._assemble_run_docx(config,pages_dir=pages,output_path=tmp_path/"out.docx",output_dir=tmp_path,run_dir=tmp_path,run_state=state,assembly_options={})
    assert _has_page(path.read_bytes()) is expected
    if version: assert primary_assembly_recipe(settings)["version"]==version


@pytest.mark.parametrize("source_evidence", [False,True])
def test_new_primary_passes_existing_automatic_writer_without_footer_hint(tmp_path,source_evidence):
    from legalpdf_translate.ordinary_auto_layout_artifacts import bind_raw_page_map,build_unreviewed_candidate,verify_unreviewed_candidate
    from legalpdf_translate.saved_docx_layout import default_decisions
    path=_raw(tmp_path,"Body.\nص. 1")
    raw=path.read_bytes();mapping=json.loads(path.with_suffix(".source_map.json").read_text(encoding="utf-8"))
    snapshot=bind_raw_page_map(raw,mapping,source_pdf_sha256="a"*64,selected_pages=(1,),target_lang="AR")
    frames=[{"page_number":1,"width_px":595,"height_px":842,"image_sha256":"b"*64}]
    decisions=default_decisions(snapshot.saved_snapshot,frames)
    for choice in decisions["paragraphs"]:choice["regions"]=[{"page_number":1,"bbox_px":[10,10,200,100]}]
    evidence={"version":"fictional_proposal_v1","document_reviewed":False,"rendered_layout_acceptance":"not_evaluated"}
    if source_evidence:evidence["source_evidence"]={"version":"ordinary_source_evidence_v1","pages":[{"page_number":1,"layout":[],"findings":[]}]}
    candidate=build_unreviewed_candidate(raw,snapshot,frames,decisions,proposal_evidence=evidence)
    verify_unreviewed_candidate(raw,snapshot,frames,decisions,candidate,proposal_evidence=evidence)
    assert not _has_page(candidate.docx_bytes)
    assert "ص. 1" in "\n".join(p.text for p in Document(BytesIO(candidate.docx_bytes)).paragraphs).replace("\u200e","")
    assert json.loads(candidate.source_map_bytes)["writer_version"].endswith("v7" if source_evidence else "v6")


def test_new_policy_differs_from_exact_historical_primary_policy():
    prior=_optional_hint_policy("ordinary_optional_layout_hint_normalization_v2",emphasis_normalization=True,metadata_normalization=True)
    assert prior != frozen_automatic_layout_policy()
