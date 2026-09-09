"""Logical boundaries and page-label regressions; never invoke native Word."""
from copy import deepcopy
import json

from docx import Document
from docx.oxml.ns import qn
import pytest

from legalpdf_translate.document_layout import derive_page_layout
from legalpdf_translate.document_structure import PageStructure, StructureBlock, text_sha256
from legalpdf_translate.docx_writer import assemble_docx, sanitize_bidi_controls
from legalpdf_translate.types import TargetLang


def _save(folder, number, rows, *, start=False, regions=False):
    folder.mkdir(exist_ok=True)
    blocks = [StructureBlock(f"p{number:04d}_b{index:04d}", **row) for index, row in enumerate(rows, 1)]
    source = PageStructure(number, text_sha256("\n".join(block.text for block in blocks)), blocks,
                           source_file_sha256="b" * 64, document_start=start)
    source.source_text_sha256 = source.source_sha256
    target = deepcopy(source)
    target.translation_sha256 = text_sha256(target.text)
    if regions:
        layout = derive_page_layout(target.to_dict())
        assert layout["status"] in {"regions", "needs_review"}
        target.metadata["layout"] = layout
    path = folder / f"page_{number:04d}.txt"
    path.write_text(target.text, encoding="utf-8")
    path.with_suffix(".structure.json").write_text(json.dumps(target.to_dict()), encoding="utf-8")
    path.with_suffix(".source_structure.json").write_text(json.dumps(source.to_dict()), encoding="utf-8")
    return target


def _opening(*, arabic=False):
    return [
        {"text": "المحكمة" if arabic else "Court", "role": "header", "bbox": (40, 30, 550, 45)},
        {"text": "الملف [[187/26]]" if arabic else "Case 187/26", "role": "reference", "bbox": (40, 60, 550, 75)},
        {"text": "قرار قضائي" if arabic else "Judicial decision", "role": "heading", "bbox": (40, 100, 550, 118),
         "document_start": True},
        {"text": "يجب الوفاء بالالتزام" if arabic else "The obligation must be fulfilled.", "bbox": (40, 150, 550, 168)},
    ]


def _break_count(document):
    return len(document._element.body.xpath(".//w:br[@w:type='page']|.//w:pPr/w:sectPr"))


def _mapping(path):
    return json.loads(path.with_suffix(".source_map.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("lang", [TargetLang.EN, TargetLang.AR])
@pytest.mark.parametrize("page_breaks", [False, True])
@pytest.mark.parametrize("page_number,table_end", [(1, False), (2, False), (2, True)])
def test_initial_page_and_title_start_are_one_boundary(tmp_path, lang, page_breaks, page_number, table_end):
    folder = tmp_path / "pages"
    if page_number == 2:
        first = [{"text": "Earlier complete document."}]
        if table_end:
            first.append({"text": "Earlier table cell", "role": "table_cell", "table_id": "t",
                          "row": 0, "col": 0})
        _save(folder, 1, first, start=True)
    current = _save(folder, page_number, _opening(arabic=lang == TargetLang.AR), start=True)
    before = {path.name: path.read_bytes() for path in folder.iterdir()}

    output = assemble_docx(folder, tmp_path / "out.docx", lang=lang, page_breaks=page_breaks)

    document = Document(output)
    assert _break_count(document) == page_number - 1
    opening_paragraphs = document.paragraphs[-4:]
    assert all(not paragraph._p.xpath(".//w:br[@w:type='page']") for paragraph in opening_paragraphs)
    assert [row["block_id"] for row in _mapping(output)["pages"][-1]["blocks"]] == [block.id for block in current.blocks]
    assert before == {path.name: path.read_bytes() for path in folder.iterdir()}


@pytest.mark.parametrize("page_start", [False, True])
def test_first_output_page_has_no_orphan_frontmatter_before_initial_title(tmp_path, page_start):
    folder = tmp_path / "pages"
    _save(folder, 1, _opening(), start=page_start)
    document = Document(assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.EN, page_breaks=False))
    assert _break_count(document) == 0


@pytest.mark.parametrize("page_breaks", [False, True])
def test_genuine_later_intrapage_start_survives_initial_boundary_deduplication(tmp_path, page_breaks):
    folder = tmp_path / "pages"
    _save(folder, 1, [{"text": "Previous document."}], start=True)
    rows = _opening() + [{"text": "Separate later decision", "role": "heading", "document_start": True},
                         {"text": "A second substantive obligation."}]
    _save(folder, 2, rows, start=True)
    document = Document(assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.EN, page_breaks=page_breaks))
    assert _break_count(document) == 2
    assert document.paragraphs[-3]._p.xpath(".//w:br[@w:type='page']")


def test_prior_substantive_content_prevents_consuming_later_start(tmp_path):
    folder = tmp_path / "pages"
    rows = [{"text": "Substantive earlier obligation."}, {"text": "New decision", "role": "heading", "document_start": True}]
    _save(folder, 1, rows, start=True)
    document = Document(assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.EN, page_breaks=False))
    assert _break_count(document) == 1


def _region_rows(*, later_start=False, footer=False):
    rows = _opening(arabic=True)[:3]
    rows += [
        {"text": "عنوان أيسر", "role": "heading", "bbox": (30, 140, 240, 157)},
        {"text": "التزام أيسر", "bbox": (30, 175, 240, 195)},
        {"text": "عنوان أيمن", "role": "heading", "bbox": (300, 140, 560, 157)},
        {"text": "التزام أيمن", "bbox": (300, 175, 560, 195)},
        {"text": "القرار اللاحق", "role": "heading", "bbox": (30, 240, 560, 258), "document_start": later_start},
    ]
    if footer:
        rows.append({"text": "الصفحة \u2066[[2]]\u2069 من [[3]]", "role": "footer", "bbox": (30, 800, 560, 820)})
    return rows


@pytest.mark.parametrize("later_start", [False, True])
def test_interior_start_regions_keep_conservative_fallback_and_single_initial_boundary(tmp_path, later_start):
    folder = tmp_path / "pages"
    _save(folder, 1, [{"text": "Earlier document."}], start=True)
    _save(folder, 2, _region_rows(later_start=later_start), start=True, regions=True)
    output = assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.AR, page_breaks=False)
    # Existing region admission rejects interior start markers. Its flow
    # fallback still needs the same deduplication; do not weaken admission.
    assert _mapping(output)["pages"][1]["layout_status"] == "needs_review_flow_fallback"
    assert _break_count(Document(output)) == 1 + int(later_start)


@pytest.mark.parametrize("label", ["الصفحة \u2066[[1]]\u2069", "صفحة [[٢]] من [[٣]]", "Page [[1]] of [[2]]", "\u200f[[1]] / [[2]]\u200e"])
@pytest.mark.parametrize("page_breaks", [False, True])
def test_protected_page_only_footers_map_to_actual_word_page_field(tmp_path, label, page_breaks):
    folder = tmp_path / "pages"
    _save(folder, 1, [{"text": "التزام كامل"}, {"text": label, "role": "footer"}])
    output = assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.AR, page_breaks=page_breaks)
    document = Document(output)
    assert [paragraph.text for paragraph in document.paragraphs] == ["التزام كامل"]
    assert _mapping(output)["pages"][0]["blocks"][1]["location"]["kind"] == "generated_footer_page_field"
    assert " PAGE " in document.sections[0].footer._element.xml


def test_protected_page_footer_in_regions_is_not_duplicate_body_text(tmp_path):
    folder = tmp_path / "pages"
    rows = [
        {"text": "عنوان أيسر", "role": "heading", "bbox": (25, 45, 225, 65)},
        {"text": "التزام أيسر", "bbox": (25, 75, 225, 100)},
        {"text": "عنوان أيمن", "role": "heading", "bbox": (285, 45, 565, 65)},
        {"text": "التزام أيمن", "bbox": (285, 75, 565, 110)},
        {"text": "الصفحة \u2066[[2]]\u2069 من [[3]]", "role": "footer", "bbox": (25, 800, 565, 820)},
    ]
    _save(folder, 1, rows, start=True, regions=True)
    output = assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.AR, page_breaks=False)
    assert _mapping(output)["pages"][0]["layout_status"] == "regions"
    assert _mapping(output)["pages"][0]["blocks"][-1]["location"]["kind"] == "generated_footer_page_field"
    visible = "".join(Document(output)._element.body.itertext())
    assert "الصفحة" not in visible


@pytest.mark.parametrize("role,text", [
    ("paragraph", "الصفحة [[1]]"),
    ("footer", "راجع الصفحة [[1]] خلال [[10]] أيام"),
    ("footer", "الصفحة [[1]] مع التزام إضافي"),
    ("footer", "[[court@example.invalid]] هاتف [[210000000]]"),
    ("footer", "الصفحة [[1]"),
])
def test_page_label_recognition_does_not_delete_substantive_or_malformed_text(tmp_path, role, text):
    folder = tmp_path / "pages"
    _save(folder, 1, [{"text": text, "role": role}])
    output = assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.AR, page_breaks=False)
    assert _mapping(output)["pages"][0]["blocks"][0]["location"]["kind"] == "body_paragraph"
    assert Document(output).paragraphs[0].text


def test_page_label_fallback_inside_active_furniture_sections_preserves_contacts(tmp_path):
    folder = tmp_path / "pages"
    for number in (1, 2):
        _save(folder, number, [
            {"text": "Tribunal Judicial de Cidade", "role": "header", "bbox": (100, 90, 500, 110)},
            {"text": f"Corpo completo {number}.", "bbox": (40, 150, 550, 170)},
            {"text": "E-mail: court@example.invalid", "role": "footer", "bbox": (100, 800, 500, 820)},
        ], start=number == 1)
    _save(folder, 3, [{"text": "مستند مستقل", "document_start": True},
                      {"text": "الصفحة [[3]]", "role": "footer"}], start=True)
    output = assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.AR, page_breaks=False)
    document = Document(output)
    mapping = _mapping(output)
    assert mapping["pages"][0]["section_furniture_adopted"] is True
    assert "court@example.invalid" in "\n".join(paragraph.text for paragraph in document.sections[0].footer.paragraphs)
    assert mapping["pages"][2]["blocks"][-1]["location"]["kind"] == "generated_footer_page_field"
    assert len(document.sections) == 2
