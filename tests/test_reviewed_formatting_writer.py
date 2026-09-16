"""Synthetic in-memory DOCX contract tests, never native/render acceptance."""
from copy import deepcopy
from dataclasses import replace
import hashlib
import io
import json
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from lxml import etree
import pytest

from tests.test_reviewed_formatting import packet, two_page_packet, validate
from legalpdf_translate.reviewed_formatting_writer import (
    ReviewedFormattingWriterError, build_reviewed_docx, validate_reviewed_docx,
)


def projection(two_pages=False, lang="FR"):
    if two_pages:
        manifest, pages = two_page_packet()
    else:
        manifest, page = packet()
        pages = [page]
    manifest["target_lang"] = lang
    return validate(manifest, pages[0], pages=pages, target_lang=lang)


def decoded(artifact):
    return Document(io.BytesIO(artifact.docx_bytes)), json.loads(artifact.source_map_bytes)


def changed_document(artifact, mutate):
    document, mapping = decoded(artifact)
    mutate(document)
    output = io.BytesIO()
    document.save(output)
    content = output.getvalue()
    mapping["docx_sha256"] = hashlib.sha256(content).hexdigest()
    return content, json.dumps(mapping).encode()


def changed_package(artifact, name, mutate):
    output = io.BytesIO()
    with ZipFile(io.BytesIO(artifact.docx_bytes)) as source, ZipFile(output, "w") as target:
        for member in source.infolist():
            raw = source.read(member.filename)
            target.writestr(member, mutate(raw) if member.filename == name else raw)
    content = output.getvalue()
    mapping = json.loads(artifact.source_map_bytes)
    mapping["docx_sha256"] = hashlib.sha256(content).hexdigest()
    return content, json.dumps(mapping).encode()


@pytest.mark.parametrize("lang", ["FR", "EN"])
def test_complete_case_has_unlinked_real_parts_exact_text_and_fixed_user_typography(lang):
    reviewed = projection(two_pages=True, lang=lang)
    before = deepcopy(reviewed)
    artifact = build_reviewed_docx(reviewed)
    document, mapping = decoded(artifact)
    validate_reviewed_docx(artifact.docx_bytes, artifact.source_map_bytes, projection=reviewed)
    assert reviewed == before
    assert len(document.sections) == 2
    assert mapping["policy"] == "source_page_matched_reviewed_v1"
    assert mapping["rendered_page_count"] is None and mapping["layout_review_required"]
    assert mapping["rendered_layout_acceptance"] == "not_evaluated"
    assert mapping["source_page_count"] == 2
    assert mapping["projection_sha256"] == artifact.projection_sha256
    parts, locators = set(), set()
    for section, page in zip(document.sections, mapping["pages"]):
        assert not section.header.is_linked_to_previous and not section.footer.is_linked_to_previous
        assert not section.different_first_page_header_footer
        assert section.page_width.twips == Cm(21).twips and section.page_height.twips == Cm(29.7).twips
        assert section.left_margin.twips == Cm(1.7).twips and section.right_margin.twips == Cm(1.7).twips
        assert section.top_margin.twips == Cm(1.5).twips and section.bottom_margin.twips == Cm(1.5).twips
        assert not list(section._sectPr.iter(qn("w:pgNumType")))
        for kind in ("header", "footer"):
            name = page["section_parts"][kind]
            assert name not in parts
            parts.add(name)
        for fragment in page["fragments"]:
            loc = fragment["location"]
            key = (loc["part_uri"], loc["paragraph_index"])
            assert key not in locators
            locators.add(key)
            paragraph = (document.paragraphs[loc["paragraph_index"]] if loc["kind"] == "body_paragraph"
                         else getattr(section, "header" if loc["kind"] == "section_header" else "footer").paragraphs[loc["paragraph_index"]])
            assert hashlib.sha256(paragraph.text.encode()).hexdigest() == fragment["display_text_sha256"]
            for run in paragraph.runs:
                assert run.font.name == "Times New Roman" and run.font.size.pt == 10.5
    assert all(p.text for p in document.paragraphs), "No generated empty body paragraph for section breaks"
    assert document.sections[0].footer.paragraphs[-1].text == "1 / 2"
    assert document.sections[1].footer.paragraphs[-1].text == "2 / 2"


def test_trailing_separators_and_display_only_transformations_are_explicit_and_reversible_from_source():
    target_parts = ["Tribunal judiciaire\r\nDivision pénale\r\n\r\n",
        "Le nom \u2066[[Cafe\u0301]]\u2069 est préservé.\r\nDeuxième ligne.\r\n \r\n",
        "Le juge\r\nNome Exemplo\r\n\r\n", "Largo do Exemplo\r\nTél. : 123456789\r\n", "1 / 1"]
    manifest, page = packet(target_parts=target_parts)
    reviewed = validate(manifest, page)
    artifact = build_reviewed_docx(reviewed)
    document, mapping = decoded(artifact)
    row = mapping["pages"][0]["fragments"][1]
    assert row["trailing_separators"] == "\r\n \r\n"
    assert row["display_policy"] == "unwrap_placeholders_strip_bidi_crlf_v1"
    assert document.paragraphs[row["location"]["paragraph_index"]].text == "Le nom Cafe\u0301 est préservé.\nDeuxième ligne."
    assert row["target_range"] == list(reviewed.pages[0].fragments[1].target_range)
    assert row["target_text_sha256"] == reviewed.pages[0].fragments[1].target_text_sha256
    validate_reviewed_docx(artifact.docx_bytes, artifact.source_map_bytes, projection=reviewed)


@pytest.mark.parametrize("field,value", [("policy", "flow"), ("layout_review_required", False),
    ("rendered_page_count", 1), ("target_lang", "HE"), ("manifest_sha256", "wrong")])
def test_unvalidated_or_unsupported_projection_state_rejected(field, value):
    with pytest.raises(ReviewedFormattingWriterError):
        build_reviewed_docx(replace(projection(), **{field: value}))


@pytest.mark.parametrize("field,value", [("target_text", "forged"), ("parent_block_id", "p0001_b9999"),
    ("target_range", (0, 1)), ("bbox_px", (0, 0, 600, 840)), ("parent_uncertain", False),
    ("source_text_sha256", "a" * 64)])
def test_forged_fragment_dataclass_rejected(field, value):
    reviewed = projection()
    page = reviewed.pages[0]
    changed = replace(page.fragments[0], **{field: value})
    reviewed = replace(reviewed, pages=(replace(page, fragments=(changed, *page.fragments[1:])),))
    with pytest.raises(ReviewedFormattingWriterError):
        build_reviewed_docx(reviewed)


@pytest.mark.parametrize("mutate", [
    lambda doc: doc.paragraphs[0].add_run("Added words"),
    lambda doc: doc.add_paragraph("Unowned words"),
    lambda doc: doc.add_paragraph(),
    lambda doc: doc.sections[0].header.paragraphs[0].add_run("Added header"),
    lambda doc: doc.sections[0].footer.paragraphs[0].add_run("Added footer"),
    lambda doc: doc.add_table(rows=1, cols=1),
    lambda doc: doc.paragraphs[0]._p.append(OxmlElement("w:hyperlink")),
    lambda doc: setattr(doc.paragraphs[0].runs[0].font, "hidden", True),
    lambda doc: setattr(doc.paragraphs[0].runs[0].font, "size", Cm(.1)),
    lambda doc: setattr(doc.sections[0], "top_margin", Cm(.2)),
    lambda doc: setattr(doc.sections[1].header, "is_linked_to_previous", True),
    lambda doc: setattr(doc.sections[0], "different_first_page_header_footer", True),
])
def test_docx_tampering_rejected_even_after_hash_rebinding(mutate):
    reviewed = projection(two_pages=True)
    artifact = build_reviewed_docx(reviewed)
    raw, mapping = changed_document(artifact, mutate)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(raw, mapping, projection=reviewed)


@pytest.mark.parametrize("element", ["w:fldChar", "w:instrText", "w:drawing", "w:pict", "w:delText", "w:object"])
def test_hidden_fields_drawings_and_other_unowned_run_content_rejected(element):
    reviewed = projection()
    artifact = build_reviewed_docx(reviewed)
    def mutate(doc):
        node = OxmlElement(element)
        doc.paragraphs[0].runs[0]._r.append(node)
    raw, mapping = changed_document(artifact, mutate)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(raw, mapping, projection=reviewed)


@pytest.mark.parametrize("change", [
    lambda m: m.update(layout_review_required=False),
    lambda m: m.update(rendered_page_count=2),
    lambda m: m.update(docx_sha256="a" * 64),
    lambda m: m["pages"][0]["fragments"][0].update(trailing_separators=""),
    lambda m: m["pages"][0]["fragments"][0]["location"].update(part_uri="/word/footer1.xml"),
    lambda m: m["pages"][0]["fragments"][1].update(location=deepcopy(m["pages"][0]["fragments"][0]["location"])),
    lambda m: m["pages"][0]["fragments"].reverse(),
    lambda m: m["pages"][0].update(source_image_sha256="a" * 64),
])
def test_map_binding_locator_ownership_and_separator_tampering_rejected(change):
    reviewed = projection()
    artifact = build_reviewed_docx(reviewed)
    mapping = json.loads(artifact.source_map_bytes)
    change(mapping)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(artifact.docx_bytes, json.dumps(mapping).encode(), projection=reviewed)


def test_duplicate_xml_or_extra_package_content_rejected():
    reviewed = projection()
    artifact = build_reviewed_docx(reviewed)
    output = io.BytesIO()
    with ZipFile(io.BytesIO(artifact.docx_bytes)) as source, ZipFile(output, "w") as target:
        for member in source.infolist():
            target.writestr(member, source.read(member.filename))
        target.writestr("word/hidden.xml", "<hidden>unmapped</hidden>")
    content = output.getvalue()
    mapping = json.loads(artifact.source_map_bytes)
    mapping["docx_sha256"] = hashlib.sha256(content).hexdigest()
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(content, json.dumps(mapping).encode(), projection=reviewed)


def test_source_gap_is_bounded_and_never_changes_font_or_text_to_force_page_count():
    reviewed = projection()
    artifact = build_reviewed_docx(reviewed)
    document, mapping = decoded(artifact)
    for row in mapping["pages"][0]["fragments"]:
        assert 0 <= row["spacing"]["space_before_pt"] <= 48
        assert row["spacing"]["basis"] == "source_region_gap_capped_48pt_v1"
    assert mapping["rendered_page_count"] is None
    assert all(run.font.size.pt == 10.5 for p in document.paragraphs for run in p.runs)


@pytest.mark.parametrize("part,path,attribute", [
    ("word/document.xml", ".//w:p", "text"),
    ("word/document.xml", ".//w:p", "tail"),
    ("word/document.xml", ".//w:r", "text"),
    ("word/document.xml", ".//w:t", "tail"),
    ("word/document.xml", ".//w:pPr", "tail"),
    ("word/document.xml", ".//w:body", "text"),
    ("word/document.xml", ".//w:sectPr", "tail"),
    ("word/header1.xml", ".", "text"),
    ("word/footer1.xml", ".", "text"),
])
def test_direct_xml_text_and_tails_cannot_hide_outside_text_nodes(part, path, attribute):
    reviewed = projection()
    artifact = build_reviewed_docx(reviewed)
    def mutate(raw):
        root = etree.fromstring(raw)
        node = root.find(path, {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
        setattr(node, attribute, "unowned XML text")
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    raw, mapping = changed_package(artifact, part, mutate)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(raw, mapping, projection=reviewed)


def test_bool_cannot_impersonate_numeric_map_count_or_uncertainty_flag():
    reviewed = projection()
    artifact = build_reviewed_docx(reviewed)
    mapping = json.loads(artifact.source_map_bytes)
    mapping["source_page_count"] = True
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(artifact.docx_bytes, json.dumps(mapping).encode(), projection=reviewed)
    page = reviewed.pages[0]
    first = replace(page.fragments[0], parent_uncertain=1)
    with pytest.raises(ReviewedFormattingWriterError):
        build_reviewed_docx(replace(reviewed, pages=(replace(page, fragments=(first, *page.fragments[1:])),)))


def test_external_relationship_is_rejected_without_following_it():
    reviewed = projection()
    artifact = build_reviewed_docx(reviewed)
    def external(raw):
        return raw.replace(b"</Relationships>", b'<Relationship Id="rIdExternal" Type="urn:unsupported" '
            b'Target="https://example.invalid/never-contact" TargetMode="External"/></Relationships>')
    content, mapping = changed_package(artifact, "word/_rels/document.xml.rels", external)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(content, mapping, projection=reviewed)


def test_non_user_text_markup_is_escaped_and_not_executable_ooxml():
    manifest, page = packet(target_parts=["Tribunal judiciaire\n", "<w:fldChar> & \"texte\"\n",
        "Le juge\nNome Exemplo\n", "Largo do Exemplo\nTél. : 123456789\n", "1 / 1"])
    reviewed = validate(manifest, page)
    artifact = build_reviewed_docx(reviewed)
    document, _ = decoded(artifact)
    assert document.paragraphs[0].text == '<w:fldChar> & "texte"'
    assert not list(document._element.iter(qn("w:fldChar")))


def test_even_page_and_first_page_furniture_cannot_be_smuggled_into_unused_parts():
    reviewed = projection()
    artifact = build_reviewed_docx(reviewed)
    def mutate(doc):
        doc.sections[0].first_page_header.is_linked_to_previous = False
        doc.sections[0].first_page_header.paragraphs[0].text = "hidden first-page header"
    content, mapping = changed_document(artifact, mutate)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(content, mapping, projection=reviewed)
