"""V7 source-owned rules/footer stories; XML proofs do not claim native acceptance."""
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
from lxml import etree
import pytest
from legalpdf_translate import saved_docx_layout as model,saved_docx_layout_writer as writer
from legalpdf_translate.docx_writer import _configure_styles,_add_page_number_footer
from tests.test_saved_docx_layout import document_bytes,reviewed,changed_part

def fixture(lang='EN',custom=False):
    def edit(doc):
        _configure_styles(doc,lang);_add_page_number_footer(doc,lang)
        if custom:doc.sections[0].footer.paragraphs[0].add_run(' custom')
        if lang=='AR':
            for p in doc.paragraphs:etree.SubElement(p._p.get_or_add_pPr(),model.W+'bidi').set(model.W+'val','1')
    text=['Recipient address','Body obligation.','________________','Reply instruction.','1 / 1'] if lang!='AR' else ['اسم المستلم Nadia Example','يبقى الالتزام القانوني قائماً.','________________','يرجى إرسال الرد.','الصفحة 1 / 1']
    raw=document_bytes(text,edit)
    snapshot,pages,decisions=reviewed(raw,lang)
    decisions['review'].update(document_reviewed=False,pages_reviewed=[],reviewer='',note='')
    ids=[p['id'] for p in snapshot['paragraphs']]
    decisions['bands']=[{'kind':'flow','groups':[{'paragraph_ids':[i],'panel':False}]} for i in ids]
    for row in decisions['paragraphs']:row.update(regions=[{'page_number':1,'bbox_px':[100,100,800,200]}])
    decisions['paragraphs'][-1].update(role='source_folio',regions=[{'page_number':1,'bbox_px':[400,1300,600,1320]}])
    evidence={'version':'ordinary_source_evidence_v1','pages':[{'page_number':1,'findings':[],'layout':[{'kind':'decorative_rule','paragraph_ids':[ids[2]],'bbox':[.1,.7,.9,.71]},{'kind':'source_footer','paragraph_ids':ids[-2:],'bbox':[.1,.9,.9,.98]}],'source_coverage_verified':False}]}
    context={'selected_pages':[1],'page_groups':[[1,ids]],'source_evidence':evidence}
    return raw,snapshot,pages,decisions,context

def build(args):return writer.build_unreviewed_docx(*args[:4],ordinary_context=args[4])
def verify(artifact,args):writer.validate_built_docx(artifact.docx_bytes,artifact.source_map,original_docx=args[0],snapshot=args[1],pages=args[2],decisions=args[3],require_review=False,ordinary_context=args[4])

def test_arabic_source_folio_numeric_expression_is_contiguous_ltr():
    args=multipage_fixture('AR');artifact=build(args);verify(artifact,args)
    with ZipFile(BytesIO(artifact.docx_bytes)) as archive:
        for name,literal in [('word/footer2.xml','2 / 7'),('word/footer3.xml','7 / 7')]:
            footer=etree.fromstring(archive.read(name));paragraph=footer.findall(model.W+'p')[-1]
            assert ''.join(n.text or '' for n in paragraph.iter(model.W+'t'))=='الصفحة '+literal
            direction=paragraph.find(model.W+'dir')
            assert direction.get(model.W+'val')=='ltr'
            assert ''.join(n.text or '' for n in direction.iter(model.W+'t'))==literal
            assert all(run.find(model.W+'rPr/'+model.W+'rtl').get(model.W+'val')=='0' for run in direction.findall(model.W+'r'))
    bad=changed_part(artifact.docx_bytes,'word/footer2.xml',lambda root:root.find('.//'+model.W+'dir').set(model.W+'val','rtl'))
    mapping=deepcopy(artifact.source_map);mapping['docx_sha256']=model._sha(bad)
    with pytest.raises(ValueError):verify(writer.SavedDocxLayoutArtifact(bad,mapping),args)

def test_folio_direction_recipe_preserves_fonts_and_declines_mixed_text():
    from legalpdf_translate.ordinary_source_layout import _folio_ltr_span
    paragraph=etree.fromstring(('<w:p xmlns:w="'+model.W[1:-1]+'"><w:r><w:rPr><w:rFonts w:ascii="Arial" w:cs="Arial"/><w:sz w:val="22"/></w:rPr><w:t>الصفحة 12 / 37</w:t></w:r></w:p>').encode())
    original=''.join(paragraph.itertext());span=_folio_ltr_span(paragraph)
    assert span['literal']=='12 / 37' and ''.join(paragraph.itertext())==original
    assert all(run.find(model.W+'rPr/'+model.W+'sz').get(model.W+'val')=='22' for run in paragraph.iter(model.W+'r'))
    mixed=deepcopy(paragraph);mixed.find('.//'+model.W+'t').text='Extra operative text '
    before=etree.tostring(mixed)
    assert _folio_ltr_span(mixed) is None and etree.tostring(mixed)==before

@pytest.mark.parametrize('lang',['EN','FR','AR'])
def test_rule_border_and_first_page_source_footer_exact_owned_text(lang):
    args=fixture(lang);original=deepcopy(args[3]);artifact=build(args);verify(artifact,args)
    assert artifact.source_map['writer_version']=='saved_docx_layout_writer_source_layout_v7'
    assert not artifact.source_map['exact_text_preserved']
    assert artifact.source_map['exact_meaningful_text_preserved']
    assert args[3]==original
    with ZipFile(BytesIO(artifact.docx_bytes)) as z:
        body=etree.fromstring(z.read('word/document.xml'))
        assert not any(n.text=='Reply instruction.' for n in body.iter(model.W+'t'))
        assert body.find('.//'+model.W+'titlePg') is not None
        footer=etree.fromstring(z.read('word/footer2.xml'))
        assert [''.join(n.text or '' for n in p.iter(model.W+'t')) for p in footer.findall(model.W+'p')]==[''.join(t['text'] for t in r['tokens'] if t['kind']=='t') for r in args[1]['paragraphs'][-2:]]
        assert b' PAGE ' not in z.read('word/footer2.xml')
        assert list(etree.fromstring(z.read('word/footer1.xml')).iter(model.W+'t'))==[]
    footer_rows=artifact.source_map['paragraphs'][-2:]
    assert all(row['part_uri']=='word/footer2.xml' for row in footer_rows)


def test_custom_footer_retains_source_flow_with_explicit_qualification():
    args=fixture(custom=True);artifact=build(args);verify(artifact,args)
    assert artifact.source_map['source_layout_plan']['source_footers']==[]
    assert artifact.source_map['source_layout_plan']['qualifications']
    assert all(row['part_uri']=='word/document.xml' for row in artifact.source_map['paragraphs'])

@pytest.mark.parametrize('mutation',['text','border','title','footer'])
def test_v7_tamper_fails_closed(mutation):
    args=fixture();artifact=build(args)
    if mutation=='footer':name='word/footer2.xml';patch=lambda root:setattr(next(root.iter(model.W+'t')),'text','changed')
    elif mutation=='title':name='word/document.xml';patch=lambda root:root.find('.//'+model.W+'titlePg').getparent().remove(root.find('.//'+model.W+'titlePg'))
    elif mutation=='border':name='word/document.xml';patch=lambda root:root.find('.//'+model.W+'bottom').set(model.W+'sz','8')
    else:name='word/document.xml';patch=lambda root:setattr(next(root.iter(model.W+'t')),'text','changed')
    raw=changed_part(artifact.docx_bytes,name,patch)
    with pytest.raises(ValueError):verify(writer.SavedDocxLayoutArtifact(raw,artifact.source_map),args)


def test_v6_remains_exact_and_manual_cannot_activate_v7():
    args=fixture();oldcontext={k:v for k,v in args[4].items() if k!='source_evidence'}
    old=writer.build_unreviewed_docx(*args[:4],ordinary_context=oldcontext)
    assert old.source_map['writer_version']==writer.ORDINARY_PRESENTATION_WRITER_VERSION
    writer.validate_built_docx(old.docx_bytes,old.source_map,original_docx=args[0],snapshot=args[1],pages=args[2],decisions=args[3],require_review=False,ordinary_context=oldcontext)
    new=build(args)
    with pytest.raises(ValueError):writer.validate_built_docx(new.docx_bytes,new.source_map,original_docx=args[0],snapshot=args[1],pages=args[2],decisions=args[3],require_review=True,ordinary_context=args[4])


@pytest.mark.parametrize('lang',['EN','FR','AR'])
def test_owned_footer_word_count_and_fee_count_once(tmp_path,lang):
    from legalpdf_translate.ordinary_source_layout import owned_story_count_descriptor
    from legalpdf_translate.joblog_flow import count_words_from_owned_story_map,count_words_from_docx
    from legalpdf_translate.pricing import translation_fee_eur
    args=fixture(lang);artifact=build(args);path=tmp_path/'selected.docx';path.write_bytes(artifact.docx_bytes)
    descriptor=owned_story_count_descriptor(artifact.source_map)
    count=count_words_from_owned_story_map(path,descriptor)
    expected=sum(len(r["text"].split()) for index,r in enumerate(args[1]["paragraphs"]) if index!=2)
    assert count==expected
    assert count_words_from_docx(path)==sum(len(r["text"].split()) for r in args[1]["paragraphs"][:2])
    assert translation_fee_eur(count,.027)==translation_fee_eur(expected,.027)
    duplicate=deepcopy(descriptor);duplicate['paragraphs'].append(deepcopy(duplicate['paragraphs'][-1]))
    with pytest.raises(ValueError):count_words_from_owned_story_map(path,duplicate)


def multipage_fixture(lang='EN',long_body=False,selected=(2,7)):
    from docx.enum.text import WD_BREAK
    def edit(doc):
        _configure_styles(doc,lang);_add_page_number_footer(doc,lang)
        doc.paragraphs[2].runs[-1].add_break(WD_BREAK.PAGE)
        if lang=='AR':
            for p in doc.paragraphs:etree.SubElement(p._p.get_or_add_pPr(),model.W+'bidi').set(model.W+'val','1')
    sentence='يبقى الالتزام القانوني قائماً في القضية REF123. ' if lang=='AR' else 'Meaningful source body sentence. '
    body=sentence*300 if long_body else sentence.strip()
    text=[body,'Reply source two.','2 / 7',body,'Reply source seven.','7 / 7'] if lang!='AR' else [body,'يرجى الرد على المصدر الثاني.','الصفحة 2 / 7',body,'يرجى الرد على المصدر السابع.','الصفحة 7 / 7']
    raw=document_bytes(text,edit)
    snapshot,frames,decisions=reviewed(raw,lang)
    pages=[]
    for page in range(1,max(selected)+1):
        frame=deepcopy(frames[0]);frame['page_number']=page;pages.append(frame)
    ids=[row['id'] for row in snapshot['paragraphs']]
    groups=[[selected[0],ids[:3]],[selected[1],ids[3:]]]
    decisions['review'].update(document_reviewed=False,pages_reviewed=[],reviewer='',note='')
    decisions['bands']=[{'kind':'flow','groups':[{'paragraph_ids':[i],'panel':False}]} for i in ids]
    for index,row in enumerate(decisions['paragraphs']):
        page=selected[index//3];row.update(regions=[{'page_number':page,'bbox_px':[100,100,800,200]}])
        if index%3==2:row.update(role='source_folio',regions=[{'page_number':page,'bbox_px':[400,1300,600,1320]}])
    evidence={'version':'ordinary_source_evidence_v1','pages':[{'page_number':page,'findings':[],'layout':[{'kind':'source_footer','paragraph_ids':owned[-2:],'bbox':[.1,.9,.9,.98]}],'source_coverage_verified':False} for page,owned in groups]}
    return raw,snapshot,pages,decisions,{'selected_pages':list(selected),'page_groups':groups,'source_evidence':evidence}

@pytest.mark.parametrize('lang',['EN','FR','AR'])
def test_sparse_multipage_source_footers_first_page_only_on_spill(lang):
    args=multipage_fixture(lang,long_body=True);artifact=build(args);verify(artifact,args)
    plan=artifact.source_map['source_layout_plan']
    assert [r['source_page_number'] for r in plan['source_footers']]==[2,7]
    with ZipFile(BytesIO(artifact.docx_bytes)) as z:
        doc=etree.fromstring(z.read('word/document.xml'))
        assert len(list(doc.iter(model.W+'sectPr')))==2
        assert len(list(doc.iter(model.W+'titlePg')))==2
        for ordinal,folio in [(1,'2 / 7'),(2,'7 / 7')]:
            footer=etree.fromstring(z.read(f'word/footer{1+ordinal}.xml'))
            assert [n.text for n in footer.iter(model.W+'t')][-1].endswith(folio)
        assert not list(doc.iter(model.W+'br'))
        assert not list(etree.fromstring(z.read('word/footer1.xml')).iter(model.W+'t'))


def test_partial_footer_coverage_keeps_original_flow():
    args=multipage_fixture();args[4]['source_evidence']['pages'].pop()
    artifact=build(args);verify(artifact,args)
    assert not artifact.source_map['source_layout_plan']['source_footers']
    assert artifact.source_map['source_layout_plan']['qualifications']


@pytest.mark.parametrize('mutation,error', [('missing_title','source_section_first_page'),('disabled_title','source_section_first_page'),('continuous','source_section_next_page')])
def test_independent_section_guard_rejects_mutated_transform(monkeypatch,mutation,error):
    import legalpdf_translate.ordinary_source_layout as layer
    original=layer.transform
    def malicious(*args):
        artifact=original(*args)
        def patch(root):
            title=root.find('.//'+model.W+'titlePg')
            if mutation=='missing_title':title.getparent().remove(title)
            elif mutation=='disabled_title':title.set(model.W+'val','0')
            else:root.find('.//'+model.W+'sectPr/'+model.W+'type').set(model.W+'val','continuous')
        raw=changed_part(artifact.docx_bytes,'word/document.xml',patch)
        mapping=deepcopy(artifact.source_map);mapping['docx_sha256']=model._sha(raw)
        return writer.SavedDocxLayoutArtifact(raw,mapping)
    monkeypatch.setattr(layer,'transform',malicious)
    with pytest.raises(ValueError,match=error):build(fixture())


def test_owned_footer_edit_recounts_current_verified_bytes(tmp_path):
    from legalpdf_translate.ordinary_source_layout import owned_story_count_descriptor
    from legalpdf_translate.joblog_flow import count_words_from_owned_story_map
    args=fixture();artifact=build(args)
    raw=changed_part(artifact.docx_bytes,'word/footer2.xml',lambda root:setattr(next(root.iter(model.W+'t')),'text','Expanded reply instruction.'))
    path=tmp_path/'approved-correction.docx';path.write_bytes(raw)
    descriptor=owned_story_count_descriptor(artifact.source_map)
    with pytest.raises(ValueError,match='identity_changed'):count_words_from_owned_story_map(path,descriptor)
    descriptor['docx_sha256']=model._sha(raw) # simulated server-approved text revision, not client admission
    assert count_words_from_owned_story_map(path,descriptor)==10


def test_independent_rule_guard_rejects_mutated_transform(monkeypatch):
    import legalpdf_translate.ordinary_source_layout as layer
    original=layer.transform
    def malicious(*args):
        artifact=original(*args)
        raw=changed_part(artifact.docx_bytes,'word/document.xml',lambda root:root.find('.//'+model.W+'pBdr/'+model.W+'bottom').set(model.W+'sz','12'))
        mapping=deepcopy(artifact.source_map);mapping['docx_sha256']=model._sha(raw)
        return writer.SavedDocxLayoutArtifact(raw,mapping)
    monkeypatch.setattr(layer,'transform',malicious)
    with pytest.raises(ValueError,match='source_rule_delta_changed'):build(fixture())


def test_v7_word_formatting_adoption_retains_sections_and_footer_ownership():
    from legalpdf_translate.ordinary_edited_revision import qualify_edited_docx
    args=fixture('AR');artifact=build(args)
    def format_only(root):
        for paragraph in root.iter(model.W+'p'):
            props=paragraph.find(model.W+'pPr')
            if props is None:props=etree.Element(model.W+'pPr');paragraph.insert(0,props)
            for name in ('bidi','jc'):
                for node in list(props.findall(model.W+name)):props.remove(node)
            etree.SubElement(props,model.W+'bidi').set(model.W+'val','1')
            etree.SubElement(props,model.W+'jc').set(model.W+'val','right')
    edited=changed_part(artifact.docx_bytes,'word/document.xml',format_only)
    qualify_edited_docx(artifact.docx_bytes,edited,source_layout_map=artifact.source_map)
    bad=changed_part(edited,'word/document.xml',lambda root:root.find('.//'+model.W+'titlePg').getparent().remove(root.find('.//'+model.W+'titlePg')))
    with pytest.raises(ValueError,match='section_ownership_changed'):qualify_edited_docx(artifact.docx_bytes,bad,source_layout_map=artifact.source_map)


def test_partition_children_preserve_exact_text_with_source_footer_sections():
    args=fixture();args[3]['version']=model.PARTITION_DECISIONS_VERSION
    identifier=args[1]['paragraphs'][1]['id'];args[3]['paragraph_partitions']=[{'paragraph_id':identifier,'offsets':[5]}]
    artifact=build(args);verify(artifact,args)
    row=next(r for r in artifact.source_map['paragraphs'] if r['paragraph_id']==identifier)
    assert len(row['parts'])==2
    assert row['part_uri']=='word/document.xml'
    with ZipFile(BytesIO(artifact.docx_bytes)) as z:
        root=etree.fromstring(z.read(row['part_uri']))
        text=''.join(''.join(n.text or '' for n in root.xpath(part['location'],namespaces={'w':model.W[1:-1]})[0].iter(model.W+'t')) for part in row['parts'])
        assert text=='Body obligation.'
