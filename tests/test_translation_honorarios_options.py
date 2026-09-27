"""Fictional per-document fee wording; no native exporter, profile or Gmail access."""
from datetime import date

import pytest
from docx import Document

from legalpdf_translate.honorarios_docx import (
    DEFAULT_TRANSLATOR_DECLARATION, build_honorarios_draft,
    build_honorarios_paragraph_texts, generate_honorarios_docx,
    validate_translation_honorarios_options,
)
from legalpdf_translate.user_profile import UserProfile


def _draft(**kwargs):
    return build_honorarios_draft(
        case_number="FICTIONAL-17", word_count=1501,
        case_entity="Procuradoria de Exemplo", case_city="Exemplo",
        profile=UserProfile(id="fiction", first_name="Ana", last_name="Exemplo",
                            postal_address="Rua Exemplo 1", iban="TEST-IBAN",
                            iva_text="23%", irs_text="Sem retenção"),
        today=date(2026, 9, 27), **kwargs,
    )


def test_default_uses_institution_and_contains_no_declaration_or_signature():
    draft = _draft()
    paragraphs = build_honorarios_paragraph_texts(draft)
    assert paragraphs[2][0] == "À Procuradoria de Exemplo"
    assert len(paragraphs) == 19
    assert not draft.include_translator_declaration
    assert draft.translator_declaration_text == ""
    assert not any("Assinatura:" in text or "Comprometo-me" in text for text, _ in paragraphs)


@pytest.mark.parametrize("custom", ["", "Declaração revista para o presente processo."])
def test_explicit_declaration_preserves_identity_and_only_adds_blank_signature(tmp_path, custom):
    draft = _draft(include_translator_declaration=True, translator_declaration_text=custom)
    path = generate_honorarios_docx(draft, tmp_path / "unsigned.docx")
    texts = [p.text for p in Document(path).paragraphs]
    expected = custom or DEFAULT_TRANSLATOR_DECLARATION
    assert texts.count(expected) == 1
    assert texts.index(expected) < texts.index("Espera deferimento,")
    assert texts[-1] == "Assinatura: ______________________________"
    assert "Nome: Ana Exemplo" in texts
    assert "Morada: Rua Exemplo 1" in texts
    assert "O Pagamento deverá ser efetuado para o seguinte IBAN: TEST-IBAN" in texts
    assert "Este serviço inclui a taxa IVA de 23% e não tem retenção de IRS." in texts
    assert "O documento traduzido contém 1501 palavras." in texts
    assert not Document(path).inline_shapes


def test_custom_recipient_exact_text_and_unchecked_declaration_are_per_document():
    custom = "À entidade indicada\nSecção Exemplo"
    draft = _draft(recipient_block=custom, translator_declaration_text="Ignored until explicitly selected")
    paragraphs = build_honorarios_paragraph_texts(draft)
    assert paragraphs[2][0] == custom
    assert draft.translator_declaration_text == ""
    assert _draft().recipient_block == ""


@pytest.mark.parametrize("options", [
    {"include_translator_declaration": "false"},
    {"include_translator_declaration": 1},
    {"recipient_block": []}, {"recipient_block": "x" * 1001},
    {"translator_declaration_text": "x" * 4001},
    {"recipient_block": "bad\x00text"}, {"translator_declaration_text": "bad\x07text"},
])
def test_invalid_options_are_rejected_before_export(options):
    with pytest.raises(ValueError):
        validate_translation_honorarios_options(**options)


def test_invalid_finalization_options_never_trigger_native_preflight(tmp_path, monkeypatch):
    from legalpdf_translate.gmail_browser_service import GmailBrowserSessionManager
    manager = GmailBrowserSessionManager()
    def forbidden(**kwargs):
        pytest.fail("Invalid wording reached native preflight")
    monkeypatch.setattr(manager, "preflight_batch_finalization", forbidden)
    with pytest.raises(ValueError):
        manager.finalize_batch(runtime_mode="shadow", workspace_id="fiction",
                               settings_path=tmp_path / "settings.json", output_filename=None,
                               profile_id=None, include_translator_declaration="true")
