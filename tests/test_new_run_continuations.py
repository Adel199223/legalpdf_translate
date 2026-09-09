"""Continuation assembly is source-derived, opt-in and never rewrites page commits."""

from dataclasses import replace
import hashlib
import json

import fitz
import pytest
from docx import Document

from legalpdf_translate import workflow as workflow_module
from legalpdf_translate.document_structure import PageStructure, StructureBlock, text_sha256
from legalpdf_translate.docx_writer import assemble_docx
from legalpdf_translate.openai_client import ApiCallResult
from legalpdf_translate.structured_artifacts import publish_structured_page
from legalpdf_translate.types import ImageMode, OcrMode, RunConfig, TargetLang
from legalpdf_translate.workflow import TranslationWorkflow


SOURCE = (
    "O arguido deve comparecer perante o tribunal e cumprir as obrigacoes impostas pela decisao, "
    "mantendo durante todo o periodo a condicao de",
    "se apresentar nos termos da notificacao e informar o tribunal sobre qualquer alteracao relevante "
    "para o cumprimento destas obrigacoes.",
)
TARGET = (
    "The defendant must attend court and comply with the obligations imposed by the decision, "
    "maintaining throughout the period the requirement to",
    "appear as stated in the notice and inform the court about any change relevant to compliance "
    "with those obligations.",
)


class FakeClient:
    def __init__(self):
        self.calls = []

    def create_page_response(self, **kwargs):
        self.calls.append(kwargs)
        payload, _ = json.JSONDecoder().raw_decode(kwargs["prompt_text"])
        blocks = [{"id": row["id"], "text": TARGET[int(row["id"][1:5]) - 1]}
                  for row in payload["blocks"]]
        return ApiCallResult(raw_output=json.dumps({"blocks": blocks}),
            usage={"input_tokens": 12, "output_tokens": 14}, response_id="offline-continuation",
            response_status="completed")


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.delenv("LEGALPDF_TRANSLATION_PROTOCOL", raising=False)
    monkeypatch.setattr(workflow_module, "load_gui_settings", lambda: {})
    monkeypatch.setattr(workflow_module, "load_environment", lambda: None)
    monkeypatch.setattr(workflow_module, "run_translation_auth_test",
                        lambda *a, **k: pytest.fail("Unexpected authentication call"))


def configuration(tmp_path, *, page_breaks=False):
    source = tmp_path / "split-source.pdf"
    with fitz.open() as document:
        for index, text in enumerate(SOURCE):
            page = document.new_page()
            y = 650 if index == 0 else 80
            assert page.insert_textbox(fitz.Rect(48, y, 540, y + 110), text, fontsize=11) >= 0
        document.save(source)
    output = tmp_path / "output"
    output.mkdir()
    return RunConfig(source, output, TargetLang.EN, image_mode=ImageMode.OFF,
                     ocr_mode=OcrMode.OFF, workers=1, resume=False, page_breaks=page_breaks)


def page_hashes(pages):
    # The existing workflow may regenerate its separate layout cache. Only the
    # source/target/TXT/commit unit is immutable translation evidence.
    return {path.name: hashlib.sha256(path.read_bytes()).hexdigest()
            for path in pages.iterdir() if path.is_file() and path.name.endswith(
                (".txt", ".source_structure.json", ".structure.json", ".commit.json"))}


def mapping(output):
    return json.loads(output.with_suffix(".source_map.json").read_text(encoding="utf-8"))["pages"]


def assert_joined_once(output):
    paragraphs = [p.text for p in Document(output).paragraphs if p.text.strip()]
    assert paragraphs == [" ".join(TARGET)]
    pages = mapping(output)
    assert [page["source_page_number"] for page in pages] == [1, 2]
    first, second = pages[0]["blocks"][0], pages[1]["blocks"][0]
    assert second["joined_to_block_id"] == first["block_id"]
    assert second["location"] == first["location"]
    for text in TARGET:
        assert paragraphs[0].count(text) == 1


def assert_no_join(output):
    assert all("joined_to_block_id" not in block
               for page in mapping(output) for block in page["blocks"])


def test_fresh_two_page_workflow_joins_once_and_preserves_all_commits_on_resume_and_rebuild(tmp_path):
    config = configuration(tmp_path)
    client = FakeClient()
    workflow = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2")
    result = workflow.run(config)
    assert result.success, result.error
    assert len(client.calls) == 2
    assert_joined_once(result.output_docx)
    pages_dir = result.run_dir / "pages"
    before = page_hashes(pages_dir)
    for path in pages_dir.glob("*.structure.json"):
        data = json.loads(path.read_text(encoding="utf-8"))
        assert not data["continuation_from_previous"] and not data["continuation_to_next"]
        assert all(block["continuation_of"] is None for block in data["blocks"])
    original_docx = result.output_docx.read_bytes()
    resumed = workflow.run(replace(config, resume=True))
    assert resumed.success, resumed.error
    assert_joined_once(resumed.output_docx)
    rebuilt = workflow.rebuild_docx(config)
    assert_joined_once(rebuilt)
    assert len(client.calls) == 2
    assert page_hashes(pages_dir) == before
    assert result.output_docx.read_bytes() == original_docx


def test_real_page_matching_run_and_rebuild_do_not_join(tmp_path):
    config = configuration(tmp_path, page_breaks=True)
    client = FakeClient()
    workflow = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2")
    result = workflow.run(config)
    assert result.success, result.error
    before = page_hashes(result.run_dir / "pages")
    assert_no_join(result.output_docx)
    assert [p.text for p in Document(result.output_docx).paragraphs if p.text.strip()] == list(TARGET)
    assert_no_join(workflow.rebuild_docx(config))
    assert page_hashes(result.run_dir / "pages") == before
    assert len(client.calls) == 2


def test_real_manual_text_rebuild_does_not_join_or_rebind_committed_evidence(tmp_path):
    config = configuration(tmp_path)
    client = FakeClient()
    workflow = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2")
    result = workflow.run(config)
    assert result.success, result.error
    pages = result.run_dir / "pages"
    edited = pages / "page_0002.txt"
    edited.write_text(TARGET[1] + " Human clarification.", encoding="utf-8")
    before = page_hashes(pages)
    output = workflow.rebuild_docx(config)
    assert_no_join(output)
    assert "Human clarification." in " ".join(p.text for p in Document(output).paragraphs)
    assert mapping(output)[1]["structure_status"] == "invalid_sidecar_txt_fallback"
    assert page_hashes(pages) == before
    assert len(client.calls) == 2


def test_cancelled_real_run_does_not_join_uncompleted_neighbor(tmp_path):
    config = configuration(tmp_path)
    client = FakeClient()
    workflow = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2")

    def cancel(done, total, status):
        if done == 1 and "finished" in status:
            workflow.cancel()

    workflow._progress_callback = cancel
    result = workflow.run(config)
    assert not result.success and result.error == "cancelled"
    assert len(client.calls) == 1
    assert_no_join(result.partial_docx)
    assert [page["source_page_number"] for page in mapping(result.partial_docx)] == [1]
    assert TARGET[1] not in " ".join(p.text for p in Document(result.partial_docx).paragraphs)


def saved_pages(tmp_path, *, change=None, furniture=False, contact=False):
    pages = tmp_path / "pages"
    sources = []
    for index, text in enumerate(SOURCE):
        number = index + 1
        y = 700 if index == 0 else 100
        blocks = [StructureBlock(f"p{number:04d}_b0001", text, bbox=(48, y, 540, y + 35))]
        if furniture:
            blocks.insert(0, StructureBlock(f"p{number:04d}_b0002", "Tribunal Judicial da Comarca",
                                           role="header", bbox=(200, 40, 400, 62)))
        if contact:
            blocks.extend([
                StructureBlock(f"p{number:04d}_b0003", "Largo da Justica, 1", role="footer",
                               bbox=(160, 800, 410, 810), alignment="center"),
                StructureBlock(f"p{number:04d}_b0004", "Telef: 210000000 - E-mail: court@example.invalid",
                               role="footer", bbox=(100, 812, 500, 822), alignment="center"),
            ])
        source = PageStructure(number, "", blocks, source_file_sha256="f" * 64,
                               provenance="digital_pdf")
        sources.append(source)
    if change:
        change(sources)
    for index, source in enumerate(sources):
        source.source_sha256 = source.source_text_sha256 = text_sha256(source.text)
        target = source.to_dict()
        for block in target["blocks"]:
            if block["role"] == "header":
                block["text"] = "Judicial Court"
            elif block["role"] != "footer":
                block["text"] = TARGET[index]
        text = "\n".join(block["text"] for block in target["blocks"])
        target["translation_sha256"] = text_sha256(text)
        publish_structured_page(pages, source_structure=source, translated_structure=target,
            translated_text=text, protocol_identity={"protocol": "legal_blocks_v2", "fingerprint": "a" * 64},
            page_fingerprint="b" * 64)
    return pages


def test_opt_in_requires_explicit_selected_pages_and_default_does_not_derive(tmp_path):
    pages = saved_pages(tmp_path)
    before = page_hashes(pages)
    with pytest.raises(ValueError, match="completed-page selection"):
        assemble_docx(pages, tmp_path / "rejected.docx", lang=TargetLang.EN, page_breaks=False,
                      derive_source_continuations=True)
    assert not (tmp_path / "rejected.docx").exists()
    output = assemble_docx(pages, tmp_path / "legacy-default.docx", lang=TargetLang.EN, page_breaks=False)
    assert_no_join(output)
    assert page_hashes(pages) == before


@pytest.mark.parametrize("selection", [[1], [2], [1, 3]])
def test_partial_selection_never_bridges_a_missing_page(tmp_path, selection):
    def change(sources):
        if selection == [1, 3]:
            sources[1].page_number = 3
            sources[1].blocks[0].id = "p0003_b0001"
    pages = saved_pages(tmp_path, change=change)
    before = page_hashes(pages)
    output = assemble_docx(pages, tmp_path / "partial.docx", lang=TargetLang.EN, page_breaks=False,
                          page_numbers=selection, partial_output=True, derive_source_continuations=True)
    assert_no_join(output)
    assert [page["source_page_number"] for page in mapping(output)] == selection
    assert page_hashes(pages) == before


@pytest.mark.parametrize("guard", ["uncertain_page", "uncertain_block", "new_document", "new_block_document",
                                   "no_geometry", "other_source", "stale_flags", "missing_source", "corrupt_source",
                                   "manual_text", "missing_target", "corrupt_target", "page_matching"])
def test_derivation_cannot_override_negative_evidence(tmp_path, guard):
    def change(sources):
        if guard == "uncertain_page": sources[1].uncertain = True
        if guard == "uncertain_block": sources[1].blocks[0].uncertain = True
        if guard == "new_document": sources[1].document_start = True
        if guard == "new_block_document": sources[1].blocks[0].document_start = True
        if guard == "no_geometry": sources[0].blocks[0].bbox = None
        if guard == "other_source": sources[1].source_file_sha256 = "d" * 64
        if guard == "stale_flags":
            sources[0].blocks[0].bbox = None
            sources[0].continuation_to_next = sources[1].continuation_from_previous = True
            sources[1].blocks[0].continuation_of = sources[0].blocks[0].id
            sources[1].metadata.update(continuation_evidence="adjacent_source_fragment", continuation_bridge={"stale": True})
    pages = saved_pages(tmp_path, change=change)
    if guard == "missing_source": (pages / "page_0002.source_structure.json").unlink()
    if guard == "corrupt_source": (pages / "page_0002.source_structure.json").write_text("{}", encoding="utf-8")
    if guard == "manual_text": (pages / "page_0002.txt").write_text(TARGET[1] + " Human edit.", encoding="utf-8")
    if guard == "missing_target": (pages / "page_0002.structure.json").unlink()
    if guard == "corrupt_target": (pages / "page_0002.structure.json").write_text("{}", encoding="utf-8")
    before = page_hashes(pages)
    output = assemble_docx(pages, tmp_path / "negative.docx", lang=TargetLang.EN,
                          page_breaks=guard == "page_matching", page_numbers=[1, 2], derive_source_continuations=True)
    assert_no_join(output)
    assert page_hashes(pages) == before


def test_repeated_header_bridge_uses_source_proof_without_rewriting_artifacts(tmp_path):
    pages = saved_pages(tmp_path, furniture=True)
    before = page_hashes(pages)
    output = assemble_docx(pages, tmp_path / "furnished.docx", lang=TargetLang.EN, page_breaks=False,
                          page_numbers=[1, 2], derive_source_continuations=True)
    page_map = mapping(output)
    body = page_map[1]["blocks"][1]
    assert body["joined_to_block_id"] == "p0001_b0001"
    assert body["continuation_evidence"] == "revalidated_identical_source_furniture"
    assert len([block for page in page_map for block in page["blocks"]]) == 4
    assert " ".join(TARGET) in [paragraph.text for paragraph in Document(output).paragraphs]
    assert page_hashes(pages) == before


def test_contact_furniture_consolidation_composes_with_derived_split_body_join(tmp_path):
    pages = saved_pages(tmp_path, furniture=True, contact=True)
    before = page_hashes(pages)
    output = assemble_docx(pages, tmp_path / "contact-and-continuation.docx", lang=TargetLang.EN,
                          page_breaks=False, page_numbers=[1, 2], derive_source_continuations=True)
    report = json.loads(output.with_suffix(".source_map.json").read_text(encoding="utf-8"))
    document = Document(output)
    sections = report["section_furniture"]["sections"]
    assert len(sections) == 1 and sections[0]["consolidated"]
    assert len(document.sections) == 1
    section = document.sections[0]
    assert [paragraph.text for paragraph in section.header.paragraphs] == ["Judicial Court"]
    assert "court@example.invalid" in section.footer._element.xml
    assert " PAGE " in section.footer._element.xml
    assert not section.header.is_linked_to_previous and not section.footer.is_linked_to_previous
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    assert paragraphs == [" ".join(TARGET)]
    assert all(paragraphs[0].count(fragment) == 1 for fragment in TARGET)
    rows = [block for page in report["pages"] for block in page["blocks"]]
    assert len(rows) == len({row["block_id"] for row in rows}) == 8
    for page in report["pages"]:
        assert page["section_furniture_adopted"]
        for block in page["blocks"]:
            if block["role"] in {"header", "footer"}:
                assert block["location"]["kind"] == f"section_{block['role']}"
    first, second = report["pages"][0]["blocks"][1], report["pages"][1]["blocks"][1]
    assert second["joined_to_block_id"] == first["block_id"]
    assert second["location"] == first["location"]
    assert second["continuation_evidence"] == "revalidated_identical_source_furniture"
    assert page_hashes(pages) == before
