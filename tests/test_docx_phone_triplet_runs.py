"""Fictional phone-shaped triplet grouping; OOXML is not native visual acceptance."""

import hashlib
import socket
import subprocess

from docx import Document
from docx.oxml.ns import qn
import pytest

from legalpdf_translate import docx_writer as writer
from legalpdf_translate.types import TargetLang

PHONE = "321 654 987"


def visible(text):
    return writer.sanitize_bidi_controls(writer.unwrap_internal_placeholders(text))


def protected(text, isolated):
    return f"\u2066[[{text}]]\u2069" if isolated else f"[[{text}]]"


def assert_unchanged(text, runs, strip):
    expected = writer.unwrap_internal_placeholders(text)
    if strip:
        expected = writer.sanitize_bidi_controls(expected)
    assert "".join(chunk for _, chunk in runs) == expected


@pytest.mark.parametrize("shape", ["plain", "whole", "split", "mixed"])
@pytest.mark.parametrize("isolated", [False, True])
@pytest.mark.parametrize("strip", [False, True])
def test_closed_phone_triplet_is_one_ltr_run(shape, isolated, strip):
    p = lambda value: protected(value, isolated)
    token = {"plain": PHONE, "whole": p(PHONE), "split": " ".join(p(x) for x in PHONE.split()),
             "mixed": p("321") + " 654 " + p("987")}[shape]
    text = f"الهاتف {token} ثم النص."
    runs, mixed = writer._segment_rtl_placeholder_aware_runs(text, strip_bidi_controls=strip)
    assert_unchanged(text, runs, strip)
    assert mixed
    assert ("ltr", PHONE) in [(kind, visible(chunk)) for kind, chunk in runs]
    assert any(kind == "rtl" and "الهاتف" in chunk for kind, chunk in runs)
    assert any(kind == "rtl" and "ثم النص" in chunk for kind, chunk in runs)


@pytest.mark.parametrize("token", [
    "[[09]] [[30]]", "[[321]] [[654]]", "[[321]] [[654]] [[987]] [[111]]",
    "[[12]] [[321]] [[654]] [[987]]", "[[321]] [[654]] [[987]] [[12]]",
    "[[12]]  [[321]] [[654]] [[987]]", "[[321]] [[654]] [[987]]  [[12]]",
    "[[32]] [[654]] [[987]]", "[[321]] [[65]] [[987]]", "[[321]] [[654]] [[98]]",
    "[[3210]] [[654]] [[987]]", "[[321]] [[6540]] [[987]]", "[[321]] [[654]] [[9870]]",
    "[[321]]\t[[654]] [[987]]", "[[321]]\u00a0[[654]] [[987]]",
    "[[321]]  [[654]] [[987]]", "[[321]]\n[[654]] [[987]]",
    "A[[321]] [[654]] [[987]]", "[[321]] [[654]] [[987]]B",
    "أ[[321]] [[654]] [[987]]", "[[321]] [[654]] [[987]]أ",
    "A\u0301[[321]] [[654]] [[987]]", "[[321]] [[654]] [[987]]\u0301",
    "A\u0301.[[321]] [[654]] [[987]]", "[[321]] [[654]] [[987]].\u0301A",
    "+[[321]] [[654]] [[987]]", "[[321]] [[654]] [[987]]-9",
    "ABC-[[321]] [[654]] [[987]]", "[[321]] [[654]] [[987]].5",
    "[[321]]/[[654]] [[987]]", "[[٣٢١]] [[٦٥٤]] [[٩٨٧]]",
    "[[[321]] [[654]] [[987]]]", "[[321]] [[654]] [[987]",
    "[[Rua Exemplo]] [[321]] [[Vila Exemplo]]",
])
@pytest.mark.parametrize("strip", [False, True])
def test_other_shapes_and_attached_values_keep_original_space_directions(token, strip):
    text = f"النص {token} نهاية"
    # Existing line segmentation is the unchanged baseline for these non-phone forms.
    baseline, _ = writer._segment_rtl_placeholder_aware_line(text, strip_bidi_controls=strip)
    baseline = writer._keep_numeric_slash_separators_ltr(baseline)
    runs, _ = writer._segment_rtl_placeholder_aware_runs(text, strip_bidi_controls=strip)
    assert_unchanged(text, runs, strip)
    assert [(kind, char) for kind, chunk in runs for char in chunk] == [
        (kind, char) for kind, chunk in baseline for char in chunk]


@pytest.mark.parametrize("opener,closer", [
    ("\u2067", "\u2069"), ("\u2068", "\u2069"), ("\u202e", "\u202c"),
    ("\u202b", "\u202c"), ("\u2066", ""),
])
@pytest.mark.parametrize("separator", [" ", "\n"])
def test_conflicting_or_unclosed_retained_scope_does_not_group(opener, closer, separator):
    token = " ".join(protected(x, True) for x in PHONE.split())
    text = opener + "نص" + separator + "الهاتف " + token + " نهاية" + closer
    runs, _ = writer._segment_rtl_placeholder_aware_runs(text, strip_bidi_controls=False)
    assert_unchanged(text, runs, False)
    assert not any(kind == "ltr" and PHONE in visible(chunk) for kind, chunk in runs)
    stripped, _ = writer._segment_rtl_placeholder_aware_runs(text, strip_bidi_controls=True)
    assert any(kind == "ltr" and visible(chunk) == PHONE for kind, chunk in stripped)


def test_unmatched_control_and_explicit_line_boundary_are_not_consumed():
    token = "[[321]] [[654]] [[987]]"
    for text in ("النص " + token + "\u2069 نهاية", "النص [[321]] [[654]]\n[[987]] نهاية"):
        runs, _ = writer._segment_rtl_placeholder_aware_runs(text, strip_bidi_controls=False)
        assert_unchanged(text, runs, False)
        assert not any(kind == "ltr" and PHONE in visible(chunk) for kind, chunk in runs)


def test_only_internal_phone_spaces_change_and_explicit_lines_remain_separate():
    segments = [("rtl", "النص ("), ("ltr", "321"), ("rtl", " "),
                ("ltr", "654"), ("rtl", " "), ("ltr", "987"), ("rtl", ") نهاية\n"),
                ("ltr", "321"), ("rtl", " \n"), ("ltr", "654"), ("rtl", " "), ("ltr", "987")]
    runs = writer._keep_phone_triplet_spaces_ltr(segments)
    before = [(kind, char) for kind, chunk in segments for char in chunk]
    after = [(kind, char) for kind, chunk in runs for char in chunk]
    assert [(a, b) for a, b in zip(before, after) if a != b] == [
        (("rtl", " "), ("ltr", " ")), (("rtl", " "), ("ltr", " "))]
    assert "".join(chunk for _, chunk in runs) == "".join(chunk for _, chunk in segments)
    assert sum(kind == "ltr" and chunk == PHONE for kind, chunk in runs) == 1


@pytest.mark.parametrize("control", ["\u061c", "\u200f", "\u202e", "\u2067", "\ufeff"])
def test_conflicting_retained_control_inside_phone_keeps_spaces(control):
    text = "الهاتف [[321]] " + control + "[[654]] [[987]] نهاية"
    runs, _ = writer._segment_rtl_placeholder_aware_runs(text, strip_bidi_controls=False)
    assert_unchanged(text, runs, False)
    assert not any(kind == "ltr" and PHONE in visible(chunk) for kind, chunk in runs)


@pytest.mark.parametrize("strip", [False, True])
def test_saved_phone_docx_has_one_ltr_number_and_preserves_source(tmp_path, monkeypatch, strip):
    def forbidden(*args, **kwargs):
        pytest.fail("Offline phone fixture must not start native or network operations")
    monkeypatch.setattr(subprocess, "Popen", forbidden)
    monkeypatch.setattr(socket, "socket", forbidden)
    pages = tmp_path / "pages"
    pages.mkdir()
    token = " ".join(protected(x, True) for x in PHONE.split())
    text = f"الهاتف {token} مكالمة إلى الشبكة الثابتة."
    source = pages / "page_0001.txt"
    source.write_text(text, encoding="utf-8")
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    output = writer.assemble_docx(pages, tmp_path / "phone.docx", lang=TargetLang.AR,
                                  page_breaks=False, strip_bidi_controls=strip)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    paragraph = next(p for p in Document(output).paragraphs if "الهاتف" in p.text)
    assert visible(paragraph.text) == visible(text)
    numbers = [r for r in paragraph.runs if PHONE in visible(r.text)]
    assert len(numbers) == 1
    assert visible(numbers[0].text) == PHONE
    assert numbers[0]._r.find(f"{qn('w:rPr')}/{qn('w:rtl')}").get(qn("w:val")) == "0"
    for run in paragraph.runs:
        if any("\u0620" <= char <= "\u064a" for char in run.text):
            assert run._r.find(f"{qn('w:rPr')}/{qn('w:rtl')}").get(qn("w:val")) == "1"
    expected = writer.unwrap_internal_placeholders(text)
    if strip:
        expected = writer.sanitize_bidi_controls(expected)
    assert paragraph.text.replace("\u200e", "") == expected
