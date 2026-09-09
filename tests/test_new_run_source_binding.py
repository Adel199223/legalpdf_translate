"""Real workflow source-winner regressions; providers/OCR are synthetic only."""
from __future__ import annotations

from dataclasses import replace
import json
from pathlib import Path
import subprocess

import fitz
import pytest

from legalpdf_translate import workflow as wf
from legalpdf_translate.checkpoint import load_run_state
from legalpdf_translate.document_structure import structure_from_text
from legalpdf_translate.ocr_engine import OcrResult
from legalpdf_translate.types import OcrMode
from legalpdf_translate.workflow import TranslationWorkflow

from .test_new_translation_blocks import FakeClient, SOURCE, configuration, offline


@pytest.fixture(autouse=True)
def no_native(monkeypatch, offline):
    assert Path(wf.__file__).resolve().is_relative_to(Path(__file__).resolve().parents[1] / "src")
    monkeypatch.setattr(subprocess, "Popen", lambda *a, **k: pytest.fail("Native process forbidden"))


def read_source(result, number=1):
    return json.loads((result.run_dir / "pages" / f"page_{number:04d}.source_structure.json").read_text("utf-8"))


def fake_ocr(monkeypatch, result):
    calls = []
    def recognize(path, number, **kwargs):
        calls.append((number, kwargs))
        return result
    monkeypatch.setattr(TranslationWorkflow, "_resolve_ocr_engine_for_reason", lambda *a, **k: (object(), True))
    monkeypatch.setattr(wf, "ocr_pdf_page_text", recognize)
    return calls


def test_stale_geometry_cannot_follow_different_selected_ocr_text(tmp_path, monkeypatch):
    config = replace(configuration(tmp_path), ocr_mode=OcrMode.ALWAYS)
    stale = structure_from_text("Old source that must not be used.", page_number=1).to_dict()
    stale["blocks"][0]["bbox"] = [20, 20, 500, 100]
    stale["provenance"], stale["uncertain"] = "local_ocr_tsv", False
    winner = SOURCE + " O prazo deve ser cumprido sem demora."
    calls = fake_ocr(monkeypatch, OcrResult(winner, "local", None, len(winner), .95, structure=stale))
    client = FakeClient()
    result = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2").run(config)
    assert result.success, result.error
    source = read_source(result)
    assert source["provenance"] == "text_fallback" and source["uncertain"]
    assert all(block["bbox"] is None for block in source["blocks"])
    assert " ".join(block["text"] for block in source["blocks"]).split() == winner.split()
    assert len(calls) == len(client.calls) == 1


def test_api_ocr_text_remains_complete_but_never_inherits_native_pdf_boxes(tmp_path, monkeypatch):
    config = replace(configuration(tmp_path), ocr_mode=OcrMode.ALWAYS)
    winner = SOURCE + " A pessoa notificada conserva o direito de resposta."
    calls = fake_ocr(monkeypatch, OcrResult(winner, "api", None, len(winner), .99))
    result = TranslationWorkflow(client=FakeClient(), translation_protocol="legal_blocks_v2").run(config)
    assert result.success, result.error
    source = read_source(result)
    assert source["provenance"] == "text_fallback" and source["uncertain"]
    assert all(block["bbox"] is None for block in source["blocks"])
    assert " ".join(block["text"] for block in source["blocks"]).split() == winner.split()
    state = load_run_state(result.run_dir / "run_state.json")
    assert state.pages["1"]["layout_review_required"] is True
    assert state.pages["1"]["fidelity_review_status"] == "not_evaluated"
    assert len(calls) == 1


def test_selected_middle_page_never_extracts_ocr_or_translates_unselected_neighbors(tmp_path, monkeypatch):
    config = replace(configuration(tmp_path, pages=3), start_page=2, end_page=2)
    original = wf.extract_ordered_page_text
    extracted = []
    def extract(path, index, **kwargs):
        extracted.append(index)
        assert index == 1, "Unselected neighbor extracted"
        return original(path, index, **kwargs)
    monkeypatch.setattr(wf, "extract_ordered_page_text", extract)
    monkeypatch.setattr(wf, "ocr_pdf_page_text", lambda *a, **k: pytest.fail("Neighbor OCR forbidden"))
    client = FakeClient()
    result = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2").run(config)
    assert result.success and len(client.calls) == 1 and extracted
    payload, _ = json.JSONDecoder().raw_decode(client.calls[0]["prompt_text"])
    assert payload["page"] == 2 and all(row["id"].startswith("p0002_") for row in payload["blocks"])
    assert "context_only_not_to_translate" not in payload
    assert not list((result.run_dir / "pages").glob("page_0001.*"))
    assert not list((result.run_dir / "pages").glob("page_0003.*"))


def test_source_mutation_after_response_blocks_publication_and_retains_usage(tmp_path):
    config = configuration(tmp_path)
    original = config.pdf_path.read_bytes()
    def mutate(response, payload):
        config.pdf_path.write_bytes(original + b"\n% synthetic source change\n")
        return response
    client = FakeClient(transform=mutate)
    result = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2").run(config)
    assert not result.success and len(client.calls) == 1
    assert not list((result.run_dir / "pages").glob("page_*.commit.json"))
    assert not list((result.run_dir / "pages").glob("page_*.txt"))
    state = load_run_state(result.run_dir / "run_state.json")
    assert state.done_count == 0 and state.pages["1"]["input_tokens"] == 12


def test_real_digital_table_and_list_source_blocks_remain_complete_and_editable(tmp_path):
    config = configuration(tmp_path)
    document = fitz.open()
    page = document.new_page()
    page.insert_textbox(fitz.Rect(40, 60, 550, 190), SOURCE, fontsize=11)
    page.insert_text((40, 220), "a) Comparecer no tribunal", fontsize=11)
    page.insert_text((40, 245), "b) Apresentar os documentos", fontsize=11)
    for x in (40, 240, 440):
        page.draw_line((x, 300), (x, 400))
    for y in (300, 350, 400):
        page.draw_line((40, y), (440, y))
    for point, text in (((50, 325), "Alpha"), ((250, 325), "Beta"), ((50, 375), "Gamma")):
        page.insert_text(point, text, fontsize=11)
    config.pdf_path.write_bytes(document.tobytes())
    document.close()
    client = FakeClient()
    def target(response, output):
        request, _ = json.JSONDecoder().raw_decode(client.calls[-1]["prompt_text"])
        translations = {"a) Comparecer no tribunal": "a) Attend court", "b) Apresentar os documentos": "b) Provide the documents",
                        "Alpha": "Alpha", "Beta": "Beta", "Gamma": "Gamma", "": ""}
        for before, after in zip(request["blocks"], output["blocks"]):
            after["text"] = translations.get(before["text"], after["text"])
        response.raw_output = json.dumps(output)
        return response
    client.transform = target
    result = TranslationWorkflow(client=client, translation_protocol="legal_blocks_v2").run(config)
    assert result.success, result.error
    source = read_source(result)
    assert source["provenance"] == "digital_pdf", source
    assert len([block for block in source["blocks"] if block["role"] == "list_item"]) == 2
    cells = [block for block in source["blocks"] if block["role"] == "table_cell"]
    assert [(cell["row"], cell["col"], cell["text"]) for cell in cells] == [
        (0, 0, "Alpha"), (0, 1, "Beta"), (1, 0, "Gamma"), (1, 1, "")]
    from docx import Document
    output = Document(result.output_docx)
    assert len(output.tables) == 1
    assert [[cell.text for cell in row.cells] for row in output.tables[0].rows] == [["Alpha", "Beta"], ["Gamma", ""]]
