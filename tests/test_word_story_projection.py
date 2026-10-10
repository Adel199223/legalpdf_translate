from io import BytesIO
from zipfile import ZipFile
from copy import deepcopy

import pytest
from docx import Document
from lxml import etree

from legalpdf_translate.ordinary_section_ownership import verified_word_story_projection, W
from legalpdf_translate.ordinary_edited_revision import qualify_edited_docx
from legalpdf_translate.ordinary_text_correction import paragraphs, word_changes, imported_word_map
from legalpdf_translate.ordinary_layout_contracts import digest


def members(raw):
    with ZipFile(BytesIO(raw)) as z:return {p:z.read(p) for p in z.namelist()}


def packed(parts):
    out=BytesIO()
    with ZipFile(out,'w') as z:
        for p,v in parts.items():z.writestr(p,v)
    return out.getvalue()


def fixture():
    d=Document();d.add_paragraph('Body reference ABC 123');d.sections[0].different_first_page_header_footer=True
    d.sections[0].footer.paragraphs[0].text=''
    d.sections[0].first_page_footer.paragraphs[0].text='Reply instruction 2 / 7'
    out=BytesIO();d.save(out);return out.getvalue()


def serialized(raw):
    parts=members(raw);rels=etree.fromstring(parts['word/_rels/document.xml.rels'])
    first=next(n for n in rels if n.get('Target')=='footer2.xml')
    first.set('Target','footer3.xml');parts['word/footer3.xml']=parts['word/footer2.xml']
    parts['word/footer2.xml']=parts['word/footer1.xml']
    parts['word/_rels/document.xml.rels']=etree.tostring(rels)
    return packed(parts)


def v7(raw):return {'writer_version':'saved_docx_layout_writer_source_layout_v7','docx_sha256':digest(raw)}


def test_word_footer_renumbering_projects_by_slot_and_preserves_ids():
    raw=fixture();saved=serialized(raw)
    projection=qualify_edited_docx(raw,saved,source_layout_map=v7(raw))
    assert projection['part_map']['word/footer2.xml']=='word/footer3.xml'
    before=paragraphs(raw);mapped=imported_word_map(raw,saved,before)
    meaningful=next(p for p in before if p['text'].startswith('Reply'))
    actual=next(p for p in mapped if p['paragraph_id']==meaningful['paragraph_id'])
    assert actual['part_uri']=='word/footer3.xml'
    assert projection['edited_sha256']==digest(saved)


@pytest.mark.parametrize('kind',['slot','text','folio','control','orphan','orphan_header','external','relationship_type','continuous','disabled'])
def test_reject_substantive_story_or_section_change(kind):
    raw=fixture();parts=members(serialized(raw))
    if kind in {'slot','continuous','disabled'}:
        doc=etree.fromstring(parts['word/document.xml']);sect=doc.find('.//'+W+'sectPr')
        if kind=='slot':
            ref=next(n for n in sect if n.tag==W+'footerReference' and n.get(W+'type')=='first');ref.set(W+'type','even')
        elif kind=='continuous':etree.SubElement(sect,W+'type').set(W+'val','continuous')
        else:sect.find(W+'titlePg').set(W+'val','0')
        parts['word/document.xml']=etree.tostring(doc)
    elif kind in {'external','relationship_type'}:
        rels=etree.fromstring(parts['word/_rels/document.xml.rels'])
        ref=next(n for n in rels if n.get('Target')=='footer3.xml')
        ref.set('TargetMode','External') if kind=='external' else ref.set('Type','urn:wrong/footer')
        parts['word/_rels/document.xml.rels']=etree.tostring(rels)
    elif kind=='orphan_header':
        root=etree.fromstring(parts['word/footer3.xml']);root.tag=W+'hdr'
        parts['word/header9.xml']=etree.tostring(root)
    else:
        root=etree.fromstring(parts['word/footer3.xml'])
        if kind=='text':next(root.iter(W+'t')).text='Changed source instruction'
        elif kind=='folio':next(root.iter(W+'t')).text='Reply instruction 2 / 8'
        elif kind=='control':etree.SubElement(next(root.iter(W+'r')),W+'fldChar').set(W+'fldCharType','begin')
        else:parts['word/footer4.xml']=parts['word/footer3.xml']
        parts['word/footer3.xml']=etree.tostring(root)
    with pytest.raises(ValueError):qualify_edited_docx(raw,packed(parts),source_layout_map=v7(raw))


def test_default_next_page_omission_is_semantic_but_footer_import_is_explicit():
    raw=fixture();a=members(raw);doc=etree.fromstring(a['word/document.xml'])
    etree.SubElement(doc.find('.//'+W+'sectPr'),W+'type').set(W+'val','nextPage');a['word/document.xml']=etree.tostring(doc);raw=packed(a)
    b=members(serialized(raw));doc=etree.fromstring(b['word/document.xml']);sect=doc.find('.//'+W+'sectPr');sect.remove(sect.find(W+'type'));b['word/document.xml']=etree.tostring(doc)
    footer=etree.fromstring(b['word/footer3.xml']);next(footer.iter(W+'t')).text='Reviewed reply instruction 2 / 7';b['word/footer3.xml']=etree.tostring(footer);saved=packed(b)
    changes=word_changes(raw,saved)
    assert len(changes)==1 and changes[0]['before']=='Reply instruction 2 / 7'
    mapping=imported_word_map(raw,saved,paragraphs(raw))
    assert any(p['part_uri']=='word/footer3.xml' and p['paragraph_id']==changes[0]['paragraph_id'] for p in mapping)


@pytest.mark.parametrize('value',['0','false','off'])
def test_unchanged_disabled_first_page_is_equivalent_to_absent(value):
    raw=fixture();parts=members(raw);root=etree.fromstring(parts['word/document.xml'])
    title=root.find('.//'+W+'titlePg');title.set(W+'val',value)
    parts['word/document.xml']=etree.tostring(root);before=packed(parts)
    title.getparent().remove(title);parts['word/document.xml']=etree.tostring(root)
    qualify_edited_docx(before,packed(parts),source_layout_map=v7(before))


def test_empty_continuation_slots_may_split_and_inherit():
    d=Document();d.add_paragraph('First section');d.sections[0].different_first_page_header_footer=True
    d.sections[0].footer.paragraphs[0].text=''
    d.sections[0].first_page_footer.paragraphs[0].text='First source footer'
    d.add_section();d.add_paragraph('Second section')
    out=BytesIO();d.save(out);raw=out.getvalue();parts=members(raw)
    doc=etree.fromstring(parts['word/document.xml']);sects=list(doc.iter(W+'sectPr'))
    rels=etree.fromstring(parts['word/_rels/document.xml.rels']);ns='{http://schemas.openxmlformats.org/package/2006/relationships}'
    empty_target=next(n.get('Target') for n in rels if n.get('Type','').endswith('/footer') and n.get('Target')=='footer1.xml')
    parts['word/footer9.xml']=parts['word/'+empty_target]
    rel=etree.SubElement(rels,ns+'Relationship',Id='rIdBlankSplit',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer',Target='footer9.xml')
    ref=etree.SubElement(sects[-1],W+'footerReference');ref.set(W+'type','even');ref.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id',rel.get('Id'))
    parts['word/document.xml']=etree.tostring(doc);parts['word/_rels/document.xml.rels']=etree.tostring(rels)
    projection=qualify_edited_docx(raw,packed(parts),source_layout_map=v7(raw))
    assert 'word/footer9.xml' in projection['edited_empty_parts']


def test_malformed_v7_edited_package_uses_controlled_inventory_refusal():
    raw=fixture()
    with pytest.raises(ValueError,match='edited_docx_invalid'):
        qualify_edited_docx(raw,b'not a ZIP',source_layout_map=v7(raw))


@pytest.mark.parametrize('merge',[False,True])
def test_identical_nonempty_story_rename_preserves_distinct_owner_sets(merge):
    d=Document();d.add_paragraph('First body');d.sections[0].different_first_page_header_footer=True
    d.sections[0].first_page_footer.paragraphs[0].text='Identical source footer'
    section=d.add_section();section.different_first_page_header_footer=True
    section.first_page_footer.is_linked_to_previous=False
    section.first_page_footer.paragraphs[0].text='Identical source footer';d.add_paragraph('Second body')
    out=BytesIO();d.save(out);raw=out.getvalue();parts=members(raw)
    rels=etree.fromstring(parts['word/_rels/document.xml.rels'])
    refs=[r for r in rels if r.get('Type','').endswith('/footer')]
    assert len(refs)==2
    left,right=[r.get('Target') for r in refs]
    refs[0].set('Target',right);refs[1].set('Target',right if merge else left)
    parts['word/_rels/document.xml.rels']=etree.tostring(rels)
    if merge:
        with pytest.raises(ValueError):qualify_edited_docx(raw,packed(parts),source_layout_map=v7(raw))
    else:
        projection=qualify_edited_docx(raw,packed(parts),source_layout_map=v7(raw))
        assert len(projection['part_map'])==2 and len(set(projection['part_map'].values()))==2
