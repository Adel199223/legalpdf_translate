from copy import deepcopy
import pytest
from legalpdf_translate.ordinary_layout_contracts import normalize_proposals, OrdinaryLayoutError
from legalpdf_translate.saved_docx_layout import inspect_docx
from tests.test_ordinary_layout_contracts import proposal
from tests.test_ordinary_layout_service import make_case


@pytest.mark.parametrize('spacing', [None, 0, 0.0, 1, True, False])
def test_page_boundary_only_null_or_numeric_zero_in_plain_flow(tmp_path, monkeypatch, spacing):
    case = make_case(tmp_path, monkeypatch)
    view = case.view['review']
    snapshot = inspect_docx(case.job.reviewed_docx, 'EN')
    snapshot['paragraphs'][0]['has_page_break'] = True
    snapshot['paragraphs'][0]['tokens'] = [{'kind':'page_break', 'text':'\f', 'start':0, 'end':1, 'run':0}]
    value = proposal(view)
    value['paragraphs'][0].update(role='body', heading_level=0, heading_size_pt=None,
        bold=False, alignment='inherit', space_before_pt=spacing, space_after_pt=spacing)
    value['bands'] = [{'kind':'flow','groups':[{'paragraph_ids':[p['paragraph_id'] for p in value['paragraphs']], 'panel':False}]}]
    original = deepcopy(value)
    if spacing is None or type(spacing) in {int,float} and spacing == 0:
        normalize_proposals(snapshot, view, [value])
        for key, changed in [('bold',True), ('alignment','center'), ('space_before_pt',1)]:
            invalid = deepcopy(value); invalid['paragraphs'][0][key] = changed
            with pytest.raises(OrdinaryLayoutError, match='page_break_requires_flow'):
                normalize_proposals(snapshot, view, [invalid])
        invalid=deepcopy(value); invalid['bands'][0]['groups'][0]['panel']=True
        with pytest.raises(OrdinaryLayoutError): normalize_proposals(snapshot,view,[invalid])
        invalid=proposal(view); invalid['paragraphs'][0]=deepcopy(value['paragraphs'][0])
        with pytest.raises(OrdinaryLayoutError): normalize_proposals(snapshot,view,[invalid])
    else:
        with pytest.raises(OrdinaryLayoutError, match='page_break_requires_flow'):
            normalize_proposals(snapshot,view,[value])
    assert value == original


@pytest.mark.parametrize('sentinel', ['', ' \u00a0', '\u200e', '\u0301', '\t\n'])
def test_control_only_break_cannot_escape_neutral_guard(tmp_path, monkeypatch, sentinel):
    case = make_case(tmp_path, monkeypatch)
    view = case.view['review']
    snapshot = inspect_docx(case.job.reviewed_docx, 'EN')
    snapshot['paragraphs'][0]['has_page_break'] = True
    snapshot['paragraphs'][0]['tokens'] = [{'kind':'t', 'text':sentinel}]
    value = proposal(view)
    value['bands'] = [{'kind':'flow','groups':[{'paragraph_ids':[p['paragraph_id'] for p in value['paragraphs']], 'panel':False}]}]
    value['paragraphs'][0].update(role='body', heading_level=0, heading_size_pt=None,
        bold=False, alignment='center', space_before_pt=0, space_after_pt=0)
    with pytest.raises(OrdinaryLayoutError, match='page_break_requires_flow'):
        normalize_proposals(snapshot,view,[value])


@pytest.mark.parametrize('text,role,alignment', [('Footer contact.', 'body', 'left'), ('4', 'source_folio', 'center'), ('Section title', 'heading', 'right')])
def test_real_text_terminal_break_uses_bounded_styles_and_preserves_exact_tokens(tmp_path, monkeypatch, text, role, alignment):
    from io import BytesIO
    from docx import Document
    from docx.enum.text import WD_BREAK
    from tests import test_ordinary_layout_service as fixtures
    from legalpdf_translate.saved_docx_layout_writer import build_docx
    doc = Document()
    doc.add_paragraph(text).add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph('Next unchanged paragraph.')
    stream = BytesIO(); doc.save(stream); raw = stream.getvalue()
    monkeypatch.setattr(fixtures, 'docx_bytes', lambda *a, **k: raw)
    case = make_case(tmp_path,monkeypatch)
    snapshot = inspect_docx(raw,'EN'); view=case.view['review']
    value=proposal(view)
    value['paragraphs'][0].update(role=role, heading_level=1 if role=='heading' else 0,
        heading_size_pt=12 if role=='heading' else None, bold=True, italic=True,
        alignment=alignment, space_before_pt=3, space_after_pt=4,
        emphasis=[{'start':0,'end':text.index(' '),'bold':True,'italic':False,'underline':True}] if ' ' in text else [])
    value['bands']=[{'kind':'flow','groups':[{'paragraph_ids':[p['paragraph_id'] for p in value['paragraphs']], 'panel':False}]}]
    normalized=normalize_proposals(snapshot,view,[value])
    normalized['review'].update(document_reviewed=True, pages_reviewed=[1],reviewer='Fictional reviewer',note='Fixture source checked')
    built=build_docx(raw,snapshot,view['pages'],normalized)
    after=inspect_docx(built.docx_bytes,'EN')
    assert [p['content_sha256'] for p in after['paragraphs']]==[p['content_sha256'] for p in snapshot['paragraphs']]
    assert after['paragraphs'][0]['has_page_break']
    assert after['paragraphs'][0]['text']==text+'\f'
    for mutation in ('panel','columns','emphasis_break','partition','bad_spacing','bad_alignment'):
        bad=deepcopy(value)
        if mutation=='panel': bad['bands'][0]['groups'][0]['panel']=True
        elif mutation=='columns':
            bad['bands']=[{'kind':'columns','widths_pct':[50,50],'gutter_pt':12,'cells':[{'groups':[{'paragraph_ids':[p['paragraph_id']], 'panel':False}]} for p in bad['paragraphs']]}]
        elif mutation=='emphasis_break': bad['paragraphs'][0]['emphasis']=[{'start':0,'end':len(text)+1,'bold':True,'italic':False,'underline':True}]
        elif mutation=='partition':
            from legalpdf_translate.ordinary_layout_contracts import PROPOSAL_VERSION_V2
            bad['version']=PROPOSAL_VERSION_V2
            bad['paragraph_partitions']=[{'paragraph_id':bad['paragraphs'][0]['paragraph_id'],'split_before':[text[-1]]}]
        elif mutation=='bad_spacing': bad['paragraphs'][0]['space_after_pt']=73
        else: bad['paragraphs'][0]['alignment']='invented'
        with pytest.raises(OrdinaryLayoutError): normalize_proposals(snapshot,view,[bad])
