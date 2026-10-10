"""Synthetic local correction lineage; no provider or native calls."""
from dataclasses import replace
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from lxml import etree
import pytest

from legalpdf_translate.ordinary_layout_contracts import OrdinaryLayoutError
from legalpdf_translate.ordinary_layout_manager import OrdinaryLayoutManager
from legalpdf_translate.ordinary_text_correction import paragraphs, apply_actions, word_changes
from tests.test_ordinary_layout_service import make_case, nonce
from legalpdf_translate.ordinary_text_correction import W


def action(pid, text, kind="replace"):
    return {"kind": kind, "paragraph_id": pid, "text": text, "source_page": 1, "source_bbox": [.1, .1, .9, .2]}


@pytest.mark.parametrize("lang,text", [("EN", "Corrected fictional address R"), ("FR", "Adresse fictive corrigée R"), ("AR", "عنوان خيالي R 123")])
def test_approval_delivery_reload_and_immutable_original(tmp_path, monkeypatch, lang, text):
    case = make_case(tmp_path, monkeypatch, lang)
    service = case.manager.service
    state = service.text_correction_state(case.job)
    pid = state["paragraphs"][0]["paragraph_id"]
    draft = service.draft_text_correction(case.job, nonce(), state["parent"], [action(pid, text)])
    assert service.state(case.job.job_id)["delivery"] is None
    selected = service.approve_text_correction(case.job, draft["draft_id"], nonce(), True, True, "Compared with source")
    assert selected["delivery"]["text_approved"] and selected["delivery"]["output_review_required"]
    artifact = case.manager.resolve_delivery(case.job.job_id, 1)
    assert paragraphs(artifact.path.read_bytes())[0]["text"] == text
    assert artifact.word_count > 0 and artifact.kind == "text_corrected"
    reopened = OrdinaryLayoutManager(case.root, mode="shadow", workspace_id="fixture", job_resolver=lambda _: case.job)
    assert reopened.resolve_delivery(case.job.job_id, 1) == artifact
    assert case.job.original_docx == case.job.reviewed_docx
    assert service.text_correction_state(case.job)["paragraphs"][0]["paragraph_id"] == pid


def test_cancel_stale_approval_owner_and_freeze(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    service = case.manager.service
    state = service.text_correction_state(case.job)
    edits = [action(state["paragraphs"][0]["paragraph_id"], "Replacement")]
    a = service.draft_text_correction(case.job, nonce(), state["parent"], edits)
    b = service.draft_text_correction(case.job, nonce(), state["parent"], edits)
    service.cancel_text_correction(case.job, b["draft_id"])
    with pytest.raises(OrdinaryLayoutError, match="cancelled"):
        service.approve_text_correction(case.job, b["draft_id"], nonce(), True, True, "Compared")
    with pytest.raises(OrdinaryLayoutError, match="owner_changed"):
        service.text_correction_state(replace(case.job, run_id="wrong-run"))
    service.approve_text_correction(case.job, a["draft_id"], nonce(), True, True, "Compared")
    with pytest.raises(OrdinaryLayoutError, match="parent_stale"):
        service.draft_text_correction(case.job, nonce(), state["parent"], edits)
    selected = service.state(case.job.job_id)["delivery"]
    service.review_text_output(case.job, selected["selection_id"], 1, True)
    case.manager.resolve_delivery(case.job.job_id, 1, nonce())
    with pytest.raises(OrdinaryLayoutError, match="frozen"):
        service.draft_text_correction(case.job, nonce(), state["parent"], edits)


def test_insert_delete_preserve_stable_ids_and_package(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    service = case.manager.service
    state = service.text_correction_state(case.job)
    first, second = state["paragraphs"]
    draft = service.draft_text_correction(case.job, nonce(), state["parent"], [
        action(first["paragraph_id"], "Inserted source annotation", "insert_before")])
    service.approve_text_correction(case.job, draft["draft_id"], nonce(), True, True, "Source annotation")
    state = service.text_correction_state(case.job)
    assert [r["paragraph_id"] for r in state["paragraphs"]][1:] == [first["paragraph_id"], second["paragraph_id"]]
    draft = service.draft_text_correction(case.job, nonce(), state["parent"], [action(first["paragraph_id"], "", "delete")])
    service.approve_text_correction(case.job, draft["draft_id"], nonce(), True, True, "Duplicate removed")
    assert service.text_correction_state(case.job)["paragraphs"][-1]["paragraph_id"] == second["paragraph_id"]


def test_word_import_requires_safe_structure_and_keeps_unapproved_bytes(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    base = case.job.reviewed_docx
    doc = Document(BytesIO(base)); doc.paragraphs[0].text = "Edited fictional text"
    stream = BytesIO(); doc.save(stream); working = stream.getvalue()
    assert len(word_changes(base, working)) == 1
    doc.add_paragraph("Unowned insertion"); stream = BytesIO(); doc.save(stream)
    with pytest.raises(OrdinaryLayoutError, match="structure_unsupported|section_changed"):
        word_changes(base, stream.getvalue())


def test_missing_region_and_controls_refuse_before_publication(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    pid = paragraphs(case.job.reviewed_docx)[0]["paragraph_id"]
    bad = action(pid, "Missing source region", "insert_after"); bad["source_bbox"] = None
    with pytest.raises(OrdinaryLayoutError, match="region_required"):
        apply_actions(case.job.reviewed_docx, [bad], "EN", (1,))
    with pytest.raises(OrdinaryLayoutError, match="invalid_text"):
        apply_actions(case.job.reviewed_docx, [action(pid, "Bad\u202e control")], "EN", (1,))


def test_literal_brackets_and_repeated_insert_ids_are_exact(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch, "AR")
    raw = case.job.reviewed_docx
    pid = paragraphs(raw)[0]["paragraph_id"]
    output, _, mapping = apply_actions(raw, [action(pid, "النص [[ABC]]")], "AR", (1,))
    assert paragraphs(output)[0]["text"] == "النص [[ABC]]"
    inserted, _, mapping = apply_actions(output, [action(pid, "نص جديد", "insert_after")], "AR", (1,), paragraph_map=mapping)
    again, _, mapping = apply_actions(inserted, [action(pid, "نص جديد", "insert_after")], "AR", (1,), paragraph_map=mapping)
    assert len(paragraphs(again)) == len({r["paragraph_id"] for r in mapping}) == 4


def test_word_section_and_noneditable_story_change_refuse(tmp_path, monkeypatch):
    from lxml import etree
    case = make_case(tmp_path, monkeypatch)
    doc = Document(BytesIO(case.job.reviewed_docx)); doc.paragraphs[0].text = "Approved wording candidate"
    props = doc.paragraphs[1]._p.get_or_add_pPr(); etree.SubElement(props, W + "sectPr")
    stream = BytesIO(); doc.save(stream)
    with pytest.raises(OrdinaryLayoutError, match="section_changed"):
        word_changes(case.job.reviewed_docx, stream.getvalue())
    def note(raw, text):
        output = BytesIO()
        with ZipFile(BytesIO(raw)) as source, ZipFile(output, "w") as target:
            for info in source.infolist():
                target.writestr(info, source.read(info))
            target.writestr("word/footnotes.xml", f'<w:footnotes xmlns:w="{W[1:-1]}"><w:footnote w:id="1"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:footnote></w:footnotes>')
        return output.getvalue()
    original = note(case.job.reviewed_docx, "Original note")
    doc = Document(BytesIO(case.job.reviewed_docx)); doc.paragraphs[0].text = "Changed body"
    stream = BytesIO(); doc.save(stream)
    with pytest.raises(OrdinaryLayoutError, match="story_unsupported"):
        word_changes(original, note(stream.getvalue(), "Hidden note change"))


@pytest.mark.parametrize("what", ["docx", "map"])
def test_selected_copy_and_rehashed_map_tampering_refuse(tmp_path, monkeypatch, what):
    from legalpdf_translate.ordinary_layout_service import _read, _write
    case = make_case(tmp_path, monkeypatch)
    service = case.manager.service
    state = service.text_correction_state(case.job)
    draft = service.draft_text_correction(case.job, nonce(), state["parent"], [action(state["paragraphs"][0]["paragraph_id"], "Corrected")])
    selected = service.approve_text_correction(case.job, draft["draft_id"], nonce(), True, True, "")
    delivery = service.resolve_delivery(case.job.job_id, 1)
    if what == "docx":
        delivery.path.write_bytes(case.job.original_docx)
    else:
        path = delivery.path.parent / "source_map.json"
        value = _read(path); value["source_character_coverage"] = "invented_complete"
        from legalpdf_translate.ordinary_layout_contracts import digest, encode
        path.write_bytes(encode({"payload": value, "sha256": digest(encode(value))}))
    with pytest.raises(OrdinaryLayoutError, match="changed"):
        service.resolve_delivery(case.job.job_id, 1)


def test_repeated_formatting_uses_one_working_root_and_imports_later_text(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    service = case.manager.service
    state = service.text_correction_state(case.job)
    draft = service.draft_text_correction(case.job, nonce(), state["parent"], [action(state["paragraphs"][0]["paragraph_id"], "Corrected fictional wording")])
    service.approve_text_correction(case.job, draft["draft_id"], nonce(), True, True, "")
    working = service.text_corrected_review_copy(case.job.job_id)
    for bold in [True, False]:
        doc = Document(working); doc.paragraphs[0].runs[0].bold = bold; doc.save(working)
        service.adopt_text_corrected_word_edit(case.job.job_id)
        assert service.text_corrected_review_copy(case.job.job_id) == working
    doc = Document(working); doc.paragraphs[0].text = "Explicit later text edit"; doc.save(working)
    with pytest.raises(OrdinaryLayoutError):
        service.adopt_text_corrected_word_edit(case.job.job_id)
    state = service.text_correction_state(case.job)
    imported = service.draft_text_correction(case.job, nonce(), state["parent"], [], import_word=True)
    assert imported["changes"][0]["after"] == "Explicit later text edit"
    assert working.exists()


def test_source_finding_requires_explicit_review_and_is_revision_bound(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch); service = case.manager.service
    original_state = service.state
    finding = {"finding_id": "a"*64, "kind": "missing_readable_literal", "literal": "R", "page_number": 1, "source_bbox": [.1,.1,.2,.2], "insertion_context": {"paragraph_id": "p1", "position": "before"}}
    def state(job_id):
        value = original_state(job_id); value["suggestions"] = [{"source_coverage_findings": [finding]}]; return value
    monkeypatch.setattr(service, "state", state)
    view = service.text_correction_state(case.job)
    assert view["source_findings"][0]["review_status"] == "unresolved"
    with pytest.raises(OrdinaryLayoutError):
        service.review_source_finding(case.job, finding["finding_id"], view["parent"], "corrected", True)
    reviewed = service.review_source_finding(case.job, finding["finding_id"], view["parent"], "reviewed_no_change", True)
    # There is no automatic selected artifact in this fixture; a new selected correction must still reset review.
    draft = service.draft_text_correction(case.job, nonce(), view["parent"], [action(view["paragraphs"][0]["paragraph_id"], "Corrected")])
    service.approve_text_correction(case.job, draft["draft_id"], nonce(), True, True, "")
    current = service.text_correction_state(case.job)
    assert current["source_findings"][0]["review_status"] == "unresolved"
    current = service.review_source_finding(case.job, finding["finding_id"], current["parent"], "corrected", True)
    assert current["source_findings"][0]["review_status"] == "corrected"


def test_correction_counts_meaningful_header_footer_text_once():
    from legalpdf_translate.ordinary_text_correction import count_correction_words
    doc = Document(); doc.add_paragraph("Body two")
    doc.sections[0].header.paragraphs[0].text = "Header three words"
    doc.sections[0].footer.paragraphs[0].text = "Source footer four words"
    stream = BytesIO(); doc.save(stream)
    assert count_correction_words(stream.getvalue()) == 9


def test_semantic_word_sections_allow_benign_metadata_but_reject_footer_repoint(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    doc = Document(BytesIO(case.job.reviewed_docx)); doc.sections[0].footer.paragraphs[0].text = "Owned footer"
    stream = BytesIO(); doc.save(stream); original = stream.getvalue()
    def edit(repoint=False):
        output = BytesIO()
        with ZipFile(BytesIO(original)) as source, ZipFile(output, "w") as target:
            for info in source.infolist():
                raw = source.read(info)
                if info.filename == "word/document.xml":
                    root = etree.fromstring(raw); next(root.iter(W+"t")).text = "Changed text"
                    section = next(root.iter(W+"sectPr")); section.set(W+"rsidR", "12345678")
                    reference = next(section.iter(W+"footerReference")); reference.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "rIdChanged")
                    raw = etree.tostring(root)
                elif info.filename == "word/_rels/document.xml.rels":
                    root = etree.fromstring(raw)
                    for rel in root:
                        if rel.get("Type").endswith("/footer"):
                            rel.set("Id", "rIdChanged")
                            if repoint: rel.set("Target", "footer999.xml")
                    raw = etree.tostring(root)
                target.writestr(info,raw)
        return output.getvalue()
    assert len(word_changes(original,edit())) == 1
    with pytest.raises(OrdinaryLayoutError,match="section_changed"):
        word_changes(original,edit(True))


def test_natural_arabic_in_latin_target_gets_native_direction_and_literal_brackets(tmp_path, monkeypatch):
    case=make_case(tmp_path,monkeypatch)
    raw=case.job.reviewed_docx
    value="نص [[ABC]]"
    edited,changes,mapping=apply_actions(raw,[action(paragraphs(raw)[0]["paragraph_id"],value)],"EN",(1,))
    with ZipFile(BytesIO(edited)) as archive:root=etree.fromstring(archive.read("word/document.xml"))
    first=next(root.iter(W+"p"))
    assert first.find(W+"pPr/"+W+"bidi").get(W+"val")=="1"
    assert "".join(n.text or "" for n in first.iter(W+"t"))==value
    replaced,_,_=apply_actions(edited,[action(mapping[0]["paragraph_id"],"Latin text")],"EN",(1,),paragraph_map=mapping)
    with ZipFile(BytesIO(replaced)) as archive:root=etree.fromstring(archive.read("word/document.xml"))
    assert next(root.iter(W+"p")).find(W+"pPr/"+W+"bidi").get(W+"val")=="0"


@pytest.mark.parametrize("language",["EN","FR"])
def test_mixed_latin_sentence_retains_ltr_paragraph_with_arabic_run(tmp_path,monkeypatch,language):
    case=make_case(tmp_path,monkeypatch,lang=language)
    value="The witness wrote محمد."
    edited,_,_=apply_actions(case.job.reviewed_docx,[action(paragraphs(case.job.reviewed_docx)[0]["paragraph_id"],value)],language,(1,))
    with ZipFile(BytesIO(edited)) as archive:root=etree.fromstring(archive.read("word/document.xml"))
    first=next(root.iter(W+"p"))
    assert first.find(W+"pPr/"+W+"bidi").get(W+"val")=="0"
    arabic=next(r for r in first.findall(W+"r") if "محمد" in "".join(n.text or "" for n in r.iter(W+"t")))
    assert arabic.find(W+"rPr/"+W+"rtl").get(W+"val")=="1"
    assert "".join(n.text or "" for n in first.iter(W+"t"))==value


def test_v7_source_footer_count_and_correction_formatting_section_guard(tmp_path):
    from tests.test_ordinary_source_layout import fixture,build,verify
    from legalpdf_translate.ordinary_layout_service import _candidate_word_count
    from legalpdf_translate.ordinary_edited_revision import qualify_edited_docx
    from legalpdf_translate.ordinary_text_correction_service import _qualify_corrected_formatting
    from legalpdf_translate.ordinary_text_correction import count_correction_words
    args=fixture();candidate=build(args);verify(candidate,args)
    path=tmp_path/"verified.docx";path.write_bytes(candidate.docx_bytes)
    assert _candidate_word_count(candidate,path)==count_correction_words(candidate.docx_bytes)==9
    doc=Document(path);doc.paragraphs[0].paragraph_format.alignment=2;doc.save(path)
    qualify_edited_docx(candidate.docx_bytes,path.read_bytes(),source_layout_map=candidate.source_map)
    assert _candidate_word_count(candidate,path)==9
    _qualify_corrected_formatting(candidate.docx_bytes,path.read_bytes())
    doc=Document(path);doc.sections[0].different_first_page_header_footer=False;doc.save(path)
    with pytest.raises(OrdinaryLayoutError,match="section_changed"):
        _qualify_corrected_formatting(candidate.docx_bytes,path.read_bytes())


def test_safe_source_folio_wrapper_remains_correctable_without_reversing_fraction(tmp_path, monkeypatch):
    from legalpdf_translate.ordinary_source_layout import _folio_ltr_span
    case=make_case(tmp_path,monkeypatch,lang='AR')
    doc=Document(BytesIO(case.job.reviewed_docx));doc.paragraphs[0].text='الصفحة 2 / 7'
    assert _folio_ltr_span(doc.paragraphs[0]._p)['literal']=='2 / 7'
    out=BytesIO();doc.save(out);raw=out.getvalue()
    row=paragraphs(raw)[0];assert row['editable']
    edited,_,_=apply_actions(raw,[action(row['paragraph_id'],'الصفحة 3 / 7')],'AR',(1,))
    with ZipFile(BytesIO(edited)) as z:root=etree.fromstring(z.read('word/document.xml'))
    wrapper=next(root.iter(W+'dir'))
    assert wrapper.get(W+'val')=='ltr'
    assert ''.join(t.text or '' for t in wrapper.iter(W+'t'))=='3 / 7'
    with pytest.raises(OrdinaryLayoutError,match='folio_structure_changed'):
        apply_actions(raw,[action(row['paragraph_id'],'Unrecognized changed paragraph')],'AR',(1,))
    wrapper.set(W+'val','rtl')
    assert not __import__('legalpdf_translate.ordinary_text_correction',fromlist=['_editable'])._editable(wrapper.getparent())


def test_word_import_refuses_mapped_decorative_rule(tmp_path, monkeypatch):
    case=make_case(tmp_path,monkeypatch);service=case.manager.service
    state=service.text_correction_state(case.job)
    draft=service.draft_text_correction(case.job,nonce(),state['parent'],[action(state['paragraphs'][0]['paragraph_id'],'First approved text')])
    service.approve_text_correction(case.job,draft['draft_id'],nonce(),True,True,'')
    working=service.text_corrected_review_copy(case.job.job_id)
    doc=Document(working);doc.paragraphs[0].text='Forbidden decorative text';doc.save(working)
    parent=service._correction_parent
    def mapped(*args,**kwargs):
        raw,identity,rows=parent(*args,**kwargs);rows[0]['decorative_rule']=True
        return raw,identity,rows
    monkeypatch.setattr(service,'_correction_parent',mapped)
    state=service.text_correction_state(case.job)
    with pytest.raises(OrdinaryLayoutError,match='paragraph_not_editable'):
        service.draft_text_correction(case.job,nonce(),state['parent'],[],import_word=True)


def test_explicit_word_open_check_and_import_preserve_owned_copy(tmp_path, monkeypatch):
    from legalpdf_translate.word_automation import WordAutomationResult
    case=make_case(tmp_path,monkeypatch);service=case.manager.service
    state=service.text_correction_state(case.job)
    draft=service.draft_text_correction(case.job,nonce(),state['parent'],[action(state['paragraphs'][0]['paragraph_id'],'Approved fictional text')])
    service.approve_text_correction(case.job,draft['draft_id'],nonce(),True,True,'')
    current=service.text_correction_state(case.job);calls=[]
    monkeypatch.setattr('legalpdf_translate.word_automation.open_docx_in_word',lambda path:calls.append(path) or WordAutomationResult(True,'open','Opened owned copy'))
    with pytest.raises(OrdinaryLayoutError,match='stale'):
        service.open_text_correction_word(case.job,state['parent'])
    assert not calls
    assert service.open_text_correction_word(case.job,current['parent'])['open_result']['ok']
    working=calls[-1];doc=Document(working);doc.paragraphs[0].text='Explicit Word text edit';doc.save(working)
    changed=working.read_bytes()
    service.open_text_correction_word(case.job,current['parent']);assert working.read_bytes()==changed
    with pytest.raises(OrdinaryLayoutError,match='text_changes_pending'):
        service.check_text_correction_word(case.job,current['parent'])
    imported=service.draft_text_correction(case.job,nonce(),current['parent'],[],import_word=True)
    service.approve_text_correction(case.job,imported['draft_id'],nonce(),True,True,'')
    current=service.text_correction_state(case.job);service.open_text_correction_word(case.job,current['parent'])
    doc=Document(calls[-1]);doc.paragraphs[0].paragraph_format.alignment=2;doc.save(calls[-1])
    formatted=service.check_text_correction_word(case.job,current['parent'])
    assert formatted['parent']['sha256']!=current['parent']['sha256']
    selected=service.state(case.job.job_id)['delivery']
    service.review_text_output(case.job,selected['selection_id'],selected['generation'],True)
    service.resolve_delivery(case.job.job_id,selected['generation'],nonce())
    before=len(calls)
    with pytest.raises(OrdinaryLayoutError,match='frozen'):
        service.open_text_correction_word(case.job,formatted['parent'])
    assert len(calls)==before
