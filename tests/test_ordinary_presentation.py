"""Ordinary V6 metadata presentation; fictional packages are not native proof."""
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
from lxml import etree
import pytest
from legalpdf_translate import saved_docx_layout as model
from legalpdf_translate import saved_docx_layout_writer as writer
from legalpdf_translate.docx_writer import _configure_styles,_add_page_number_footer
from tests.test_saved_docx_layout import document_bytes,reviewed,changed_part

def packet(lang='EN',footer='generated'):
    def edit(doc):
        _configure_styles(doc,lang)
        _add_page_number_footer(doc,lang)
        if footer=='custom':doc.sections[0].footer.paragraphs[0].add_run(' custom')
        if footer=='numpages':
            for node in doc.sections[0].footer._element.iter(model.W+'instrText'):node.text=' NUMPAGES '
    raw=document_bytes(['Reference 12','Recipient address','Obligation remains.','1 / 1'],edit)
    snapshot,pages,dec=reviewed(raw,lang)
    ids=[p['id'] for p in snapshot['paragraphs']]
    dec['review'].update(document_reviewed=False,pages_reviewed=[],reviewer='',note='')
    dec['bands']=[{'kind':'flow','groups':[{'paragraph_ids':[i],'panel':False}]} for i in ids]
    for c,role,box in zip(dec['paragraphs'],['reference','recipient','body','source_folio'],[[20,500,400,530],[20,300,400,350],[20,700,800,750],[400,1300,600,1320]]):
        c.update(role=role,regions=[{'page_number':1,'bbox_px':box}])
    context={'selected_pages':[1],'page_groups':[[1,ids]]}
    return raw,snapshot,pages,dec,context

def build(args):
    raw,snapshot,pages,dec,context=args
    return writer.build_unreviewed_docx(raw,snapshot,pages,dec,ordinary_context=context)

def validate(artifact,args):
    raw,snapshot,pages,dec,context=args
    writer.validate_built_docx(artifact.docx_bytes,artifact.source_map,original_docx=raw,snapshot=snapshot,pages=pages,decisions=dec,require_review=False,ordinary_context=context)

def footer_text(raw):
    with ZipFile(BytesIO(raw)) as z:return z.read('word/footer1.xml')

@pytest.mark.parametrize('lang',['EN','FR','AR'])
def test_metadata_adjacent_inversion_exact_text_and_generated_footer_only(lang):
    args=packet(lang);before=deepcopy(args[3]);art=build(args);validate(art,args)
    ids=[p['id'] for p in args[1]['paragraphs']]
    assert art.source_map['writer_version']==writer.ORDINARY_PRESENTATION_WRITER_VERSION
    assert art.source_map['rendered_paragraph_ids']==[ids[1],ids[0],*ids[2:]]
    assert art.source_map['original_paragraph_ids']==ids
    assert art.source_map['exact_text_preserved'] and not art.source_map['logical_order_preserved']
    assert art.source_map['document_reviewed'] is False
    assert art.source_map['suppressed_generated_page_footer_parts']==['word/footer1.xml']
    assert b' PAGE ' not in footer_text(art.docx_bytes)
    assert args[3]==before
    # Legacy default still has raw order and PAGE, with historical verification.
    old=writer.build_unreviewed_docx(*args[:4]);assert old.source_map['writer_version']==writer.AUTOMATIC_SEPARATED_WRITER_VERSION
    writer.validate_built_docx(old.docx_bytes,old.source_map,original_docx=args[0],snapshot=args[1],pages=args[2],decisions=args[3],require_review=False)
    assert b' PAGE ' in footer_text(old.docx_bytes)

@pytest.mark.parametrize('barrier',['body','signature','list','panel','overlap','equal','full','unmapped','partition'])
def test_metadata_barriers_keep_original_order(barrier):
    args=list(packet());dec=args[3]
    if barrier in {'body','signature','list'}:dec['paragraphs'][1]['role']=barrier
    elif barrier=='panel':dec['bands'][1]['groups'][0]['panel']=True
    elif barrier in {'overlap','equal'}:dec['paragraphs'][1]['regions'][0]['bbox_px']=[20,500,400,550]
    elif barrier=='full':dec['paragraphs'][1]['regions'][0]['bbox_px']=[0,0,1000,1400]
    elif barrier=='unmapped':dec['paragraphs'][1].update(regions=[],unmapped_reason='No reliable region')
    else:
        dec['version']=model.PARTITION_DECISIONS_VERSION;dec['paragraph_partitions']=[{'paragraph_id':dec['paragraphs'][1]['paragraph_id'],'offsets':[10]}];dec['paragraphs'][1]['role']='body'
    art=build(args);assert art.source_map['original_order_preserved'];validate(art,args)

@pytest.mark.parametrize('kind',['custom','missing_page','midpage','prose','body_number'])
def test_folio_and_footer_strict_non_suppression(kind):
    args=list(packet(footer=kind if kind in {'custom','numpages'} else 'generated'))
    if kind=='missing_page':args[4]['selected_pages']=[1,7]
    elif kind=='midpage':args[3]['paragraphs'][3]['regions'][0]['bbox_px']=[400,500,600,520]
    elif kind=='body_number':args[3]['paragraphs'][3]['role']='body'
    elif kind=='prose':
        raw=document_bytes(['Reference 12','Recipient address','Obligation remains.','Page 1 / 1 operative reply'],lambda d:(_configure_styles(d,'EN'),_add_page_number_footer(d,'EN')))
        args[0]=raw;args[1]=model.inspect_docx(raw,'EN')
    art=build(args);validate(art,args)
    assert footer_text(art.docx_bytes)==footer_text(args[0])
    assert not art.source_map['suppressed_generated_page_footer_parts']

def test_v6_map_and_footer_tamper_and_missing_context_reject():
    args=packet();art=build(args)
    mapping=deepcopy(art.source_map);mapping['rendered_paragraph_ids'].reverse()
    with pytest.raises(model.SavedDocxLayoutError):writer.validate_built_docx(art.docx_bytes,mapping,original_docx=args[0],snapshot=args[1],pages=args[2],decisions=args[3],require_review=False,ordinary_context=args[4])
    with pytest.raises(model.SavedDocxLayoutError,match='context_required'):writer.validate_built_docx(art.docx_bytes,art.source_map,original_docx=args[0],snapshot=args[1],pages=args[2],decisions=args[3],require_review=False)
    def inject(root):etree.SubElement(root[0],model.W+'r').append(etree.Element(model.W+'t'));root[0][-1][0].text='hidden change'
    bad=changed_part(art.docx_bytes,'word/footer1.xml',inject)
    with pytest.raises(model.SavedDocxLayoutError,match='unaffected_package_changed'):writer.validate_built_docx(bad,art.source_map,original_docx=args[0],snapshot=args[1],pages=args[2],decisions=args[3],require_review=False,ordinary_context=args[4])


def test_numpages_footer_helper_retains_unsupported_custom_field():
    from legalpdf_translate.ordinary_presentation import suppress_footer
    raw=packet()[0]
    def edit(root):
        for node in root.iter(model.W+'instrText'):node.text=' NUMPAGES '
    raw=changed_part(raw,'word/footer1.xml',edit)
    with ZipFile(BytesIO(raw)) as z:members={n:z.read(n) for n in z.namelist()}
    output,changes=suppress_footer(members,raw,'EN',True)
    assert output==members and not changes

@pytest.mark.parametrize('box',[[30,400,20,450],[20,400,400,400]])
def test_reversed_empty_geometry_rejected_by_shared_validator(box):
    args=list(packet());args[3]['paragraphs'][1]['regions'][0]['bbox_px']=box
    with pytest.raises(model.SavedDocxLayoutError):build(args)


def test_cross_page_regions_and_substantive_intervening_group_never_swapped():
    args=list(packet());ids=args[4]['page_groups'][0][1]
    args[2].append(dict(args[2][0],page_number=2));args[4]={'selected_pages':[1,2],'page_groups':[[1,[ids[0],*ids[2:]]],[2,[ids[1]]]]}
    args[3]['paragraphs'][1]['regions'][0]['page_number']=2
    assert build(args).source_map['original_order_preserved']
    args=list(packet());dec=args[3]
    dec['bands']=[dec['bands'][0],dec['bands'][2],dec['bands'][1],dec['bands'][3]]
    # Strict manual/raw decision ordering remains mandatory, not repaired here.
    with pytest.raises(model.SavedDocxLayoutError):build(args)

def test_malicious_derivation_cannot_reorder_body(monkeypatch):
    args=list(packet());args[3]['paragraphs'][1]['role']='body'
    real=writer._ordinary_presentation
    def corrupt(*values):
        dec,plan=real(*values);dec['bands'][0],dec['bands'][1]=dec['bands'][1],dec['bands'][0]
        plan['rendered_paragraph_ids'][0:2]=reversed(plan['rendered_paragraph_ids'][0:2]);plan['swaps']=[{'forged':True}]
        return dec,plan
    monkeypatch.setattr(writer,'_ordinary_presentation',corrupt)
    with pytest.raises(model.SavedDocxLayoutError,match='unsafe_inversion'):build(args)


def test_malicious_suppression_cannot_remove_custom_footer(monkeypatch):
    args=packet(footer='custom')
    def corrupt(members,raw,lang,enabled):
        output=dict(members);root=model._xml(output['word/footer1.xml'])
        for child in list(root[0]):
            if child.tag!=model.W+'pPr':root[0].remove(child)
        output['word/footer1.xml']=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
        return output,['word/footer1.xml']
    monkeypatch.setattr(writer,'_ordinary_footer',corrupt)
    with pytest.raises(model.SavedDocxLayoutError,match='not_generated_page'):build(args)

@pytest.mark.parametrize('label',['الصفحة 1','الصفحة ١','pág. 1','Page 1 of 1'])
def test_definite_arabic_and_source_page_labels(label):
    args=list(packet('AR'));raw=document_bytes(['Reference 12','Recipient address','Obligation remains.',label],lambda d:(_configure_styles(d,'AR'),_add_page_number_footer(d,'AR')))
    args[0]=raw;args[1]=model.inspect_docx(raw,'AR');art=build(args)
    assert art.source_map['suppressed_generated_page_footer_parts']==['word/footer1.xml'];validate(art,args)


def test_complete_two_page_folios_with_terminal_pagebreak_preserved():
    from docx.enum.text import WD_BREAK
    def edit(doc):
        _configure_styles(doc,'AR');_add_page_number_footer(doc,'AR');doc.paragraphs[1].runs[0].add_break(WD_BREAK.PAGE)
    raw=document_bytes(['First obligation.','الصفحة ١','Second obligation.','الصفحة ٢'],edit)
    snapshot,pages,dec=reviewed(raw,'AR');pages.append(dict(pages[0],page_number=2));ids=[r['id'] for r in snapshot['paragraphs']]
    dec['review'].update(document_reviewed=False,pages_reviewed=[],reviewer='',note='')
    for index,c in enumerate(dec['paragraphs']):
        c.update(role='source_folio' if index%2 else 'body',regions=[{'page_number':1 if index<2 else 2,'bbox_px':[400,1300,600,1320] if index%2 else [20,400,800,600]}])
    dec['bands']=[{'kind':'flow','groups':[{'paragraph_ids':ids,'panel':False}]}]
    context={'selected_pages':[1,2],'page_groups':[[1,ids[:2]],[2,ids[2:]]]};args=raw,snapshot,pages,dec,context;art=build(args);validate(art,args)
    assert art.source_map['suppressed_generated_page_footer_parts']==['word/footer1.xml']
    assert snapshot['paragraphs'][1]['has_page_break']
    dec['paragraphs'][3]['role']='body';assert not build(args).source_map['suppressed_generated_page_footer_parts']
