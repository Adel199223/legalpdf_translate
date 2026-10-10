from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
import hashlib
import json
from pathlib import Path
import uuid

import fitz
import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from io import BytesIO

from legalpdf_translate.docx_writer import assemble_docx
from legalpdf_translate.ordinary_auto_layout_artifacts import (
    OrdinaryAutoArtifactError, bind_raw_page_map, build_unreviewed_candidate,
    verify_unreviewed_candidate,
)
from legalpdf_translate.saved_docx_layout_service import (
    SavedDocxLayoutService, SavedDocxLayoutServiceError,
)
from legalpdf_translate.saved_docx_layout import default_decisions
from legalpdf_translate.types import TargetLang


def _raw(tmp_path: Path):
    pages = tmp_path / "pages"
    pages.mkdir()
    (pages / "page_0002.txt").write_text("First fictional paragraph\nSecond fictional paragraph", encoding="utf-8")
    (pages / "page_0003.txt").write_text("Third fictional paragraph", encoding="utf-8")
    output = assemble_docx(pages, tmp_path / "raw.docx", lang=TargetLang.EN,
                           page_breaks=False, page_numbers=[2, 3])
    raw = output.read_bytes()
    mapping = json.loads(output.with_suffix(".source_map.json").read_text(encoding="utf-8"))
    return raw, mapping


def _source_pdf():
    document = fitz.open()
    for number in range(1, 4):
        page = document.new_page(width=595, height=842)
        page.insert_text((72, 90), f"Fictional source page {number}")
    raw = document.tobytes()
    document.close()
    return raw


def _candidate_fixture(tmp_path):
    raw, mapping = _raw(tmp_path)
    snapshot = bind_raw_page_map(raw, mapping, source_pdf_sha256="a" * 64,
                                 selected_pages=(2, 3), target_lang="EN")
    frames = [{"page_number": n, "width_px": 595, "height_px": 842,
               "image_sha256": f"{n}" * 64} for n in (1, 2, 3)]
    decisions = default_decisions(snapshot.saved_snapshot, frames)
    for choice, row in zip(decisions["paragraphs"], snapshot.paragraphs):
        if row.page_number is not None:
            choice["regions"] = [{"page_number": row.page_number,
                                  "bbox_px": [10, 10, 200, 100]}]
    evidence = {"version": "fictional_proposal_v1", "document_reviewed": False,
                "rendered_layout_acceptance": "not_evaluated"}
    artifact = build_unreviewed_candidate(raw, snapshot, frames, decisions,
                                          proposal_evidence=evidence)
    return raw, snapshot, frames, decisions, evidence, artifact


def test_retained_v6_candidate_with_v3_source_evidence_verifies(tmp_path):
    from legalpdf_translate import ordinary_auto_layout_artifacts as artifacts
    from legalpdf_translate.saved_docx_layout_writer import build_unreviewed_docx
    raw, snapshot, frames, decisions, evidence, _ = _candidate_fixture(tmp_path)
    evidence['source_evidence'] = {'version': 'ordinary_source_evidence_v1', 'pages': []}
    context = {'selected_pages': list(snapshot.selected_pages),
               'page_groups': [[page, list(ids)] for page, ids in snapshot.page_groups]}
    checked = artifacts._checked_automatic_decisions(snapshot, frames, decisions)
    rendered = build_unreviewed_docx(raw, snapshot.saved_snapshot, frames, checked,
                                    ordinary_context=context)
    assert rendered.source_map['writer_version'] == 'saved_docx_layout_writer_ordinary_presentation_v6'
    mapping = artifacts._json(artifacts._extended_source_map(rendered.source_map, snapshot, evidence))
    receipt = artifacts._json(artifacts._candidate_receipt(snapshot, checked, evidence,
                                                          rendered.docx_bytes, mapping))
    candidate = artifacts.AutoCandidateArtifact(rendered.docx_bytes, mapping, receipt,
        artifacts._sha(rendered.docx_bytes), artifacts._sha(mapping), artifacts._sha(receipt))
    before = (candidate.docx_bytes, candidate.source_map_bytes, candidate.receipt_bytes)
    verify_unreviewed_candidate(raw, snapshot, frames, decisions, candidate, proposal_evidence=evidence)
    assert before == (candidate.docx_bytes, candidate.source_map_bytes, candidate.receipt_bytes)


def test_raw_binding_preserves_physical_pages_and_rejects_ambiguous_owner(tmp_path):
    raw, mapping = _raw(tmp_path)
    source_hash = hashlib.sha256(_source_pdf()).hexdigest()
    snapshot = bind_raw_page_map(raw, mapping, source_pdf_sha256=source_hash,
                                 selected_pages=(2, 3), target_lang="EN")
    assert snapshot.page_ids == {2: ("p000001", "p000002"), 3: ("p000003",)}
    assert snapshot.saved_snapshot is not None
    changed = deepcopy(mapping)
    changed["pages"][1]["blocks"][0]["location"] = changed["pages"][0]["blocks"][0]["location"]
    with pytest.raises(OrdinaryAutoArtifactError, match="ambiguous_raw_page"):
        bind_raw_page_map(raw, changed, source_pdf_sha256=source_hash,
                          selected_pages=(2, 3), target_lang="EN")


def test_raw_binding_reads_nested_table_paths_and_section_controls():
    document = Document()
    document.add_paragraph("Opening unit")
    outer = document.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).paragraphs[0].text = "Nested unit"
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph("Closing unit")
    stream = BytesIO()
    document.save(stream)
    raw = stream.getvalue()
    mapping = {"version": 1, "docx_sha256": hashlib.sha256(raw).hexdigest(),
               "source_page_count": 2, "pages": [
                   {"source_page_number": 2, "blocks": [
                       {"location": {"kind": "body_paragraph", "paragraph_index": 0}},
                       {"location": {"kind": "layout_region_paragraph", "paragraph_index": 0,
                           "table_path": [{"table_index": 0, "row": 0, "col": 0},
                                          {"table_index": 0, "row": 0, "col": 0}]}}]},
                   {"source_page_number": 3, "blocks": [
                       {"location": {"kind": "body_paragraph", "paragraph_index": 2}}]}]}
    snapshot = bind_raw_page_map(raw, mapping, source_pdf_sha256="a" * 64,
                                 selected_pages=(2, 3), target_lang="EN")
    assert [row.text for row in snapshot.paragraphs if not row.control] == [
        "Opening unit", "Nested unit", "Closing unit"]
    assert snapshot.page_ids[2] == ("p000001", "p000003")
    assert snapshot.page_ids[3] == ("p000006",)
    assert any(row.control for row in snapshot.paragraphs)
    assert snapshot.saved_snapshot is None


def test_unreviewed_candidate_service_preserves_review_gate_and_provenance(tmp_path):
    raw, mapping = _raw(tmp_path)
    source = _source_pdf()
    service = SavedDocxLayoutService(tmp_path / "app", mode="shadow", workspace_id="fixture")
    imported = service.import_document(source, raw, "EN", uuid.uuid4().hex)
    review_id = imported["review_id"]
    generation = imported["generation"]
    evidence = {"version": "fictional_proposal_v1", "document_reviewed": False,
                "rendered_layout_acceptance": "not_evaluated"}
    policy = "c" * 64
    service.attach_ordinary_binding(review_id, mapping, (2, 3), policy, evidence,
                                    expected_generation=generation)
    with pytest.raises(SavedDocxLayoutServiceError, match="review_required"):
        service.build(review_id, generation, uuid.uuid4().hex, review_confirmed=False)
    built = service.build_unreviewed_candidate(review_id, generation, uuid.uuid4().hex)
    verified = service.verified_unreviewed_candidate(review_id, built["candidate_id"])
    assert verified.docx_sha256 == built["docx_sha256"]
    assert verified.raw_docx_sha256 == hashlib.sha256(raw).hexdigest()
    assert verified.source_pdf_sha256 == hashlib.sha256(source).hexdigest()
    assert verified.selected_pages == (2, 3)
    assert verified.policy_fingerprint == policy
    assert verified.receipt["document_reviewed"] is False
    assert verified.source_map["geometry_basis"] == "model_proposal_unreviewed"
    assert verified.receipt["rendered_layout_acceptance"] == "not_evaluated"
    assert (tmp_path / "pages" / "page_0002.txt").read_text(encoding="utf-8").startswith("First")
    with pytest.raises(SavedDocxLayoutServiceError, match="ordinary_binding_changed"):
        service.attach_ordinary_binding(review_id, mapping, (2, 3), "d" * 64, evidence,
                                        expected_generation=generation)


@pytest.mark.parametrize(("field", "wrong"), [
    ("docx_sha256", "b" * 64), ("raw_docx_sha256", "b" * 64),
    ("raw_source_map_sha256", "b" * 64), ("source_pdf_sha256", "b" * 64),
    ("decisions_sha256", "b" * 64), ("target_lang", "AR"),
    ("source_character_coverage", "verified"), ("exact_text_preserved", False),
])
def test_rehashed_false_receipt_claims_are_rejected(tmp_path, field, wrong):
    raw, snapshot, frames, decisions, evidence, artifact = _candidate_fixture(tmp_path)
    changed = artifact.receipt
    changed[field] = wrong
    encoded = json.dumps(changed, sort_keys=True, ensure_ascii=False, allow_nan=False,
                         separators=(",", ":")).encode("utf-8")
    forged = replace(artifact, receipt_bytes=encoded,
                     receipt_sha256=hashlib.sha256(encoded).hexdigest())
    with pytest.raises(OrdinaryAutoArtifactError, match="candidate_provenance_changed"):
        verify_unreviewed_candidate(raw, snapshot, frames, decisions, forged,
                                    proposal_evidence=evidence)


@pytest.mark.parametrize(("field", "wrong"), [
    ("raw_source_map_sha256", "b" * 64), ("source_pdf_sha256", "b" * 64),
    ("raw_snapshot_fingerprint", "b" * 64), ("proposal_evidence_sha256", "b" * 64),
    ("source_association_basis", "operator_verified"),
])
def test_rehashed_false_source_map_claims_are_rejected(tmp_path, field, wrong):
    raw, snapshot, frames, decisions, evidence, artifact = _candidate_fixture(tmp_path)
    mapping = artifact.source_map
    mapping[field] = wrong
    map_bytes = json.dumps(mapping, sort_keys=True, ensure_ascii=False, allow_nan=False,
                           separators=(",", ":")).encode("utf-8")
    receipt = artifact.receipt
    receipt["source_map_sha256"] = hashlib.sha256(map_bytes).hexdigest()
    receipt_bytes = json.dumps(receipt, sort_keys=True, ensure_ascii=False, allow_nan=False,
                               separators=(",", ":")).encode("utf-8")
    forged = replace(artifact, source_map_bytes=map_bytes,
        source_map_sha256=hashlib.sha256(map_bytes).hexdigest(),
        receipt_bytes=receipt_bytes,
        receipt_sha256=hashlib.sha256(receipt_bytes).hexdigest())
    with pytest.raises(OrdinaryAutoArtifactError, match="candidate_provenance_changed"):
        verify_unreviewed_candidate(raw, snapshot, frames, decisions, forged,
                                    proposal_evidence=evidence)


@pytest.mark.parametrize("wrong_page", [1, 3])
def test_proposal_region_must_match_raw_owner_not_other_or_excluded_page(tmp_path, wrong_page):
    raw, snapshot, frames, decisions, evidence, artifact = _candidate_fixture(tmp_path)
    assert artifact.source_map["paragraphs"][0]["raw_source_page_number"] == 2
    changed = deepcopy(decisions)
    changed["paragraphs"][0]["regions"][0]["page_number"] = wrong_page
    with pytest.raises(OrdinaryAutoArtifactError, match="proposal_page_owner_mismatch"):
        build_unreviewed_candidate(raw, snapshot, frames, changed,
                                   proposal_evidence=evidence)
