"""Server-owned selected metrics survive Save and staged Gmail validation."""
from dataclasses import replace
from hashlib import sha256
from pathlib import Path
from types import SimpleNamespace
from copy import deepcopy

from docx import Document
import pytest

from legalpdf_translate import translation_service as service
from legalpdf_translate.ordinary_layout_contracts import DeliveryArtifact
from legalpdf_translate.ordinary_text_correction import count_correction_words
from tests.test_translation_completion_metrics import completion_seed, completion_job


def fixture_delivery(tmp_path, kind="text_corrected"):
    path = tmp_path / "selected.docx"
    doc = Document()
    doc.add_paragraph(" ".join(["body"] * 147))
    doc.sections[0].footer.paragraphs[0].text = " ".join(["source"] * 12)
    doc.save(path)
    count = count_correction_words(path.read_bytes())
    assert count == 159
    artifact = DeliveryArtifact("fictional-job", "fictional-run", path,
        sha256(path.read_bytes()).hexdigest(), count, "AR", "a" * 64,
        1, "b" * 32, kind, False)
    seed = completion_seed(path)
    seed.update(word_count=159, rate_per_word=.027, expected_total=4.29,
                api_cost=.33127325)
    return path, artifact, seed


@pytest.fixture(autouse=True)
def no_capability_probes(monkeypatch):
    monkeypatch.setattr(service, "build_translation_capability_flags", lambda **_: {})


def test_save_preserves_verified_footer_count_and_auto_fee(tmp_path):
    path, artifact, seed = fixture_delivery(tmp_path)
    assert service.reviewed_translation_word_count(path) == 147
    response = service.save_translation_row(settings_path=tmp_path / "settings.json",
        job_log_db_path=tmp_path / "jobs.sqlite", form_values={**seed, "word_count": 1, "expected_total_mode": "auto"},
        seed_payload=seed, word_count_docx=path, verified_delivery=artifact)
    row = response["saved_result"]
    assert row["word_count"] == 159
    assert response["normalized_payload"]["expected_total"] == 4.29
    assert response["normalized_payload"]["api_cost"] == .33127325
    assert sha256(path.read_bytes()).hexdigest() == artifact.sha256


@pytest.mark.parametrize("mutation", ["hash", "run", "language", "type", "bytes", "missing"])
def test_invalid_selected_context_refuses_before_writes(tmp_path, mutation):
    path, artifact, seed = fixture_delivery(tmp_path)
    bad = {"hash": lambda: replace(artifact, sha256="0" * 64),
           "run": lambda: replace(artifact, run_id="other"),
           "language": lambda: replace(artifact, target_lang="EN"),
           "type": lambda: {"word_count": 159}}.get(mutation, lambda: artifact)()
    if mutation == "bytes":
        path.write_bytes(path.read_bytes() + b"changed")
    if mutation == "missing":
        path.unlink()
    with pytest.raises(ValueError, match="verified|could not be read"):
        service.save_translation_row(settings_path=tmp_path / "settings.json",
            job_log_db_path=tmp_path / "jobs.sqlite", form_values=seed,
            seed_payload=seed, word_count_docx=path, verified_delivery=bad)
    assert not (tmp_path / "jobs.sqlite").exists()
    assert not (tmp_path / "settings.json").exists()


def test_client_seed_and_form_cannot_supply_count_context(tmp_path):
    path, artifact, seed = fixture_delivery(tmp_path)
    spoof = {**seed, "verified_delivery": artifact, "owned_story_count": 159}
    _, row = service.validate_translation_row(job_log_db_path=tmp_path / "jobs.sqlite",
        form_values=spoof, seed_payload=seed, word_count_docx=path)
    assert row["word_count"] == 147


@pytest.mark.parametrize("mutation", [None, "locked_count", "locked_selection", "locked_generation", "staged_bytes"])
def test_gmail_prospective_and_staged_save_reuse_verified_count(tmp_path, monkeypatch, mutation):
    from tests.test_gmail_browser_service import _load_result
    from legalpdf_translate.gmail_browser_service import GmailBrowserSessionManager
    from legalpdf_translate.gmail_batch import DownloadedGmailAttachment, GmailBatchSession
    import legalpdf_translate.gmail_browser_service as gmail
    import legalpdf_translate.ordinary_layout_integration as integration
    monkeypatch.setattr(gmail, "build_gmail_browser_capability_flags", lambda **_: {})
    path, artifact, seed = fixture_delivery(tmp_path)
    loaded = _load_result(message_id="message", thread_id="thread", subject="Fictional",
        account_email="user@example.test", attachment_ids=("attachment",))
    source = tmp_path / "source.pdf"; source.write_bytes(b"fictional")
    attachment = DownloadedGmailAttachment(candidate=loaded.message.attachments[0], saved_path=source, start_page=1, page_count=1)
    session = GmailBatchSession(intake_context=loaded.intake_context, message=loaded.message,
        gog_path=loaded.gog_path, account_email=loaded.account_email,
        downloaded_attachments=(attachment,), download_dir=tmp_path, selected_target_lang="AR")
    manager = GmailBrowserSessionManager()
    manager._store_loaded_result(runtime_mode="shadow", workspace_id="fictional", result=loaded)
    manager._workspace(runtime_mode="shadow", workspace_id="fictional").batch_session = session
    job = completion_job(path); job["result"]["save_seed"] = seed
    job["config"] = {"source_path": str(source), "start_page": 1}
    monkeypatch.setattr(integration, "delivery_job_snapshot", lambda *a, **k: (job, artifact))
    locked = {"locked_count": replace(artifact, word_count=147),
              "locked_selection": replace(artifact, selection_id="d" * 32),
              "locked_generation": replace(artifact, generation=2)}.get(mutation, artifact)
    ordinary = SimpleNamespace(state=lambda _: {"status": "ready"}, resolve_delivery=lambda *a, **k: locked)
    staged = []
    real_stage = gmail.stage_gmail_batch_translated_docx
    def stage(**kwargs):
        result = real_stage(**kwargs); staged.append(result)
        if mutation == "staged_bytes": result.write_bytes(result.read_bytes() + b"changed")
        return result
    monkeypatch.setattr(gmail, "stage_gmail_batch_translated_docx", stage)
    if mutation:
        with pytest.raises(ValueError, match="delivery_changed|staged_delivery_changed"):
            manager.confirm_current_batch_translation(runtime_mode="shadow", workspace_id="fictional",
                settings_path=tmp_path / "settings.json", job_log_db_path=tmp_path / "jobs.sqlite",
                translation_jobs=SimpleNamespace(get_job=lambda _: job), job_id="fictional-job",
                form_values=seed, ordinary_layout_manager=ordinary, baseline_id="c" * 32,
                expected_delivery_generation=1)
        assert not (tmp_path / "jobs.sqlite").exists() and session.confirmed_items == []
        if mutation != "staged_bytes": assert staged == []
        return
    result = manager.confirm_current_batch_translation(runtime_mode="shadow", workspace_id="fictional",
        settings_path=tmp_path / "settings.json", job_log_db_path=tmp_path / "jobs.sqlite",
        translation_jobs=SimpleNamespace(get_job=lambda _: job), job_id="fictional-job",
        form_values=seed, ordinary_layout_manager=ordinary, baseline_id="c" * 32,
        expected_delivery_generation=1)
    item = session.confirmed_items[0]
    assert item.translated_word_count == 159
    assert sha256(item.staged_translated_docx_path.read_bytes()).hexdigest() == artifact.sha256
    history = service.list_translation_history(db_path=tmp_path / "jobs.sqlite")
    assert history[0]["row"]["word_count"] == 159
    assert history[0]["row"]["expected_total"] == 4.29
    # Stop at the real fee builder boundary; no export or draft dispatch occurs.
    monkeypatch.setattr(manager, "preflight_batch_finalization", lambda **_: {"normalized_payload": {"finalization_preflight": {"finalization_ready": True}}})
    monkeypatch.setattr(gmail, "write_gmail_batch_session_report", lambda _: None)
    import legalpdf_translate.interpretation_service as interpretation
    monkeypatch.setattr(interpretation, "_current_profile", lambda **_: ([], "fictional", SimpleNamespace()))
    monkeypatch.setattr(interpretation, "_profile_missing_fields", lambda _: [])
    captured = {}
    class FeeBoundary(Exception): pass
    def capture_fee(**kwargs):
        captured.update(kwargs); raise FeeBoundary()
    monkeypatch.setattr("legalpdf_translate.honorarios_docx.build_honorarios_draft", capture_fee)
    with pytest.raises(FeeBoundary):
        manager.finalize_batch(runtime_mode="shadow", workspace_id="fictional",
            settings_path=tmp_path / "settings.json", output_filename="", profile_id=None)
    assert captured["word_count"] == 159


def test_browser_save_route_passes_server_selected_count(tmp_path, monkeypatch):
    from tests.test_shadow_web_api import _build_app, _completed_ar_job
    import legalpdf_translate.shadow_web.app as web
    path, artifact, seed = fixture_delivery(tmp_path)
    with _build_app(tmp_path, monkeypatch) as client:
        job = _completed_ar_job(path)
        job["config"]["target_lang"] = "EN"
        seed["target_lang"] = "EN"
        job["result"]["save_seed"] = seed
        artifact = replace(artifact, job_id=job["job_id"], target_lang="EN")
        monkeypatch.setattr(client.app.state.shadow_context.translation_jobs, "get_job", lambda _: job)
        monkeypatch.setattr(web, "delivery_job_snapshot", lambda *a, **k: (job, artifact))
        shown = client.get("/api/translation/jobs/" + job["job_id"])
        assert shown.status_code == 200
        assert shown.json()["normalized_payload"]["job"]["result"]["save_seed"]["word_count"] == 159
        saved = client.post("/api/translation/save-row", json={"job_id":job["job_id"],
            "form_values":{**seed,"word_count":147,"expected_total_mode":"auto"},
            "seed_payload":{**seed,"word_count":999}})
        assert saved.status_code == 200, saved.text
        assert saved.json()["normalized_payload"]["word_count"] == 159
        assert saved.json()["normalized_payload"]["expected_total"] == 4.29


def test_real_correction_resolver_count_is_used_for_save(tmp_path, monkeypatch):
    from io import BytesIO
    from tests.test_ordinary_layout_service import make_case, nonce
    from tests.test_ordinary_text_correction import action
    case = make_case(tmp_path, monkeypatch)
    doc = Document(BytesIO(case.job.original_docx))
    doc.sections[0].footer.paragraphs[0].text = " ".join(["source"] * 12)
    buf = BytesIO(); doc.save(buf)
    job = replace(case.job, original_docx=buf.getvalue(), reviewed_docx=buf.getvalue(),
        mapping_docx_sha256=sha256(buf.getvalue()).hexdigest())
    # Prepare the changed fictional source under a fresh independently owned service.
    from legalpdf_translate.ordinary_layout_manager import OrdinaryLayoutManager
    manager = OrdinaryLayoutManager(tmp_path / "new-owner", mode="shadow", workspace_id="fixture",
        job_resolver=lambda _: job)
    manager.prepare(job.job_id, nonce())
    state = manager.service.text_correction_state(job)
    draft = manager.service.draft_text_correction(job, nonce(), state["parent"],
        [action(state["paragraphs"][0]["paragraph_id"], "Corrected fictional wording")])
    manager.service.approve_text_correction(job, draft["draft_id"], nonce(), True, True, "Compared source")
    row = manager.service.state(job.job_id)["delivery"]
    manager.service.review_text_output(job, row["selection_id"], row["generation"], True)
    delivery = manager.resolve_delivery(job.job_id, row["generation"], require_settled=True)
    assert delivery.word_count == service.reviewed_translation_word_count(delivery.path) + 12
    seed = completion_seed(delivery.path)
    seed.update(run_id=delivery.run_id, target_lang=delivery.target_lang, word_count=delivery.word_count,
        rate_per_word=.027, expected_total_mode="auto")
    # The mode is a form option, never a seed field.
    form = dict(seed); seed.pop("expected_total_mode")
    result = service.save_translation_row(settings_path=tmp_path / "saved-settings.json",
        job_log_db_path=tmp_path / "saved.sqlite", form_values=form, seed_payload=seed,
        word_count_docx=delivery.path, verified_delivery=delivery)
    assert result["saved_result"]["word_count"] == delivery.word_count


@pytest.mark.parametrize("lang", ["EN", "FR", "AR"])
def test_verified_automatic_v7_owned_map_count_survives_save(tmp_path, lang):
    from tests.test_ordinary_source_layout import fixture, build, verify
    from legalpdf_translate.ordinary_layout_service import _candidate_word_count
    args = fixture(lang); candidate = build(args); verify(candidate, args)
    path = tmp_path / "automatic.docx"; path.write_bytes(candidate.docx_bytes)
    count = _candidate_word_count(candidate, path)
    assert count > service.reviewed_translation_word_count(path)
    artifact = DeliveryArtifact("fictional-job", "fictional-run", path, sha256(path.read_bytes()).hexdigest(),
        count, lang, "a" * 64, 0, "b" * 32, "automatic_unreviewed", False)
    seed = completion_seed(path); seed.update(target_lang=lang, word_count=count, rate_per_word=.027)
    _, row = service.validate_translation_row(job_log_db_path=tmp_path / "unused.sqlite",
        form_values={**seed,"expected_total_mode":"auto"}, seed_payload=seed,
        word_count_docx=path, verified_delivery=artifact)
    assert row["word_count"] == count
    from legalpdf_translate.joblog_flow import translation_fee_eur
    assert row["expected_total"] == translation_fee_eur(count,.027)
