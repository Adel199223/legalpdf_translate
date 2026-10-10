from copy import deepcopy

import pytest

from legalpdf_translate.ordinary_layout_contracts import (
    PROPOSAL_VERSION, PROPOSAL_VERSION_V2, PROPOSAL_VERSION_V3, OrdinaryLayoutError,
    _resolve_partition_anchors, normalize_proposals, proposal_schema,
)
from legalpdf_translate.saved_docx_layout import inspect_docx
from tests.test_ordinary_layout_contracts import proposal
from tests.test_ordinary_layout_service import make_case


def test_fresh_schema_requires_bounded_literal_anchors_and_preserves_historical_version():
    schema = proposal_schema(1, ["p000001"])["schema"]
    assert PROPOSAL_VERSION == "ordinary_layout_proposal_v1"
    assert schema["properties"]["version"]["enum"] == [PROPOSAL_VERSION_V3]
    historical = proposal_schema(1, ["p000001"], version=PROPOSAL_VERSION_V2)["schema"]
    assert historical["properties"]["version"]["enum"] == [PROPOSAL_VERSION_V2]
    assert historical["properties"]["paragraph_partitions"] == schema["properties"]["paragraph_partitions"]
    assert "paragraph_partitions" in schema["required"]
    item = schema["properties"]["paragraph_partitions"]["items"]
    assert item["additionalProperties"] is False
    anchors = item["properties"]["split_before"]
    assert anchors["minItems"] == 1 and anchors["maxItems"] == 7
    assert anchors["items"]["maxLength"] == 160


def test_anchor_resolution_preserves_unicode_and_boundary_spaces_without_numeric_model_offsets():
    text = "First sentence. Second sentence. Third sentence."
    snapshot = {"paragraphs": [{"id": "p000001", "text": text}]}
    value = {"paragraph_partitions": [{"paragraph_id": "p000001",
              "split_before": ["Second sentence.", "Third sentence."]}]}
    result = _resolve_partition_anchors(value, snapshot, ["p000001"])
    offsets = result[0]["offsets"]
    assert offsets == [16, 33]
    children = [text[a:b] for a, b in zip([0] + offsets, offsets + [len(text)])]
    assert "".join(children) == text and children[0].endswith(" ")


@pytest.mark.parametrize("anchors", [[], ["missing"], ["word"], ["other", "word"],
                                    ["other", "other"], [" "], ["x" * 161], [1]])
def test_anchors_fail_closed_for_missing_ambiguous_unordered_or_unbounded_values(anchors):
    value = {"paragraph_partitions": [{"paragraph_id": "p000001", "split_before": anchors}]}
    snapshot = {"paragraphs": [{"id": "p000001", "text": "word word other"}]}
    with pytest.raises(OrdinaryLayoutError):
        _resolve_partition_anchors(value, snapshot, ["p000001"])


def test_overlapping_occurrences_are_ambiguous():
    with pytest.raises(OrdinaryLayoutError, match="anchor_ambiguous"):
        _resolve_partition_anchors({"paragraph_partitions": [{"paragraph_id": "p000001",
            "split_before": ["aaa"]}]}, {"paragraphs": [{"id": "p000001", "text": "aaaa"}]}, ["p000001"])


def test_parent_order_identity_and_page_capacity_are_enforced():
    ids = [f"p{i:06d}" for i in range(1, 201)]
    snapshot = {"paragraphs": [{"id": pid, "text": "one two"} for pid in ids]}
    cut = {"paragraph_id": ids[0], "split_before": ["two"]}
    with pytest.raises(OrdinaryLayoutError, match="page_too_large"):
        _resolve_partition_anchors({"paragraph_partitions": [cut]}, snapshot, ids)
    for rows in ([cut, cut], [{**cut, "paragraph_id": "unknown"}],
                 [{**cut, "paragraph_id": ids[1]}, cut], [{**cut, "offsets": [4]}]):
        with pytest.raises(OrdinaryLayoutError):
            _resolve_partition_anchors({"paragraph_partitions": rows}, snapshot, ids[:2])


def test_empty_v2_partitions_normalize_identically_to_retained_v1(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    view = case.view["review"]
    old = proposal(view)
    new = deepcopy(old)
    new.update(version=PROPOSAL_VERSION_V2, paragraph_partitions=[])
    snapshot = inspect_docx(case.job.reviewed_docx, "EN")
    assert normalize_proposals(snapshot, view, [old]) == normalize_proposals(snapshot, view, [new])
    assert old["version"] == PROPOSAL_VERSION and "paragraph_partitions" not in old


def test_retained_v2_response_crosses_fresh_v3_sdk_and_accounting_boundary(tmp_path, monkeypatch):
    import json
    import tests.test_ordinary_layout_manager as manager_fixture
    case = make_case(tmp_path, monkeypatch)
    old_proposal = manager_fixture.proposal
    def response(view):
        value = old_proposal(view)
        value.update(version=PROPOSAL_VERSION_V2, paragraph_partitions=[])
        return value
    monkeypatch.setattr(manager_fixture, "proposal", response)
    paid = manager_fixture.enable_paid(case)
    result = manager_fixture.suggest(case)
    assert result["status"] == "applied_unreviewed"
    assert len(paid.calls) == 1
    request = paid.calls[0]
    assert request["text"]["format"]["schema"]["properties"]["version"]["enum"] == [PROPOSAL_VERSION_V3]
    prompt = json.loads(request["input"][0]["content"][0]["text"])
    assert prompt["version"] == PROPOSAL_VERSION_V3
    assert result["accounting"]["complete"] is True


@pytest.mark.parametrize("version,include_partitions", [(PROPOSAL_VERSION, True), (PROPOSAL_VERSION_V2, False)])
def test_version_specific_response_shape_is_exact(tmp_path, monkeypatch, version, include_partitions):
    case = make_case(tmp_path, monkeypatch)
    value = proposal(case.view["review"])
    value["version"] = version
    if include_partitions:
        value["paragraph_partitions"] = []
    with pytest.raises(OrdinaryLayoutError, match="invalid_proposal"):
        normalize_proposals(inspect_docx(case.job.reviewed_docx, "EN"), case.view["review"], [value])


@pytest.mark.parametrize('language', ['EN', 'FR', 'AR'])
def test_v2_nonempty_partition_uses_shared_validator_without_changing_parent_text(tmp_path, monkeypatch, language):
    from io import BytesIO
    from docx import Document
    import tests.test_ordinary_layout_service as service_fixture
    texts = {'EN': 'First phrase. Second phrase. Third phrase.',
             'FR': 'Première phrase. Deuxième phrase. Troisième phrase.',
             'AR': 'الجملة الأولى. الجملة الثانية. الجملة الثالثة.'}
    text = texts[language]
    anchors = {'EN': ['Second phrase.', 'Third phrase.'],
               'FR': ['Deuxième phrase.', 'Troisième phrase.'],
               'AR': ['الجملة الثانية.', 'الجملة الثالثة.']}[language]
    def raw(*args, **kwargs):
        doc = Document();doc.add_paragraph('Fictional introduction');doc.add_paragraph(text)
        stream = BytesIO();doc.save(stream);return stream.getvalue()
    monkeypatch.setattr(service_fixture, 'docx_bytes', raw)
    case = make_case(tmp_path, monkeypatch, language)
    value = proposal(case.view['review'])
    value.update(version=PROPOSAL_VERSION_V2,
                 paragraph_partitions=[{'paragraph_id': 'p000002', 'split_before': anchors}])
    value['bands'] = [{'kind': 'flow', 'groups': [{'paragraph_ids': ['p000001', 'p000002'], 'panel': False}]}]
    snapshot = inspect_docx(case.job.reviewed_docx, language)
    retained = deepcopy(snapshot)
    result = normalize_proposals(snapshot, case.view['review'], [value])
    assert result['version'] == 'saved_docx_layout_decisions_v2'
    assert result['paragraph_partitions'] == [{'paragraph_id': 'p000002',
                                              'offsets': [text.index(anchor) for anchor in anchors]}]
    assert snapshot == retained
    assert [row['paragraph_id'] for row in result['paragraphs']] == ['p000001', 'p000002']
