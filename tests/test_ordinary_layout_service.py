"""Fictional documents only: local review ownership and immutable delivery."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
from types import SimpleNamespace
import uuid

from docx import Document
import fitz
from PIL import Image
import pytest

from legalpdf_translate import saved_docx_layout_service as storage
from legalpdf_translate.ordinary_layout_contracts import OrdinaryLayoutJob, OrdinaryLayoutError, digest
from legalpdf_translate.ordinary_layout_manager import OrdinaryLayoutManager
from legalpdf_translate.saved_docx_layout import inspect_docx


def nonce():
    return uuid.uuid4().hex


def docx_bytes(lang="EN", suffix="", pages=1):
    doc = Document()
    for _ in range(pages):
        doc.add_paragraph({"EN": "Fictional court", "FR": "Tribunal fictif", "AR": "محكمة خيالية"}[lang] + suffix)
        doc.add_paragraph({"EN": "Preserve all words.", "FR": "Préserver tous les mots.", "AR": "الحفاظ على كل الكلمات."}[lang])
    stream = BytesIO(); doc.save(stream)
    return stream.getvalue()


def pdf_bytes(pages=1):
    with fitz.open() as pdf:
        for _ in range(pages):
            pdf.new_page(width=240, height=320).insert_text((20, 40), "Fictional court")
        return pdf.tobytes()


def fake_prepare(folder, language):
    snapshot = inspect_docx(storage._read(folder / "original.docx", storage.DOCX_MAX_BYTES), language)
    storage._write(folder / "snapshot.json", snapshot, storage.SNAPSHOT_MAX_BYTES)
    stream = BytesIO(); Image.new("RGB", (240, 320), "white").save(stream, format="PNG")
    raw = stream.getvalue()
    storage._mkdir(folder / "images")
    with fitz.open(folder / "source.pdf") as source:
        count = len(source)
    for number in range(1, count + 1):
        storage._atomic(folder / "images" / f"page-{number:04d}.png", raw)
    storage._write(folder / "pages.json", [{"page_number": number, "width_px": 240, "height_px": 320,
        "image_sha256": digest(raw), "rotation": 0, "render_dpi": 144, "backend": "fictional", "backend_version": "1"}
        for number in range(1, count + 1)])


def make_case(tmp_path, monkeypatch, lang="EN", mapping=True, pages=1):
    monkeypatch.setattr(storage, "_prepare_import", fake_prepare)
    raw = docx_bytes(lang, pages=pages)
    job = OrdinaryLayoutJob("tx-fictional", "shadow", "fixture", "run-fictional", pdf_bytes(pages), raw, raw,
        lang, tuple(range(1, pages + 1)), {"session": "fictional", "attachment": "fictional"},
        {p: (f"p{2*p-1:06d}", f"p{2*p:06d}") for p in range(1, pages + 1)} if mapping else {}, digest(raw) if mapping else "")
    state = {"job": job}
    manager = OrdinaryLayoutManager(tmp_path / "isolated", mode="shadow", workspace_id="fixture",
                                    job_resolver=lambda _: state["job"])
    view = manager.prepare(job.job_id, nonce())
    return SimpleNamespace(job=job, state=state, manager=manager, view=view, root=tmp_path / "isolated")


def built(case):
    saved = case.manager.service.saved
    view = saved.read(case.view["review"]["review_id"])
    decisions = deepcopy(view["decisions"])
    decisions["paragraphs"][0].update(role="heading", heading_level=1, heading_size_pt=12, bold=True)
    decisions["review"].update(reviewer="Fictional operator", note="All fictional source pages checked.",
                                document_reviewed=True, pages_reviewed=[1])
    view = saved.save_decisions(view["review_id"], view["generation"], nonce(), decisions)
    return saved.build(view["review_id"], view["generation"], nonce(), True)


def select(case, artifact, *, accepted=True):
    manager, baseline = case.manager, case.view["baseline_id"]
    if accepted:
        manager.accept_output(case.job.job_id, artifact["artifact_id"], artifact["generation"], nonce(), True,
                              expected_baseline_id=baseline)
    return manager.select_delivery(case.job.job_id, 0, nonce(), "reviewed", artifact["review_id"],
        artifact["artifact_id"], artifact["generation"], expected_baseline_id=baseline)


@pytest.mark.parametrize("language", ["EN", "FR", "AR"])
def test_auto_bound_import_build_accept_select_and_restart(tmp_path, monkeypatch, language):
    case = make_case(tmp_path, monkeypatch, language)
    assert case.view["review"]["decisions"]["review"]["document_reviewed"] is False
    assert all(p["regions"] for p in case.view["review"]["decisions"]["paragraphs"])
    artifact = built(case)
    with pytest.raises(OrdinaryLayoutError, match="output_review_required"):
        select(case, artifact, accepted=False)
    state = select(case, artifact)
    delivery = case.manager.resolve_delivery(case.job.job_id, state["delivery_generation"])
    assert delivery.path.read_bytes() == case.manager.service.saved.artifact(artifact["review_id"], artifact["artifact_id"], "docx")
    assert delivery.word_count > 0 and not delivery.frozen
    assert [r["text"] for r in inspect_docx(delivery.path.read_bytes(), language)["paragraphs"]] == [
        r["text"] for r in inspect_docx(case.job.reviewed_docx, language)["paragraphs"]]
    reopened = OrdinaryLayoutManager(case.root, mode="shadow", workspace_id="fixture", job_resolver=lambda _: case.job)
    assert reopened.resolve_delivery(case.job.job_id, 1) == delivery
    folder = reopened.service.folder(case.job.job_id) / "baselines" / case.view["baseline_id"]
    assert (folder / "provider.docx").read_bytes() == case.job.original_docx
    assert (folder / "reviewed.docx").read_bytes() == case.job.reviewed_docx


def test_build_not_acceptance_and_original_requires_explicit_choice(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    with pytest.raises(OrdinaryLayoutError, match="ordinary_layout_confirmation_required"):
        case.manager.select_delivery(case.job.job_id, 0, nonce(), "original", expected_baseline_id=case.view["baseline_id"])
    view = case.manager.select_delivery(case.job.job_id, 0, nonce(), "original", keep_ordinary_confirmed=True,
                                        expected_baseline_id=case.view["baseline_id"])
    assert view["delivery"]["sha256"] == digest(case.job.reviewed_docx)


def test_owner_source_current_word_and_mapping_guards(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    case.state["job"] = replace(case.job, reviewed_docx=docx_bytes(suffix=" changed"))
    assert case.manager.state(case.job.job_id)["stale"]
    with pytest.raises(OrdinaryLayoutError, match="baseline_stale"):
        case.manager.select_delivery(case.job.job_id, 0, nonce(), "original", keep_ordinary_confirmed=True,
                                        expected_baseline_id=case.view["baseline_id"])
    with pytest.raises(OrdinaryLayoutError, match="page_mapping_stale"):
        case.manager.prepare(case.job.job_id, nonce())
    case.state["job"] = replace(case.job, mode="live")
    with pytest.raises(OrdinaryLayoutError, match="job_owner_mismatch"):
        case.manager.prepare(case.job.job_id, nonce())


def test_generations_nonces_and_frozen_copy_survive_later_independent_edits(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    artifact = built(case)
    state = select(case, artifact)
    selected = case.manager.resolve_delivery(case.job.job_id, 1, nonce())
    old = selected.path.read_bytes()
    view = case.manager.service.saved.read(artifact["review_id"])
    changed = deepcopy(view["decisions"]); changed["paragraphs"][1]["bold"] = True
    changed["review"].update(document_reviewed=False, pages_reviewed=[])
    case.manager.service.saved.save_decisions(view["review_id"], view["generation"], nonce(), changed)
    case.state["job"] = replace(case.job, reviewed_docx=docx_bytes(suffix=" later interactive edit"))
    assert case.manager.resolve_delivery(case.job.job_id, 1).path.read_bytes() == old
    with pytest.raises(OrdinaryLayoutError, match="delivery_frozen"):
        case.manager.prepare(case.job.job_id, nonce())
    with pytest.raises(OrdinaryLayoutError):
        case.manager.resolve_delivery(case.job.job_id, 2)
    selected.path.write_bytes(old + b"changed")
    with pytest.raises(OrdinaryLayoutError, match="delivery_changed"):
        case.manager.resolve_delivery(case.job.job_id, 1)


def test_unfrozen_selection_rejects_new_review_generation(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch); artifact = built(case); select(case, artifact)
    view = case.manager.service.saved.read(artifact["review_id"])
    changed = deepcopy(view["decisions"]); changed["paragraphs"][1]["italic"] = True
    case.manager.service.saved.save_decisions(view["review_id"], view["generation"], nonce(), changed)
    with pytest.raises(OrdinaryLayoutError, match="delivery_stale"):
        case.manager.resolve_delivery(case.job.job_id, 1)


def test_new_baseline_rejects_old_browser_generation_even_when_number_matches(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    newer = case.manager.prepare(case.job.job_id, nonce())
    assert newer["generation"] == case.view["generation"]
    with pytest.raises(OrdinaryLayoutError, match="baseline_stale"):
        case.manager.select_delivery(case.job.job_id, 0, nonce(), "original", keep_ordinary_confirmed=True,
                                        expected_baseline_id=case.view["baseline_id"])


def test_selection_retry_precedes_generation_and_body_change_conflicts(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    chosen = nonce()
    first = case.manager.select_delivery(case.job.job_id, 0, chosen, "original", keep_ordinary_confirmed=True,
                                         expected_baseline_id=case.view["baseline_id"])
    again = case.manager.select_delivery(case.job.job_id, 0, chosen, "original", keep_ordinary_confirmed=True,
                                         expected_baseline_id=case.view["baseline_id"])
    assert first == again
    with pytest.raises(OrdinaryLayoutError, match="nonce_conflict"):
        case.manager.select_delivery(case.job.job_id, 1, chosen, "original", keep_ordinary_confirmed=True,
                                     expected_baseline_id=case.view["baseline_id"])


def test_no_current_handle_after_restart_keeps_read_only_state(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    def missing(_):
        raise ValueError("private path")
    reopened = OrdinaryLayoutManager(case.root, mode="shadow", workspace_id="fixture", job_resolver=missing)
    assert reopened.state(case.job.job_id)["attached"] is False
    with pytest.raises(OrdinaryLayoutError, match="job_unavailable"):
        reopened.prepare(case.job.job_id, nonce())


def test_open_review_guard_prevents_accepting_changed_review_copy(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch); artifact = built(case)
    calls = []
    def guard(*values):
        calls.append(values)
        raise OrdinaryLayoutError("review_copy_changed", 409)
    case.manager.output_review_guard = guard
    with pytest.raises(OrdinaryLayoutError, match="review_copy_changed"):
        select(case, artifact)
    assert calls == [(case.job.job_id, artifact["artifact_id"], artifact["generation"], case.view["baseline_id"])]
    assert case.manager.state(case.job.job_id)["output_reviews"] == []
    checked = case.manager.review_artifact(case.job.job_id, artifact["artifact_id"], artifact["generation"],
        expected_baseline_id=case.view["baseline_id"])
    assert checked.docx_bytes == case.manager.service.saved.artifact(artifact["review_id"], artifact["artifact_id"], "docx")


def test_saved_service_wrong_owner_is_rejected(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    with pytest.raises(OrdinaryLayoutError, match="review_owner_mismatch"):
        OrdinaryLayoutManager(case.root, mode="live", workspace_id="fixture", job_resolver=lambda _: case.job,
            saved_service=case.manager.service.saved)


def test_prepare_recovers_only_unpublished_immediate_pointer(tmp_path, monkeypatch):
    case = make_case(tmp_path, monkeypatch)
    operation = nonce()
    atomic = storage._atomic
    failed = False
    def fail_pointer(path, raw):
        nonlocal failed
        if path.parent.name == "prepared" and not failed:
            failed = True
            raise OSError("fictional interrupted publication")
        return atomic(path, raw)
    monkeypatch.setattr(storage, "_atomic", fail_pointer)
    with pytest.raises(OrdinaryLayoutError):
        case.manager.prepare(case.job.job_id, operation)
    recovered = case.manager.prepare(case.job.job_id, operation)
    assert recovered["baseline_id"] == operation
    assert recovered["review"]["decisions"]["review"]["document_reviewed"] is False
