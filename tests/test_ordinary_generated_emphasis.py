from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from legalpdf_translate.ordinary_layout_contracts import _generated_emphasis, OrdinaryLayoutError
from legalpdf_translate.saved_docx_layout import inspect_docx, _emphasis, SavedDocxLayoutError

def row(text="Alpha beta gamma."):
    document=Document();document.add_paragraph(text);stream=BytesIO();document.save(stream)
    return inspect_docx(stream.getvalue(),"EN")["paragraphs"][0]

def span(start,end,**changes):
    return {"start":start,"end":end,"bold":True,"italic":False,"underline":False}|changes

def test_drops_only_word_interior_span_and_preserves_valid_styles():
    original=[span(1,5),span(6,10,italic=True)]
    before=deepcopy(original)
    assert _generated_emphasis(row(),original)==[original[1]]
    assert original==before
    with pytest.raises(SavedDocxLayoutError,match="unsafe_emphasis_boundary"):_emphasis(row(),original)

@pytest.mark.parametrize("invalid",[
    span(1,3,unknown=True),span(True,3),span(1,3,bold=1),span(1,3,bold=False),
    span(-1,3),span(1,999),span(3,3),span(5,3),
])
def test_entire_original_shape_is_rejected_before_filtering(invalid):
    with pytest.raises(OrdinaryLayoutError):_generated_emphasis(row(),[invalid])

@pytest.mark.parametrize("spans",[[span(6,8),span(1,3)],[span(1,8),span(6,10)]])
def test_bad_order_or_overlap_is_never_hidden_by_dropping(spans):
    with pytest.raises(OrdinaryLayoutError):_generated_emphasis(row(),spans)

@pytest.mark.parametrize("text,start,end",[("A\u200dB value",1,3),("A\u2066B value",1,3)])
def test_nonlexical_or_control_edges_remain_rejected(text,start,end):
    with pytest.raises(OrdinaryLayoutError):_generated_emphasis(row(text),[span(start,end)])

def test_protected_identifier_interior_remains_rejected():
    # Existing literal-token protection remains strict even for lexical endpoints.
    document=Document();p=document.add_paragraph("نص ");run=p.add_run("[[AB123456789CD]]")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rtl=OxmlElement("w:rtl");rtl.set(qn("w:val"),"0");run._r.get_or_add_rPr().append(rtl)
    stream=BytesIO();document.save(stream);snapshot=inspect_docx(stream.getvalue(),"AR")
    with pytest.raises(OrdinaryLayoutError):_generated_emphasis(snapshot["paragraphs"][0],[span(4,8)])

def test_nontext_token_overlap_remains_rejected():
    value=row();value["tokens"].append({"kind":"tab","start":2,"end":3})
    with pytest.raises(OrdinaryLayoutError):_generated_emphasis(value,[span(1,5)])

@pytest.mark.parametrize("text",["Alpha beta gamma.","École café fin.","نص عربي طويل"])
def test_optional_interior_drop_never_expands_or_rewrites_text(text):
    value=row(text);before=deepcopy(value)
    assert _generated_emphasis(value,[span(1,len(text))])==[]
    assert value==before


@pytest.mark.parametrize("text,cut", [("الجهة：عنوان", 6), ("Label:Value", 6), ("A/B value", 1), ("A—B value", 1)])
def test_safe_punctuation_edges_drop_optional_styles_without_changing_text(text, cut):
    value = row(text); before = deepcopy(value)
    styles = [span(0, cut), span(cut, len(text), bold=False, underline=True)]
    rejected = []
    assert _generated_emphasis(value, styles, rejected=rejected) == []
    assert len(rejected) == 2
    assert all(item["reason"] == "unsupported_optional_emphasis_edge" for item in rejected)
    assert value == before
    with pytest.raises(OrdinaryLayoutError):
        _generated_emphasis(value, styles, drop_unsupported_edges=False)
    with pytest.raises(SavedDocxLayoutError):
        _emphasis(value, styles)


def test_unsupported_safe_span_does_not_hide_later_control_or_overlap():
    value = row("Label:Value")
    with pytest.raises(OrdinaryLayoutError):
        _generated_emphasis(value, [span(0, 6), span(5, 11)])
    value = row("Label:Value A\u200dB")
    with pytest.raises(OrdinaryLayoutError):
        _generated_emphasis(value, [span(0, 6), span(13, 15)])


def test_emphasis_evidence_is_versioned_reviewable_and_legacy_absence_preserved():
    from legalpdf_translate.ordinary_layout_contracts import proposal_source_evidence, source_evidence_review_fields
    from tests.test_ordinary_source_evidence import fixture
    snapshot, view, proposal = fixture()
    actual = row("Label:Value"); actual["id"] = "p000001"
    snapshot["paragraphs"] = [actual]
    proposal["paragraphs"][0]["emphasis"] = [span(0, 6), span(6, 11)]
    original = deepcopy(proposal)
    result = proposal_source_evidence(snapshot, view, proposal)
    evidence = result["optional_emphasis_normalization"]
    assert evidence["version"] == "ordinary_optional_emphasis_normalization_v1"
    assert len(evidence["rejected_spans"]) == 2
    assert source_evidence_review_fields([result])["source_layout_review_required"] is True
    assert "optional_emphasis_normalization" not in proposal_source_evidence(
        snapshot, view, proposal, emphasis_normalization_version=None)
    assert proposal == original



def test_settled_optional_edge_recovery_has_no_dispatch_and_binds_warning(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from tests import test_ordinary_auto_layout_workflow as flow
    from legalpdf_translate import ordinary_layout_manager as manager_module, ordinary_auto_layout as automatic
    from legalpdf_translate.ordinary_layout_contracts import PROPOSAL_VERSION_V3, normalize_proposals, OrdinaryLayoutError
    from legalpdf_translate.ordinary_layout_service import _read, _write
    original = flow._proposal
    def proposal(view, page, columns=0):
        result = original(view, page, columns)
        result.update(version=PROPOSAL_VERSION_V3, paragraph_partitions=[], source_coverage_findings=[], source_layout_evidence=[])
        result['paragraphs'][0]['emphasis'] = [span(0, 6), span(6, 11)]
        return result
    monkeypatch.setattr(flow, '_proposal', proposal)
    source = tmp_path/'source.pdf'; flow._source(source, ('digital',))
    output = tmp_path/'output'; output.mkdir()
    app, jobs, sdk, requests = flow._app(tmp_path, monkeypatch, ('Label:Value\nOther body.',), (1,))
    scope = '?mode=shadow&workspace=fictional'
    with TestClient(app) as client:
        with monkeypatch.context() as historical:
            historical.setattr(manager_module, 'normalize_proposals',
                lambda *args: normalize_proposals(*args, emphasis_normalization_version=None))
            job_id = flow._start(client, source, output, 'EN', start=1, end=1, image_mode='off')
            initial = flow._wait(jobs, job_id)
        assert initial['result']['automatic_layout']['status'] == 'raw_fallback'
        manager = app.state.ordinary_layouts.manager_for_context(app.state.shadow_context, 'shadow', 'fictional')
        failed = [_read(path) for path in manager.service.root.rglob('result.json')]
        assert any(item.get('error_code') == 'ordinary_layout_invalid_proposal_decisions' for item in failed)
        prior_cost = manager.layout_costs(job_id)['known_cost_usd']
        provider = manager.provider_factory
        def forbidden(*args):
            value = provider(*args)
            value._client.responses.create = lambda **kwargs: pytest.fail('Retained style recovery must not dispatch')
            return value
        manager.provider_factory = forbidden
        checker = automatic._proposal_evidence
        tampered = []
        def missing_marker(manager, origin, baseline, review, nonce, *args, **kwargs):
            path = manager.service.root/origin/'baselines'/baseline/'suggestions'/nonce/'source_evidence.json'
            before = path.read_bytes(); changed = _read(path)
            changed['pages'][0].pop('optional_emphasis_normalization')
            path.unlink(); _write(path, changed)
            result_path = path.with_name('result.json'); before_result = result_path.read_bytes()
            result = _read(result_path)
            import hashlib
            from legalpdf_translate.ordinary_layout_contracts import encode
            result['source_layout_evidence_sha256'] = hashlib.sha256(encode(changed['pages'])).hexdigest()
            result_path.unlink(); _write(result_path, result)
            try:
                with pytest.raises(OrdinaryLayoutError, match='source_evidence_response_changed'):
                    checker(manager, origin, baseline, review, nonce, *args, **kwargs)
                tampered.append(True)
                raise OrdinaryLayoutError('ordinary_layout_source_evidence_response_changed', 409)
            finally:
                path.write_bytes(before)
                result_path.write_bytes(before_result)
        with monkeypatch.context() as mutation:
            mutation.setattr(automatic, '_proposal_evidence', missing_marker)
            assert client.post(f'/api/translation/jobs/{job_id}/layout/recover'+scope).status_code == 200
            assert flow._wait(jobs, job_id)['result']['automatic_layout']['status'] == 'raw_fallback'
        assert tampered == [True]
        assert client.post(f'/api/translation/jobs/{job_id}/layout/recover'+scope).status_code == 200
        final = flow._wait(jobs, job_id)['result']['automatic_layout']
        assert final['status'] == 'automatic_unreviewed', final
        assert final['source_layout_review_required'] is True
        assert len(final['source_layout_rejected_hints']) == 2
        assert all(item['kind'] == 'optional_emphasis' for item in final['source_layout_rejected_hints'])
        from decimal import Decimal
        assert Decimal(final['layout_costs']['known_cost_usd']) == Decimal(prior_cost)
        assert final['layout_costs']['complete'] is True
        assert client.post(f'/api/translation/jobs/{job_id}/layout/recover'+scope).status_code == 200
        again = flow._wait(jobs, job_id)['result']['automatic_layout']
        assert again['sha256'] == final['sha256']
        assert again['layout_costs'] == final['layout_costs']
        assert again['source_layout_rejected_hints'] == final['source_layout_rejected_hints']
        assert len(requests) == len(sdk.requests) == 1


@pytest.mark.parametrize("version", ["ordinary_layout_proposal_v1", "ordinary_layout_proposal_v2", "ordinary_layout_proposal_v3"])
def test_only_v3_can_drop_new_edges_with_bound_evidence(tmp_path, monkeypatch, version):
    from tests import test_ordinary_layout_service as service
    from tests.test_ordinary_layout_contracts import proposal
    from legalpdf_translate.ordinary_layout_contracts import normalize_proposals
    document = Document(); document.add_paragraph("Label:Value"); document.add_paragraph("Other body.")
    stream = BytesIO(); document.save(stream); raw = stream.getvalue()
    monkeypatch.setattr(service, 'docx_bytes', lambda *args, **kwargs: raw)
    case = service.make_case(tmp_path, monkeypatch)
    view = case.view['review']; value = proposal(view)
    value['version'] = version; value['paragraphs'][0]['emphasis'] = [span(0, 6), span(6, 11)]
    if version != 'ordinary_layout_proposal_v1': value['paragraph_partitions'] = []
    if version == 'ordinary_layout_proposal_v3':
        value.update(source_coverage_findings=[], source_layout_evidence=[])
        result = normalize_proposals(inspect_docx(raw, 'EN'), view, [value])
        assert result['paragraphs'][0]['emphasis'] == []
    else:
        with pytest.raises(OrdinaryLayoutError, match='invalid_proposal_decisions'):
            normalize_proposals(inspect_docx(raw, 'EN'), view, [value])
