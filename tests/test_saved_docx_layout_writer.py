"""Fictional source-layout derivatives; package tests are not visual acceptance."""
from copy import deepcopy
import io
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree
import pytest

from legalpdf_translate import saved_docx_layout as model
from legalpdf_translate import saved_docx_layout_writer as writer
from legalpdf_translate.docx_writer import _add_rtl_flags, _set_rtl_run_props, _set_ltr_run_props
from tests.test_saved_docx_layout import document_bytes, changed_part, reviewed


def packet(lang="EN"):
    texts = {
        "EN": ["Example Public Office", "Reference 12/2026/AB", "A clear notice", "Keep this whole phrase unchanged.", "Further information", "Final original paragraph"],
        "FR": ["Service public exemple", "Référence 12/2026/AB", "Un avis clair", "Conserver la de\u0301claration complète intacte.", "Informations complémentaires", "Dernier paragraphe original"],
        "AR": ["مكتب عام تجريبي", "مرجع محفوظ", "إشعار واضح", "يُحفظ النَصّ الكامل دون تغيير", "معلومات إضافية", "الفقرة الأخيرة الأصلية"],
    }[lang]
    def edit(doc):
        if lang == "AR":
            for paragraph in doc.paragraphs:
                _add_rtl_flags(paragraph)
                for run in paragraph.runs:
                    _set_rtl_run_props(run, bidi_lang="ar-SA")
            paragraph = doc.paragraphs[1]
            literal = paragraph.add_run(" \u200e12/2026/AB\u200e")
            _set_ltr_run_props(literal)
        doc.paragraphs[-1].runs[-1].add_break(WD_BREAK.PAGE)
        doc.sections[0].footer.paragraphs[0].text = "Inherited footer"
    raw = document_bytes(texts, edit)
    snapshot, pages, dec = reviewed(raw, lang)
    ids = [p["id"] for p in snapshot["paragraphs"]]
    group = lambda identifiers, panel=False: {"paragraph_ids": identifiers, "panel": panel}
    dec["bands"] = [
        {"kind": "columns", "widths_pct": [50, 50], "gutter_pt": 12,
         "cells": [{"groups": [group(ids[:1])]}, {"groups": [group(ids[1:2])]}]},
        {"kind": "columns", "widths_pct": [40, 60], "gutter_pt": 16,
         "cells": [{"groups": [group(ids[2:3], True)]}, {"groups": [group(ids[3:5])]}]},
        {"kind": "flow", "groups": [group(ids[5:])]},
    ]
    dec["paragraphs"][0].update(bold=True, space_after_pt=8)
    dec["paragraphs"][2].update(role="heading", heading_level=1, heading_size_pt=16,
                                 bold=True, underline=True, space_before_pt=4, space_after_pt=6)
    dec["paragraphs"][4].update(role="heading", heading_level=2, heading_size_pt=12, bold=True)
    phrase = {"EN": "this whole phrase", "FR": "la de\u0301claration complète", "AR": "النَصّ الكامل"}[lang]
    start = snapshot["paragraphs"][3]["text"].index(phrase)
    dec["paragraphs"][3]["emphasis"] = [{"start": start, "end": start + len(phrase), "bold": True, "italic": True, "underline": False}]
    for choice in dec["paragraphs"]:
        choice["alignment"] = "right" if lang == "AR" else "left"
    return raw, snapshot, pages, dec


def xml(raw):
    with ZipFile(io.BytesIO(raw)) as package:
        return etree.fromstring(package.read("word/document.xml"))


def located(root, path):
    return root.getroottree().xpath(path, namespaces=root.nsmap)[0]


@pytest.mark.parametrize("lang", ["EN", "FR", "AR"])
def test_structural_derivative_preserves_text_order_packages_and_editability(lang):
    raw, snapshot, pages, dec = packet(lang)
    before = deepcopy((snapshot, pages, dec))
    artifact = writer.build_docx(raw, snapshot, pages, dec)
    assert (snapshot, pages, dec) == before
    writer.validate_built_docx(artifact.docx_bytes, artifact.source_map, original_docx=raw,
                              snapshot=snapshot, pages=pages, decisions=dec)
    root = xml(artifact.docx_bytes)
    with ZipFile(io.BytesIO(raw)) as original, ZipFile(io.BytesIO(artifact.docx_bytes)) as output:
        assert original.namelist() == output.namelist()
        assert all(original.read(name) == output.read(name) for name in original.namelist() if name != "word/document.xml")
    mapping = artifact.source_map
    assert mapping["exact_text_preserved"] and mapping["logical_order_preserved"]
    assert mapping["source_character_coverage"] == "not_proven"
    assert mapping["rendered_layout_acceptance"] == "not_evaluated" and mapping["rendered_page_count"] is None
    assert not list(root.iter(qn("w:trHeight"))) and not list(root.iter(qn("w:cantSplit")))
    assert len(list(root.iter(qn("w:tbl")))) == 3
    assert all(n.get(qn("w:val")) == "0" for n in root.iter(qn("w:bidiVisual")))
    assert [n.get(qn("w:fill")) for n in root.iter(qn("w:shd"))] == ["E7E7E7"]
    output_rows = [located(root, row["location"]) for row in mapping["paragraphs"]]
    assert [p for p in root[0].iter(qn("w:p")) if p in output_rows] == output_rows
    assert output_rows[-1].getparent().tag == qn("w:body")
    assert output_rows[-1].find(".//" + qn("w:br")).get(qn("w:type")) == "page"
    heading = output_rows[2]
    assert heading.find(qn("w:pPr") + "/" + qn("w:outlineLvl")).get(qn("w:val")) == "0"
    assert heading.find(qn("w:pPr") + "/" + qn("w:keepNext")).get(qn("w:val")) == "1"
    assert all(n.get(qn("w:val")) == "32" for n in heading.iter(qn("w:sz")))
    assert all(n.get(qn("w:val")) == "32" for n in heading.iter(qn("w:szCs")))
    for name in ("b", "bCs", "u"):
        assert list(heading.iter(qn("w:" + name)))
    phrase_p = output_rows[3]
    emphasized = [r for r in phrase_p.findall(qn("w:r")) if r.find(qn("w:rPr") + "/" + qn("w:b")) is not None]
    assert emphasized and len(phrase_p.findall(qn("w:r"))) >= 3
    span = dec["paragraphs"][3]["emphasis"][0]
    assert "".join(n.text or "" for r in emphasized for n in r.findall(qn("w:t"))) == snapshot["paragraphs"][3]["text"][span["start"]:span["end"]]
    if lang == "AR":
        for paragraph in output_rows:
            assert paragraph.find(qn("w:pPr") + "/" + qn("w:bidi")).get(qn("w:val")) == "1"
            assert paragraph.find(qn("w:pPr") + "/" + qn("w:jc")).get(qn("w:val")) == "start"
        literal_runs = [r for r in output_rows[1].findall(qn("w:r")) if "12/2026/AB" in "".join(r.itertext())]
        assert len(literal_runs) == 1
        assert literal_runs[0].find(qn("w:rPr") + "/" + qn("w:rtl")).get(qn("w:val")) == "0"


@pytest.mark.parametrize("baseline,direct", [(30, True), (15, True), (15, False)])
def test_heading_size_never_shrinks_effective_original_or_changes_style_part(baseline, direct):
    def edit(doc):
        if direct:
            doc.paragraphs[0].runs[0].font.size = Pt(baseline)
        else:
            doc.paragraphs[0].runs[0].font.size = None
            doc.styles["Normal"].font.size = Pt(baseline)
    raw = document_bytes(("A heading",), edit)
    snapshot, pages, dec = reviewed(raw)
    dec["paragraphs"][0].update(role="heading", heading_level=1, heading_size_pt=12)
    artifact = writer.build_docx(raw, snapshot, pages, dec)
    assert next(xml(artifact.docx_bytes).iter(qn("w:sz"))).get(qn("w:val")) == str(baseline * 2)


@pytest.mark.parametrize("lang,bidi,expected", [("AR", False, "left"), ("AR", True, "end"), ("EN", True, "end")])
@pytest.mark.parametrize("inherited", [False, True])
def test_physical_alignment_uses_preserved_effective_bidi_not_requested_language(lang, bidi, expected, inherited):
    def edit(doc):
        parent = doc.styles["Normal"]._element.get_or_add_pPr() if inherited else doc.paragraphs[0]._p.get_or_add_pPr()
        node = OxmlElement("w:bidi")
        node.set(qn("w:val"), "1" if bidi else "0")
        parent.append(node)
    raw = document_bytes(("Visible text",), edit)
    snapshot, pages, dec = reviewed(raw, lang)
    dec["paragraphs"][0]["alignment"] = "left"
    result = writer.build_docx(raw, snapshot, pages, dec)
    assert next(xml(result.docx_bytes).iter(qn("w:jc"))).get(qn("w:val")) == expected


def test_empty_whitespace_and_typed_tokens_survive_new_structural_empty_paragraphs():
    raw = document_bytes(("", "  ", "Line one\nLine two\tTail"))
    snapshot, pages, dec = reviewed(raw)
    ids = [p["id"] for p in snapshot["paragraphs"]]
    dec["bands"][0]["groups"] = [{"paragraph_ids": ids, "panel": True}]
    result = writer.build_docx(raw, snapshot, pages, dec)
    root = xml(result.docx_bytes)
    assert len(result.source_map["structural_paragraphs"]) == 1
    assert len(list(root[0].iter(qn("w:p")))) == 4
    assert len(list(root.iter(qn("w:tab")))) == len(list(root.iter(qn("w:br")))) == 1


def test_source_map_retains_explicit_unmapped_exception_without_provenance_fiction():
    raw, snapshot, pages, dec = packet()
    dec["paragraphs"][1].update(regions=[], unmapped_reason="No visible source counterpart; saved text retained.")
    result = writer.build_docx(raw, snapshot, pages, dec)
    assert result.source_map["qualifications"] == [{"paragraph_id": "p000002", "reason": dec["paragraphs"][1]["unmapped_reason"]}]
    assert "commit_file_sha256" not in str(result.source_map)
    assert result.source_map["review"]["reviewer_kind"] == "operator_review"


@pytest.mark.parametrize("partial", [False, True])
def test_adding_underline_retains_inherited_double_underline(partial):
    def edit(doc):
        props = doc.styles["Normal"]._element.get_or_add_rPr()
        props.append(OxmlElement("w:u"))
        props[-1].set(qn("w:val"), "double")
    raw = document_bytes(("Keep this phrase unchanged.",), edit)
    snapshot, pages, dec = reviewed(raw)
    if partial:
        dec["paragraphs"][0]["emphasis"] = [{"start": 5, "end": 16, "bold": False, "italic": False, "underline": True}]
    else:
        dec["paragraphs"][0]["underline"] = True
    result = writer.build_docx(raw, snapshot, pages, dec)
    assert not list(xml(result.docx_bytes).iter(qn("w:u")))
    with ZipFile(io.BytesIO(raw)) as before, ZipFile(io.BytesIO(result.docx_bytes)) as after:
        assert before.read("word/styles.xml") == after.read("word/styles.xml")


def test_independent_checker_detects_assembler_loss_of_xml_space(monkeypatch):
    raw = document_bytes(("  Retain edge whitespace  ",))
    snapshot, pages, dec = reviewed(raw)
    real = writer._styled_paragraph
    def faulty(*args, **kwargs):
        p = real(*args, **kwargs)
        for node in p.iter(qn("w:t")):
            node.attrib.pop(model.XML_SPACE, None)
        return p
    monkeypatch.setattr(writer, "_styled_paragraph", faulty)
    with pytest.raises(model.SavedDocxLayoutError, match="output_text_space_semantics_changed"):
        writer.build_docx(raw, snapshot, pages, dec)


@pytest.mark.parametrize("change", ["text", "rtl", "lang", "font", "extra_bold", "heading_size", "paragraph_bidi", "paragraph_indent"])
def test_independent_baseline_checker_catches_assembler_fault_repeated_by_expected_scaffold(monkeypatch, change):
    raw, snapshot, pages, dec = packet("AR")
    real = writer._styled_paragraph
    def faulty(*args, **kwargs):
        paragraph = real(*args, **kwargs)
        run = paragraph.find(qn("w:r"))
        props = run.find(qn("w:rPr"))
        if change == "text":
            run.find(qn("w:t")).text += " corruption"
        elif change in {"rtl", "lang", "font"}:
            tag, attr, value = {"rtl": ("rtl", "val", "0"), "lang": ("lang", "val", "en-US"), "font": ("rFonts", "ascii", "Courier New")}[change]
            node = props.find(qn("w:" + tag))
            if node is None:
                node = OxmlElement("w:" + tag)
                props.append(node)
            node.set(qn("w:" + attr), value)
        elif change in {"extra_bold", "heading_size"}:
            node = props.find(qn("w:b" if change == "extra_bold" else "w:sz"))
            if node is None:
                node = OxmlElement("w:b" if change == "extra_bold" else "w:sz")
                props.append(node)
            node.set(qn("w:val"), "0" if change == "extra_bold" else "4")
        else:
            ppr = paragraph.find(qn("w:pPr"))
            if change == "paragraph_bidi":
                ppr.find(qn("w:bidi")).set(qn("w:val"), "0")
            else:
                node = OxmlElement("w:ind")
                node.set(qn("w:left"), "1000")
                ppr.append(node)
        return paragraph
    monkeypatch.setattr(writer, "_styled_paragraph", faulty)
    with pytest.raises(model.SavedDocxLayoutError, match="output_(content_changed|baseline_run_semantics_changed|unauthorized_emphasis|unauthorized_heading_size|baseline_paragraph_semantics_changed)"):
        writer.build_docx(raw, snapshot, pages, dec)


@pytest.mark.parametrize("mutation", ["table_bidi", "panel", "extra_paragraph", "changed_footer", "changed_map", "changed_snapshot"])
def test_reopened_artifact_and_mapping_tampering_is_rejected(mutation):
    raw, snapshot, pages, dec = packet()
    artifact = writer.build_docx(raw, snapshot, pages, dec)
    mapping, output = deepcopy(artifact.source_map), artifact.docx_bytes
    if mutation == "changed_map": mapping["rendered_layout_acceptance"] = "accepted"
    elif mutation == "changed_snapshot": snapshot["paragraphs"][0]["text"] += " changed"
    else:
        def change(root):
            if mutation == "table_bidi": next(root.iter(qn("w:bidiVisual"))).set(qn("w:val"), "1")
            elif mutation == "panel": next(root.iter(qn("w:shd"))).set(qn("w:fill"), "FFFFFF")
            elif mutation == "extra_paragraph": root[0].insert(0, etree.Element(qn("w:p")))
            else: next(root.iter(qn("w:t"))).text += " altered"
        output = changed_part(output, "word/footer1.xml" if mutation == "changed_footer" else "word/document.xml", change)
        mapping["docx_sha256"] = model._sha(output)
    with pytest.raises(model.SavedDocxLayoutError):
        writer.validate_built_docx(output, mapping, original_docx=raw, snapshot=snapshot, pages=pages, decisions=dec)


@pytest.mark.parametrize("lang", ["EN", "FR", "AR"])
def test_successive_equal_column_bands_pair_headings_and_bodies_without_reordering(lang):
    texts = {
        "EN": ["Notice", "Response", "Keep the notice condition.", "Response remains optional."],
        "FR": ["Avis", "Réponse", "Conserver la condition de l'avis.", "La réponse reste facultative."],
        "AR": ["إشعار", "رد", "يُحفظ شرط الإشعار", "يبقى الرد اختيارياً"],
    }[lang]
    def edit(doc):
        if lang == "AR":
            for paragraph in doc.paragraphs:
                _add_rtl_flags(paragraph)
                for run in paragraph.runs:
                    _set_rtl_run_props(run, bidi_lang="ar-SA")
    raw = document_bytes(texts, edit)
    snapshot, pages, decisions = reviewed(raw, lang)
    ids = [row["id"] for row in snapshot["paragraphs"]]
    decisions["bands"] = [{"kind": "columns", "widths_pct": [50, 50], "gutter_pt": 12,
        "cells": [{"groups": [{"paragraph_ids": [identifier], "panel": False}]}
                  for identifier in pair]} for pair in (ids[:2], ids[2:])]
    for choice in decisions["paragraphs"][:2]:
        choice.update(role="heading", heading_level=1, heading_size_pt=12, bold=True)
    artifact = writer.build_docx(raw, snapshot, pages, decisions)
    root = xml(artifact.docx_bytes)
    tables = root.find(qn("w:body")).findall(qn("w:tbl"))
    assert len(tables) == 2
    grids = [[n.get(qn("w:w")) for n in table.find(qn("w:tblGrid"))] for table in tables]
    assert grids[0] == grids[1] and grids[0][0] == grids[0][1]
    cells = [table.find(qn("w:tr")).findall(qn("w:tc")) for table in tables]
    assert [["".join(cell.itertext()) for cell in row] for row in cells] == [texts[:2], texts[2:]]
    assert artifact.source_map["exact_text_preserved"] and artifact.source_map["logical_order_preserved"]
    output_rows = [located(root, row["location"]) for row in artifact.source_map["paragraphs"]]
    assert ["".join(p.itertext()) for p in output_rows] == texts
    if lang == "AR":
        assert all(p.find(qn("w:pPr") + "/" + qn("w:bidi")).get(qn("w:val")) == "1" for p in output_rows)
