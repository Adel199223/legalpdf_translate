from copy import deepcopy
from dataclasses import replace
import hashlib
import json

from docx import Document
import pytest

from legalpdf_translate import acceptance_assembly as assembly
from legalpdf_translate.document_structure import PageStructure, StructureBlock
from legalpdf_translate.docx_writer import assemble_docx
from legalpdf_translate.formatting_support import fingerprint, digest_text
from legalpdf_translate.structured_artifacts import publish_structured_page, StructuredArtifactError
from legalpdf_translate.types import TargetLang


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture(params=['legal_blocks_reviewed_recovery_v1', 'legal_blocks_legacy_reviewed_recovery_v1'])
def case(tmp_path, request):
    refs = []
    for n, protocol in enumerate(("legal_blocks_v2", "legal_blocks_recovery_v1",
                                 request.param), 1):
        source_text = f"Decisão {n}\nPedido admitido."
        source = PageStructure(page_number=n, source_sha256=digest_text(source_text),
            source_text_sha256=digest_text(source_text), source_file_sha256="a" * 64,
            uncertain=True, blocks=[StructureBlock(f"p{n:04d}_b0001", source_text, uncertain=True)])
        target = source.to_dict()
        text = f"Décision {n}\nDemande admise."
        target["blocks"][0]["text"] = text
        target["translation_sha256"] = digest_text(text)
        identity = {"protocol": protocol, "fingerprint": "b" * 64}
        page_fingerprint = "c" * 64
        page_result = {"usage": {}, "metadata": {}, "image_used": False, "retry_used": False}
        if protocol != "legal_blocks_v2":
            provenance = {"version": protocol, "lang": "FR", "preferences_sha256": "d" * 64}
            identity["fingerprint"] = fingerprint(provenance)
            page_fingerprint = fingerprint({"recovery": identity, "page": n})
            target["metadata"]["recovery"] = provenance
            page_result = None
        folder = tmp_path / f"original_{n}"
        record = publish_structured_page(folder, source_structure=source, translated_structure=target,
            translated_text=text, protocol_identity=identity, page_fingerprint=page_fingerprint,
            page_result=page_result)
        refs.append(assembly.PageBundleRef(folder, record, digest(folder / f"page_{n:04d}.commit.json"), "e" * 64))
    args = dict(page_bundles=refs, full_case_pages=[1, 2, 3], source_file_sha256="a" * 64,
                lang=TargetLang.FR, preferences_sha256="d" * 64, output_dir=tmp_path / "assembled")
    binding = assembly.assembly_binding(**{k: v for k, v in args.items() if k != "output_dir"})
    args["evidence_guard"] = lambda: deepcopy(binding)
    return args


def snapshot(case):
    return {p: p.read_bytes() for ref in case["page_bundles"] for p in ref.pages_dir.iterdir()}


def read_map(output):
    return json.loads(output.with_suffix(".source_map.json").read_text(encoding="utf-8"))


def write_map(output, value):
    output.with_suffix(".source_map.json").write_text(json.dumps(value), encoding="utf-8")


def wrapped_writer(change):
    def writer(folder, output, **kwargs):
        result = assemble_docx(folder, output, **kwargs)
        change(folder, output, kwargs)
        return result
    return writer


def test_mixed_case_uses_exact_original_bytes_once_and_verifies_real_docx(case):
    before, calls = snapshot(case), []
    def writer(folder, output, **kw):
        calls.append((folder, kw.copy()))
        return assemble_docx(folder, output, **kw)
    output = assembly.assemble_mixed_case(**case, writer=writer)
    assert len(calls) == 1
    assert calls[0][0] == output.parent / "pages"
    assert calls[0][1]["page_numbers"] == [1, 2, 3]
    assert calls[0][1]["partial_output"] is False
    assert calls[0][1]["derive_source_continuations"] is True
    assert calls[0][1]["page_breaks"] is False
    for path, content in before.items():
        assert path.read_bytes() == content
        assert (output.parent / "pages" / path.name).read_bytes() == content
    doc = Document(output)
    assert [p.text for p in doc.paragraphs if p.text.strip()] == [
        f"Décision {n}\nDemande admise." for n in range(1, 4)]
    manifest = json.loads((output.parent / "assembly.json").read_bytes())
    assert manifest["original_commits"] == [r.expected_commit for r in case["page_bundles"]]
    assert manifest["provider_dispatch_count"] == 0
    assert manifest["assembly_cost_usd"] == "0"
    assert manifest["workflow_resumed"] is False
    assert manifest["fidelity_acceptance"] == manifest["rendered_layout_acceptance"] == "not_evaluated"
    assert manifest["visible_text_coverage"] == "exact_unique_body_locators"


@pytest.mark.parametrize("pages", [[], [2, 3], [1, 1, 2], [1, 3, 2], [True, 2, 3], [1, 2]])
def test_incomplete_duplicate_or_unordered_selection_rejected(case, pages):
    case["full_case_pages"] = pages
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_mixed_case(**case)
    assert not case["output_dir"].exists()


@pytest.mark.parametrize("field,value", [("source_file_sha256", "f" * 64),
    ("lang", TargetLang.AR), ("preferences_sha256", "f" * 64), ("evidence_guard", None)])
def test_wrong_case_identity_or_missing_guard_rejected(case, field, value):
    case[field] = value
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_mixed_case(**case)
    assert not case["output_dir"].exists()


@pytest.mark.parametrize("field", ["commit_file_sha256", "evidence_identity_sha256"])
def test_changed_reference_binding_rejected(case, field):
    case["page_bundles"][0] = replace(case["page_bundles"][0], **{field: "f" * 64})
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_mixed_case(**case)


@pytest.mark.parametrize("suffix", ["txt", "source_structure.json", "structure.json", "commit.json"])
@pytest.mark.parametrize("missing", [False, True])
def test_original_missing_or_modified_rejected_before_new_output(case, suffix, missing):
    path = case["page_bundles"][0].pages_dir / f"page_0001.{suffix}"
    if missing:
        path.unlink()
    else:
        path.write_bytes(path.read_bytes() + b" ")
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_mixed_case(**case)
    assert not case["output_dir"].exists()


@pytest.mark.parametrize("suffix", ["layout.json", "layout_eligibility.json"])
def test_unbound_original_layout_rejected(case, suffix):
    (case["page_bundles"][0].pages_dir / f"page_0001.{suffix}").write_text("{}")
    with pytest.raises(StructuredArtifactError, match="unbound_layout"):
        assembly.assemble_mixed_case(**case)
    assert not case["output_dir"].exists()


@pytest.mark.parametrize("call_number", [1, 3, 6, 8])
def test_original_mutation_across_guard_phases_fails(case, call_number):
    original_guard = case["evidence_guard"]
    count = 0
    def guard():
        nonlocal count
        count += 1
        if count == call_number:
            (case["page_bundles"][0].pages_dir / "page_0001.txt").write_text("changed")
        return original_guard()
    case["evidence_guard"] = guard
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_mixed_case(**case)
    assert not (case["output_dir"] / "assembly.json").exists()


@pytest.mark.parametrize("extra", ["page_0001.layout.json", "page_0004.txt", "unexpected.json"])
def test_writer_cannot_add_unbound_staged_inputs(case, extra):
    def change(folder, output, kw):
        (folder / extra).write_text("{}")
    with pytest.raises(StructuredArtifactError, match="unexpected_staged"):
        assembly.assemble_mixed_case(**case, writer=wrapped_writer(change))
    assert not (case["output_dir"] / "assembly.json").exists()


def test_writer_cannot_mutate_staged_text(case):
    def change(folder, output, kw):
        (folder / "page_0001.txt").write_text("changed")
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_mixed_case(**case, writer=wrapped_writer(change))


@pytest.mark.parametrize("field,value", [("structured_page_count", 2), ("structure_fallback_count", 1)])
def test_writer_fallback_rejected(case, field, value):
    with pytest.raises(StructuredArtifactError, match="writer_fallback"):
        assembly.assemble_mixed_case(**case, writer=wrapped_writer(lambda f, o, kw: kw["stats"].update({field: value})))


@pytest.mark.parametrize("mutation", [
    lambda m: m.update(docx_sha256="f" * 64),
    lambda m: m.update(source_page_count=2),
    lambda m: m.update(rendered_page_count=3),
    lambda m: m["pages"].reverse(),
    lambda m: m["pages"][0].update(source_file_sha256="f" * 64),
    lambda m: m["pages"][0].update(translation_sha256="f" * 64),
    lambda m: m["pages"][0].update(source_block_coverage_status="incomplete"),
    lambda m: m["pages"][0]["blocks"].clear(),
    lambda m: m["pages"][0]["blocks"][0].update(block_id="wrong"),
    lambda m: m["pages"][0]["blocks"][0].update(location={}),
    lambda m: m["pages"][0]["blocks"][0].update(furniture_alias={}),
])
def test_source_map_must_bind_docx_source_and_each_block(case, mutation):
    def change(folder, output, kw):
        mapping = read_map(output)
        mutation(mapping)
        write_map(output, mapping)
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_mixed_case(**case, writer=wrapped_writer(change))
    assert not (case["output_dir"] / "assembly.json").exists()


@pytest.mark.parametrize("mode", ["alter_text", "extra_paragraph", "extra_header", "extra_table", "swap_order"])
def test_actual_docx_content_is_verified_not_only_claimed_mapping(case, mode):
    def change(folder, output, kw):
        doc, mapping = Document(output), read_map(output)
        if mode == "alter_text":
            doc.paragraphs[0].text += " changed"
        elif mode == "extra_paragraph":
            doc.add_paragraph("Unmapped addition")
        elif mode == "extra_header":
            doc.sections[0].header.paragraphs[0].text = "Unmapped header"
        elif mode == "extra_table":
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "Unmapped table"
        else:
            left, right = [page["blocks"][0]["location"]["paragraph_index"] for page in mapping["pages"][:2]]
            a, b = doc.paragraphs[left].text, doc.paragraphs[right].text
            doc.paragraphs[left].text, doc.paragraphs[right].text = b, a
            mapping["pages"][0]["blocks"][0]["location"]["paragraph_index"] = right
            mapping["pages"][1]["blocks"][0]["location"]["paragraph_index"] = left
        doc.save(output)
        mapping["docx_sha256"] = digest(output)  # Rehashed lies still cannot pass.
        write_map(output, mapping)
    with pytest.raises(StructuredArtifactError):
        assembly.assemble_mixed_case(**case, writer=wrapped_writer(change))
    assert not (case["output_dir"] / "assembly.json").exists()


def test_existing_output_refused_without_overwriting(case):
    case["output_dir"].mkdir()
    existing = case["output_dir"] / "keep.txt"
    existing.write_text("existing user data")
    with pytest.raises(FileExistsError):
        assembly.assemble_mixed_case(**case)
    assert existing.read_text() == "existing user data"


def test_output_cannot_be_nested_with_originals(case):
    case["output_dir"] = case["page_bundles"][0].pages_dir / "child"
    with pytest.raises(StructuredArtifactError, match="overlaps_evidence"):
        assembly.assemble_mixed_case(**case)


def test_parent_traversal_cannot_hide_output_overlap(case):
    original = case["page_bundles"][0].pages_dir
    sibling = original.parent / "sibling"
    sibling.mkdir()
    case["output_dir"] = sibling / ".." / original.name / "new"
    with pytest.raises(StructuredArtifactError, match="absolute_path"):
        assembly.assemble_mixed_case(**case)
    assert not (original / "new").exists()
