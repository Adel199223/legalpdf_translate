"""Synthetic, offline regressions for source-owned isolated contact footers.

DOCX/XML assertions prove editable placement and isolation, not Word pagination.
No private document, provider, credential or native Word operation is used.
"""
from __future__ import annotations

from copy import deepcopy
import json
import socket

from docx import Document
from docx.oxml.ns import qn
import pytest

from legalpdf_translate.document_layout import derive_page_layout
from legalpdf_translate.document_structure import PageStructure, text_sha256
import legalpdf_translate.docx_writer as writer
from legalpdf_translate.layout_integration import prepare_layout_rebuild
from legalpdf_translate.section_furniture import plan_section_furniture
from legalpdf_translate.types import TargetLang
from tests.test_source_folio_layout import synthetic_pair


@pytest.fixture(autouse=True)
def _offline_only(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Isolated-footer fixtures must not use providers, credentials or network")

    monkeypatch.setattr(socket, "create_connection", forbidden)
    monkeypatch.setattr(socket.socket, "connect", forbidden)
    for path in (
        "legalpdf_translate.openai_client.resolve_openai_key_with_source",
        "legalpdf_translate.openai_client.run_translation_auth_test",
        "legalpdf_translate.openai_client.OpenAI",
        "legalpdf_translate.ocr_engine.build_ocr_engine",
        "legalpdf_translate.ocr_engine.resolve_ocr_api_key_source",
        "legalpdf_translate.ocr_engine.OpenAI",
    ):
        monkeypatch.setattr(path, forbidden)


def _rehash_source(source):
    source["source_sha256"] = source["source_text_sha256"] = text_sha256(
        "\n".join(block["text"] for block in source["blocks"])
    )


def _retarget(source, previous=None):
    target = deepcopy(source)
    if previous is not None:
        texts = {block["id"]: block["text"] for block in previous["blocks"]}
        for block in target["blocks"]:
            block["text"] = texts.get(block["id"], block["text"])
    target["translation_sha256"] = text_sha256(
        "\n".join(block["text"] for block in target["blocks"])
    )
    return target


def _pair(number=5, lang=TargetLang.AR):
    source, target = synthetic_pair(number)
    source["blocks"][6]["text"] = "a) João Exemplo;"
    _rehash_source(source)
    if lang == TargetLang.AR:
        target["blocks"][6]["text"] = "أ) João Exemplo؛"
    else:
        for block in target["blocks"]:
            block["text"] = next(item["text"] for item in source["blocks"] if item["id"] == block["id"])
        target["blocks"][6]["text"] = "a) João Exemplo;"
    target = _retarget(source, target)
    target["metadata"]["layout"] = derive_page_layout(source)
    return source, target


def _save_pairs(folder, pairs):
    folder.mkdir(exist_ok=True)
    for source, target in pairs:
        path = folder / f"page_{source['page_number']:04d}.txt"
        path.write_text("\n".join(block["text"] for block in target["blocks"]), encoding="utf-8")
        path.with_suffix(".source_structure.json").write_text(json.dumps(source), encoding="utf-8")
        path.with_suffix(".structure.json").write_text(json.dumps(target), encoding="utf-8")


def _read_output(path):
    return Document(path), json.loads(path.with_suffix(".source_map.json").read_text(encoding="utf-8"))


def _assert_not_adopted(plan, *, review):
    assert not any(section["consolidated"] for section in plan["sections"])
    assert all(not page["adopted_footer_ids"] and not page["adopted_header_ids"] for page in plan["pages"])
    if review:
        assert plan["review_required"]


@pytest.mark.parametrize("number", [1, 5, 7])
def test_only_footer_is_adopted_from_one_complete_source_page_without_mutation(number):
    pair = _pair(number)
    # The source-shaped helper's fraction must remain a valid printed folio.
    pair[0]["blocks"][18]["text"] = "3/5"
    _rehash_source(pair[0])
    pair = (pair[0], _retarget(pair[0], pair[1]))
    pair[1]["blocks"][18]["text"] = "3/5"
    pair[1]["translation_sha256"] = text_sha256("\n".join(b["text"] for b in pair[1]["blocks"]))
    before = deepcopy(pair)
    plan = plan_section_furniture([pair])
    section, = plan["sections"]
    page, = plan["pages"]
    assert section["consolidated"]
    assert section["adoption_basis"] == "isolated_source_contact_footer_v1"
    assert section["parts"]["header"] is None
    assert section["parts"]["footer"]
    assert section["page_numbers"] == [number]
    assert section["source_file_sha256"] == pair[0]["source_file_sha256"]
    assert page["adopted_header_ids"] == []
    assert page["adopted_footer_ids"] == [f"p{number:04d}_b0018", f"p{number:04d}_b0020"]
    assert len(section["parts"]["footer"]["aliases"]) == 2
    assert not plan["review_required"]
    assert pair == before


@pytest.mark.parametrize("guard", [
    "missing_body_geometry", "uncertain_body", "continued_body", "body_outside_page",
    "body_overlaps_footer", "body_below_footer", "footer_gap_under_six_points",
    "overlapping_contacts", "uncertain_footer", "continued_footer", "footer_not_at_bottom",
    "lone_address", "operative_footer", "table_footer", "no_substantive_body", "only_separators",
    "page_continues_from_previous", "page_continues_to_next", "uncertain_page",
    "uncertain_boundary", "interior_document_start", "stale_source", "stale_target",
    "partial_target", "oversized_target_footer", "invalid_layout",
])
def test_unsafe_or_incomplete_single_page_stays_inline_and_review_required(guard):
    source, target = _pair()
    assert plan_section_furniture([(source, target)])["sections"][0]["consolidated"]
    body, address, contact = source["blocks"][10], source["blocks"][17], source["blocks"][19]
    if guard == "missing_body_geometry":
        body["bbox"] = None
    elif guard == "uncertain_body":
        body["uncertain"] = True
    elif guard == "continued_body":
        body["continuation_of"] = "p0004_b0001"
    elif guard == "body_outside_page":
        body["bbox"] = [-1, 373, 524, 425]
    elif guard == "body_overlaps_footer":
        body["bbox"] = [70, 798, 524, 807]
    elif guard == "body_below_footer":
        body["bbox"] = [70, 824, 524, 835]
    elif guard == "footer_gap_under_six_points":
        body["bbox"] = [70, 780, 524, 798]
    elif guard == "overlapping_contacts":
        contact["bbox"][1] = address["bbox"][3] - 1
    elif guard == "uncertain_footer":
        contact["uncertain"] = True
    elif guard == "continued_footer":
        contact["continuation_of"] = "p0004_b0001"
    elif guard == "footer_not_at_bottom":
        contact["bbox"] = [185, 450, 402, 460]
    elif guard == "lone_address":
        source["blocks"].remove(contact)
    elif guard == "operative_footer":
        contact["text"] += "; deve pagar no prazo de cinco dias."
    elif guard == "table_footer":
        contact.update(role="table_cell", table_id="contact_table", row=0, col=0)
    elif guard == "no_substantive_body":
        source["blocks"] = [b for b in source["blocks"] if b["role"] in {"header", "footer"}]
    elif guard == "only_separators":
        source["blocks"] = [b for b in source["blocks"] if b["role"] in {"header", "footer"} or b["text"] == "*"]
    elif guard == "page_continues_from_previous":
        source["continuation_from_previous"] = True
    elif guard == "page_continues_to_next":
        source["continuation_to_next"] = True
    elif guard == "uncertain_page":
        source["uncertain"] = True
    elif guard == "uncertain_boundary":
        source["metadata"]["document_boundary_review_required"] = True
    elif guard == "interior_document_start":
        body["document_start"] = True
    elif guard == "invalid_layout":
        source["metadata"]["layout"] = {"status": "needs_review", "review_required": True, "bands": []}
    _rehash_source(source)
    target = _retarget(source, target)
    # Ensure negative geometry/content fixtures do not pass merely because the
    # source/target pair is stale or malformed for an unrelated reason.
    PageStructure.from_dict(source)
    PageStructure.from_dict(target)
    if guard == "stale_source":
        source["blocks"][0]["text"] += " changed without rehash"
    elif guard == "stale_target":
        target["blocks"][0]["text"] += " changed without rehash"
    elif guard == "partial_target":
        target["blocks"].pop()
        target["translation_sha256"] = text_sha256("\n".join(b["text"] for b in target["blocks"]))
    elif guard == "oversized_target_footer":
        target["blocks"][19]["text"] = "x" * 601
        target["translation_sha256"] = text_sha256("\n".join(b["text"] for b in target["blocks"]))
    before = deepcopy((source, target))
    _assert_not_adopted(plan_section_furniture([(source, target)]), review=True)
    assert (source, target) == before


@pytest.mark.parametrize("barrier", ["changed_contact", "changed_header", "document_start", "missing_before", "missing_after"])
def test_singleton_intervals_inside_larger_plans_do_not_gain_isolated_footers(barrier):
    first, second = _pair(5), _pair(6)
    if barrier.startswith("missing"):
        pairs = [None, first] if barrier == "missing_before" else [first, None]
    else:
        source, target = second
        if barrier == "changed_contact":
            source["blocks"][19]["text"] = "Telef: 220000000 - E-mail: other@example.invalid"
        elif barrier == "changed_header":
            source["blocks"][0]["text"] = "Tribunal Judicial de Outro Exemplo"
        else:
            source["document_start"] = True
        _rehash_source(source)
        pairs = [first, (source, _retarget(source, target))]
    _assert_not_adopted(plan_section_furniture(pairs), review=barrier.startswith("missing"))


def test_two_page_recurrence_still_adopts_both_parts_not_isolated_basis():
    plan = plan_section_furniture([_pair(5), _pair(6)])
    section, = plan["sections"]
    assert section["consolidated"]
    assert section.get("adoption_basis") != "isolated_source_contact_footer_v1"
    assert section["parts"]["header"] and section["parts"]["footer"]
    assert all(page["adopted_header_ids"] and page["adopted_footer_ids"] for page in plan["pages"])


def test_isolated_rejection_preserves_detailed_plan_and_warning_in_saved_source_map(tmp_path):
    source, target = _pair()
    source["blocks"][10]["bbox"] = None
    target = _retarget(source, target)
    folder = tmp_path / "pages"
    _save_pairs(folder, [(source, target)])
    before = {p.name: p.read_bytes() for p in folder.iterdir()}
    document, mapping = _read_output(writer.assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.AR,
                                                       page_breaks=False))
    page, = mapping["pages"]
    warning = "section_furniture_isolated_contact_evidence_unavailable"
    assert mapping["layout_review_required"] and page["layout_review_required"]
    assert warning in page["layout_warnings"]
    plan = mapping["section_furniture"]
    assert plan["review_required"]
    assert warning in plan["pages"][0]["warnings"]
    assert not plan["sections"][0]["consolidated"]
    assert len(page["blocks"]) == 20
    assert all(row["location"]["kind"] == "body_paragraph" for row in page["blocks"]
               if not row["block_id"].endswith("_b0019"))
    assert sum("E-mail" in p.text for p in document.paragraphs) == 1
    assert "E-mail" not in document.sections[0].footer._element.xml
    assert before == {p.name: p.read_bytes() for p in folder.iterdir()}


@pytest.mark.parametrize("lang,font_size,font", [
    (TargetLang.AR, 11, "Arial"), (TargetLang.EN, 10.5, "Times New Roman"),
    (TargetLang.FR, 10.5, "Times New Roman"),
])
def test_docx_footer_has_contacts_page_and_complete_editable_mapping_without_body_duplicates(tmp_path, lang, font_size, font):
    pair = _pair(lang=lang)
    folder = tmp_path / "pages"
    _save_pairs(folder, [pair])
    before = {p.name: p.read_bytes() for p in folder.iterdir()}
    document, mapping = _read_output(writer.assemble_docx(folder, tmp_path / "out.docx", lang=lang, page_breaks=False))
    section, = document.sections
    page, = mapping["pages"]
    plan, = mapping["section_furniture"]["sections"]
    assert plan["adoption_basis"] == "isolated_source_contact_footer_v1"
    assert plan["parts"]["header"] is None
    assert len(section.footer.paragraphs) == 3
    assert " PAGE " in section.footer._element.xml
    assert section._sectPr.find(qn("w:pgNumType")) is None
    assert not section.different_first_page_header_footer
    assert not section.footer.is_linked_to_previous
    assert all(not p.text for p in section.header.paragraphs)
    expected = {b["id"]: b for b in pair[1]["blocks"]}
    assert len(page["blocks"]) == len(expected) == 20
    assert {b["block_id"] for b in page["blocks"]} == set(expected)
    for row in page["blocks"]:
        location = row["location"]
        if row["block_id"].endswith("_b0019"):
            assert location["kind"] == "generated_footer_page_field"
            continue
        is_contact = row["block_id"].endswith(("_b0018", "_b0020"))
        assert location["kind"] == ("section_footer" if is_contact else "body_paragraph")
        paragraph = (section.footer.paragraphs if is_contact else document.paragraphs)[location["paragraph_index"]]
        assert writer.sanitize_bidi_controls(paragraph.text) == expected[row["block_id"]]["text"]
        if is_contact:
            assert location["section_id"] == plan["section_id"]
            assert str(section.footer.part.partname) == location["part_uri"]
            assert paragraph._p.pPr.find(qn("w:jc")).get(qn("w:val")) == "center"
            assert not any(expected[row["block_id"]]["text"] in p.text for p in document.paragraphs)
        for run in paragraph.runs:
            if not run._r.xpath("./w:t"):
                continue
            assert run.font.size.pt == font_size
            fonts = run._r.rPr.find(qn("w:rFonts"))
            assert all(fonts.get(qn(f"w:{slot}")) == font for slot in ("ascii", "hAnsi", "cs"))
    gap = next(item for item in page["source_spacing"]["applied"] if item["block_id"] == "p0005_b0003")
    assert gap["applied_space_before_pt"] == 24
    assert any("João Exemplo" in run.text for p in document.paragraphs for run in p.runs)
    assert "[[" not in document._element.xml
    assert plan["render_reserve"]["footer_body_gap_pt"] == 6
    assert plan["render_reserve"]["word_render_verified"] is False
    assert not document._element.xpath(".//w:br[@w:type='page']")
    assert not document._element.xpath(".//w:pPr/w:sectPr")
    assert not mapping["layout_review_required"]
    assert mapping["rendered_page_count"] is None
    assert before == {p.name: p.read_bytes() for p in folder.iterdir()}


@pytest.mark.parametrize("lang", [TargetLang.AR, TargetLang.EN, TargetLang.FR])
def test_real_footer_reserve_overflow_keeps_all_contacts_in_body_at_readable_size(tmp_path, lang):
    source, target = _pair(lang=lang)
    # Scale actual synthetic geometry to a physically short source page. The
    # contact pair still qualifies but the existing 144pt body reserve cannot.
    scale = 180 / source["height_pt"]
    source["height_pt"] = 180
    for block in source["blocks"]:
        block["bbox"][1] *= scale
        block["bbox"][3] *= scale
    target = _retarget(source, target)
    assert plan_section_furniture([(source, target)])["sections"][0]["consolidated"]
    folder = tmp_path / "pages"
    _save_pairs(folder, [(source, target)])
    document, mapping = _read_output(writer.assemble_docx(folder, tmp_path / "out.docx", lang=lang, page_breaks=False))
    assert mapping["layout_review_required"]
    plan, = mapping["section_furniture"]["sections"]
    assert not plan["consolidated"] and plan["reserve_rejected"]
    assert "section_furniture_reserve_exceeds_page" in mapping["pages"][0]["layout_warnings"]
    assert all(row["location"]["kind"] == "body_paragraph" for row in mapping["pages"][0]["blocks"]
               if not row["block_id"].endswith("_b0019"))
    assert "E-mail" not in document.sections[0].footer._element.xml
    assert sum("E-mail" in p.text for p in document.paragraphs) == 1
    size = 11 if lang == TargetLang.AR else 10.5
    assert all(run.font.size.pt == size for p in document.paragraphs for run in p.runs if run._r.xpath("./w:t"))


def test_explicit_page_matching_keeps_contact_in_body(tmp_path):
    pair = _pair()
    _assert_not_adopted(plan_section_furniture([pair], page_breaks=True), review=False)
    folder = tmp_path / "pages"
    _save_pairs(folder, [pair])
    document, mapping = _read_output(writer.assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.AR, page_breaks=True))
    assert "section_furniture" not in mapping
    assert sum("E-mail" in p.text for p in document.paragraphs) == 1
    assert "E-mail" not in document.sections[0].footer._element.xml


def test_legacy_contact_text_does_not_become_source_proven_footer(tmp_path):
    folder = tmp_path / "pages"
    folder.mkdir()
    text = "Legacy body\nLargo do Exemplo, 1\nTelef: 210000000 - E-mail: court@example.invalid"
    (folder / "page_0005.txt").write_text(text, encoding="utf-8")
    document, mapping = _read_output(writer.assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.EN, page_breaks=False))
    assert "section_furniture" not in mapping
    assert sum("E-mail" in p.text for p in document.paragraphs) == 1
    assert "E-mail" not in document.sections[0].footer._element.xml


def test_up_to_page_does_not_turn_larger_saved_run_into_isolated_footer_output(tmp_path):
    folder = tmp_path / "pages"
    _save_pairs(folder, [_pair(5), _pair(6)])
    before = {p.name: p.read_bytes() for p in folder.iterdir()}
    document, mapping = _read_output(writer.assemble_docx(folder, tmp_path / "out.docx", lang=TargetLang.AR,
                                                       page_breaks=False, up_to_page=5))
    assert [page["source_page_number"] for page in mapping["pages"]] == [5]
    assert sum("E-mail" in p.text for p in document.paragraphs) == 1
    assert "E-mail" not in document.sections[0].footer._element.xml
    assert not mapping.get("section_furniture", {}).get("sections", []) or not any(
        section["consolidated"] for section in mapping["section_furniture"]["sections"])
    assert before == {p.name: p.read_bytes() for p in folder.iterdir()}


def test_cancelled_layout_preparation_writes_no_derivative_or_source_change(tmp_path):
    folder = tmp_path / "pages"
    _save_pairs(folder, [_pair()])
    before = {p.name: p.read_bytes() for p in folder.iterdir()}
    # Cancellation occurs before any page/source rendering. No actual PDF or
    # OCR/credential access is required merely to preserve a cancelled rebuild.
    result = prepare_layout_rebuild(folder, tmp_path / "not-opened.pdf", cancelled=lambda: True)
    assert result["cancelled"] and not result["prepared_pages"]
    assert before == {p.name: p.read_bytes() for p in folder.iterdir()}
