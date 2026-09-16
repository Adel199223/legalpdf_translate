"""Source-bound editable column gutters; synthetic data, no native rendering."""
from copy import deepcopy
from dataclasses import replace
import hashlib
import json

from docx.oxml.ns import qn
import pytest

from legalpdf_translate.reviewed_regions import GUTTER_POLICY, GUTTER_VERSION, ReviewedRegionsError
from legalpdf_translate.reviewed_formatting_writer import (
    ReviewedFormattingWriterError, build_reviewed_docx, validate_reviewed_docx,
)
from legalpdf_translate.reviewed_region_writer import GUTTER_TABLE_POLICY, TABLE_POLICY
from tests.test_reviewed_regions import packet, validate, fid
from tests.test_reviewed_formatting import validate as validate_formatting
from tests.test_reviewed_formatting_v2 import region_packet
from tests.test_reviewed_formatting_writer import changed_document, decoded
from tests.test_reviewed_region_writer import located


def enable(layout, gaps):
    layout.update(version=GUTTER_VERSION, gutter_policy=GUTTER_POLICY)
    for block in layout["body"]:
        if block["kind"] == "table":
            block["column_gaps_px"] = deepcopy(gaps)


def test_gutter_is_explicit_and_bound_to_all_adjacent_source_cell_pairs():
    layout, fragments = packet()
    table = layout["body"][0]
    table["rows"] = [{"cells": [{"fragment_ids": [fid(2)]}, {"fragment_ids": [fid(3)]}]},
                     {"cells": [{"fragment_ids": [fid(4)]}, {"fragment_ids": [fid(5)]}]}]
    fragments[4] = replace(fragments[4], bbox_px=(480, 130, 950, 150))
    enable(layout, [80])
    before = deepcopy((layout, fragments))
    assert json.loads(validate(layout, fragments)) == layout
    assert (layout, fragments) == before
    table["column_gaps_px"] = [80.01]
    with pytest.raises(ReviewedRegionsError, match="region_unsupported_column_gap"):
        validate(layout, fragments)


@pytest.mark.parametrize("gaps", [None, (), "40", [], [0, 0], [True], [-1], [101],
                                  [float("nan")], [float("inf")], [10**1000], ["40"], [{}]])
def test_invalid_or_unsupported_gap_cannot_change_column_layout(gaps):
    layout, fragments = packet()
    enable(layout, gaps)
    with pytest.raises(ReviewedRegionsError):
        validate(layout, fragments)


@pytest.mark.parametrize("mutation", [
    lambda l: l.pop("gutter_policy"),
    lambda l: l.update(gutter_policy="automatic"),
    lambda l: l["body"][0].pop("column_gaps_px"),
    lambda l: l["body"][0].update(cell_padding=5),
    lambda l: l.update(version="reviewed_region_layout_v1"),
])
def test_version_and_policy_never_silently_promote_gutters(mutation):
    layout, fragments = packet()
    enable(layout, [40])
    mutation(layout)
    with pytest.raises(ReviewedRegionsError):
        validate(layout, fragments)


def test_empty_cells_do_not_invent_support_for_a_gutter():
    layout, fragments = packet()
    table = layout["body"][0]
    table["column_widths"] = [35, 10, 55]
    table["rows"][0]["cells"].insert(1, {"fragment_ids": []})
    enable(layout, [0, 0])
    assert json.loads(validate(layout, fragments)) == layout
    table["column_gaps_px"] = [0, 1]
    with pytest.raises(ReviewedRegionsError, match="region_unsupported_column_gap"):
        validate(layout, fragments)


def test_sparse_later_rows_preserve_the_supported_gutter_and_empty_cell():
    layout, fragments = packet()
    table = layout["body"][0]
    # Source-rank preserving populated row followed by a sparse signature row.
    table["rows"].append({"cells": [{"fragment_ids": []}, {"fragment_ids": [fid(6)]}]})
    layout["body"].pop()
    enable(layout, [40])
    result = json.loads(validate(layout, fragments))
    assert result["body"][0]["rows"][1]["cells"][0] == {"fragment_ids": []}
    assert result["body"][0]["column_gaps_px"] == [40]


def formatting_projection(lang="FR", *, gutters=True, widths=None):
    parts = None
    if lang == "AR":
        parts = ["المحكمة القضائية\nالدائرة الجنائية\n\n",
                 "يُحافَظ على مهلة عشرة أيام.\nيحضر [[João Guerreiro]] في [[09]]:[[30]].\n\n",
                 "القاضية\n[[Nome Exemplo]]\n\n",
                 "[[Largo do Exemplo - 1234-567 Exemplo]]\nالهاتف: [[123456789]]\n", "1 / 1"]
    manifest, page = region_packet(columns=True, target_parts=parts)
    manifest["target_lang"] = lang
    layout = manifest["pages"][0]["region_layout"]
    if gutters:
        enable(layout, [20])
    if widths:
        layout["body"][0]["column_widths"] = widths
    return validate_formatting(manifest, page, target_lang=lang)


@pytest.mark.parametrize("lang", ["FR", "EN", "AR"])
def test_docx_gutter_reserves_space_without_text_font_or_outer_width_changes(lang):
    projection = formatting_projection(lang)
    artifact = build_reviewed_docx(projection)
    doc, mapping = decoded(artifact)
    validate_reviewed_docx(artifact.docx_bytes, artifact.source_map_bytes, projection=projection)
    assert mapping["table_policy"] == GUTTER_TABLE_POLICY
    table_map = mapping["pages"][0]["tables"][0]
    # Fixture: 20px on a600px source page, mapped to595.276pt, rounded to397twips.
    assert table_map["column_gaps_px"] == [20]
    assert table_map["column_gaps_twips"] == [397]
    assert table_map["cell_margins_twips"] == [{"left": 0, "right": 198}, {"left": 199, "right": 0}]
    assert sum(table_map["column_widths_twips"]) == 9978
    for column, expected in enumerate(((0, 198), (199, 0))):
        props = doc.tables[0].cell(0, column)._tc.tcPr
        margins = props.find(qn("w:tcMar"))
        assert [int(margins.find(qn("w:" + side)).get(qn("w:w"))) for side in ("left", "right")] == list(expected)
        assert list(props).index(margins) < list(props).index(props.find(qn("w:vAlign")))
    for row in mapping["pages"][0]["fragments"]:
        p = located(doc, row)
        text = p.text.replace("\u200e", "") if lang == "AR" else p.text
        assert hashlib.sha256(text.encode()).hexdigest() == row["display_text_sha256"]
        assert all(run.font.size.pt == (11 if lang == "AR" else 10.5) for run in p.runs)
    assert mapping["rendered_layout_acceptance"] == "not_evaluated"


def test_existing_region_version_keeps_zero_margins_and_map_policy():
    doc, mapping = decoded(build_reviewed_docx(formatting_projection(gutters=False)))
    assert mapping["table_policy"] == TABLE_POLICY
    assert "column_gaps_px" not in mapping["pages"][0]["tables"][0]
    assert not list(doc._element.iter(qn("w:tcMar")))


def test_large_gutter_cannot_consume_a_narrow_column_text_area():
    projection = formatting_projection(widths=[1, 99])
    with pytest.raises(ReviewedFormattingWriterError, match="column_gutter_leaves_no_text_width"):
        build_reviewed_docx(projection)


@pytest.mark.parametrize("side", ["left", "right"])
def test_modified_physical_cell_margins_are_rejected_after_rehash(side):
    projection = formatting_projection()
    artifact = build_reviewed_docx(projection)
    def mutate(doc):
        node = doc.tables[0].cell(0, 0)._tc.tcPr.find(qn("w:tcMar")).find(qn("w:" + side))
        node.set(qn("w:w"), "1")
    raw, mapping = changed_document(artifact, mutate)
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(raw, mapping, projection=projection)


def test_modified_gutter_source_map_is_rejected():
    projection = formatting_projection()
    artifact = build_reviewed_docx(projection)
    mapping = json.loads(artifact.source_map_bytes)
    mapping["pages"][0]["tables"][0]["column_gaps_twips"] = [0]
    with pytest.raises(ReviewedFormattingWriterError):
        validate_reviewed_docx(artifact.docx_bytes, json.dumps(mapping).encode(), projection=projection)
