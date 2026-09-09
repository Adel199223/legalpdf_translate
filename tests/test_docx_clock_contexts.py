"""Offline saved-DOCX clock direction checks, not native visual acceptance."""

from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import pytest

from legalpdf_translate.document_structure import PageStructure, StructureBlock, text_sha256
from legalpdf_translate.docx_writer import assemble_docx, sanitize_bidi_controls, unwrap_internal_placeholders
from legalpdf_translate.types import TargetLang


@pytest.fixture(autouse=True)
def _no_native_processes(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Offline clock fixture must not launch a native process.")
    monkeypatch.setattr(subprocess, "Popen", forbidden)


def _visible(text: str) -> str:
    return sanitize_bidi_controls(unwrap_internal_placeholders(text))


def _snapshot(folder: Path) -> dict[str, str]:
    return {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in folder.iterdir()}


def _saved_pages(folder: Path, clock: str) -> dict[str, str]:
    """Bound synthetic source/target pairs; no OCR or provider invocation."""
    folder.mkdir()
    time = _visible(clock)
    target_texts = {
        "header": f"المحكمة القضائية - الساعة {clock}",
        "heading": f"موعد الجلسة {clock}",
        "body": f"تبدأ الجلسة في الساعة {clock} أمام المحكمة.",
        "list": f"1) الحضور عند الساعة {clock} دون تأخير.",
        "table": f"وقت الجلسة {clock}",
        "footer": f"هاتف: [[210000000]] - ساعات الاتصال {clock}",
    }
    definitions = [
        ("header", dict(text=f"Tribunal Judicial - Horário {time}", role="header", bbox=(80, 70, 510, 90), alignment="center", bold=True)),
        ("heading", dict(text=f"Hora da audiência {time}", role="heading", bbox=(80, 160, 510, 180), alignment="center", bold=True)),
        ("body", dict(text=f"A audiência começa às {time} no tribunal.", role="paragraph", bbox=(50, 210, 540, 245), alignment="justify")),
        ("list", dict(text=f"1) Comparecer às {time} sem atraso.", role="list_item", bbox=(50, 280, 540, 315), alignment="right")),
        ("table", dict(text=f"Hora {time}", role="table_cell", bbox=(50, 350, 540, 380), table_id="clock_table", row=0, col=0, alignment="left", italic=True)),
        ("footer", dict(text=f"Telef: 210000000 - Horário {time}", role="footer", bbox=(80, 800, 510, 820), alignment="center")),
    ]
    for number in (1, 2):
        source = PageStructure(
            page_number=number, source_sha256="a" * 64, source_file_sha256="b" * 64,
            document_start=number == 1,
            blocks=[StructureBlock(f"p{number:04d}_b{index:04d}", **values)
                    for index, (_, values) in enumerate(definitions, 1)],
        )
        source.source_sha256 = source.source_text_sha256 = text_sha256(source.text)
        target = deepcopy(source.to_dict())
        for block, (key, _) in zip(target["blocks"], definitions):
            block["text"] = target_texts[key]
        translated = "\n".join(block["text"] for block in target["blocks"])
        target["translation_sha256"] = text_sha256(translated)
        page = folder / f"page_{number:04d}.txt"
        page.write_text(translated, encoding="utf-8")
        page.with_suffix(".source_structure.json").write_text(json.dumps(source.to_dict(), ensure_ascii=False), encoding="utf-8")
        page.with_suffix(".structure.json").write_text(json.dumps(target, ensure_ascii=False), encoding="utf-8")
    return target_texts


def _build(tmp_path: Path, clock: str, strip: bool):
    pages = tmp_path / "pages"
    texts = _saved_pages(pages, clock)
    before = _snapshot(pages)
    output = assemble_docx(pages, tmp_path / "clocks.docx", lang=TargetLang.AR,
                           page_breaks=False, page_numbers=[1, 2], strip_bidi_controls=strip)
    assert _snapshot(pages) == before
    document = Document(output)
    mapping = json.loads(output.with_suffix(".source_map.json").read_text(encoding="utf-8"))
    assert len(document.sections) == 1
    assert all(page["section_furniture_adopted"] for page in mapping["pages"])
    assert len([block for page in mapping["pages"] for block in page["blocks"]]) == 12
    assert len(document.tables) == 2
    paragraphs = list(document.paragraphs)
    paragraphs += [p for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
    paragraphs += list(document.sections[0].header.paragraphs) + list(document.sections[0].footer.paragraphs)
    by_context = {key: [p for p in paragraphs if _visible(p.text) == _visible(text)]
                  for key, text in texts.items()}
    assert {key: len(value) for key, value in by_context.items()} == {
        "header": 1, "heading": 2, "body": 2, "list": 2, "table": 2, "footer": 1,
    }
    with ZipFile(output) as archive:
        for part in archive.namelist():
            if part.startswith("word/") and part.endswith(".xml"):
                xml = archive.read(part).decode("utf-8")
                assert "[[" not in xml and "]]" not in xml
    return document, by_context


def _assert_clock_run(paragraph, clock: str):
    runs = [run for run in paragraph.runs if clock in _visible(run.text)]
    assert len(runs) == 1, [(run.text, run._r.xml) for run in paragraph.runs]
    assert _visible(runs[0].text) == clock
    assert runs[0]._r.get_or_add_rPr().find(qn("w:rtl")).get(qn("w:val")) == "0"
    # Adjacent Arabic text must not be pulled into the LTR clock run.
    assert all(not ("\u0600" <= c <= "\u06ff") for c in _visible(runs[0].text))


@pytest.mark.parametrize("context", ["body", "list", "table", "heading", "header", "footer"])
@pytest.mark.parametrize("clock", ["09:30", "09:30:05"])
@pytest.mark.parametrize("strip", [True, False])
def test_split_clock_is_one_ltr_run_in_every_saved_document_context(tmp_path, context, clock, strip):
    wrapped = ":".join(f"[[{piece}]]" for piece in clock.split(":"))
    document, paragraphs = _build(tmp_path, wrapped, strip)
    for paragraph in paragraphs[context]:
        _assert_clock_run(paragraph, clock)
    assert " PAGE " in document.sections[0].footer._element.xml


@pytest.mark.parametrize("clock", ["00:00", "09:30", "23:59:05"])
@pytest.mark.parametrize("shape", ["plain", "whole", "isolated_split"])
@pytest.mark.parametrize("strip", [True, False])
def test_plain_whole_and_isolated_clock_shapes_preserve_all_context_text(tmp_path, clock, shape, strip):
    text = clock if shape == "plain" else f"[[{clock}]]" if shape == "whole" else ":".join(
        f"\u2066[[{piece}]]\u2069" for piece in clock.split(":"))
    _, paragraphs = _build(tmp_path, text, strip)
    for context in paragraphs.values():
        for paragraph in context:
            _assert_clock_run(paragraph, clock)


@pytest.mark.parametrize("strip", [True, False])
def test_clock_fix_preserves_readable_font_emphasis_and_context_alignment(tmp_path, strip):
    document, paragraphs = _build(tmp_path, "[[09]]:[[30]]", strip)
    expected_alignment = {"body": "both", "list": "start", "table": "end",
                          "heading": "center", "header": "center", "footer": "center"}
    for key, items in paragraphs.items():
        for paragraph in items:
            assert paragraph._p.get_or_add_pPr().find(qn("w:jc")).get(qn("w:val")) == expected_alignment[key]
            assert paragraph._p.get_or_add_pPr().find(qn("w:bidi")) is not None
            for run in paragraph.runs:
                if not run.text:
                    continue
                assert run.font.size.pt == 11
                props = run._r.get_or_add_rPr()
                assert props.find(qn("w:szCs")).get(qn("w:val")) == "22"
                fonts = props.find(qn("w:rFonts"))
                assert all(fonts.get(qn("w:" + field)) == "Arial" for field in ("ascii", "hAnsi", "cs"))
                if key in {"heading", "header"}:
                    assert run.bold is True and run.font.cs_bold is True
                if key == "table":
                    assert run.italic is True and run.font.cs_italic is True
    assert document.styles["Normal"].font.size.pt == 11


@pytest.mark.parametrize("lang", [TargetLang.EN, TargetLang.FR])
def test_clock_fix_does_not_change_non_arabic_default_font_size(tmp_path, lang):
    pages = tmp_path / "pages"
    pages.mkdir()
    (pages / "page_0001.txt").write_text("09:30 23:59:05", encoding="utf-8")
    output = assemble_docx(pages, tmp_path / "latin.docx", lang=lang, page_breaks=False)
    document = Document(output)
    assert document.paragraphs[0].text == "09:30 23:59:05"
    for run in document.paragraphs[0].runs:
        assert run.font.name == "Times New Roman" and run.font.size.pt == 10.5
