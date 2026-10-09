"""Fictional ordinary completion through bounded proposals and normal delivery."""
from __future__ import annotations

from copy import deepcopy
import hashlib
from io import BytesIO
import json
from pathlib import Path
from types import SimpleNamespace
import uuid
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from legalpdf_translate.budget_reservations import ReservationBudget
from legalpdf_translate.openai_client import OpenAIResponsesClient
from legalpdf_translate.ordinary_auto_layout import frozen_automatic_layout_policy, run_automatic_layout
from legalpdf_translate.ordinary_auto_layout_artifacts import bind_raw_page_map
from legalpdf_translate.ordinary_layout_accounting import layout_accounting_policy
from legalpdf_translate.ordinary_layout_contracts import LayoutSuggestionPolicy, OrdinaryLayoutJob, PROPOSAL_VERSION
from legalpdf_translate.ordinary_layout_contracts import OrdinaryLayoutError
from legalpdf_translate.ordinary_layout_manager import OrdinaryLayoutManager
from legalpdf_translate.ordinary_layout_integration import delivery_job_snapshot
from legalpdf_translate.usage_accounting import DispatchAccounting
from tests.test_ordinary_auto_layout_artifacts import _raw, _source_pdf


def _proposal(view, page):
    rows = []
    for choice in view["decisions"]["paragraphs"]:
        if {region["page_number"] for region in choice["regions"]} != {page}:
            continue
        row = deepcopy(choice)
        row.pop("regions"); row.pop("unmapped_reason")
        row["bbox"] = [0.08, 0.08, 0.92, 0.9]
        if not rows:
            row.update(role="heading", heading_level=1, heading_size_pt=12, bold=True)
        rows.append(row)
    assert rows
    return {"version": PROPOSAL_VERSION, "page_number": page, "paragraphs": rows,
        "bands": [{"kind": "flow", "groups": [{"paragraph_ids": [row["paragraph_id"]],
            "panel": index == 0} for index, row in enumerate(rows)]}]}


def test_two_physical_pages_publish_unreviewed_default_and_reuse_without_dispatch(tmp_path):
    raw, mapping = _raw(tmp_path)
    source = _source_pdf()
    source_hash = hashlib.sha256(source).hexdigest()
    mapped = bind_raw_page_map(raw, mapping, source_pdf_sha256=source_hash,
                               selected_pages=(2, 3), target_lang="EN")
    canonical_map_hash = mapped.raw_source_map_sha256
    byte_map_hash = hashlib.sha256(json.dumps(mapping).encode("utf-8")).hexdigest()
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    policy = frozen_automatic_layout_policy()
    calls = []
    manager_slot = {}
    budget = ReservationBudget(tmp_path / "budget.json", cap_usd="2",
                               identity={"fixture": "two-page-auto"})
    accounting_args = layout_accounting_policy().accounting_arguments()

    def make_job(job_id):
        return OrdinaryLayoutJob(job_id, "shadow", "fixture", "run-fictional", source, raw, raw,
            "EN", (2, 3), {"source_sha256": source_hash,
                "original_sha256": hashlib.sha256(raw).hexdigest(),
                "raw_source_map_sha256": canonical_map_hash}, mapped.page_ids,
            hashlib.sha256(raw).hexdigest())

    jobs = {key: make_job(key) for key in ("tx-first", "tx-second", "tx-third")}

    class TrustedJobs:
        def trusted_ordinary_raw_map(self, job_id, **_scope):
            return deepcopy(mapping)

        def trusted_ordinary_baseline_identity(self, job_id, **_scope):
            return {"run_dir": str(run_dir), "run_id": "run-fictional",
                "source_pdf_sha256": source_hash,
                "raw_docx_sha256": hashlib.sha256(raw).hexdigest(),
                "raw_source_map_bytes_sha256": byte_map_hash,
                "selected_pages": [2, 3], "target_lang": "EN",
                "runtime_mode": "shadow", "workspace_id": "fixture"}

    class Accounting:
        jobs = TrustedJobs()

        def authorize_automatic(self, job, policy_fingerprint):
            assert job.job_id in jobs and policy_fingerprint == policy.split(":", 1)[1]

    def provider(_job, selected_policy):
        def create(**request):
            page = (2, 3)[len(calls)]
            calls.append(request)
            view = manager_slot["manager"].service.state(_job.job_id)["review"]
            return SimpleNamespace(id=f"fictional-{page}", model="gpt-5.2",
                status="completed", service_tier="default", output=[],
                output_text=json.dumps(_proposal(view, page)),
                usage={"input_tokens": 80, "output_tokens": 60, "total_tokens": 140,
                    "input_tokens_details": {"cached_tokens": 0},
                    "output_tokens_details": {"reasoning_tokens": 0}})
        return OpenAIResponsesClient(model=selected_policy.model,
            max_transport_retries=0, pre_call_jitter_seconds=0,
            sdk_client=SimpleNamespace(base_url="https://api.openai.com/v1/",
                                       responses=SimpleNamespace(create=create)))

    def accountant(job, operation_nonce, operation_dir, _policy):
        return DispatchAccounting(operation_dir / "accounting",
            run_identity={"job_id": job.job_id, "operation_nonce": operation_nonce},
            budget_context=budget, **accounting_args)

    def make_manager(job_id):
        manager = OrdinaryLayoutManager(tmp_path / "app", mode="shadow", workspace_id="fixture",
            job_resolver=lambda _: jobs[job_id], provider_factory=provider,
            accounting_factory=accountant,
            suggestion_policy=LayoutSuggestionPolicy("gpt-5.2", ".812", "1.624"))
        manager_slot["manager"] = manager
        return manager

    first = make_manager("tx-first")
    result = run_automatic_layout(first, Accounting(), "tx-first", policy)
    assert len(calls) == 2 and result["delivery_kind"] == "automatic_unreviewed"
    assert result["layout_costs"]["complete"] is True
    assert Path(result["output_path"]).read_bytes()
    state = first.state("tx-first")
    assert state["delivery"]["kind"] == "automatic_unreviewed"
    assert state["delivery"]["document_reviewed"] is False
    assert first.resolve_delivery("tx-first", 0).sha256 == result["sha256"]
    job_snapshot = {"job_id": "tx-first", "result": {"save_seed": {
        "run_id": "run-fictional", "output_docx": str(tmp_path / "raw.docx"),
        "estimated_api_cost": 0.0, "api_cost": 0.0}},
        "artifacts": {"output_docx": str(tmp_path / "raw.docx")}}
    selected, delivered = delivery_job_snapshot(first, job_snapshot)
    assert delivered.kind == "automatic_unreviewed"
    assert selected["result"]["save_seed"]["output_docx"] == str(delivered.path)
    assert selected["delivery"]["sha256"] == result["sha256"]
    retained = delivered.path.read_bytes()
    delivered.path.write_bytes(retained + b"user edit")
    with pytest.raises(OrdinaryLayoutError, match="automatic_candidate_changed"):
        delivery_job_snapshot(first, job_snapshot)
    delivered.path.write_bytes(retained)
    marker_path = run_dir / "ordinary_auto_layout" / "candidate.json"
    marker_path.unlink()
    early_pointer = json.loads((run_dir / "ordinary_auto_layout" / "operation.json").read_text(encoding="utf-8"))
    early_operation = (first.service.root / early_pointer["origin_job_id"] / "baselines" /
        early_pointer["baseline_id"] / "suggestions" / early_pointer["operation_nonce"])
    (early_operation / "result.json").unlink()
    recovered = run_automatic_layout(make_manager("tx-second"), Accounting(), "tx-second", policy)
    assert recovered["sha256"] == result["sha256"] and len(calls) == 2
    reused = run_automatic_layout(make_manager("tx-second"), Accounting(), "tx-second", policy)
    assert reused["reused_durable_candidate"] is True
    assert reused["sha256"] == result["sha256"] and len(calls) == 2
    second = make_manager("tx-second")
    alias = second.state("tx-second")
    assert alias["automatic_alias"] is True and alias["editor_rebase_required"] is True
    assert alias["delivery"]["kind"] == "automatic_unreviewed"
    assert second.resolve_delivery("tx-second", 0).sha256 == result["sha256"]
    alias_job = deepcopy(job_snapshot)
    alias_job["job_id"] = "tx-second"
    alias_snapshot, alias_delivery = delivery_job_snapshot(second, alias_job, mutation=True,
        baseline_id=alias["baseline_id"], expected_delivery_generation=0)
    assert alias_delivery.sha256 == result["sha256"]
    assert alias_snapshot["layout_costs"]["cost_usd"] == result["layout_costs"]["cost_usd"]
    working_path = first.service.automatic_review_copy("tx-first")
    edited_doc = Document(BytesIO(retained))
    edited_doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    edited_bytes = BytesIO()
    edited_doc.save(edited_bytes)
    working_path.write_bytes(edited_bytes.getvalue())
    with pytest.raises(OrdinaryLayoutError, match="edited_revision_changes_detected"):
        first.service.adopt_automatic_word_edit("tx-first", without_changes=True)
    edited_view = first.service.adopt_automatic_word_edit("tx-first")
    assert edited_view["delivery"]["kind"] == "automatic_unreviewed_edited"
    selected_edit, edited_delivery = delivery_job_snapshot(first, job_snapshot, mutation=True,
        baseline_id=state["baseline_id"], expected_delivery_generation=0)
    assert edited_delivery.path.read_bytes() == edited_bytes.getvalue()
    assert edited_delivery.sha256 == hashlib.sha256(edited_bytes.getvalue()).hexdigest()
    assert selected_edit["result"]["save_seed"]["output_docx"] == str(edited_delivery.path)
    altered_doc = Document(BytesIO(edited_bytes.getvalue()))
    altered_doc.paragraphs[0].add_run(" changed legal text")
    altered_bytes = BytesIO()
    altered_doc.save(altered_bytes)
    working_path.write_bytes(altered_bytes.getvalue())
    with pytest.raises(OrdinaryLayoutError, match="edited_layout_rebase_required"):
        first.service.adopt_automatic_word_edit("tx-first")
    assert working_path.read_bytes() == altered_bytes.getvalue()
    header_doc = Document(BytesIO(edited_bytes.getvalue()))
    header_doc.sections[0].header.paragraphs[0].text = "Injected visible header"
    header_bytes = BytesIO()
    header_doc.save(header_bytes)
    working_path.write_bytes(header_bytes.getvalue())
    with pytest.raises(OrdinaryLayoutError, match="edited_layout_rebase_required"):
        first.service.adopt_automatic_word_edit("tx-first")
    field_doc = Document(BytesIO(edited_bytes.getvalue()))
    first_run = field_doc.paragraphs[0].runs[0]._r
    paragraph_node = first_run.getparent()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DATE")
    paragraph_node.insert(paragraph_node.index(first_run), field)
    field.append(first_run)
    field_bytes = BytesIO()
    field_doc.save(field_bytes)
    working_path.write_bytes(field_bytes.getvalue())
    with pytest.raises(OrdinaryLayoutError, match="edited_layout_rebase_required"):
        first.service.adopt_automatic_word_edit("tx-first")
    working_path.write_bytes(edited_bytes.getvalue())
    frozen_edit = first.resolve_delivery("tx-first", 0, uuid.uuid4().hex)
    assert frozen_edit.sha256 == edited_delivery.sha256 and frozen_edit.frozen
    later_doc = Document(BytesIO(edited_bytes.getvalue()))
    later_doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    later_bytes = BytesIO()
    later_doc.save(later_bytes)
    working_path.write_bytes(later_bytes.getvalue())
    with pytest.raises(OrdinaryLayoutError, match="delivery_frozen"):
        first.service.adopt_automatic_word_edit("tx-first")
    assert first.resolve_delivery("tx-first", 0).sha256 == frozen_edit.sha256
    assert frozen_edit.path.read_bytes() == edited_bytes.getvalue()
    assert (run_dir / "ordinary_auto_layout" / "intent.json").is_file()
    marker_path.unlink()
    pointer = json.loads((run_dir / "ordinary_auto_layout" / "operation.json").read_text(encoding="utf-8"))
    operation = (first.service.root / pointer["origin_job_id"] / "baselines" /
        pointer["baseline_id"] / "suggestions" / pointer["operation_nonce"])
    (operation / "result.json").unlink()
    (operation / "page-0003.response.json").unlink()
    with pytest.raises(OrdinaryLayoutError, match="pending_or_uncertain"):
        run_automatic_layout(make_manager("tx-third"), Accounting(), "tx-third", policy)
    assert len(calls) == 2
