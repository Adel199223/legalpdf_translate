"""Offline, source-bound rebuild integration; synthetic text and no Word/API."""
from __future__ import annotations

from copy import deepcopy
import hashlib
import json
import socket
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import fitz
import pytest

from legalpdf_translate.docx_writer import assemble_docx, sanitize_bidi_controls
from legalpdf_translate.layout_integration import prepare_layout_rebuild
from legalpdf_translate.types import TargetLang
from tests.test_source_folio_layout import synthetic_pair


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Formatting-only rebuild must never use the network")

    monkeypatch.setattr(socket, "create_connection", forbidden)
    monkeypatch.setattr(socket.socket, "connect", forbidden)


def _saved_case(tmp_path, numbers=(5,), *, source_layout=None):
    pdf = tmp_path / "synthetic-original.pdf"
    with fitz.open() as document:
        for _ in range(max(numbers)):
            document.new_page(width=595.44, height=841.68)
        document.save(pdf)
    file_hash = hashlib.sha256(pdf.read_bytes()).hexdigest()
    folder = tmp_path / "pages"
    folder.mkdir()
    pairs = [synthetic_pair(number) for number in numbers]
    for source, target in pairs:
        assert "layout" not in source["metadata"]
        assert "layout" not in target["metadata"]
        source["source_file_sha256"] = target["source_file_sha256"] = file_hash
        if source_layout is not None:
            source["metadata"]["layout"] = deepcopy(source_layout)
        path = folder / f"page_{source['page_number']:04d}.txt"
        path.write_text("\n".join(row["text"] for row in target["blocks"]), encoding="utf-8")
        path.with_suffix(".source_structure.json").write_text(json.dumps(source, ensure_ascii=False), encoding="utf-8")
        path.with_suffix(".structure.json").write_text(json.dumps(target, ensure_ascii=False), encoding="utf-8")
    originals = {path: path.read_bytes() for path in (pdf, *folder.iterdir())}
    return pdf, folder, pairs, originals


def _rebuild(tmp_path, numbers=(5,), *, page_breaks=False, source_layout=None):
    pdf, folder, pairs, originals = _saved_case(tmp_path, numbers, source_layout=source_layout)
    pair_snapshot = deepcopy(pairs)
    prepared = prepare_layout_rebuild(folder, pdf, cache_dir=tmp_path / "layout-cache")
    assert prepared == {"prepared_pages": list(numbers), "review_required_pages": [],
                        "legacy_pages": [], "cancelled": False}
    for number in numbers:
        derivative = json.loads((folder / f"page_{number:04d}.layout.json").read_text(encoding="utf-8"))
        assert derivative["layout"]["status"] == "flow"
        assert not derivative["layout"]["review_required"]
        assert derivative["layout"]["warnings"] == []
    output = assemble_docx(folder, tmp_path / "rebuilt.docx", lang=TargetLang.AR, page_breaks=page_breaks)
    assert all(path.read_bytes() == value for path, value in originals.items())
    assert pairs == pair_snapshot
    mapping = json.loads(output.with_suffix(".source_map.json").read_text(encoding="utf-8"))
    assert mapping["docx_sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert mapping["source_page_count"] == len(numbers)
    assert mapping["rendered_page_count"] is None  # XML is not native Word acceptance.
    return output, Document(output), mapping, pairs


def _mapped_paragraph(document, row):
    location = row["location"]
    if location["kind"] == "body_paragraph":
        return document.paragraphs[location["paragraph_index"]]
    if location["kind"] in {"section_header", "section_footer"}:
        part = getattr(document.sections[location["section_index"]], location["kind"][8:])
        assert str(part.part.partname) == location["part_uri"]
        return part.paragraphs[location["paragraph_index"]]
    assert location["kind"] == "generated_footer_page_field"
    return None


def _assert_complete_editable_mapping(document, mapping, pairs):
    expected = {row["id"]: row for _, target in pairs for row in target["blocks"]}
    rows = [row for page in mapping["pages"] for row in page["blocks"]]
    assert len(rows) == len(expected) == 20 * len(pairs)
    assert len({row["block_id"] for row in rows}) == len(rows)
    assert {row["block_id"] for row in rows} == set(expected)
    for row in rows:
        paragraph = _mapped_paragraph(document, row)
        if paragraph is None:
            assert row["block_id"].endswith("_b0019")
            continue
        text = expected[row["block_id"]]["text"]
        if "furniture_alias" in row:
            assert row["furniture_alias"]["original_target_text"] == text
        assert text in sanitize_bidi_controls(paragraph.text)
        for run in paragraph.runs:
            if not run._r.xpath("./w:t"):
                continue  # The existing explicit page-break run has no glyphs.
            assert run.font.size.pt == 11
            fonts = run._r.rPr.find(qn("w:rFonts"))
            assert all(fonts.get(qn(f"w:{slot}")) == "Arial" for slot in ("ascii", "hAnsi", "cs"))
            assert run._r.rPr.find(qn("w:szCs")).get(qn("w:val")) == "22"
    for page in mapping["pages"]:
        assert page["source_block_coverage_status"] == "complete"
        assert page["layout_status"] == "flow"
        assert "ambiguous_side_by_side_geometry" not in page["layout_warnings"]
        assert "section_furniture_source_evidence_unavailable" not in page["layout_warnings"]
    assert "[[" not in document._element.xml
    for section in document.sections:
        assert " PAGE " in section.footer._element.xml


@pytest.mark.parametrize("page_breaks", [False, True])
def test_single_retained_page_rebuild_keeps_twenty_ids_and_body_header(tmp_path, page_breaks):
    output, document, mapping, pairs = _rebuild(tmp_path, page_breaks=page_breaks)
    _assert_complete_editable_mapping(document, mapping, pairs)
    if page_breaks:
        assert "section_furniture" not in mapping
        assert all(row["location"]["kind"] == "body_paragraph"
                   for row in mapping["pages"][0]["blocks"] if not row["block_id"].endswith("_b0019"))
        with ZipFile(output) as archive:
            assert not any(name.startswith("word/header") and name.endswith(".xml") for name in archive.namelist())
    else:
        section, = mapping["section_furniture"]["sections"]
        assert section["adoption_basis"] == "isolated_source_contact_footer_v1"
        assert section["parts"]["header"] is None
        for row in mapping["pages"][0]["blocks"]:
            suffix = row["block_id"].split("_b")[1]
            expected = ("section_footer" if suffix in {"0018", "0020"} else
                        "generated_footer_page_field" if suffix == "0019" else "body_paragraph")
            assert row["location"]["kind"] == expected
        assert not any(p.text for p in document.sections[0].header.paragraphs)
    assert not document._element.xpath(".//w:br[@w:type='page']")
    page = mapping["pages"][0]
    if page_breaks:
        assert "source_spacing" not in page
    else:
        spacing = page["source_spacing"]
        assert spacing["status"] == "source_bound"
        assert spacing["line_height_evidence"] == "unavailable"
        gap = next(row for row in spacing["applied"] if row["block_id"] == "p0005_b0003")
        assert gap["basis"] == "source_header_body_gap"
        assert gap["desired_gap_pt"] == gap["applied_space_before_pt"] == 24
        assert gap["source_gap_pt"] == pytest.approx(28.8)
        paragraph = _mapped_paragraph(document, page["blocks"][2])
        assert paragraph.paragraph_format.space_before.pt >= 24
        body = _mapped_paragraph(document, page["blocks"][10])
        assert body.paragraph_format.line_spacing == pytest.approx(1.1)


@pytest.mark.parametrize("page_breaks", [False, True])
def test_changing_source_folios_preserve_two_pages_and_every_contact(tmp_path, page_breaks):
    output, document, mapping, pairs = _rebuild(tmp_path, (5, 6), page_breaks=page_breaks)
    _assert_complete_editable_mapping(document, mapping, pairs)
    assert [source["blocks"][18]["text"] for source, _ in pairs] == ["3/5", "4/5"]
    assert not any(paragraph.text in {"3/5", "4/5"} for paragraph in document.paragraphs)
    if page_breaks:
        assert "section_furniture" not in mapping
        assert len(document._element.xpath(".//w:br[@w:type='page']")) == 1
        assert not any("furniture_alias" in row for page in mapping["pages"] for row in page["blocks"])
        assert all("source_spacing" not in page for page in mapping["pages"])
    else:
        section, = mapping["section_furniture"]["sections"]
        assert section["consolidated"] and section["page_numbers"] == [5, 6]
        assert not mapping["layout_review_required"]
        assert len(document.sections) == 1
        assert len(document.sections[0].header.paragraphs) == 2
        assert len(document.sections[0].footer.paragraphs) == 3
        assert not document._element.xpath(".//w:br[@w:type='page']")
        assert all(page["section_furniture_adopted"] for page in mapping["pages"])
        for page in mapping["pages"]:
            rows = {row["block_id"]: row for row in page["blocks"]}
            prefix = f"p{page['source_page_number']:04d}_b"
            assert rows[prefix + "0001"]["location"]["kind"] == "section_header"
            assert rows[prefix + "0002"]["location"]["kind"] == "section_header"
            assert rows[prefix + "0018"]["location"]["kind"] == "section_footer"
            assert rows[prefix + "0020"]["location"]["kind"] == "section_footer"
            assert rows[prefix + "0019"]["location"]["kind"] == "generated_footer_page_field"
            assert not any(row["block_id"] == prefix + "0003" for row in page["source_spacing"]["applied"])


def test_present_uncertain_source_layout_is_never_replaced_by_flow_for_spacing(tmp_path):
    _, document, mapping, pairs = _rebuild(tmp_path, source_layout={
        "status": "needs_review", "review_required": True, "bands": [], "warnings": ["synthetic_review"]})
    # The new target derivative can be valid, but it cannot erase an explicit
    # uncertain source-layout decision for this spacing inference.
    page = mapping["pages"][0]
    assert page["source_spacing"]["status"] == "unavailable"
    assert page["source_spacing"]["applied"] == []
    assert pairs[0][0]["metadata"]["layout"]["review_required"] is True
    assert _mapped_paragraph(document, page["blocks"][2]).paragraph_format.space_before is None
